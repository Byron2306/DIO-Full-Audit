from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
import subprocess
import tempfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PptxInches, Pt as PptxPt


ROOT = Path(__file__).resolve().parents[2]
PROFILE_PATH = ROOT / "config" / "format_profiles.json"
SUPPORTED_CHANNELS = {"docx", "pdf", "pptx", "html", "vtt"}
APPROVED_LANGUAGE_STATES = {"human_approved", "human_approved_crystallized"}
TEXT_BLOCKS = {
    "title", "heading", "paragraph", "learning_objective", "teacher_instruction",
    "learner_instruction", "question", "expected_response", "caption", "equation",
    "reference", "annexure", "transcript",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _content_digest(value: Any) -> str:
    encoded = json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _hex_rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _pptx_rgb(value: str) -> PptxRGBColor:
    value = value.lstrip("#")
    return PptxRGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-") or "asset"


def _localizable(block: dict[str, Any]) -> dict[str, Any]:
    return {
        key: deepcopy(block[key])
        for key in ("text", "title", "items", "headers", "rows", "caption", "alt_text", "criteria")
        if key in block
    }


def _block_strings(block: dict[str, Any]) -> list[str]:
    values: list[str] = []
    for key in ("text", "title", "caption"):
        if str(block.get(key) or "").strip():
            values.append(str(block[key]).strip())
    values.extend(str(item).strip() for item in block.get("items") or [] if str(item).strip())
    values.extend(str(item).strip() for item in block.get("headers") or [] if str(item).strip())
    for row in block.get("rows") or []:
        values.extend(str(item).strip() for item in row if str(item).strip())
    for row in block.get("criteria") or []:
        if isinstance(row, dict):
            values.extend(str(item).strip() for item in row.values() if str(item).strip())
    return values


def load_profiles() -> dict[str, Any]:
    return json.loads(PROFILE_PATH.read_text(encoding="utf-8"))


def validate_semantic_content(content: dict[str, Any]) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    if content.get("schema") != "dio.semantic_content.v1":
        errors.append("schema must be dio.semantic_content.v1")
    for key in ("object_id", "version", "title", "source_language"):
        if not str(content.get(key) or "").strip():
            errors.append(f"{key} is required")
    context = content.get("context") or {}
    for key in ("product", "artifact_type", "audience"):
        if not str(context.get(key) or "").strip():
            errors.append(f"context.{key} is required")
    blocks = content.get("blocks") or []
    if not blocks:
        errors.append("at least one semantic block is required")
    seen: set[str] = set()
    previous_heading = 0
    for index, block in enumerate(blocks, 1):
        block_id = str(block.get("block_id") or "")
        block_type = str(block.get("type") or "")
        if not block_id or block_id in seen:
            errors.append(f"block {index} has a missing or duplicate block_id")
        seen.add(block_id)
        if block_type not in TEXT_BLOCKS | {"bullet_list", "table", "rubric", "figure", "diagram", "page_break"}:
            errors.append(f"{block_id or index} has unsupported type {block_type or '(missing)'}")
        if block_type in TEXT_BLOCKS and not str(block.get("text") or "").strip():
            errors.append(f"{block_id} requires text")
        if block_type == "heading":
            level = int(block.get("level") or 1)
            if level > previous_heading + 1 and previous_heading:
                warnings.append(f"{block_id} skips a heading level")
            previous_heading = level
        if block_type in {"table", "rubric"}:
            headers = block.get("headers") or []
            rows = block.get("rows") or []
            if not headers or not rows:
                errors.append(f"{block_id} requires headers and rows")
            if headers and any(len(row) != len(headers) for row in rows):
                errors.append(f"{block_id} row width does not match its headers")
        if block_type in {"figure", "diagram"}:
            if not str(block.get("alt_text") or "").strip():
                errors.append(f"{block_id} requires alt_text")
            if not str(block.get("caption") or "").strip():
                errors.append(f"{block_id} requires a caption")
        if block_type == "figure" and not str(block.get("media_path") or "").strip():
            errors.append(f"{block_id} requires media_path")
        if block_type == "question" and float(block.get("marks") or 0) < 0:
            errors.append(f"{block_id} has a negative mark allocation")
    return {"passed": not errors, "errors": errors, "warnings": warnings, "block_count": len(blocks)}


def project_language(content: dict[str, Any], language: str, *, release_mode: bool) -> tuple[dict[str, Any], dict[str, Any]]:
    projected = deepcopy(content)
    source_language = str(content["source_language"])
    if language == source_language:
        projected["render_language"] = language
        return projected, {"language": language, "lane": "source", "status": "canonical", "stale_units": []}
    lane = (content.get("translations") or {}).get(language)
    if not lane:
        raise ValueError(f"No translation lane exists for {language}")
    if str(lane.get("source_version") or "") != str(content["version"]):
        raise ValueError(f"{language} translation lane is stale for source version {content['version']}")
    status = str(lane.get("status") or "human_review_required")
    if release_mode and status not in APPROVED_LANGUAGE_STATES:
        raise ValueError(f"{language} translation is not human-approved for release")
    overlays = {str(row.get("block_id") or ""): row for row in lane.get("blocks") or []}
    missing: list[str] = []
    stale: list[str] = []
    for block in projected["blocks"]:
        if not _localizable(block):
            continue
        overlay = overlays.get(block["block_id"])
        if not overlay:
            missing.append(block["block_id"])
            continue
        expected = _content_digest(_localizable(block))
        if str(overlay.get("source_hash") or "") != expected:
            stale.append(block["block_id"])
            continue
        for key, value in _localizable(overlay).items():
            block[key] = value
    if missing or stale:
        raise ValueError(f"Incomplete {language} translation lane; missing={missing}, stale={stale}")
    projected["title"] = str(lane.get("title") or projected["title"])
    projected["render_language"] = language
    projected["translation_status"] = status
    return projected, {"language": language, "lane": "translation", "status": status, "stale_units": stale}


def build_paragraph_semantic_content(
    *,
    object_id: str,
    version: str,
    title: str,
    source_language: str,
    source_rows: list[dict[str, str]],
    context: dict[str, Any],
    translations: dict[str, Any] | None = None,
    target_language: str | None = None,
    translation_status: str = "human_review_required",
) -> dict[str, Any]:
    blocks: list[dict[str, Any]] = []
    for index, row in enumerate(source_rows):
        blocks.append({
            "block_id": str(row.get("paragraph_id") or f"P{index + 1}"),
            "type": "title" if index == 0 else "paragraph",
            "text": str(row["text"]),
        })
    content: dict[str, Any] = {
        "schema": "dio.semantic_content.v1",
        "object_id": object_id,
        "version": version,
        "title": title,
        "source_language": source_language,
        "context": context,
        "blocks": blocks,
        "translations": {},
    }
    if translations and target_language:
        rows = []
        for block in blocks:
            translation = translations[block["block_id"]]
            rows.append({
                "block_id": block["block_id"],
                "source_hash": _content_digest(_localizable(block)),
                "text": str(translation.get("translated") or translation.get("target_text") or ""),
            })
        content["translations"][target_language] = {
            "source_version": version,
            "status": translation_status,
            "blocks": rows,
        }
    return content


def semantic_content_from_lingua_object(lingua: dict[str, Any]) -> dict[str, Any]:
    if lingua.get("schema") != "dio.lingua.semantic_object.v1":
        raise ValueError("Expected a dio.lingua.semantic_object.v1 object")
    origin = lingua.get("origin") or {}
    source = lingua.get("source") or {}
    source_units = source.get("units") or []
    if not source_units:
        raise ValueError("Lingua object contains no source units")
    type_map = {
        "title": "title", "heading": "heading", "learning_objective": "learning_objective",
        "teacher_instruction": "teacher_instruction", "learner_instruction": "learner_instruction",
        "question": "question", "expected_response": "expected_response", "caption": "caption",
        "reference": "reference", "transcript": "transcript", "hook": "heading",
        "labelled_instruction": "teacher_instruction", "campaign_body": "paragraph", "body": "paragraph",
    }
    blocks: list[dict[str, Any]] = []
    for index, unit in enumerate(source_units):
        unit_type = str(unit.get("unit_type") or "body")
        block_type = type_map.get(unit_type, "paragraph")
        text = str(unit.get("source_text") or "").strip()
        if index == 0 and block_type == "paragraph" and len(text) < 140:
            block_type = "title"
        blocks.append({
            "block_id": str(unit["unit_id"]),
            "type": block_type,
            "level": 1 if block_type == "heading" else None,
            "text": text,
        })
        if blocks[-1]["level"] is None:
            blocks[-1].pop("level")
    title_unit = next((block for block in blocks if block["type"] == "title"), None)
    title = str((title_unit or {}).get("text") or origin.get("artifact_id") or lingua["object_id"])
    content: dict[str, Any] = {
        "schema": "dio.semantic_content.v1",
        "object_id": str(lingua["object_id"]),
        "version": str(source.get("version") or "1.0.0"),
        "title": title.lstrip("# "),
        "source_language": str(source.get("language") or "English"),
        "context": {
            "product": str(origin.get("product") or "document_studio"),
            "artifact_type": str(origin.get("artifact_type") or "document"),
            "audience": str(origin.get("audience") or "controlled audience"),
            "subject": lingua.get("subject"),
            "grade": lingua.get("grade"),
            "curriculum_concept": lingua.get("curriculum_concept"),
        },
        "blocks": blocks,
        "translations": {},
    }
    source_by_id = {str(row["unit_id"]): row for row in source_units}
    for language, lane in (lingua.get("translations") or {}).items():
        overlays = []
        for translated in lane.get("units") or []:
            unit_id = str(translated.get("unit_id") or "")
            source_unit = source_by_id.get(unit_id)
            block = next((item for item in blocks if item["block_id"] == unit_id), None)
            if not source_unit or not block:
                raise ValueError(f"{language} contains unknown Lingua unit {unit_id}")
            if str(translated.get("source_hash") or "") != str(source_unit.get("source_hash") or ""):
                raise ValueError(f"{language} Lingua unit {unit_id} is stale")
            if str(translated.get("status") or "").startswith("stale"):
                raise ValueError(f"{language} Lingua unit {unit_id} is stale")
            overlays.append({
                "block_id": unit_id,
                "source_hash": _content_digest(_localizable(block)),
                "text": str(translated.get("target_text") or ""),
            })
        content["translations"][language] = {
            "source_version": content["version"],
            "status": str(lane.get("status") or "human_review_required"),
            "blocks": overlays,
            "title": next(
                (row["text"].lstrip("# ") for row in overlays if row["block_id"] == (title_unit or {}).get("block_id")),
                content["title"],
            ),
            "lingua_authority": {
                "semantic_object_id": lingua["object_id"],
                "source_document_hash": source.get("document_hash"),
            },
        }
    return content


def _set_repeat_table_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _shade_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_text(cell: Any, value: str, *, bold: bool = False, colour: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(value)
    run.bold = bold
    if colour:
        run.font.color.rgb = colour
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _add_answer_line(document: Document) -> None:
    paragraph = document.add_paragraph(" ")
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "666666")
    borders.append(bottom)


def _configure_docx(
    document: Document,
    projected: dict[str, Any],
    style: dict[str, Any],
    delivery: dict[str, Any],
    logo_path: Path | None = None,
) -> None:
    section = document.sections[0]
    page = style["page"]
    if page.get("orientation") == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(297), Mm(210)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
    top, right, bottom, left = page["margins_mm"]
    section.top_margin, section.right_margin = Mm(top), Mm(right)
    section.bottom_margin, section.left_margin = Mm(bottom), Mm(left)
    fonts, sizes, colours = style["fonts"], style["sizes_pt"], style["colours"]
    for name in ("Normal", "Body Text"):
        document.styles[name].font.name = fonts["body"]
        document.styles[name].font.size = Pt(sizes["body"])
    for name, key in (("Title", "title"), ("Subtitle", "subtitle"), ("Heading 1", "h1"), ("Heading 2", "h2"), ("Heading 3", "h3")):
        document.styles[name].font.name = fonts["display"]
        document.styles[name].font.size = Pt(sizes[key])
        document.styles[name].font.color.rgb = _hex_rgb(colours["accent"] if name != "Subtitle" else colours["muted"])
    document.styles["Caption"].font.name = fonts["body"]
    document.styles["Caption"].font.size = Pt(sizes["caption"])
    document.styles["Caption"].font.color.rgb = _hex_rgb(colours["muted"])
    if style.get("header"):
        header = section.header.paragraphs[0]
        header.text = style["header"]
        header.runs[0].font.name = fonts["display"]
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = _hex_rgb(colours["muted"])
    footer = section.footer.paragraphs[0]
    footer.add_run(str(style.get("footer") or ""))
    footer.add_run("    ")
    _add_page_number(footer)
    if style.get("cover"):
        if logo_path:
            logo = document.add_paragraph()
            logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            logo.add_run().add_picture(str(logo_path), width=Inches(1.35))
        document.add_paragraph(projected["title"], style="Title")
        context = projected.get("context") or {}
        subtitle = " | ".join(str(context.get(key)) for key in ("artifact_type", "audience") if context.get(key))
        document.add_paragraph(subtitle, style="Subtitle")
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in (
            ("Product", context.get("product")), ("Subject", context.get("subject")),
            ("Grade", context.get("grade")), ("Curriculum focus", context.get("curriculum_concept")),
            ("Language", projected.get("render_language")), ("Version", projected.get("version")),
        ):
            if value in (None, ""):
                continue
            cells = table.add_row().cells
            _set_cell_text(cells[0], label, bold=True, colour=_hex_rgb(colours["accent"]))
            _set_cell_text(cells[1], str(value))
        if delivery.get("review_label"):
            notice = document.add_paragraph("CONTROLLED REVIEW ASSET. Human approval is required before external release.")
            notice.paragraph_format.space_before = Pt(14)
            notice.runs[0].bold = True
            notice.runs[0].font.color.rgb = _hex_rgb(colours["accent_2"])
        document.add_page_break()
        heading_count = sum(block.get("type") == "heading" for block in projected["blocks"])
        if heading_count >= 2:
            document.add_heading("Contents", level=1)
            paragraph = document.add_paragraph()
            run = paragraph.add_run()
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
            separate = OxmlElement("w:fldChar")
            separate.set(qn("w:fldCharType"), "separate")
            placeholder = OxmlElement("w:t")
            placeholder.text = "Update this field in Word or LibreOffice to populate the table of contents."
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            run._r.extend([begin, instruction, separate, placeholder, end])
            document.add_page_break()


def _make_diagram(block: dict[str, Any], path: Path, style: dict[str, Any]) -> Path:
    width, height = 1400, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"
    bold_path = "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"
    font = ImageFont.truetype(font_path, 29)
    bold = ImageFont.truetype(bold_path, 31)
    colours = style["colours"]
    ink, accent, fill = "#" + colours["ink"], "#" + colours["accent"], "#" + colours["fill"]
    nodes = block.get("nodes") or []
    if not nodes:
        nodes = [{"id": "A", "label": block.get("caption") or "Diagram", "x": 0.5, "y": 0.5}]
    positions: dict[str, tuple[int, int, int, int]] = {}
    for index, node in enumerate(nodes):
        x = int(float(node.get("x", (index + 1) / (len(nodes) + 1))) * width)
        y = int(float(node.get("y", 0.5)) * height)
        box_w, box_h = 310, 120
        positions[str(node.get("id") or index)] = (x - box_w // 2, y - box_h // 2, x + box_w // 2, y + box_h // 2)
    for edge in block.get("edges") or []:
        source = positions.get(str(edge.get("from")))
        target = positions.get(str(edge.get("to")))
        if not source or not target:
            continue
        start = ((source[0] + source[2]) // 2, (source[1] + source[3]) // 2)
        end = ((target[0] + target[2]) // 2, (target[1] + target[3]) // 2)
        draw.line((start, end), fill=accent, width=5)
        if edge.get("label"):
            middle = ((start[0] + end[0]) // 2, (start[1] + end[1]) // 2)
            draw.text(middle, str(edge["label"]), font=font, fill=ink, anchor="mm", stroke_width=5, stroke_fill="white")
    for index, node in enumerate(nodes):
        box = positions[str(node.get("id") or index)]
        draw.rounded_rectangle(box, radius=12, fill=fill, outline=accent, width=4)
        draw.multiline_text(((box[0] + box[2]) // 2, (box[1] + box[3]) // 2), str(node.get("label") or node.get("id") or ""), font=bold, fill=ink, anchor="mm", align="center", spacing=6)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, "PNG", optimize=True)
    return path


def _resolve_media(block: dict[str, Any], source_root: Path, asset_dir: Path, style: dict[str, Any]) -> Path:
    if block["type"] == "diagram":
        return _make_diagram(block, asset_dir / f"{_slug(block['block_id'])}.png", style)
    media = Path(str(block["media_path"])).expanduser()
    media = media if media.is_absolute() else source_root / media
    if not media.is_file():
        raise FileNotFoundError(f"Figure source not found: {media}")
    return media.resolve()


def _set_picture_alt_text(inline_shape: Any, alt_text: str) -> None:
    doc_properties = inline_shape._inline.docPr
    doc_properties.set("descr", alt_text)
    doc_properties.set("title", alt_text[:120])


def _render_docx(
    projected: dict[str, Any],
    path: Path,
    style: dict[str, Any],
    delivery: dict[str, Any],
    source_root: Path,
    asset_dir: Path,
    template_path: Path | None = None,
    logo_path: Path | None = None,
) -> None:
    document = Document(str(template_path)) if template_path else Document()
    _configure_docx(document, projected, style, delivery, logo_path)
    colours = style["colours"]
    first_title_skipped = False
    question_number = 0
    for block in projected["blocks"]:
        kind = block["type"]
        if kind == "title" and style.get("cover") and not first_title_skipped:
            first_title_skipped = True
            continue
        if kind == "page_break":
            document.add_page_break()
        elif kind == "title":
            document.add_paragraph(block["text"], style="Title")
        elif kind == "heading":
            document.add_heading(block["text"], level=int(block.get("level") or 1))
        elif kind in {"paragraph", "reference", "transcript"}:
            paragraph = document.add_paragraph(block["text"])
            paragraph.paragraph_format.space_after = Pt(7)
        elif kind in {"learning_objective", "teacher_instruction", "learner_instruction", "expected_response"}:
            labels = {
                "learning_objective": "Learning objective", "teacher_instruction": "Teacher note",
                "learner_instruction": "Learner instruction", "expected_response": "Expected response",
            }
            table = document.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            _shade_cell(table.cell(0, 0), colours["fill"])
            paragraph = table.cell(0, 0).paragraphs[0]
            run = paragraph.add_run(labels[kind] + ": ")
            run.bold = True
            run.font.color.rgb = _hex_rgb(colours["accent"])
            paragraph.add_run(block["text"])
            document.add_paragraph().paragraph_format.space_after = Pt(1)
        elif kind == "question":
            question_number += 1
            table = document.add_table(rows=1, cols=2)
            table.autofit = False
            table.columns[0].width = Inches(6.25)
            table.columns[1].width = Inches(0.65)
            left, right = table.rows[0].cells
            number = str(block.get("number") or question_number)
            _set_cell_text(left, f"{number}. {block['text']}")
            _set_cell_text(right, f"({block.get('marks', 0):g})", bold=True, colour=_hex_rgb(colours["accent"]))
            right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for _ in range(max(0, int(block.get("answer_lines") or 0))):
                _add_answer_line(document)
        elif kind == "bullet_list":
            for item in block.get("items") or []:
                document.add_paragraph(str(item), style="List Bullet")
        elif kind in {"table", "rubric"}:
            if block.get("title"):
                heading = document.add_paragraph(str(block["title"]))
                heading.runs[0].bold = True
                heading.runs[0].font.color.rgb = _hex_rgb(colours["accent"])
            table = document.add_table(rows=1, cols=len(block["headers"]))
            table.style = "Table Grid"
            for index, header in enumerate(block["headers"]):
                _shade_cell(table.rows[0].cells[index], colours["fill"])
                _set_cell_text(table.rows[0].cells[index], str(header), bold=True, colour=_hex_rgb(colours["accent"]))
            if style.get("table_header_repeat"):
                _set_repeat_table_header(table.rows[0])
            for values in block["rows"]:
                cells = table.add_row().cells
                for index, value in enumerate(values):
                    _set_cell_text(cells[index], str(value))
            document.add_paragraph().paragraph_format.space_after = Pt(1)
        elif kind in {"figure", "diagram"}:
            media = _resolve_media(block, source_root, asset_dir, style)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture = paragraph.add_run().add_picture(str(media), width=Inches(float(block.get("width_inches") or 6.3)))
            _set_picture_alt_text(picture, str(block["alt_text"]))
            caption = document.add_paragraph(str(block["caption"]), style="Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "caption":
            document.add_paragraph(block["text"], style="Caption")
        elif kind == "equation":
            paragraph = document.add_paragraph(block["text"])
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.runs[0].font.name = style["fonts"]["mono"]
        elif kind == "annexure":
            document.add_page_break()
            document.add_heading(block["text"], level=1)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _add_pptx_textbox(slide: Any, text: str, *, x: float, y: float, w: float, h: float, size: float, colour: str, bold: bool = False) -> Any:
    shape = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(w), PptxInches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = PptxInches(0.08)
    frame.margin_top = frame.margin_bottom = PptxInches(0.05)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = PptxPt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = _pptx_rgb(colour)
    paragraph.font.name = "Liberation Sans"
    return shape


def _pptx_background(slide: Any, style: dict[str, Any], label: str) -> None:
    colours = style["colours"]
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = _pptx_rgb(colours["paper"])
    _add_pptx_textbox(slide, label, x=0.55, y=7.08, w=12.2, h=0.25, size=8, colour=colours["muted"])


def _set_pptx_alt_text(shape: Any, alt_text: str) -> None:
    properties = shape._element.xpath(".//p:cNvPr")
    if properties:
        properties[0].set("descr", alt_text)


def _render_pptx(
    projected: dict[str, Any],
    path: Path,
    style: dict[str, Any],
    delivery: dict[str, Any],
    source_root: Path,
    asset_dir: Path,
    template_path: Path | None = None,
    logo_path: Path | None = None,
) -> None:
    deck = Presentation(str(template_path)) if template_path else Presentation()
    deck.slide_width, deck.slide_height = PptxInches(13.333), PptxInches(7.5)
    blank = next((layout for layout in deck.slide_layouts if str(layout.name).casefold() == "blank"), deck.slide_layouts[-1])
    colours = style["colours"]
    title_slide = deck.slides.add_slide(blank)
    _pptx_background(title_slide, style, projected["object_id"])
    _add_pptx_textbox(title_slide, projected["title"], x=0.7, y=1.25, w=11.9, h=1.8, size=30, colour=colours["accent"], bold=True)
    context = projected.get("context") or {}
    subtitle = "\n".join(str(context.get(key)) for key in ("subject", "grade", "curriculum_concept", "audience") if context.get(key))
    _add_pptx_textbox(title_slide, subtitle, x=0.72, y=3.25, w=10.8, h=1.7, size=16, colour=colours["muted"])
    if logo_path:
        title_slide.shapes.add_picture(str(logo_path), PptxInches(11.0), PptxInches(0.45), width=PptxInches(1.55))
    for block in projected["blocks"]:
        if block["type"] in {"title", "page_break"}:
            continue
        slide = deck.slides.add_slide(blank)
        _pptx_background(slide, style, projected["object_id"])
        kind = block["type"]
        heading = str(block.get("title") or block.get("label") or kind.replace("_", " ").title())
        if kind == "heading":
            heading = block["text"]
        _add_pptx_textbox(slide, heading, x=0.62, y=0.35, w=11.9, h=0.72, size=22, colour=colours["accent"], bold=True)
        if kind in {"figure", "diagram"}:
            media = _resolve_media(block, source_root, asset_dir, style)
            with Image.open(media) as image:
                ratio = image.width / max(image.height, 1)
            height = 5.2
            width = min(10.8, height * ratio)
            x = (13.333 - width) / 2
            picture = slide.shapes.add_picture(str(media), PptxInches(x), PptxInches(1.18), width=PptxInches(width), height=PptxInches(height))
            picture.name = str(block["alt_text"])[:250]
            _set_pptx_alt_text(picture, str(block["alt_text"]))
            _add_pptx_textbox(slide, str(block["caption"]), x=0.9, y=6.45, w=11.5, h=0.45, size=10, colour=colours["muted"])
        elif kind in {"table", "rubric"}:
            rows, cols = len(block["rows"]) + 1, len(block["headers"])
            table = slide.shapes.add_table(rows, cols, PptxInches(0.72), PptxInches(1.28), PptxInches(11.9), PptxInches(5.55)).table
            for col in range(cols):
                table.columns[col].width = PptxInches(11.9 / cols)
            for col, value in enumerate(block["headers"]):
                cell = table.cell(0, col)
                cell.text = str(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _pptx_rgb(colours["fill"])
            for row_index, values in enumerate(block["rows"], 1):
                for col, value in enumerate(values):
                    table.cell(row_index, col).text = str(value)
            for row in range(rows):
                for col in range(cols):
                    frame = table.cell(row, col).text_frame
                    frame.word_wrap = True
                    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    for paragraph in frame.paragraphs:
                        paragraph.font.name = style["fonts"]["body"]
                        paragraph.font.size = PptxPt(max(9, min(14, 26 - rows - cols)))
                        paragraph.font.bold = row == 0
                        paragraph.font.color.rgb = _pptx_rgb(colours["ink"])
        else:
            if kind == "bullet_list":
                text = "\n".join(f"• {item}" for item in block.get("items") or [])
            else:
                text = str(block.get("text") or "")
                if kind == "question":
                    text += f"\n\nMarks: {block.get('marks', 0):g}"
            font_size = 22 if len(text) < 420 else 18 if len(text) < 850 else 14
            _add_pptx_textbox(slide, text, x=0.75, y=1.35, w=11.7, h=5.5, size=font_size, colour=colours["ink"])
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)


def _html_block(block: dict[str, Any], source_root: Path, asset_dir: Path, style: dict[str, Any], media_dir: Path) -> str:
    kind = block["type"]
    identifier = html.escape(str(block["block_id"]))
    if kind == "page_break":
        return '<hr class="page-break" aria-hidden="true">'
    if kind == "title":
        return f'<h1 id="{identifier}">{html.escape(block["text"])}</h1>'
    if kind == "heading":
        level = min(4, int(block.get("level") or 1) + 1)
        return f'<h{level} id="{identifier}">{html.escape(block["text"])}</h{level}>'
    if kind in {"paragraph", "reference", "transcript"}:
        return f'<p id="{identifier}">{html.escape(block["text"])}</p>'
    if kind in {"learning_objective", "teacher_instruction", "learner_instruction", "expected_response"}:
        label = kind.replace("_", " ").title()
        return f'<aside id="{identifier}" class="callout"><strong>{label}:</strong> {html.escape(block["text"])}</aside>'
    if kind == "question":
        return f'<section id="{identifier}" class="question"><p>{html.escape(block["text"])}</p><strong class="marks">({float(block.get("marks") or 0):g})</strong></section>'
    if kind == "bullet_list":
        return '<ul id="{}">{}</ul>'.format(identifier, "".join(f"<li>{html.escape(str(item))}</li>" for item in block.get("items") or []))
    if kind in {"table", "rubric"}:
        caption = f"<caption>{html.escape(str(block.get('title') or kind.title()))}</caption>"
        headers = "".join(f'<th scope="col">{html.escape(str(item))}</th>' for item in block["headers"])
        rows = "".join("<tr>" + "".join(f"<td>{html.escape(str(item))}</td>" for item in row) + "</tr>" for row in block["rows"])
        return f'<table id="{identifier}">{caption}<thead><tr>{headers}</tr></thead><tbody>{rows}</tbody></table>'
    if kind in {"figure", "diagram"}:
        media = _resolve_media(block, source_root, asset_dir, style)
        media_dir.mkdir(parents=True, exist_ok=True)
        target = media_dir / f"{_slug(block['block_id'])}{media.suffix.casefold()}"
        if media.resolve() != target.resolve():
            shutil.copy2(media, target)
        return f'<figure id="{identifier}"><img src="media/{html.escape(target.name)}" alt="{html.escape(block["alt_text"])}"><figcaption>{html.escape(block["caption"])}</figcaption></figure>'
    if kind == "caption":
        return f'<p id="{identifier}" class="caption">{html.escape(block["text"])}</p>'
    if kind == "equation":
        return f'<pre id="{identifier}" class="equation">{html.escape(block["text"])}</pre>'
    if kind == "annexure":
        return f'<h2 id="{identifier}" class="annexure">{html.escape(block["text"])}</h2>'
    return ""


def _render_html(projected: dict[str, Any], path: Path, style: dict[str, Any], delivery: dict[str, Any], source_root: Path, asset_dir: Path) -> None:
    colours = style["colours"]
    body = "\n".join(_html_block(block, source_root, asset_dir, style, path.parent / "media") for block in projected["blocks"])
    context = projected.get("context") or {}
    metadata = " · ".join(str(context.get(key)) for key in ("subject", "grade", "curriculum_concept", "audience") if context.get(key))
    review = '<p class="review">CONTROLLED REVIEW ASSET. Human approval is required before external release.</p>' if delivery.get("review_label") else ""
    document = f'''<!doctype html>
<html lang="{html.escape(projected.get('render_language', projected['source_language']))}">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(projected['title'])}</title>
<style>
:root{{--ink:#{colours['ink']};--accent:#{colours['accent']};--muted:#{colours['muted']};--line:#{colours['line']};--fill:#{colours['fill']};}}
*{{box-sizing:border-box}} body{{margin:0;color:var(--ink);font:16px/1.55 system-ui,sans-serif;background:#fff}} main{{max-width:920px;margin:auto;padding:42px 28px 70px}} h1,h2,h3,h4{{color:var(--accent);line-height:1.2}} .meta{{color:var(--muted);border-bottom:2px solid var(--accent);padding-bottom:18px}} .review{{border:1px solid var(--line);padding:10px 12px;font-weight:700}} .callout{{display:block;background:var(--fill);border-left:4px solid var(--accent);padding:12px 14px;margin:16px 0}} .question{{display:grid;grid-template-columns:1fr auto;gap:18px;border-bottom:1px solid var(--line)}} .marks{{align-self:center}} table{{display:block;width:100%;max-width:100%;overflow-x:auto;border-collapse:collapse;margin:18px 0}} th,td{{border:1px solid var(--line);padding:8px;text-align:left;vertical-align:top}} th{{background:var(--fill)}} caption{{font-weight:700;text-align:left;margin-bottom:6px}} figure{{margin:24px 0}} img{{display:block;max-width:100%;height:auto;margin:auto}} figcaption,.caption{{color:var(--muted);font-size:.9rem}} .equation{{max-width:100%;overflow:auto;padding:14px;background:var(--fill)}} .page-break{{border:0;break-after:page}} @media print{{main{{max-width:none;padding:0}} table{{display:table;overflow:visible}} .page-break{{break-after:page}}}} @media(max-width:600px){{main{{padding:24px 16px}} .question{{grid-template-columns:1fr}}}}
</style></head><body><main><header><h1>{html.escape(projected['title'])}</h1><p class="meta">{html.escape(metadata)}</p>{review}</header>{body}</main></body></html>'''
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(document, encoding="utf-8")


def _render_vtt(projected: dict[str, Any], path: Path) -> bool:
    transcript = [block for block in projected["blocks"] if block["type"] == "transcript"]
    if not transcript:
        return False
    lines = ["WEBVTT", ""]
    cursor = 0.0
    for index, block in enumerate(transcript, 1):
        start = float(block.get("start_seconds", cursor))
        end = float(block.get("end_seconds", start + max(2.0, len(block["text"].split()) / 2.3)))
        cursor = end
        def stamp(value: float) -> str:
            hours, remainder = divmod(int(value * 1000), 3600000)
            minutes, remainder = divmod(remainder, 60000)
            seconds, millis = divmod(remainder, 1000)
            return f"{hours:02}:{minutes:02}:{seconds:02}.{millis:03}"
        lines.extend([str(index), f"{stamp(start)} --> {stamp(end)}", block["text"], ""])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return True


def _convert_docx_to_pdf(docx_path: Path, pdf_dir: Path) -> Path:
    from adapters.libreoffice_low_memory import convert_many_to_pdf

    produced = convert_many_to_pdf(
        [docx_path],
        pdf_dir,
        profile_prefix="dio-format-lo-",
    )
    if len(produced) != 1:
        raise RuntimeError(
            f"Format Core expected one PDF, received {len(produced)}"
        )
    return produced[0]

def _normalise_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _docx_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        values.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(values)


def _pptx_text(path: Path) -> str:
    deck = Presentation(path)
    values: list[str] = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                values.append(str(shape.text))
            if getattr(shape, "has_table", False):
                values.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(values)


def _qa_outputs(projected: dict[str, Any], outputs: dict[str, Path], warnings: list[str]) -> dict[str, Any]:
    errors: list[str] = []
    text_outputs: dict[str, str] = {}
    if "docx" in outputs:
        text_outputs["docx"] = _normalise_text(_docx_text(outputs["docx"]))
    if "pptx" in outputs:
        text_outputs["pptx"] = _normalise_text(_pptx_text(outputs["pptx"]))
    if "html" in outputs:
        raw = outputs["html"].read_text(encoding="utf-8")
        text_outputs["html"] = _normalise_text(html.unescape(re.sub(r"<[^>]+>", " ", raw)))
    expected = []
    for block in projected["blocks"]:
        if block["type"] in {"figure", "diagram"}:
            expected.append(str(block.get("caption") or ""))
        else:
            expected.extend(_block_strings(block))
    expected = [value for value in expected if len(_normalise_text(value)) >= 3]
    completeness: dict[str, Any] = {}
    for channel, rendered in text_outputs.items():
        missing = [value for value in expected if _normalise_text(value) not in rendered]
        completeness[channel] = {"expected_strings": len(expected), "missing_count": len(missing), "missing": missing[:12]}
        if missing:
            errors.append(f"{channel} omitted {len(missing)} semantic strings")
    if "pdf" in outputs:
        completed = subprocess.run(["pdftotext", str(outputs["pdf"]), "-"], text=True, capture_output=True, check=False, timeout=60)
        if completed.returncode != 0 or len(_normalise_text(completed.stdout)) < 30:
            errors.append("pdf text extraction was empty or failed")
    return {"passed": not errors, "errors": errors, "warnings": warnings, "semantic_completeness": completeness}


def render_semantic_asset(
    content: dict[str, Any],
    out_dir: Path,
    *,
    style_profile: str = "dio_professional",
    delivery_profile: str = "editable_review",
    language: str | None = None,
    channels: list[str] | None = None,
    release_mode: bool = False,
    source_root: Path | None = None,
    template_paths: dict[str, str | Path] | None = None,
) -> dict[str, Any]:
    validation = validate_semantic_content(content)
    if not validation["passed"]:
        raise ValueError("Invalid semantic content: " + "; ".join(validation["errors"]))
    profiles = load_profiles()
    if style_profile not in profiles["styles"]:
        raise ValueError(f"Unknown style profile: {style_profile}")
    if delivery_profile not in profiles["deliveries"]:
        raise ValueError(f"Unknown delivery profile: {delivery_profile}")
    style = profiles["styles"][style_profile]
    delivery = profiles["deliveries"][delivery_profile]
    selected_channels = list(dict.fromkeys(channels or delivery["channels"]))
    unknown = set(selected_channels) - SUPPORTED_CHANNELS
    if unknown:
        raise ValueError(f"Unsupported format channels: {', '.join(sorted(unknown))}")
    language = language or str(content["source_language"])
    projected, language_receipt = project_language(content, language, release_mode=release_mode)
    out_dir = out_dir.resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    source_root = (source_root or ROOT).resolve()
    resolved_templates: dict[str, Path] = {}
    for kind, raw_path in (template_paths or {}).items():
        if kind not in {"docx", "pptx", "logo"}:
            raise ValueError(f"Unsupported template kind: {kind}")
        template = Path(raw_path).expanduser()
        template = template if template.is_absolute() else source_root / template
        if not template.is_file():
            raise FileNotFoundError(f"Format template not found: {template}")
        resolved_templates[kind] = template.resolve()
    assets = out_dir / "assets"
    stem = f"{_slug(content['object_id'])}-{_slug(language)}"
    outputs: dict[str, Path] = {}
    docx_path = out_dir / f"{stem}.docx"
    if "docx" in selected_channels or "pdf" in selected_channels:
        _render_docx(
            projected, docx_path, style, delivery, source_root, assets,
            resolved_templates.get("docx"), resolved_templates.get("logo"),
        )
        if "docx" in selected_channels:
            outputs["docx"] = docx_path
    if "pdf" in selected_channels:
        outputs["pdf"] = _convert_docx_to_pdf(docx_path, out_dir)
        if "docx" not in selected_channels:
            docx_path.unlink(missing_ok=True)
    if "pptx" in selected_channels:
        outputs["pptx"] = out_dir / f"{stem}.pptx"
        _render_pptx(
            projected, outputs["pptx"], style, delivery, source_root, assets,
            resolved_templates.get("pptx"), resolved_templates.get("logo"),
        )
    if "html" in selected_channels:
        outputs["html"] = out_dir / f"{stem}.html"
        _render_html(projected, outputs["html"], style, delivery, source_root, assets)
    warnings = list(validation["warnings"])
    if "vtt" in selected_channels:
        vtt = out_dir / f"{stem}.vtt"
        if _render_vtt(projected, vtt):
            outputs["vtt"] = vtt
        else:
            warnings.append("vtt requested but the semantic object contains no transcript blocks")
    source_chars = sum(len(value) for block in content["blocks"] for value in _block_strings(block))
    target_chars = sum(len(value) for block in projected["blocks"] for value in _block_strings(block))
    expansion_ratio = round(target_chars / max(source_chars, 1), 3)
    if language != content["source_language"] and expansion_ratio > 1.45:
        warnings.append(f"translated content expansion ratio {expansion_ratio} requires layout review")
    qa = _qa_outputs(projected, outputs, warnings)
    receipt = {
        "schema": "dio.format_core.receipt.v1",
        "object_id": content["object_id"],
        "source_version": content["version"],
        "rendered_at": utc_now(),
        "status": ("rendered_release_candidate" if release_mode else "rendered_review_candidate") if qa["passed"] else "format_validation_failed",
        "style_profile": style_profile,
        "style_profile_hash": _content_digest(style),
        "delivery_profile": delivery_profile,
        "delivery_profile_hash": _content_digest(delivery),
        "language": language_receipt,
        "release_mode": release_mode,
        "semantic_source_hash": _content_digest(content),
        "semantic_blocks": validation["block_count"],
        "translation_expansion_ratio": expansion_ratio,
        "templates": {
            kind: {"path": str(path), "sha256": _sha256(path)}
            for kind, path in sorted(resolved_templates.items())
        },
        "qa": qa,
        "claim_boundary": "Format Core rebuilds typed semantic content through controlled profiles; it does not claim pixel-identical reconstruction of arbitrary source files.",
        "outputs": [
            {"channel": channel, "path": str(path.relative_to(out_dir)), "sha256": _sha256(path), "bytes": path.stat().st_size}
            for channel, path in sorted(outputs.items())
        ],
    }
    _write_json(out_dir / "FORMAT_CORE_RECEIPT.json", receipt)
    if not qa["passed"]:
        raise ValueError("Format Core output validation failed: " + "; ".join(qa["errors"]))
    return receipt
