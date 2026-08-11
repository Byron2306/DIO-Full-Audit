from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document


DOCUMENT_INPUTS = {".docx", ".odt", ".rtf", ".txt", ".md", ".html", ".htm", ".pdf", ".pptx", ".xlsx"}
DOCUMENT_OUTPUTS = {".docx", ".odt", ".rtf", ".txt", ".md", ".html", ".pdf", ".pptx", ".xlsx"}
DIRECT_TEXT_INPUTS = {".txt", ".md", ".html", ".htm"}
OFFICE_INPUTS = {".docx", ".odt", ".rtf", ".pptx", ".xlsx"}
OFFICE_OUTPUTS = {".docx", ".odt", ".rtf", ".pdf", ".pptx", ".xlsx"}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_hash(value: Any) -> str:
    raw = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def command_path(*names: str) -> str | None:
    for name in names:
        resolved = shutil.which(name)
        if resolved:
            return resolved
    return None


def _plain_text_from_markup(text: str, suffix: str) -> str:
    if suffix in {".html", ".htm"}:
        text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
        text = re.sub(r"(?i)<br\s*/?>", "\n", text)
        text = re.sub(r"(?i)</p\s*>", "\n\n", text)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        text = html.unescape(text)
    if suffix == ".md":
        text = re.sub(r"(?m)^#{1,6}\s+", "", text)
        text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
        text = re.sub(r"[*_`~]+", "", text)
    return re.sub(r"[ \t]+", " ", text).strip()


def _text_to_docx(source: Path, output: Path) -> None:
    text = _plain_text_from_markup(source.read_text(encoding="utf-8", errors="replace"), source.suffix.lower())
    document = Document()
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if block:
            document.add_paragraph(block)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def _docx_to_text(source: Path, output: Path, markdown: bool = False) -> None:
    document = Document(source)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    value = "\n\n".join(paragraphs) + ("\n" if paragraphs else "")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(value, encoding="utf-8")


def _text_to_html(source: Path, output: Path) -> None:
    text = _plain_text_from_markup(source.read_text(encoding="utf-8", errors="replace"), source.suffix.lower())
    blocks = [f"<p>{html.escape(block.strip())}</p>" for block in re.split(r"\n\s*\n", text) if block.strip()]
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("<!doctype html><meta charset=\"utf-8\">\n" + "\n".join(blocks) + "\n", encoding="utf-8")


@dataclass(frozen=True)
class ConversionPlan:
    source: Path
    output: Path
    source_format: str
    output_format: str
    strategy: str
    converter: str
    command: list[str] | None
    lossy: bool
    warnings: list[str]

    def as_dict(self) -> dict[str, Any]:
        return {
            "schema": "dio.document_conversion_plan.v1",
            "source": str(self.source),
            "output": str(self.output),
            "source_format": self.source_format,
            "output_format": self.output_format,
            "strategy": self.strategy,
            "converter": self.converter,
            "command": self.command,
            "lossy": self.lossy,
            "warnings": self.warnings,
        }


def plan_conversion(source: Path, output: Path) -> ConversionPlan:
    source = source.expanduser().resolve()
    output = output.expanduser().resolve()
    source_format = source.suffix.lower()
    output_format = output.suffix.lower()
    if source_format not in DOCUMENT_INPUTS:
        raise ValueError(f"Unsupported Document Studio source format: {source_format or '(none)'}")
    if output_format not in DOCUMENT_OUTPUTS:
        raise ValueError(f"Unsupported Document Studio output format: {output_format or '(none)'}")
    if source == output:
        raise ValueError("Conversion output must differ from the source path.")

    warnings: list[str] = []
    lossy = False
    if source_format in DIRECT_TEXT_INPUTS and output_format == ".docx":
        return ConversionPlan(source, output, source_format, output_format, "direct_text_to_docx", "python-docx", None, source_format in {".html", ".htm", ".md"}, ["Rich markup is reduced to paragraph content in the direct converter."] if source_format != ".txt" else [])
    if source_format == ".docx" and output_format in {".txt", ".md"}:
        return ConversionPlan(source, output, source_format, output_format, "direct_docx_to_text", "python-docx", None, True, ["Layout, images, comments, footnotes and advanced formatting are not represented in plain-text output."])
    if source_format in DIRECT_TEXT_INPUTS and output_format in {".html", ".htm"}:
        return ConversionPlan(source, output, source_format, output_format, "direct_text_to_html", "python-stdlib", None, source_format != ".txt", warnings)
    if source_format == ".pdf" and output_format in {".txt", ".md"}:
        converter = command_path("pdftotext")
        if not converter:
            raise RuntimeError("pdftotext is required for PDF-to-text conversion.")
        return ConversionPlan(source, output, source_format, output_format, "pdftotext", converter, [converter, "-layout", str(source), str(output)], True, ["PDF text extraction is structural extraction, not layout-preserving conversion."])

    if source_format in OFFICE_INPUTS and output_format in OFFICE_OUTPUTS:
        converter = command_path("libreoffice", "soffice")
        if not converter:
            raise RuntimeError("LibreOffice/soffice is required for this office-document conversion route.")
        warnings = []
        lossy = source_format != output_format
        return ConversionPlan(source, output, source_format, output_format, "libreoffice_headless", converter, None, lossy, warnings)

    pandoc = command_path("pandoc")
    if pandoc and source_format in {".txt", ".md", ".html", ".htm", ".docx"} and output_format in {".txt", ".md", ".html", ".docx", ".pdf"}:
        command = [pandoc, str(source), "-o", str(output)]
        return ConversionPlan(source, output, source_format, output_format, "pandoc", pandoc, command, source_format != output_format, warnings)

    raise ValueError(f"No governed conversion route is available for {source_format} -> {output_format}.")


def _run(command: list[str], *, timeout: int = 180) -> subprocess.CompletedProcess[str]:
    completed = subprocess.run(command, text=True, capture_output=True, check=False, timeout=timeout)
    if completed.returncode != 0:
        raise RuntimeError((completed.stderr or completed.stdout or "Document conversion failed.")[-2000:])
    return completed


def execute_conversion(source: Path, output: Path, *, dry_run: bool = False, force: bool = False) -> dict[str, Any]:
    plan = plan_conversion(source, output)
    if not plan.source.is_file():
        raise FileNotFoundError(f"Source document does not exist: {plan.source}")
    if plan.output.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing output: {plan.output}")
    source_hash = sha256(plan.source)
    plan_payload = plan.as_dict()
    plan_payload["source_sha256"] = source_hash
    plan_payload["plan_hash"] = canonical_hash(plan_payload)
    if dry_run:
        return {"status": "planned", "plan": plan_payload}

    plan.output.parent.mkdir(parents=True, exist_ok=True)
    if force and plan.output.exists():
        plan.output.unlink()

    if plan.strategy == "direct_text_to_docx":
        _text_to_docx(plan.source, plan.output)
    elif plan.strategy == "direct_docx_to_text":
        _docx_to_text(plan.source, plan.output, markdown=plan.output_format == ".md")
    elif plan.strategy == "direct_text_to_html":
        _text_to_html(plan.source, plan.output)
    elif plan.strategy in {"pdftotext", "pandoc"}:
        assert plan.command is not None
        _run(plan.command)
    elif plan.strategy == "libreoffice_headless":
        with tempfile.TemporaryDirectory(prefix="dio-doc-convert-") as temp_dir:
            temp = Path(temp_dir)
            _run([plan.converter, "--headless", "--convert-to", plan.output_format.lstrip("."), "--outdir", str(temp), str(plan.source)])
            produced = temp / f"{plan.source.stem}{plan.output_format}"
            if not produced.is_file():
                candidates = list(temp.glob(f"*{plan.output_format}"))
                if len(candidates) != 1:
                    raise RuntimeError("LibreOffice conversion completed without one unambiguous output artifact.")
                produced = candidates[0]
            shutil.move(str(produced), str(plan.output))
    else:
        raise RuntimeError(f"Unknown conversion strategy: {plan.strategy}")

    if not plan.output.is_file() or plan.output.stat().st_size == 0:
        raise RuntimeError("Conversion did not produce a non-empty output artifact.")
    receipt = {
        "schema": "dio.document_conversion_receipt.v1",
        "status": "complete",
        "source": {"path": str(plan.source), "format": plan.source_format, "sha256": source_hash, "size_bytes": plan.source.stat().st_size},
        "output": {"path": str(plan.output), "format": plan.output_format, "sha256": sha256(plan.output), "size_bytes": plan.output.stat().st_size},
        "strategy": plan.strategy,
        "converter": plan.converter,
        "lossy": plan.lossy,
        "warnings": plan.warnings,
        "source_preserved": True,
        "semantic_editing_performed": False,
        "plan_hash": plan_payload["plan_hash"],
        "completed_at": utc_now(),
    }
    receipt_path = Path(f"{plan.output}.conversion.json")
    receipt_path.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return {"status": "complete", "receipt": receipt, "receipt_path": str(receipt_path)}
