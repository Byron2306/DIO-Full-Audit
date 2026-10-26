from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def _add_inline(paragraph, text: str) -> None:
    """Render the small Markdown subset emitted by Evidex templates."""
    cursor = 0
    for match in _INLINE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def write_markdown_docx(path: Path, markdown_like_text: str) -> None:
    doc = Document()
    for raw_line in markdown_like_text.splitlines():
        line = raw_line.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            paragraph = doc.add_heading("", level=3)
            _add_inline(paragraph, line[4:])
        elif line.startswith("## "):
            paragraph = doc.add_heading("", level=2)
            _add_inline(paragraph, line[3:])
        elif line.startswith("# "):
            paragraph = doc.add_heading("", level=1)
            _add_inline(paragraph, line[2:])
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            _add_inline(paragraph, line[2:])
        else:
            paragraph = doc.add_paragraph()
            _add_inline(paragraph, line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
