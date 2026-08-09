#!/usr/bin/env python3
from __future__ import annotations

import argparse
import html
import json
import re
import shutil
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DESIGN_LAW = ROOT / "config" / "homs_design_law.json"
DEFAULT_ASSESSMENT_DESIGN = ROOT / "deliverables" / "caps_assessment_design" / "caps_assessment_design.json"


PALETTE = {
    "ink": "#18212b",
    "muted": "#5b6570",
    "paper": "#f8faf7",
    "line": "#c9d3cc",
    "green": "#1f6b3a",
    "blue": "#225c7a",
    "gold": "#f3b64b",
    "red": "#b64a3d",
    "white": "#ffffff",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: object) -> None:
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def clean(value: str, limit: int = 120) -> str:
    value = re.sub(r"\s+", " ", str(value or "")).strip()
    return value[:limit].rstrip()


def esc(value: str) -> str:
    return html.escape(str(value or ""), quote=True)


def wrap_text(value: str, width: int = 72, max_lines: int = 4) -> list[str]:
    words = clean(value, 500).split()
    lines: list[str] = []
    current: list[str] = []
    for word in words:
        candidate = " ".join([*current, word])
        if len(candidate) > width and current:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
        if len(lines) >= max_lines:
            break
    if current and len(lines) < max_lines:
        lines.append(" ".join(current))
    return lines


def text_block(lines: list[str], x: int, y: int, size: int = 22, weight: str = "400", fill: str = "#18212b", gap: int = 28) -> str:
    return "\n".join(
        f'<text x="{x}" y="{y + (index * gap)}" font-family="Arial, sans-serif" font-size="{size}" font-weight="{weight}" fill="{fill}">{esc(line)}</text>'
        for index, line in enumerate(lines)
    )


def svg_base(title: str, subtitle: str, body: str, width: int = 1280, height: int = 720) -> str:
    return f"""<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">
  <rect width="{width}" height="{height}" fill="{PALETTE['paper']}"/>
  <rect x="48" y="42" width="{width - 96}" height="{height - 84}" rx="18" fill="{PALETTE['white']}" stroke="{PALETTE['line']}" stroke-width="3"/>
  <rect x="48" y="42" width="{width - 96}" height="92" fill="{PALETTE['blue']}"/>
  <text x="82" y="96" font-family="Arial, sans-serif" font-size="34" font-weight="800" fill="{PALETTE['white']}">{esc(title)}</text>
  <text x="82" y="124" font-family="Arial, sans-serif" font-size="18" font-weight="700" fill="#dbeafe">{esc(subtitle)}</text>
  {body}
</svg>
"""


def assessment_sheet(title: str, subtitle: str, sheet_label: str, body: str, width: int = 1280, height: int = 720) -> str:
    title_lines = wrap_text(title, 52, 1)
    subtitle_lines = wrap_text(subtitle, 88, 1)
    return f"""<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">
  <rect width="{width}" height="{height}" fill="#ffffff"/>
  <rect x="54" y="42" width="{width - 108}" height="{height - 84}" fill="#ffffff" stroke="#18212b" stroke-width="2"/>
  <line x1="54" y1="146" x2="{width - 54}" y2="146" stroke="#18212b" stroke-width="2"/>
  <text x="82" y="82" font-family="Arial, sans-serif" font-size="28" font-weight="800" fill="#18212b">{esc(title_lines[0] if title_lines else title)}</text>
  {text_block(subtitle_lines, 82, 118, 15, "600", "#5b6570", 18)}
  <text x="{width - 82}" y="82" text-anchor="end" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">{esc(sheet_label)}</text>
  {body}
</svg>
"""


def render_png(svg_path: Path, png_path: Path) -> bool:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        return False
    result = subprocess.run(
        [
            ffmpeg,
            "-hide_banner",
            "-loglevel",
            "error",
            "-y",
            "-i",
            str(svg_path),
            "-frames:v",
            "1",
            "-pix_fmt",
            "rgb24",
            str(png_path),
        ],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    return result.returncode == 0 and png_path.exists()


def question_texts(pack: dict[str, Any], limit: int = 4) -> list[dict[str, Any]]:
    questions = []
    for section in pack.get("sections") or []:
        for question in section.get("questions") or []:
            questions.append(question)
            if len(questions) >= limit:
                return questions
    return questions


def rubric_items(pack: dict[str, Any], limit: int = 6) -> list[dict[str, Any]]:
    return list(pack.get("rubric") or [])[:limit]


def pack_questions(pack: dict[str, Any]) -> list[dict[str, Any]]:
    questions = []
    for section in pack.get("sections") or []:
        for question in section.get("questions") or []:
            item = dict(question)
            item["_section_title"] = section.get("title", "")
            item["_section_mode"] = section.get("mode", "")
            item["_section_stimulus"] = section.get("stimulus", "")
            questions.append(item)
    return questions


def visual_family_for_pack(pack: dict[str, Any], assessment_design: dict[str, Any] | None = None) -> str:
    assessment_design = assessment_design or {}
    visual_blueprint = pack.get("visual_blueprint") or {}
    family = str(visual_blueprint.get("assessment_family") or assessment_design.get("assessment_family") or pack.get("blueprint") or "")
    if family:
        return family
    modes = {str(section.get("mode") or "") for section in pack.get("sections", [])}
    if "performance_task_sheet" in modes or "reflection" in modes:
        return "practical_performance_or_portfolio"
    if "practical_project_design_task" in modes:
        return "practical_project_design_task"
    return "structured_test_or_task"


def planned_visual_kinds(pack: dict[str, Any]) -> set[str]:
    visual_blueprint = pack.get("visual_blueprint") or {}
    return {
        str(item.get("visual_kind") or "")
        for item in visual_blueprint.get("required_visuals") or []
        if item.get("visual_kind")
    }


def number_line_visual(pack: dict[str, Any]) -> tuple[str, str]:
    ticks = []
    labels = []
    for i in range(0, 31):
        x = 126 + i * 32
        h = 34 if i % 5 == 0 else 18
        ticks.append(f'<line x1="{x}" y1="310" x2="{x}" y2="{310 - h}" stroke="{PALETTE["ink"]}" stroke-width="3"/>')
        if i % 5 == 0:
            labels.append(f'<text x="{x}" y="350" text-anchor="middle" font-family="Arial" font-size="21" font-weight="800" fill="{PALETTE["ink"]}">{i}</text>')
    rows = []
    for index, question in enumerate(question_texts(pack, 3), start=1):
        y = 410 + (index - 1) * 72
        rows.append(f'<rect x="96" y="{y - 28}" width="1088" height="54" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        rows.append(f'<text x="112" y="{y + 6}" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">{esc(question.get("number", index))}</text>')
        rows.append(text_block(wrap_text(str(question.get("question", "")), 68, 1), 170, y + 6, 17, "500", "#18212b", 20))
        rows.append(f'<text x="1088" y="{y + 6}" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">Ans:</text>')
        rows.append(f'<line x1="1130" y1="{y + 8}" x2="1170" y2="{y + 8}" stroke="#18212b" stroke-width="2"/>')
    body = "\n".join(
        [
            '<text x="82" y="178" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Learner workspace: number line evidence</text>',
            '<text x="82" y="212" font-family="Arial, sans-serif" font-size="18" fill="#5b6570">Learner shows jumps on the line, then records the landing number.</text>',
            f'<line x1="126" y1="310" x2="{126 + 30 * 32}" y2="310" stroke="{PALETTE["ink"]}" stroke-width="5"/>',
            *ticks,
            *labels,
            '<rect x="96" y="238" width="248" height="36" fill="#f3f4f6" stroke="#18212b"/>',
            '<text x="112" y="262" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">Start</text>',
            '<rect x="380" y="238" width="248" height="36" fill="#f3f4f6" stroke="#18212b"/>',
            '<text x="396" y="262" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">Jump(s)</text>',
            '<rect x="664" y="238" width="248" height="36" fill="#f3f4f6" stroke="#18212b"/>',
            '<text x="680" y="262" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">Landing number</text>',
            *rows,
        ]
    )
    return "Learner Number Line Workspace", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "LEARNER WORKSPACE", body)


def foundation_observation_visual(pack: dict[str, Any]) -> tuple[str, str]:
    rows = []
    headers = ["Evidence", "Observed", "Notes"]
    xs = [92, 664, 842]
    widths = [572, 178, 342]
    for i, header in enumerate(headers):
        rows.append(f'<rect x="{xs[i]}" y="190" width="{widths[i]}" height="42" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        rows.append(f'<text x="{xs[i] + 14}" y="218" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">{header}</text>')
    for row_idx, question in enumerate(question_texts(pack, 5), start=1):
        y = 190 + row_idx * 58
        for i in range(3):
            rows.append(f'<rect x="{xs[i]}" y="{y}" width="{widths[i]}" height="58" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
        rows.append(text_block(wrap_text(str(question.get("question", "")), 58, 2), 106, y + 24, 15, "500", "#18212b", 18))
        rows.append(f'<rect x="724" y="{y + 17}" width="22" height="22" fill="#ffffff" stroke="#18212b" stroke-width="2"/>')
    body = "\n".join(
        [
            '<text x="82" y="174" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Teacher observation record</text>',
            *rows,
            '<text x="92" y="628" font-family="Arial, sans-serif" font-size="16" fill="#5b6570">Record evidence while the learner works with manipulatives, drawings, spoken explanation, or written response.</text>',
        ]
    )
    return "Teacher Observation Record", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "TEACHER RECORD", body)


def design_cycle_visual(pack: dict[str, Any]) -> tuple[str, str]:
    steps = ["Problem", "Inputs", "Logic", "Prototype", "Test", "Improve"]
    x_positions = [142, 316, 490, 664, 838, 1012]
    parts = []
    for i, (step, x) in enumerate(zip(steps, x_positions)):
        parts.append(f'<rect x="{x - 68}" y="214" width="136" height="68" fill="#ffffff" stroke="#18212b" stroke-width="2"/>')
        parts.append(f'<text x="{x}" y="256" text-anchor="middle" font-family="Arial, sans-serif" font-size="19" font-weight="800" fill="#18212b">{step}</text>')
        if i < len(steps) - 1:
            parts.append(f'<path d="M{x + 74} 248 L{x_positions[i + 1] - 78} 248" stroke="#18212b" stroke-width="3" marker-end="url(#arrow)"/>')
    table_rows = []
    for i, label in enumerate(["Material", "Purpose", "Constraint", "Test evidence", "Improvement"]):
        y = 384 + i * 44
        table_rows.append(f'<rect x="94" y="{y}" width="1092" height="44" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
        table_rows.append(f'<text x="112" y="{y + 29}" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">{label}</text>')
        table_rows.append(f'<line x1="372" y1="{y}" x2="372" y2="{y + 44}" stroke="#18212b" stroke-width="1"/>')
    body = "\n".join(
        [
            '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="3" orient="auto"><path d="M0,0 L0,6 L8,3 z" fill="#18212b"/></marker></defs>',
            '<text x="82" y="176" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Design and testing evidence log</text>',
            *parts,
            '<text x="94" y="360" font-family="Arial, sans-serif" font-size="20" font-weight="800" fill="#18212b">Complete during the build and testing process</text>',
            *table_rows,
        ]
    )
    return "Design And Test Evidence Log", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "PROJECT EVIDENCE", body)


def robot_maze_visual(pack: dict[str, Any]) -> tuple[str, str]:
    maze = [
        '<rect x="92" y="176" width="610" height="402" fill="#ffffff" stroke="#18212b" stroke-width="4"/>',
        '<path d="M132 532 L132 222 L312 222 L312 342 L220 342 L220 438 L440 438 L440 250 L610 250 L610 520 L514 520 L514 342 L686 342" fill="none" stroke="#18212b" stroke-width="28" stroke-linejoin="round"/>',
        '<text x="112" y="628" font-family="Arial, sans-serif" font-size="17" font-weight="800">START</text>',
        '<text x="620" y="628" font-family="Arial, sans-serif" font-size="17" font-weight="800">FINISH</text>',
        '<circle cx="132" cy="532" r="16" fill="#ffffff" stroke="#18212b" stroke-width="4"/>',
        '<rect x="662" y="326" width="48" height="32" fill="#ffffff" stroke="#18212b" stroke-width="3"/>',
    ]
    constraints = [
        ("Line rule", "Follow the black path"),
        ("Junctions", "State turn logic"),
        ("Sensor", "Show position"),
        ("Evidence", "Record tests"),
        ("Improve", "Change one thing"),
    ]
    rows = []
    for i, (label, value) in enumerate(constraints):
        y = 196 + i * 70
        rows.append(f'<rect x="760" y="{y}" width="360" height="52" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        rows.append(f'<text x="778" y="{y + 22}" font-family="Arial, sans-serif" font-size="16" font-weight="800" fill="#18212b">{label}</text>')
        rows.append(f'<text x="778" y="{y + 43}" font-family="Arial, sans-serif" font-size="15" fill="#5b6570">{value}</text>')
    body = "\n".join(
        [
            '<text x="82" y="158" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Maze stimulus and constraints</text>',
            *maze,
            *rows,
            text_block(
                wrap_text("Learner design must explain how the robot detects, turns, tests, and improves.", 48, 2),
                760,
                586,
                15,
                "400",
                "#5b6570",
                20,
            ),
        ]
    )
    return "Maze Stimulus And Constraints", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "TASK STIMULUS", body)


def performance_visual(pack: dict[str, Any]) -> tuple[str, str]:
    motifs = ["Motif A", "Motif B", "Transition", "Climax", "Resolution"]
    body_parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Performance floor plan and sequence map</text>',
        '<rect x="96" y="206" width="540" height="362" fill="#ffffff" stroke="#18212b" stroke-width="4"/>',
        '<line x1="96" y1="316" x2="636" y2="316" stroke="#9ca3af" stroke-width="1.5" stroke-dasharray="8 8"/>',
        '<line x1="96" y1="442" x2="636" y2="442" stroke="#9ca3af" stroke-width="1.5" stroke-dasharray="8 8"/>',
        '<line x1="276" y1="206" x2="276" y2="568" stroke="#9ca3af" stroke-width="1.5" stroke-dasharray="8 8"/>',
        '<line x1="456" y1="206" x2="456" y2="568" stroke="#9ca3af" stroke-width="1.5" stroke-dasharray="8 8"/>',
        '<text x="350" y="236" text-anchor="middle" font-family="Arial, sans-serif" font-size="16" font-weight="800">UPSTAGE</text>',
        '<text x="350" y="552" text-anchor="middle" font-family="Arial, sans-serif" font-size="16" font-weight="800">DOWNSTAGE / AUDIENCE</text>',
        '<path d="M150 512 C236 450 248 366 350 386 C454 406 478 296 586 256" fill="none" stroke="#18212b" stroke-width="5"/>',
        '<circle cx="150" cy="512" r="14" fill="#ffffff" stroke="#18212b" stroke-width="4"/><text x="174" y="518" font-family="Arial, sans-serif" font-size="15" font-weight="800">Entry</text>',
        '<circle cx="586" cy="256" r="14" fill="#18212b"/><text x="606" y="262" font-family="Arial, sans-serif" font-size="15" font-weight="800">Exit</text>',
    ]
    for i, motif in enumerate(motifs, start=1):
        y = 214 + (i - 1) * 64
        body_parts.append(f'<rect x="724" y="{y}" width="400" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        body_parts.append(f'<text x="742" y="{y + 20}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{motif}</text>')
        body_parts.append(f'<text x="742" y="{y + 40}" font-family="Arial, sans-serif" font-size="14" fill="#5b6570">Counts: ____   Space: ____   Dynamic: ____</text>')
    body_parts.append('<text x="724" y="554" font-family="Arial, sans-serif" font-size="15" fill="#5b6570">Learner completes this map before performance.</text>')
    body_parts.append('<text x="724" y="578" font-family="Arial, sans-serif" font-size="15" fill="#5b6570">Educator retains it with process evidence for moderation.</text>')
    return "Performance Floor Plan And Sequence Map", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "PERFORMANCE EVIDENCE", "\n".join(body_parts))


def performance_observation_visual(pack: dict[str, Any]) -> tuple[str, str]:
    xs = [80, 500, 618, 766, 920, 1070]
    widths = [420, 118, 148, 154, 150, 92]
    headers = ["Criterion", "Marks", "Evidence seen", "Concern", "Moderator note", "Score"]
    parts = ['<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Teacher observation instrument</text>']
    for i, header in enumerate(headers):
        parts.append(f'<rect x="{xs[i]}" y="194" width="{widths[i]}" height="40" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{xs[i] + 10}" y="220" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">{header}</text>')
    for row_idx, item in enumerate(rubric_items(pack, 6), start=1):
        y = 194 + row_idx * 58
        for i in range(len(headers)):
            parts.append(f'<rect x="{xs[i]}" y="{y}" width="{widths[i]}" height="58" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
        criterion = re.sub(r":.*$", "", str(item.get("criterion", "")))
        parts.append(text_block(wrap_text(criterion, 42, 2), xs[0] + 10, y + 22, 13, "700", "#18212b", 16))
        parts.append(f'<text x="{xs[1] + 14}" y="{y + 34}" font-family="Arial, sans-serif" font-size="16" font-weight="800" fill="#18212b">{esc(item.get("marks", ""))}</text>')
        for col in [2, 3, 4, 5]:
            parts.append(f'<line x1="{xs[col] + 12}" y1="{y + 34}" x2="{xs[col] + widths[col] - 12}" y2="{y + 34}" stroke="#9ca3af" stroke-width="1.5"/>')
    parts.append('<text x="80" y="654" font-family="Arial, sans-serif" font-size="15" fill="#5b6570">Use during live performance; retain with video/process evidence where moderation is required.</text>')
    return "Teacher Observation Instrument", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "MARKING EVIDENCE", "\n".join(parts))


def calculation_workspace_visual(pack: dict[str, Any]) -> tuple[str, str]:
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Working grid and method marks</text>',
    ]
    for i, question in enumerate(pack_questions(pack)[:4], start=1):
        y = 204 + (i - 1) * 108
        parts.append(f'<rect x="88" y="{y}" width="1096" height="88" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="108" y="{y + 26}" font-family="Arial, sans-serif" font-size="16" font-weight="800" fill="#18212b">{esc(question.get("number", i))} ({esc(question.get("marks", ""))})</text>')
        parts.append(text_block(wrap_text(str(question.get("question", "")), 82, 2), 188, y + 26, 15, "500", "#18212b", 18))
        for gx in range(650, 1148, 32):
            parts.append(f'<line x1="{gx}" y1="{y + 12}" x2="{gx}" y2="{y + 76}" stroke="#d1d5db" stroke-width="1"/>')
        for gy in range(y + 28, y + 76, 16):
            parts.append(f'<line x1="650" y1="{gy}" x2="1148" y2="{gy}" stroke="#d1d5db" stroke-width="1"/>')
    parts.append('<text x="90" y="656" font-family="Arial, sans-serif" font-size="15" fill="#5b6570">Award method marks from visible working before final-answer marks.</text>')
    return "Working Grid And Method Marks", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "CALCULATION WORKSPACE", "\n".join(parts))


def geography_weather_map_visual(pack: dict[str, Any]) -> tuple[str, str]:
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Synoptic weather map and data</text>',
        '<rect x="88" y="196" width="610" height="378" fill="#f8faf7" stroke="#18212b" stroke-width="3"/>',
        '<path d="M142 504 C228 422 304 394 390 334 C486 270 560 250 662 228" fill="none" stroke="#225c7a" stroke-width="3"/>',
        '<path d="M120 294 C236 236 366 222 486 246 C574 264 620 312 670 358" fill="none" stroke="#225c7a" stroke-width="3"/>',
        '<path d="M136 386 C240 332 336 326 434 356 C530 384 594 432 660 498" fill="none" stroke="#225c7a" stroke-width="3"/>',
        '<text x="210" y="282" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#225c7a">1016</text>',
        '<text x="430" y="338" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#225c7a">1012</text>',
        '<text x="540" y="488" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#225c7a">1008</text>',
        '<circle cx="260" cy="396" r="42" fill="#ffffff" stroke="#18212b" stroke-width="3"/>',
        '<text x="260" y="407" text-anchor="middle" font-family="Arial, sans-serif" font-size="34" font-weight="800" fill="#18212b">H</text>',
        '<circle cx="536" cy="286" r="42" fill="#ffffff" stroke="#b64a3d" stroke-width="3"/>',
        '<text x="536" y="297" text-anchor="middle" font-family="Arial, sans-serif" font-size="34" font-weight="800" fill="#b64a3d">L</text>',
        '<path d="M470 494 C500 442 542 414 602 394" fill="none" stroke="#b64a3d" stroke-width="5"/>',
        '<path d="M494 452 l-18 12 m46-38 l-18 12 m54-30 l-18 12" stroke="#b64a3d" stroke-width="4"/>',
        '<text x="510" y="528" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#b64a3d">Cold front</text>',
        '<path d="M610 230 L610 280 M592 248 L610 230 L628 248" stroke="#18212b" stroke-width="3" fill="none"/>',
        '<text x="610" y="300" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">N</text>',
        '<rect x="106" y="210" width="208" height="86" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
        '<text x="120" y="234" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">Legend</text>',
        '<line x1="122" y1="256" x2="166" y2="256" stroke="#225c7a" stroke-width="3"/><text x="178" y="261" font-family="Arial, sans-serif" font-size="13" fill="#18212b">Isobar</text>',
        '<line x1="122" y1="278" x2="166" y2="278" stroke="#b64a3d" stroke-width="5"/><text x="178" y="283" font-family="Arial, sans-serif" font-size="13" fill="#18212b">Cold front</text>',
        '<text x="96" y="604" font-family="Arial, sans-serif" font-size="14" fill="#5b6570">Use the pressure cells, isobars, front and wind direction to infer likely weather.</text>',
        '<text x="744" y="210" font-family="Arial, sans-serif" font-size="19" font-weight="800" fill="#18212b">Station weather data</text>',
    ]
    cols = [("Station", 116), ("Temp", 94), ("Pressure", 116), ("Weather", 126)]
    table_x = 744
    cx = table_x
    for label, width in cols:
        parts.append(f'<rect x="{cx}" y="230" width="{width}" height="38" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{cx + 10}" y="255" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">{label}</text>')
        cx += width
    rows = [("A", "18 C", "1016 hPa", "Clear"), ("B", "22 C", "1012 hPa", "Cloudy"), ("C", "16 C", "1008 hPa", "Rain likely")]
    for row_index, row in enumerate(rows, start=1):
        y = 230 + row_index * 38
        cx = table_x
        for value, (_label, width) in zip(row, cols):
            parts.append(f'<rect x="{cx}" y="{y}" width="{width}" height="38" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
            parts.append(f'<text x="{cx + 10}" y="{y + 25}" font-family="Arial, sans-serif" font-size="14" fill="#18212b">{value}</text>')
            cx += width
    for label, y in [("Pressure pattern:", 426), ("Weather evidence:", 486), ("Forecast inference:", 546)]:
        parts.append(f'<rect x="744" y="{y}" width="452" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="760" y="{y + 30}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{label}</text>')
        parts.append(f'<line x1="916" y1="{y + 30}" x2="1174" y2="{y + 30}" stroke="#9ca3af" stroke-width="1.5"/>')
    return "Synoptic Weather Map And Data", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "MAP SOURCE", "\n".join(parts))


def geography_thematic_map_visual(pack: dict[str, Any], theme: str) -> tuple[str, str]:
    configs = {
        "tectonics": {
            "title": "Plate Boundary Map And Hazard Data",
            "lead": "Use boundary type, volcanic belts and earthquake frequency to explain tectonic risk.",
            "legend": [("Convergent", "#b64a3d"), ("Divergent", "#225c7a"), ("Volcanoes", "#f3b64b")],
            "table_title": "Plate boundary evidence",
            "cols": [("Zone", 108), ("Boundary", 126), ("Quakes", 104), ("Hazard", 114)],
            "rows": [("A", "Convergent", "High", "Volcanic"), ("B", "Divergent", "Medium", "Rift"), ("C", "Transform", "High", "Fault")],
            "prompts": ["Boundary evidence:", "Hazard pattern:", "Reasoned explanation:"],
            "extra": [
                '<path d="M128 476 C230 390 290 306 410 290 C520 276 582 224 656 204" fill="none" stroke="#b64a3d" stroke-width="6"/>',
                '<path d="M132 252 C250 310 340 338 468 332 C548 328 606 380 664 472" fill="none" stroke="#225c7a" stroke-width="6" stroke-dasharray="14 10"/>',
                '<polygon points="240,350 256,382 224,382" fill="#f3b64b" stroke="#18212b" stroke-width="2"/>',
                '<polygon points="456,286 472,318 440,318" fill="#f3b64b" stroke="#18212b" stroke-width="2"/>',
                '<polygon points="570,410 586,442 554,442" fill="#f3b64b" stroke="#18212b" stroke-width="2"/>',
                '<text x="394" y="524" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#b64a3d">Active plate boundary zone</text>',
            ],
        },
        "topographic": {
            "title": "Topographic Map Extract And Orthophoto Cues",
            "lead": "Use contour spacing, grid reference, land use and route evidence to answer mapwork questions.",
            "legend": [("Contour", "#8b5e34"), ("Road", "#6b7280"), ("River", "#225c7a")],
            "table_title": "Mapwork evidence",
            "cols": [("Feature", 120), ("Grid", 104), ("Evidence", 132), ("Inference", 96)],
            "rows": [("Steep slope", "A3", "Close contours", "Hard route"), ("River", "B2", "Blue line", "Valley"), ("Road", "C4", "Dashed line", "Access")],
            "prompts": ["Grid reference:", "Map evidence:", "Inference:"],
            "extra": [
                '<path d="M150 278 C250 220 382 228 506 284 C580 318 624 388 650 468" fill="none" stroke="#8b5e34" stroke-width="3"/>',
                '<path d="M172 318 C274 276 380 280 488 324 C556 352 594 408 612 470" fill="none" stroke="#8b5e34" stroke-width="3"/>',
                '<path d="M204 362 C296 328 382 332 464 366 C514 386 548 428 562 472" fill="none" stroke="#8b5e34" stroke-width="3"/>',
                '<path d="M118 244 L664 516" fill="none" stroke="#6b7280" stroke-width="6" stroke-dasharray="18 12"/>',
                '<path d="M212 510 C282 420 358 360 456 322 C522 298 590 264 656 220" fill="none" stroke="#225c7a" stroke-width="8" opacity="0.75"/>',
                '<g stroke="#d1d5db" stroke-width="1"><path d="M188 196 L188 556"/><path d="M288 196 L288 556"/><path d="M388 196 L388 556"/><path d="M488 196 L488 556"/><path d="M588 196 L588 556"/><path d="M88 286 L688 286"/><path d="M88 376 L688 376"/><path d="M88 466 L688 466"/></g>',
            ],
        },
        "water": {
            "title": "Water Resources Map And Availability Data",
            "lead": "Use river systems, storage, demand and water-stress data to explain availability.",
            "legend": [("River", "#225c7a"), ("Dam", "#1f6b3a"), ("High stress", "#b64a3d")],
            "table_title": "Water availability",
            "cols": [("Catchment", 124), ("Rainfall", 104), ("Demand", 104), ("Stress", 120)],
            "rows": [("Upper", "620 mm", "Medium", "Moderate"), ("Middle", "480 mm", "High", "High"), ("Lower", "310 mm", "High", "Severe")],
            "prompts": ["Spatial pattern:", "Data evidence:", "Management response:"],
            "extra": [
                '<path d="M124 226 C206 306 270 334 356 382 C450 436 552 476 662 516" fill="none" stroke="#225c7a" stroke-width="18" opacity="0.72"/>',
                '<ellipse cx="330" cy="368" rx="52" ry="30" fill="#dbeafe" stroke="#1f6b3a" stroke-width="4"/>',
                '<ellipse cx="530" cy="462" rx="62" ry="34" fill="#dbeafe" stroke="#1f6b3a" stroke-width="4"/>',
                '<path d="M410 284 L642 236 L654 410 L438 438 Z" fill="#b64a3d" opacity="0.18" stroke="#b64a3d" stroke-width="3" stroke-dasharray="9 7"/>',
                '<text x="520" y="274" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#b64a3d">High water stress zone</text>',
            ],
        },
        "economic": {
            "title": "Economic Geography Map And Resource Data",
            "lead": "Use resource location, transport links and sector data to explain economic patterns.",
            "legend": [("Mining", "#f3b64b"), ("Industry", "#b64a3d"), ("Transport", "#6b7280")],
            "table_title": "Regional economic data",
            "cols": [("Region", 112), ("Resource", 112), ("Output", 104), ("Jobs", 124)],
            "rows": [("A", "Coal", "High", "42 000"), ("B", "Ports", "Medium", "18 000"), ("C", "Tourism", "Growing", "26 000")],
            "prompts": ["Location factor:", "Data trend:", "Development link:"],
            "extra": [
                '<circle cx="244" cy="338" r="42" fill="#f3b64b" stroke="#18212b" stroke-width="3"/><text x="244" y="346" text-anchor="middle" font-family="Arial" font-size="16" font-weight="800">Mine</text>',
                '<rect x="432" y="278" width="116" height="72" fill="#fee2e2" stroke="#b64a3d" stroke-width="3"/><text x="490" y="320" text-anchor="middle" font-family="Arial" font-size="16" font-weight="800">Industry</text>',
                '<path d="M120 496 C248 434 356 416 484 370 C560 342 610 298 664 236" fill="none" stroke="#6b7280" stroke-width="8" stroke-dasharray="18 10"/>',
                '<path d="M368 210 L636 260 L612 490 L338 514 Z" fill="#1f6b3a" opacity="0.12" stroke="#1f6b3a" stroke-width="3"/>',
            ],
        },
        "population": {
            "title": "Population Distribution Map And Density Data",
            "lead": "Use settlement pattern and density data to compare population distribution.",
            "legend": [("River", "#225c7a"), ("Road", "#6b7280"), ("Dense area", "#f3b64b")],
            "table_title": "Population change",
            "cols": [("Year", 104), ("Rural", 116), ("Urban", 116), ("Total", 116)],
            "rows": [("2000", "3 200", "5 600", "8 800"), ("2010", "2 900", "8 400", "11 300"), ("2025", "2 300", "13 700", "16 000")],
            "prompts": ["Map pattern:", "Data trend:", "Inference:"],
            "extra": [
                '<path d="M126 504 C214 452 226 364 326 344 C444 320 486 240 650 220" fill="none" stroke="#225c7a" stroke-width="18" opacity="0.72"/>',
                '<path d="M118 258 C238 286 330 274 450 314 C536 342 600 414 664 504" fill="none" stroke="#6b7280" stroke-width="7" stroke-dasharray="16 10"/>',
                '<rect x="426" y="286" width="124" height="84" fill="#f3b64b" stroke="#18212b" stroke-width="2"/>',
                '<text x="488" y="333" text-anchor="middle" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Urban core</text>',
                '<path d="M386 248 L594 256 L630 404 L384 412 Z" fill="#f3b64b" opacity="0.22" stroke="#b64a3d" stroke-width="3" stroke-dasharray="9 7"/>',
                '<text x="520" y="426" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#b64a3d">Urban expansion</text>',
                '<circle cx="202" cy="424" r="10" fill="#1f6b3a"/><circle cx="246" cy="394" r="8" fill="#1f6b3a"/><circle cx="284" cy="462" r="7" fill="#1f6b3a"/>',
            ],
        },
    }
    config = configs.get(theme, configs["population"])
    parts = [
        f'<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">{esc(config["title"])}</text>',
        '<rect x="88" y="196" width="600" height="360" fill="#f8faf7" stroke="#18212b" stroke-width="3"/>',
        *config["extra"],
        '<path d="M610 230 L610 280 M592 248 L610 230 L628 248" stroke="#18212b" stroke-width="3" fill="none"/>',
        '<text x="610" y="300" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">N</text>',
        '<line x1="116" y1="532" x2="236" y2="532" stroke="#18212b" stroke-width="5"/>',
        '<text x="176" y="550" text-anchor="middle" font-family="Arial, sans-serif" font-size="13" font-weight="700" fill="#18212b">0    5 km</text>',
        '<rect x="106" y="210" width="210" height="96" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
        '<text x="120" y="234" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">Legend</text>',
    ]
    for index, (label, color) in enumerate(config["legend"]):
        y = 256 + index * 23
        parts.append(f'<line x1="122" y1="{y}" x2="166" y2="{y}" stroke="{color}" stroke-width="6"/>')
        parts.append(f'<text x="178" y="{y + 5}" font-family="Arial, sans-serif" font-size="13" fill="#18212b">{esc(label)}</text>')
    parts.append(text_block(wrap_text(config["lead"], 66, 2), 96, 584, 14, "500", "#5b6570", 18))

    table_x = 744
    parts.append(f'<text x="744" y="210" font-family="Arial, sans-serif" font-size="19" font-weight="800" fill="#18212b">{esc(config["table_title"])}</text>')
    cx = table_x
    for label, width in config["cols"]:
        parts.append(f'<rect x="{cx}" y="230" width="{width}" height="38" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{cx + 10}" y="255" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">{esc(label)}</text>')
        cx += width
    for row_index, row in enumerate(config["rows"], start=1):
        y = 230 + row_index * 38
        cx = table_x
        for value, (_label, width) in zip(row, config["cols"]):
            parts.append(f'<rect x="{cx}" y="{y}" width="{width}" height="38" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
            parts.append(f'<text x="{cx + 10}" y="{y + 25}" font-family="Arial, sans-serif" font-size="14" fill="#18212b">{esc(value)}</text>')
            cx += width
    for label, y in zip(config["prompts"], [426, 486, 546]):
        parts.append(f'<rect x="744" y="{y}" width="452" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="760" y="{y + 30}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{esc(label)}</text>')
        parts.append(f'<line x1="930" y1="{y + 30}" x2="1174" y2="{y + 30}" stroke="#9ca3af" stroke-width="1.5"/>')
    return str(config["title"]), assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "MAP SOURCE", "\n".join(parts))


def geography_map_source_visual(pack: dict[str, Any]) -> tuple[str, str]:
    visual_blueprint = pack.get("visual_blueprint") or {}
    focus_blob = " ".join(
        [
            str(pack.get("assessment_title") or ""),
            " ".join(str(section.get("stimulus") or "") for section in pack.get("sections") or []),
            " ".join(visual_blueprint.get("assessment_content_focus") or []),
        ]
    ).lower()
    caps_blob = " ".join(visual_blueprint.get("caps_theme_candidates") or []).lower()
    theme_blob = f"{focus_blob} {caps_blob}"
    primary_blob = " ".join(
        [
            str(pack.get("assessment_title") or ""),
            " ".join((visual_blueprint.get("assessment_content_focus") or [])[:2]),
        ]
    ).lower()
    quality = visual_blueprint.get("theme_quality") or {}
    primary_content_line = str((quality.get("content_lines") or [""])[0])
    first_topic_blob = " ".join(
        [
            str(pack.get("assessment_title") or ""),
            primary_content_line,
            str((visual_blueprint.get("assessment_content_focus") or [""])[0]),
        ]
    ).lower()
    if primary_content_line:
        first_topic_blob = primary_content_line.lower()
    if any(token in first_topic_blob for token in ["water", "river", "dam", "catchment", "availability"]):
        return geography_thematic_map_visual(pack, "water")
    if any(token in first_topic_blob for token in ["population", "density", "distribution", "settlement", "urban", "rural"]):
        return geography_thematic_map_visual(pack, "population")
    if any(token in first_topic_blob for token in ["plate", "tectonic", "earthquake", "volcanic", "volcano"]):
        return geography_thematic_map_visual(pack, "tectonics")
    if any(token in first_topic_blob for token in ["economic", "energy", "development", "industry"]):
        return geography_thematic_map_visual(pack, "economic")
    if any(token in first_topic_blob for token in ["weather", "climate", "atmosphere", "synoptic", "isobar", "pressure"]):
        return geography_weather_map_visual(pack)
    if any(token in primary_blob for token in ["topographic", "orthophoto", "mapwork", "aerial"]):
        return geography_thematic_map_visual(pack, "topographic")

    if any(token in theme_blob for token in ["weather", "climate", "atmosphere", "synoptic", "isobar", "pressure"]):
        return geography_weather_map_visual(pack)
    if any(token in theme_blob for token in ["plate", "tectonic", "earthquake", "volcanic", "volcano"]):
        return geography_thematic_map_visual(pack, "tectonics")
    if any(token in theme_blob for token in ["topographic", "orthophoto", "mapwork", "aerial"]):
        return geography_thematic_map_visual(pack, "topographic")
    if any(token in theme_blob for token in ["water", "river", "dam", "catchment", "availability"]):
        return geography_thematic_map_visual(pack, "water")
    if any(token in theme_blob for token in ["economic", "energy", "development", "industry"]):
        return geography_thematic_map_visual(pack, "economic")
    if any(token in theme_blob for token in ["settlement", "urban", "rural", "population", "density", "distribution"]):
        return geography_thematic_map_visual(pack, "population")

    questions = pack_questions(pack)
    source_note = ""
    for question in questions:
        source_note = clean(question.get("_section_stimulus", ""), 180)
        if source_note:
            break
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Map extract and settlement data</text>',
        '<rect x="88" y="196" width="600" height="360" fill="#f8faf7" stroke="#18212b" stroke-width="3"/>',
        '<path d="M126 504 C214 452 226 364 326 344 C444 320 486 240 650 220" fill="none" stroke="#225c7a" stroke-width="18" opacity="0.72"/>',
        '<path d="M118 258 C238 286 330 274 450 314 C536 342 600 414 664 504" fill="none" stroke="#6b7280" stroke-width="7" stroke-dasharray="16 10"/>',
        '<rect x="426" y="286" width="124" height="84" fill="#f3b64b" stroke="#18212b" stroke-width="2"/>',
        '<text x="488" y="333" text-anchor="middle" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Urban core</text>',
        '<path d="M386 248 L594 256 L630 404 L384 412 Z" fill="#f3b64b" opacity="0.22" stroke="#b64a3d" stroke-width="3" stroke-dasharray="9 7"/>',
        '<text x="520" y="426" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#b64a3d">Urban expansion 2000-2025</text>',
        '<circle cx="202" cy="424" r="10" fill="#1f6b3a"/><circle cx="246" cy="394" r="8" fill="#1f6b3a"/><circle cx="284" cy="462" r="7" fill="#1f6b3a"/>',
        '<circle cx="184" cy="322" r="7" fill="#1f6b3a"/><circle cx="336" cy="408" r="8" fill="#1f6b3a"/>',
        '<text x="242" y="494" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#1f6b3a">Rural settlements</text>',
        '<path d="M610 230 L610 280 M592 248 L610 230 L628 248" stroke="#18212b" stroke-width="3" fill="none"/>',
        '<text x="610" y="300" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">N</text>',
        '<line x1="116" y1="532" x2="236" y2="532" stroke="#18212b" stroke-width="5"/>',
        '<text x="176" y="550" text-anchor="middle" font-family="Arial, sans-serif" font-size="13" font-weight="700" fill="#18212b">0    5 km</text>',
        '<rect x="106" y="210" width="184" height="86" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
        '<text x="120" y="234" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">Legend</text>',
        '<line x1="122" y1="256" x2="166" y2="256" stroke="#225c7a" stroke-width="8"/><text x="178" y="261" font-family="Arial, sans-serif" font-size="13" fill="#18212b">River</text>',
        '<line x1="122" y1="278" x2="166" y2="278" stroke="#6b7280" stroke-width="5" stroke-dasharray="10 6"/><text x="178" y="283" font-family="Arial, sans-serif" font-size="13" fill="#18212b">Road</text>',
    ]
    if source_note:
        parts.append(text_block(wrap_text(source_note, 66, 2), 96, 584, 14, "500", "#5b6570", 18))

    table_x = 744
    parts.extend(
        [
            '<text x="744" y="210" font-family="Arial, sans-serif" font-size="19" font-weight="800" fill="#18212b">Population change</text>',
        ]
    )
    cols = [("Year", 104), ("Rural", 116), ("Urban", 116), ("Total", 116)]
    cx = table_x
    for label, width in cols:
        parts.append(f'<rect x="{cx}" y="230" width="{width}" height="38" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{cx + 12}" y="255" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">{label}</text>')
        cx += width
    rows = [("2000", "3 200", "5 600", "8 800"), ("2010", "2 900", "8 400", "11 300"), ("2025", "2 300", "13 700", "16 000")]
    for row_index, row in enumerate(rows, start=1):
        y = 230 + row_index * 38
        cx = table_x
        for value, (_label, width) in zip(row, cols):
            parts.append(f'<rect x="{cx}" y="{y}" width="{width}" height="38" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
            parts.append(f'<text x="{cx + 12}" y="{y + 25}" font-family="Arial, sans-serif" font-size="14" fill="#18212b">{value}</text>')
            cx += width
    parts.extend(
        [
            '<text x="744" y="404" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">Learner evidence notes</text>',
            '<rect x="744" y="426" width="452" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
            '<text x="760" y="456" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Map pattern:</text>',
            '<line x1="880" y1="456" x2="1174" y2="456" stroke="#9ca3af" stroke-width="1.5"/>',
            '<rect x="744" y="486" width="452" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
            '<text x="760" y="516" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Data trend:</text>',
            '<line x1="868" y1="516" x2="1174" y2="516" stroke="#9ca3af" stroke-width="1.5"/>',
            '<rect x="744" y="546" width="452" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
            '<text x="760" y="576" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Inference:</text>',
            '<line x1="850" y1="576" x2="1174" y2="576" stroke="#9ca3af" stroke-width="1.5"/>',
        ]
    )
    return "Map Extract And Settlement Data", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "MAP SOURCE", "\n".join(parts))


def source_panel_visual(pack: dict[str, Any]) -> tuple[str, str]:
    planned_kinds = planned_visual_kinds(pack)
    subject_blob = " ".join(
        [
            str(pack.get("subject") or ""),
            str(pack.get("assessment_title") or ""),
            " ".join(str(q.get("_section_stimulus") or "") for q in pack_questions(pack)),
        ]
    ).lower()
    if "map_extract" in planned_kinds or ("geography" in subject_blob and "map" in subject_blob):
        return geography_map_source_visual(pack)

    questions = pack_questions(pack)
    stimuli = []
    for question in questions:
        stimulus = clean(question.get("_section_stimulus", ""), 240)
        if stimulus and stimulus not in stimuli:
            stimuli.append(stimulus)
    if not stimuli:
        stimuli = [clean(q.get("question", ""), 180) for q in questions[:2]]
    parts = ['<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Source response panel</text>']
    for i in range(2):
        x = 90 + i * 560
        label = f"Source {chr(65 + i)}"
        text = stimuli[i] if i < len(stimuli) else "Teacher inserts approved source extract, visual, map, table, or document here."
        parts.append(f'<rect x="{x}" y="204" width="500" height="180" fill="#ffffff" stroke="#18212b" stroke-width="2"/>')
        parts.append(f'<text x="{x + 18}" y="232" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">{label}</text>')
        parts.append(text_block(wrap_text(text, 52, 5), x + 18, 262, 15, "500", "#18212b", 18))
    headers = ["Question", "Source used", "Evidence to quote or describe", "Inference / answer"]
    xs = [90, 246, 440, 800]
    widths = [156, 194, 360, 322]
    for i, header in enumerate(headers):
        parts.append(f'<rect x="{xs[i]}" y="420" width="{widths[i]}" height="40" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{xs[i] + 10}" y="446" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">{header}</text>')
    for row, question in enumerate(questions[:3], start=1):
        y = 420 + row * 48
        for i in range(len(headers)):
            parts.append(f'<rect x="{xs[i]}" y="{y}" width="{widths[i]}" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
        parts.append(f'<text x="104" y="{y + 29}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{esc(question.get("number", row))}</text>')
        for col in [1, 2, 3]:
            parts.append(f'<line x1="{xs[col] + 12}" y1="{y + 29}" x2="{xs[col] + widths[col] - 12}" y2="{y + 29}" stroke="#9ca3af" stroke-width="1.5"/>')
    return "Source Panel And Evidence Table", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "SOURCE EVIDENCE", "\n".join(parts))


def source_extract_visual(pack: dict[str, Any]) -> tuple[str, str]:
    questions = pack_questions(pack)
    stimuli = []
    for question in questions:
        stimulus = clean(question.get("_section_stimulus", ""), 260)
        if stimulus and stimulus not in stimuli:
            stimuli.append(stimulus)
    if not stimuli:
        stimuli = [
            "Teacher inserts approved DBE-style photograph, infographic, extract or case source here.",
            "Learners must cite visible/source evidence before explaining.",
        ]
    parts = ['<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Source extract and visual evidence</text>']
    labels = ["Source A", "Source B"]
    for i, label in enumerate(labels):
        x = 90 + i * 560
        text = stimuli[i] if i < len(stimuli) else "Teacher inserts the second approved source or visual here."
        parts.append(f'<rect x="{x}" y="204" width="500" height="220" fill="#ffffff" stroke="#18212b" stroke-width="2"/>')
        parts.append(f'<rect x="{x + 20}" y="242" width="180" height="132" fill="#eef6f3" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{x + 110}" y="310" text-anchor="middle" font-family="Arial, sans-serif" font-size="16" font-weight="800" fill="#225c7a">visual cue</text>')
        parts.append(f'<text x="{x + 18}" y="232" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">{label}</text>')
        parts.append(text_block(wrap_text(text, 31, 5), x + 222, 266, 14, "500", "#18212b", 18))
        parts.append(f'<text x="{x + 20}" y="402" font-family="Arial, sans-serif" font-size="13" fill="#5b6570">Source note / caption:</text>')
        parts.append(f'<line x1="{x + 154}" y1="402" x2="{x + 480}" y2="402" stroke="#9ca3af" stroke-width="1.5"/>')
    for i, prompt in enumerate(["Visible evidence:", "Source evidence:", "Geographic inference:"]):
        y = 470 + i * 52
        parts.append(f'<rect x="90" y="{y}" width="1030" height="42" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="108" y="{y + 27}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{prompt}</text>')
        parts.append(f'<line x1="292" y1="{y + 27}" x2="1096" y2="{y + 27}" stroke="#9ca3af" stroke-width="1.5"/>')
    return "Source Extract And Visual Evidence", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "SOURCE EVIDENCE", "\n".join(parts))


def claim_evidence_table_visual(pack: dict[str, Any]) -> tuple[str, str]:
    questions = pack_questions(pack)
    headers = ["Question", "Claim / answer", "Evidence from source", "Reasoning"]
    xs = [84, 230, 532, 868]
    widths = [146, 302, 336, 250]
    parts = ['<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Claim-evidence-reasoning table</text>']
    for i, header in enumerate(headers):
        parts.append(f'<rect x="{xs[i]}" y="206" width="{widths[i]}" height="42" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{xs[i] + 10}" y="233" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">{header}</text>')
    for row, question in enumerate(questions[:6], start=1):
        y = 206 + row * 58
        for i in range(len(headers)):
            parts.append(f'<rect x="{xs[i]}" y="{y}" width="{widths[i]}" height="58" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
        parts.append(f'<text x="100" y="{y + 34}" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">{esc(question.get("number", row))}</text>')
        for col in [1, 2, 3]:
            parts.append(f'<line x1="{xs[col] + 12}" y1="{y + 34}" x2="{xs[col] + widths[col] - 12}" y2="{y + 34}" stroke="#9ca3af" stroke-width="1.5"/>')
    parts.append('<text x="88" y="640" font-family="Arial, sans-serif" font-size="15" fill="#5b6570">Use this to force every explanation back to the map, graph, table, photograph, diagram or extract.</text>')
    return "Claim Evidence Reasoning Table", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "EVIDENCE TABLE", "\n".join(parts))


def geography_process_diagram_visual(pack: dict[str, Any]) -> tuple[str, str]:
    title = clean(pack.get("assessment_title", "Geography process"), 120)
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Geography process diagram</text>',
        '<rect x="90" y="204" width="650" height="360" fill="#f8faf7" stroke="#18212b" stroke-width="2"/>',
        '<path d="M142 474 C236 394 320 366 420 340 C520 314 602 276 690 228" fill="none" stroke="#225c7a" stroke-width="5"/>',
        '<path d="M154 500 C248 456 354 438 468 446 C578 454 642 496 706 538" fill="none" stroke="#8b5e34" stroke-width="4" stroke-dasharray="10 8"/>',
        '<circle cx="242" cy="418" r="34" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
        '<circle cx="468" cy="342" r="34" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
        '<circle cx="620" cy="274" r="34" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
        '<text x="242" y="424" text-anchor="middle" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">A</text>',
        '<text x="468" y="348" text-anchor="middle" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">B</text>',
        '<text x="620" y="280" text-anchor="middle" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">C</text>',
        '<path d="M278 414 L430 352" stroke="#18212b" stroke-width="3" marker-end="url(#arrow)"/>',
        '<path d="M504 336 L584 286" stroke="#18212b" stroke-width="3" marker-end="url(#arrow)"/>',
        '<text x="106" y="588" font-family="Arial, sans-serif" font-size="14" fill="#5b6570">Label the process stages and explain how evidence changes from A to C.</text>',
        '<text x="790" y="214" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">Use with topic</text>',
        *text_block(wrap_text(title, 38, 3), 790, 248, 16, "600", "#18212b", 20).splitlines(),
    ]
    labels = ["Feature / stage", "Evidence in diagram", "Explanation"]
    for i, label in enumerate(labels):
        y = 344 + i * 72
        parts.append(f'<rect x="790" y="{y}" width="360" height="52" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="808" y="{y + 31}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{label}</text>')
        parts.append(f'<line x1="968" y1="{y + 31}" x2="1128" y2="{y + 31}" stroke="#9ca3af" stroke-width="1.5"/>')
    body = '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="3" orient="auto"><path d="M0,0 L0,6 L8,3 z" fill="#18212b"/></marker></defs>' + "\n".join(parts)
    return "Geography Process Diagram", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "PROCESS DIAGRAM", body)


def science_labelled_diagram_visual(pack: dict[str, Any]) -> tuple[str, str]:
    subject = str(pack.get("subject") or "Science")
    title = clean(pack.get("assessment_title", "Science diagram interpretation"), 120)
    visual_blueprint = pack.get("visual_blueprint") or {}
    topic_blob = " ".join(
        [
            title,
            " ".join(str(section.get("title") or "") for section in pack.get("sections", [])),
            " ".join(str(section.get("stimulus") or "") for section in pack.get("sections", [])),
            " ".join(str(question.get("question") or "") for question in pack_questions(pack)[:6]),
            " ".join(visual_blueprint.get("assessment_content_focus") or []),
            " ".join(visual_blueprint.get("caps_theme_candidates") or []),
        ]
    ).lower()
    if "evolution" in topic_blob or "homin" in topic_blob or "skull" in topic_blob:
        parts = [
            '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="3" orient="auto"><path d="M0,0 L0,6 L8,3 z" fill="#18212b"/></marker></defs>',
            '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Hominin cranial profile comparison</text>',
            '<rect x="88" y="204" width="1030" height="382" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
            '<line x1="410" y1="220" x2="410" y2="558" stroke="#cbd5e1" stroke-width="2" stroke-dasharray="8 8"/>',
            '<line x1="796" y1="220" x2="796" y2="558" stroke="#cbd5e1" stroke-width="2" stroke-dasharray="8 8"/>',
            '<text x="248" y="242" text-anchor="middle" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Skull A: earlier hominin profile</text>',
            '<path d="M208 350 C208 292 250 260 304 266 C356 272 386 310 380 358 C374 404 342 430 294 436 C244 442 210 408 208 350 Z" fill="#f8faf7" stroke="#18212b" stroke-width="3"/>',
            '<path d="M296 352 C360 354 400 384 406 428 C366 442 328 438 296 418 Z" fill="#ffffff" stroke="#18212b" stroke-width="3"/>',
            '<path d="M232 390 C260 422 312 430 356 410" fill="none" stroke="#18212b" stroke-width="4"/>',
            '<path d="M346 326 C362 326 374 338 374 354" fill="none" stroke="#18212b" stroke-width="3"/>',
            '<ellipse cx="282" cy="444" rx="17" ry="8" fill="#ffffff" stroke="#18212b" stroke-width="3"/>',
            '<line x1="218" y1="470" x2="392" y2="470" stroke="#18212b" stroke-width="2"/>',
            '<text x="642" y="242" text-anchor="middle" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Skull B: later hominin profile</text>',
            '<path d="M560 344 C560 268 620 230 696 240 C770 250 812 302 796 366 C784 418 740 448 682 452 C610 456 560 410 560 344 Z" fill="#f8faf7" stroke="#18212b" stroke-width="3"/>',
            '<path d="M694 358 C738 364 764 386 762 414 C730 424 700 420 678 402 Z" fill="#ffffff" stroke="#18212b" stroke-width="3"/>',
            '<path d="M594 392 C628 422 684 428 732 406" fill="none" stroke="#18212b" stroke-width="4"/>',
            '<path d="M700 324 C720 324 732 338 732 354" fill="none" stroke="#18212b" stroke-width="3"/>',
            '<ellipse cx="666" cy="454" rx="18" ry="8" fill="#ffffff" stroke="#18212b" stroke-width="3"/>',
            '<line x1="600" y1="480" x2="772" y2="480" stroke="#18212b" stroke-width="2"/>',
            '<circle cx="316" cy="300" r="20" fill="#ffffff" stroke="#225c7a" stroke-width="3"/><text x="316" y="307" text-anchor="middle" font-family="Arial, sans-serif" font-size="17" font-weight="900" fill="#225c7a">A</text>',
            '<circle cx="374" cy="392" r="20" fill="#ffffff" stroke="#225c7a" stroke-width="3"/><text x="374" y="399" text-anchor="middle" font-family="Arial, sans-serif" font-size="17" font-weight="900" fill="#225c7a">B</text>',
            '<circle cx="282" cy="444" r="19" fill="#ffffff" stroke="#225c7a" stroke-width="3"/><text x="282" y="451" text-anchor="middle" font-family="Arial, sans-serif" font-size="16" font-weight="900" fill="#225c7a">C</text>',
            '<circle cx="694" cy="286" r="20" fill="#ffffff" stroke="#225c7a" stroke-width="3"/><text x="694" y="293" text-anchor="middle" font-family="Arial, sans-serif" font-size="17" font-weight="900" fill="#225c7a">A</text>',
            '<circle cx="742" cy="394" r="20" fill="#ffffff" stroke="#225c7a" stroke-width="3"/><text x="742" y="401" text-anchor="middle" font-family="Arial, sans-serif" font-size="17" font-weight="900" fill="#225c7a">B</text>',
            '<circle cx="666" cy="454" r="19" fill="#ffffff" stroke="#225c7a" stroke-width="3"/><text x="666" y="461" text-anchor="middle" font-family="Arial, sans-serif" font-size="16" font-weight="900" fill="#225c7a">C</text>',
            '<rect x="844" y="238" width="232" height="58" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
            '<text x="862" y="262" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">A: cranial capacity</text>',
            '<text x="862" y="282" font-family="Arial, sans-serif" font-size="12" fill="#5b6570">compare brain-case volume</text>',
            '<path d="M844 266 L716 286" fill="none" stroke="#18212b" stroke-width="2" marker-end="url(#arrow)"/>',
            '<rect x="844" y="346" width="232" height="58" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
            '<text x="862" y="370" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">B: jaw projection</text>',
            '<text x="862" y="390" font-family="Arial, sans-serif" font-size="12" fill="#5b6570">reduced projection in later forms</text>',
            '<path d="M844 374 L764 392" fill="none" stroke="#18212b" stroke-width="2" marker-end="url(#arrow)"/>',
            '<rect x="844" y="454" width="232" height="58" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>',
            '<text x="862" y="478" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">C: foramen magnum</text>',
            '<text x="862" y="498" font-family="Arial, sans-serif" font-size="12" fill="#5b6570">more central with bipedalism</text>',
            '<path d="M844 482 L688 454" fill="none" stroke="#18212b" stroke-width="2" marker-end="url(#arrow)"/>',
            '<text x="92" y="622" font-family="Arial, sans-serif" font-size="14" fill="#5b6570">Compare labelled features only: cranial capacity, jaw projection and foramen magnum position.</text>',
        ]
        return f"{subject} Hominin Skull Diagram", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "DIAGRAM", "\n".join(parts))
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Labelled science diagram stimulus</text>',
        '<rect x="90" y="204" width="620" height="360" fill="#f8faf7" stroke="#18212b" stroke-width="2"/>',
        '<ellipse cx="382" cy="374" rx="190" ry="110" fill="#eef6f3" stroke="#18212b" stroke-width="3"/>',
        '<ellipse cx="382" cy="374" rx="86" ry="48" fill="#ffffff" stroke="#225c7a" stroke-width="3"/>',
        '<circle cx="292" cy="336" r="24" fill="#ffffff" stroke="#b64a3d" stroke-width="3"/>',
        '<circle cx="478" cy="416" r="24" fill="#ffffff" stroke="#1f6b3a" stroke-width="3"/>',
        '<path d="M382 264 C440 292 484 326 526 374" fill="none" stroke="#225c7a" stroke-width="4" marker-end="url(#arrow)"/>',
        '<path d="M238 444 C294 484 388 500 500 466" fill="none" stroke="#1f6b3a" stroke-width="4" marker-end="url(#arrow)"/>',
        '<line x1="292" y1="336" x2="150" y2="274" stroke="#18212b" stroke-width="2"/>',
        '<line x1="382" y1="374" x2="168" y2="386" stroke="#18212b" stroke-width="2"/>',
        '<line x1="478" y1="416" x2="612" y2="500" stroke="#18212b" stroke-width="2"/>',
        '<rect x="104" y="246" width="152" height="44" fill="#ffffff" stroke="#18212b"/>',
        '<text x="118" y="274" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Label A</text>',
        '<rect x="104" y="364" width="152" height="44" fill="#ffffff" stroke="#18212b"/>',
        '<text x="118" y="392" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Label B</text>',
        '<rect x="604" y="480" width="152" height="44" fill="#ffffff" stroke="#18212b"/>',
        '<text x="618" y="508" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">Label C</text>',
        '<text x="98" y="596" font-family="Arial, sans-serif" font-size="14" fill="#5b6570">Use this as the supplied diagram; replace labels with topic-specific structures during review.</text>',
        '<text x="780" y="214" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">Assessment actions</text>',
        *text_block(wrap_text(title, 38, 3), 780, 248, 16, "600", "#18212b", 20).splitlines(),
    ]
    prompts = ["Identify labelled structures", "Explain the process", "Relate structure to function", "Predict effect of a change"]
    for i, prompt in enumerate(prompts):
        y = 350 + i * 54
        parts.append(f'<rect x="780" y="{y}" width="360" height="42" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="798" y="{y + 27}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{esc(prompt)}</text>')
    body = '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="3" orient="auto"><path d="M0,0 L0,6 L8,3 z" fill="#18212b"/></marker></defs>' + "\n".join(parts)
    return f"{subject} Labelled Diagram Stimulus", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "DIAGRAM", body)


def science_data_graph_visual(pack: dict[str, Any]) -> tuple[str, str]:
    topic_blob = " ".join(
        [
            str(pack.get("assessment_title") or ""),
            " ".join(str(section.get("title") or "") for section in pack.get("sections", [])),
            " ".join(str(question.get("question") or "") for question in pack_questions(pack)[:6]),
        ]
    ).lower()
    if "evolution" in topic_blob or "natural selection" in topic_blob or "beak" in topic_blob:
        parts = [
            '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Finch beak depth data</text>',
            '<rect x="90" y="204" width="510" height="306" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
            '<text x="112" y="236" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">Island population sample</text>',
        ]
        cols = [("Year", 88), ("Rainfall", 112), ("Hard seeds", 118), ("Mean beak", 130)]
        rows = [("2019", "820 mm", "24%", "8.7 mm"), ("2020", "610 mm", "37%", "9.1 mm"), ("2021", "430 mm", "58%", "9.8 mm"), ("2022", "390 mm", "64%", "10.2 mm")]
        x = 112
        for label, width in cols:
            parts.append(f'<rect x="{x}" y="258" width="{width}" height="36" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
            parts.append(f'<text x="{x + 8}" y="282" font-family="Arial, sans-serif" font-size="12" font-weight="800" fill="#18212b">{label}</text>')
            x += width
        for row_idx, row in enumerate(rows, start=1):
            y = 258 + row_idx * 36
            x = 112
            for value, (_label, width) in zip(row, cols):
                parts.append(f'<rect x="{x}" y="{y}" width="{width}" height="36" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
                parts.append(f'<text x="{x + 8}" y="{y + 24}" font-family="Arial, sans-serif" font-size="12" fill="#18212b">{value}</text>')
                x += width
        parts.extend([
            '<rect x="676" y="204" width="470" height="306" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
            '<line x1="744" y1="456" x2="1094" y2="456" stroke="#18212b" stroke-width="3"/>',
            '<line x1="744" y1="456" x2="744" y2="246" stroke="#18212b" stroke-width="3"/>',
            '<polyline points="754,426 840,398 928,334 1020,286" fill="none" stroke="#225c7a" stroke-width="5"/>',
            '<circle cx="754" cy="426" r="6" fill="#225c7a"/><circle cx="840" cy="398" r="6" fill="#225c7a"/><circle cx="928" cy="334" r="6" fill="#225c7a"/><circle cx="1020" cy="286" r="6" fill="#225c7a"/>',
            '<text x="894" y="490" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">Year</text>',
            '<text x="710" y="350" transform="rotate(-90 710 350)" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">Mean beak depth</text>',
            '<text x="90" y="552" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">Natural selection link</text>',
        ])
        for i, prompt in enumerate(["Environmental pressure:", "Advantageous variation:", "Population change:"]):
            y = 574 + i * 38
            parts.append(f'<text x="110" y="{y}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{prompt}</text>')
            parts.append(f'<line x1="342" y1="{y}" x2="1110" y2="{y}" stroke="#9ca3af" stroke-width="1.5"/>')
        return "Finch Beak Data Table And Graph", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "DATA", "\n".join(parts))
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Science data table and graph stimulus</text>',
        '<rect x="90" y="204" width="510" height="300" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
        '<text x="112" y="236" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">Experimental results</text>',
    ]
    cols = [("Trial", 88), ("Variable X", 134), ("Response Y", 134), ("Unit", 92)]
    rows = [("1", "0", "2.1", "units"), ("2", "5", "5.8", "units"), ("3", "10", "9.4", "units"), ("4", "15", "12.7", "units")]
    x = 112
    for label, width in cols:
        parts.append(f'<rect x="{x}" y="258" width="{width}" height="36" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{x + 10}" y="282" font-family="Arial, sans-serif" font-size="13" font-weight="800" fill="#18212b">{label}</text>')
        x += width
    for row_idx, row in enumerate(rows, start=1):
        y = 258 + row_idx * 36
        x = 112
        for value, (_label, width) in zip(row, cols):
            parts.append(f'<rect x="{x}" y="{y}" width="{width}" height="36" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
            parts.append(f'<text x="{x + 10}" y="{y + 24}" font-family="Arial, sans-serif" font-size="13" fill="#18212b">{value}</text>')
            x += width
    parts.extend([
        '<rect x="676" y="204" width="470" height="300" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
        '<line x1="744" y1="456" x2="1094" y2="456" stroke="#18212b" stroke-width="3"/>',
        '<line x1="744" y1="456" x2="744" y2="246" stroke="#18212b" stroke-width="3"/>',
        '<polyline points="754,430 840,386 928,326 1020,276" fill="none" stroke="#225c7a" stroke-width="5"/>',
        '<circle cx="754" cy="430" r="6" fill="#225c7a"/><circle cx="840" cy="386" r="6" fill="#225c7a"/><circle cx="928" cy="326" r="6" fill="#225c7a"/><circle cx="1020" cy="276" r="6" fill="#225c7a"/>',
        '<text x="894" y="490" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">Variable X</text>',
        '<text x="710" y="350" transform="rotate(-90 710 350)" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" font-weight="800" fill="#18212b">Response Y</text>',
        '<text x="90" y="552" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">Interpretation prompts</text>',
    ])
    for i, prompt in enumerate(["Describe the trend:", "Calculate the change:", "State a conclusion supported by data:"]):
        y = 574 + i * 38
        parts.append(f'<text x="110" y="{y}" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{prompt}</text>')
        parts.append(f'<line x1="398" y1="{y}" x2="1110" y2="{y}" stroke="#9ca3af" stroke-width="1.5"/>')
    return "Science Data Table And Graph Stimulus", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "DATA", "\n".join(parts))


def case_decision_visual(pack: dict[str, Any]) -> tuple[str, str]:
    questions = pack_questions(pack)
    scenario = clean((questions[0].get("_section_stimulus") if questions else "") or pack.get("assessment_title", ""), 260)
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Case facts and decision table</text>',
        '<rect x="90" y="204" width="480" height="280" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
        '<text x="112" y="234" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">Scenario summary</text>',
        text_block(wrap_text(scenario or "Teacher inserts the approved case scenario here.", 48, 7), 112, 266, 15, "500", "#18212b", 18),
    ]
    labels = ["Stakeholder", "Problem", "Evidence", "Recommendation"]
    for i, label in enumerate(labels):
        y = 204 + i * 70
        parts.append(f'<rect x="640" y="{y}" width="420" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="660" y="{y + 30}" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">{label}</text>')
        parts.append(f'<line x1="840" y1="{y + 30}" x2="1038" y2="{y + 30}" stroke="#9ca3af" stroke-width="1.5"/>')
    headers = ["Claim", "Case evidence", "Reasoned action"]
    xs = [90, 430, 770]
    for i, header in enumerate(headers):
        parts.append(f'<rect x="{xs[i]}" y="532" width="310" height="40" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{xs[i] + 12}" y="558" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{header}</text>')
        parts.append(f'<rect x="{xs[i]}" y="572" width="310" height="58" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
    return "Case Facts And Decision Table", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "CASE STUDY", "\n".join(parts))


def investigation_visual(pack: dict[str, Any]) -> tuple[str, str]:
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Investigation method, data and conclusion</text>',
    ]
    steps = ["Aim", "Variables", "Method", "Results", "Conclusion"]
    for i, step in enumerate(steps):
        x = 100 + i * 210
        parts.append(f'<rect x="{x}" y="214" width="156" height="56" fill="#ffffff" stroke="#18212b" stroke-width="2"/>')
        parts.append(f'<text x="{x + 78}" y="249" text-anchor="middle" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">{step}</text>')
        if i < len(steps) - 1:
            parts.append(f'<line x1="{x + 156}" y1="242" x2="{x + 204}" y2="242" stroke="#18212b" stroke-width="3"/>')
    cols = ["Reading", "Observation", "Unit", "Pattern / anomaly"]
    xs = [92, 312, 596, 764]
    widths = [220, 284, 168, 340]
    for i, col in enumerate(cols):
        parts.append(f'<rect x="{xs[i]}" y="332" width="{widths[i]}" height="42" fill="#e5e7eb" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{xs[i] + 12}" y="360" font-family="Arial, sans-serif" font-size="15" font-weight="800" fill="#18212b">{col}</text>')
    for row in range(4):
        y = 374 + row * 48
        for i in range(len(cols)):
            parts.append(f'<rect x="{xs[i]}" y="{y}" width="{widths[i]}" height="48" fill="#ffffff" stroke="#18212b" stroke-width="1"/>')
    parts.append('<rect x="92" y="594" width="1012" height="46" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
    parts.append('<text x="110" y="624" font-family="Arial, sans-serif" font-size="16" font-weight="800" fill="#18212b">Conclusion supported by data:</text>')
    return "Investigation Data And Method Sheet", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "INVESTIGATION", "\n".join(parts))


def language_planner_visual(pack: dict[str, Any]) -> tuple[str, str]:
    parts = [
        '<text x="82" y="166" font-family="Arial, sans-serif" font-size="24" font-weight="800" fill="#18212b">Text response and language planner</text>',
        '<rect x="92" y="204" width="470" height="330" fill="#ffffff" stroke="#18212b" stroke-width="2"/>',
        '<text x="112" y="234" font-family="Arial, sans-serif" font-size="18" font-weight="800" fill="#18212b">Reading / viewing text panel</text>',
    ]
    prompts = ["Main idea", "Tone / purpose", "Language feature", "Evidence"]
    for i, prompt in enumerate(prompts):
        y = 278 + i * 58
        parts.append(f'<text x="112" y="{y}" font-family="Arial, sans-serif" font-size="16" font-weight="800" fill="#18212b">{prompt}</text>')
        parts.append(f'<line x1="276" y1="{y}" x2="530" y2="{y}" stroke="#9ca3af" stroke-width="1.5"/>')
    boxes = [("Plan", 640, 214), ("Draft", 640, 326), ("Edit", 640, 438)]
    for label, x, y in boxes:
        parts.append(f'<rect x="{x}" y="{y}" width="420" height="74" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
        parts.append(f'<text x="{x + 18}" y="{y + 30}" font-family="Arial, sans-serif" font-size="17" font-weight="800" fill="#18212b">{label}</text>')
        parts.append(f'<line x1="{x + 110}" y1="{y + 30}" x2="{x + 390}" y2="{y + 30}" stroke="#9ca3af" stroke-width="1.5"/>')
        parts.append(f'<line x1="{x + 110}" y1="{y + 54}" x2="{x + 390}" y2="{y + 54}" stroke="#9ca3af" stroke-width="1.5"/>')
    parts.append('<rect x="92" y="574" width="968" height="44" fill="#ffffff" stroke="#18212b" stroke-width="1.5"/>')
    parts.append('<text x="110" y="602" font-family="Arial, sans-serif" font-size="16" font-weight="800" fill="#18212b">Language structures checked:</text>')
    return "Text Response And Language Planner", assessment_sheet(pack["subject"], clean(pack["assessment_title"]), "LANGUAGE TASK", "\n".join(parts))


def generic_visual(pack: dict[str, Any]) -> tuple[str, str]:
    sections = pack.get("sections") or []
    rows = []
    for i, section in enumerate(sections[:5]):
        y = 210 + i * 66
        rows.append(f'<rect x="112" y="{y}" width="1000" height="52" rx="8" fill="{"#ffffff" if i % 2 else "#eef6f3"}" stroke="{PALETTE["line"]}"/>')
        rows.append(f'<text x="136" y="{y + 33}" font-family="Arial" font-size="22" font-weight="800" fill="{PALETTE["ink"]}">{esc(clean(section.get("title"), 70))}</text>')
        rows.append(f'<text x="900" y="{y + 33}" font-family="Arial" font-size="20" font-weight="800" fill="{PALETTE["blue"]}">{esc(clean(section.get("mode"), 28))}</text>')
    body = "\n".join(
        [
            '<text x="82" y="184" font-family="Arial" font-size="27" font-weight="800" fill="#18212b">FIGURE 1: Assessment structure</text>',
            *rows,
        ]
    )
    return "Assessment Structure Map", svg_base(pack["subject"], clean(pack["assessment_title"]), body)


def choose_visuals(pack: dict[str, Any], design_law: dict[str, Any], assessment_design: dict[str, Any] | None = None) -> list[tuple[str, str]]:
    blueprint = visual_family_for_pack(pack, assessment_design)
    if blueprint == "language_integrated_assessment" and evidence_cards(pack):
        return []
    modes = {section.get("mode") for section in pack.get("sections", [])}
    planned_kinds = planned_visual_kinds(pack)
    if planned_kinds:
        visuals: list[tuple[str, str]] = []
        is_geography = "geography" in str(pack.get("subject", "")).lower()
        is_science = any(token in str(pack.get("subject", "")).lower() for token in ["life sciences", "natural sciences", "physical sciences", "agricultural sciences", "marine sciences"])
        if "map_extract" in planned_kinds or ("data_table" in planned_kinds and is_geography):
            visuals.append(geography_map_source_visual(pack))
        elif "source_extract" in planned_kinds:
            visuals.append(source_panel_visual(pack))
        if "data_table" in planned_kinds and is_science:
            visuals.append(science_data_graph_visual(pack))
        if "source_extract" in planned_kinds and is_geography:
            visuals.append(source_extract_visual(pack))
        if "evidence_table" in planned_kinds:
            visuals.append(claim_evidence_table_visual(pack))
        if "case_table" in planned_kinds:
            visuals.append(case_decision_visual(pack))
        if "investigation_sheet" in planned_kinds:
            visuals.append(investigation_visual(pack))
        if "language_planner" in planned_kinds:
            visuals.append(language_planner_visual(pack))
        if "axis_diagram" in planned_kinds and is_geography:
            visuals.append(geography_process_diagram_visual(pack))
        elif "axis_diagram" in planned_kinds and is_science:
            visuals.append(science_labelled_diagram_visual(pack))
        elif "working_grid" in planned_kinds or "axis_diagram" in planned_kinds:
            visuals.append(calculation_workspace_visual(pack))
        if "design_log" in planned_kinds:
            visuals.extend([robot_maze_visual(pack), design_cycle_visual(pack)])
        if "performance_map" in planned_kinds:
            visuals.append(performance_visual(pack))
        if "rubric_observation_table" in planned_kinds:
            visuals.append(performance_observation_visual(pack))
        if visuals:
            deduped: list[tuple[str, str]] = []
            seen = set()
            for title, svg in visuals:
                if title not in seen:
                    deduped.append((title, svg))
                    seen.add(title)
            return deduped
    if blueprint == "foundation_activity_assessment" and pack.get("subject") == "Mathematics":
        return [number_line_visual(pack), foundation_observation_visual(pack)]
    if blueprint == "calculation_problem_solving":
        return [calculation_workspace_visual(pack)]
    if blueprint in {"source_based_plus_extended_response", "source_based_plus_essay"}:
        return [source_panel_visual(pack)]
    if blueprint == "case_study_structured_questions":
        return [case_decision_visual(pack)]
    if blueprint == "data_diagram_practical_investigation":
        return [investigation_visual(pack)]
    if blueprint == "language_integrated_assessment":
        return [language_planner_visual(pack)]
    if blueprint == "practical_project_design_task" or "practical_project_design_task" in modes:
        return [robot_maze_visual(pack), design_cycle_visual(pack)]
    if blueprint == "practical_performance_or_portfolio" or "reflection" in modes:
        return [performance_visual(pack), performance_observation_visual(pack)]
    return [generic_visual(pack)]


def build_assets(job_dir: Path, pack: dict[str, Any], design_law: dict[str, Any], assessment_design: dict[str, Any] | None = None) -> list[dict[str, str]]:
    visual_dir = job_dir / "assessment_visuals"
    visual_dir.mkdir(exist_ok=True)
    for existing in visual_dir.glob("*"):
        if existing.is_file():
            existing.unlink()
    assets = []
    for index, (title, svg) in enumerate(choose_visuals(pack, design_law, assessment_design), start=1):
        stem = f"{index:02d}_{re.sub(r'[^a-z0-9]+', '_', title.lower()).strip('_')}"
        svg_path = visual_dir / f"{stem}.svg"
        png_path = visual_dir / f"{stem}.png"
        svg_path.write_text(svg, encoding="utf-8")
        rendered = render_png(svg_path, png_path)
        quality_flags = []
        title_blob = title.lower()
        if any(token in title_blob for token in ["skull", "diagram", "model"]):
            quality_flags.append("generated_schematic_not_authoritative_source")
        if "skull" in title_blob:
            quality_flags.append("anatomy_requires_curated_or_educator_verified_asset")
        assets.append(
            {
                "title": title,
                "svg": str(svg_path),
                "png": str(png_path) if rendered else "",
                "status": "png_rendered" if rendered else "svg_only",
                "quality_flags": quality_flags,
            }
        )
    return assets


def load_assessment_design(path: Path | None) -> dict[str, Any]:
    if not path or not path.exists():
        return {}
    return load_json(path)


def assessment_design_for_pack(pack: dict[str, Any], assessment_design: dict[str, Any]) -> dict[str, Any]:
    profile_id = pack.get("canonical_profile_id")
    profiles_by_id = assessment_design.get("profiles_by_id") or {}
    if profile_id and profile_id in profiles_by_id:
        return profiles_by_id[profile_id]
    for profile in assessment_design.get("profiles") or []:
        if profile.get("profile_id") == profile_id:
            return profile
    return {}


def add_paragraph(doc: Any, text: str, bold_label: str | None = None) -> Any:
    p = doc.add_paragraph()
    if bold_label:
        p.add_run(bold_label).bold = True
    p.add_run(text)
    return p


def is_afrikaans_pack(pack: dict[str, Any]) -> bool:
    return (
        str(pack.get("language_of_assessment") or "").lower() == "afrikaans"
        or "afrikaans" in str(pack.get("canonical_profile_id") or "").lower()
        or "afrikaans" in str(pack.get("subject") or "").lower()
    )


def localized_source_type(value: Any, afrikaans: bool) -> str:
    text = str(value or "").replace("_", " ").strip()
    if not afrikaans:
        return text.title()
    return {
        "text extract": "teksuittreksel",
        "cartoon": "spotprent",
        "graph or chart": "grafiek of diagram",
        "photograph or image": "foto of beeld",
        "data table": "datatabel",
        "diagram or model": "diagram of model",
        "map extract": "kaart-uittreksel",
        "synoptic weather map": "sinoptiese weerkaart",
    }.get(text.lower(), text)


def localized_visual_title(value: Any, afrikaans: bool) -> str:
    text = str(value or "").strip()
    if not afrikaans:
        return text
    return {
        "Text Response And Language Planner": "Teksrespons- en taalbeplanner",
        "Language Structures Table": "Taalstrukturetabel",
        "Response Planner": "Antwoordbeplanner",
    }.get(text, text)


def add_cover_table(doc: Any, pack: dict[str, Any]) -> None:
    afrikaans = is_afrikaans_pack(pack)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = (
        [
            ("VAK", str(pack.get("subject", "")).upper()),
            ("GRAAD", str(pack.get("grade", ""))),
            ("ASSESSERINGSTAAK", clean(pack.get("assessment_title", ""), 160).upper()),
            ("PUNTE", str(pack.get("total_marks", ""))),
            ("TYD", str(pack.get("duration", ""))),
        ]
        if afrikaans
        else [
            ("SUBJECT", str(pack.get("subject", "")).upper()),
            ("GRADE", str(pack.get("grade", ""))),
            ("ASSESSMENT TASK", clean(pack.get("assessment_title", ""), 160).upper()),
            ("MARKS", str(pack.get("total_marks", ""))),
            ("TIME", str(pack.get("duration", ""))),
        ]
    )
    for index, (label, value) in enumerate(rows):
        table.rows[index].cells[0].text = label
        table.rows[index].cells[1].text = value
        for cell in table.rows[index].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True if cell is table.rows[index].cells[0] else run.bold


def default_instructions(pack: dict[str, Any]) -> list[str]:
    if is_afrikaans_pack(pack):
        return [
            "Lees alle instruksies noukeurig.",
            "Beantwoord alle vrae of voltooi alle aktiwiteite soos aangedui.",
            "Wys alle berekeninge, beplanning, diagramme of verduidelikings waar dit vereis word.",
            "Gebruik die ruimte wat deur die opvoeder, antwoordboek of leerdertaakboek verskaf word.",
            "Skryf netjies en duidelik.",
        ]
    subject = str(pack.get("subject") or "this assessment")
    instructions = [
        "Read all instructions carefully.",
        "Answer all questions or complete all activities as instructed.",
        "Show all working where calculations, planning, diagrams or explanations are required.",
        "Use the space provided by the educator, answer booklet or learner task book.",
        "Write neatly and clearly.",
    ]
    if "Dance" in subject:
        instructions.extend(
            [
                "Use correct subject terminology where appropriate.",
                "Follow all safety instructions during practical performance.",
            ]
        )
    if "Coding" in subject or "Robotics" in subject:
        instructions.extend(
            [
                "Label diagrams clearly.",
                "Record testing results honestly, including what did not work.",
            ]
        )
    if "Mathematics" in subject:
        instructions.append("Do not use a calculator unless the educator allows it.")
    return instructions


def add_answer_lines(doc: Any, count: int = 3) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.shared import Pt

    table = doc.add_table(rows=max(1, count), cols=1)
    table.autofit = True
    for row in table.rows:
        row.height = Pt(21)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        cell = row.cells[0]
        cell.text = "\u00a0"
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ["top", "left", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "nil")
            borders.append(border)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")
        borders.append(bottom)
        tc_pr.append(borders)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)


def add_question_paragraph(doc: Any, question: dict[str, Any], keep_with_next: bool = False) -> Any:
    from docx.enum.text import WD_TAB_ALIGNMENT

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = keep_with_next
    section = doc.sections[-1]
    text_width = section.page_width - section.left_margin - section.right_margin
    paragraph.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run(f"{question.get('number', '')}. ").bold = True
    paragraph.add_run(str(question.get("question", "")))
    paragraph.add_run("\t")
    paragraph.add_run(f"({question.get('marks', 0)})").bold = True
    return paragraph


def set_source_cell_margins(cell: Any, twips: int = 120) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ["top", "left", "bottom", "right"]:
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")


def add_visual_material(doc: Any, assets: list[dict[str, str]], pack: dict[str, Any] | None = None) -> None:
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if not assets:
        return
    source_policy = ((pack or {}).get("source_embedding") or {}).get("generated_visual_policy")
    if source_policy == "suppress_generated_visuals_when_source_assets_present" and normalized_source_assets(pack or {}):
        return
    if source_policy == "suppress_generated_visuals_when_evidence_cards_present" and evidence_cards(pack or {}):
        return
    afrikaans = is_afrikaans_pack(pack or {})
    doc.add_heading("VISUELE MATERIAAL" if afrikaans else "VISUAL MATERIAL", level=1)
    for index, asset in enumerate(assets, start=1):
        figure_label = "FIGUUR" if afrikaans else "FIGURE"
        caption = doc.add_paragraph(f"{figure_label} {index}: {localized_visual_title(asset['title'], afrikaans)}")
        caption.runs[0].bold = True
        caption.paragraph_format.space_before = 0
        caption.paragraph_format.space_after = 0
        caption.paragraph_format.keep_with_next = True
        if asset.get("png"):
            image_paragraph = doc.add_paragraph()
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_paragraph.paragraph_format.space_before = 0
            image_paragraph.paragraph_format.space_after = Pt(6)
            image_paragraph.paragraph_format.keep_together = True
            image_run = image_paragraph.add_run()
            image_run.add_picture(asset["png"], width=Inches(6.2))
        else:
            doc.add_paragraph("Visuele materiaal word afsonderlik verskaf." if afrikaans else "Visual supplied separately.")


def pack_source_assets(pack: dict[str, Any]) -> list[dict[str, Any]]:
    assets = []
    assets.extend(pack.get("source_assets") or [])
    visual_blueprint = pack.get("visual_blueprint") or {}
    assets.extend(visual_blueprint.get("source_assets") or [])
    return [asset for asset in assets if asset.get("png") or asset.get("preferred_png") or asset.get("curated_png") or asset.get("bank_png")]


def normalized_source_assets(pack: dict[str, Any]) -> list[dict[str, Any]]:
    afrikaans = is_afrikaans_pack(pack)
    normalized = []
    for index, asset in enumerate(pack_source_assets(pack), start=1):
        image_path = asset.get("png") or asset.get("preferred_png") or asset.get("curated_png") or asset.get("bank_png")
        resolved = Path(str(image_path))
        if not resolved.is_absolute():
            resolved = ROOT / resolved
        normalized.append(
            {
                "source_number": index,
                "title": asset.get("title") or localized_source_type(asset.get("source_type"), afrikaans) or f"Source {index}",
                "source_type": asset.get("source_type"),
                "paper": asset.get("paper"),
                "page": asset.get("page"),
                "bank_png": asset.get("bank_png"),
                "curated_png": asset.get("curated_png"),
                "preferred_png": asset.get("preferred_png"),
                "png": str(resolved) if resolved.exists() else str(image_path),
                "asset_status": asset.get("asset_status"),
                "quality_flags": asset.get("quality_flags") or [],
                "source_quality_level": asset.get("source_quality_level"),
                "release_gate": asset.get("release_gate"),
                "embedding_policy": asset.get("embedding_policy"),
            }
        )
    return normalized


def add_source_material(doc: Any, pack: dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    source_assets = normalized_source_assets(pack)
    if not source_assets:
        return
    afrikaans = is_afrikaans_pack(pack)
    source_label = "BRON" if afrikaans else "SOURCE"
    doc.add_heading("BRONMATERIAAL" if afrikaans else "SOURCE MATERIAL", level=1)
    for index, asset in enumerate(source_assets, start=1):
        title = asset.get("title") or localized_source_type(asset.get("source_type"), afrikaans) or f"{source_label} {index}"
        if afrikaans and str(title).lower().startswith("official-paper exemplar"):
            title = f"Amptelike vraestelvoorbeeld: {localized_source_type(asset.get('source_type'), True)}"
        caption = doc.add_paragraph(f"{source_label} {index}: {title}")
        caption.runs[0].bold = True
        caption.paragraph_format.space_before = 0
        caption.paragraph_format.space_after = 0
        caption.paragraph_format.keep_with_next = True
        image_path = asset.get("png")
        if image_path:
            image_path = Path(str(image_path))
            image_paragraph = doc.add_paragraph()
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_paragraph.paragraph_format.space_before = 0
            image_paragraph.paragraph_format.space_after = Pt(6)
            image_paragraph.paragraph_format.keep_together = True
            image_paragraph.add_run().add_picture(str(image_path), width=Inches(5.2))
        details = []
        detail_labels = (
            {"paper": "Vraestel", "page": "Bladsy", "source_type": "Brontipe", "snippet": "Uittreksel"}
            if afrikaans
            else {"paper": "Paper", "page": "Page", "source_type": "Source Type", "snippet": "Snippet"}
        )
        for key in ["paper", "page", "source_type", "snippet"]:
            if asset.get(key):
                if key == "source_type":
                    value = localized_source_type(asset[key], True) if afrikaans else str(asset[key]).replace("_", " ").title()
                else:
                    value = asset[key]
                details.append(f"{detail_labels[key]}: {value}")
        if details:
            note = doc.add_paragraph(" | ".join(details[:4]))
            note.paragraph_format.space_after = Pt(6)


def evidence_cards(pack: dict[str, Any]) -> list[dict[str, Any]]:
    cards = []
    cards.extend(pack.get("evidence_cards") or [])
    cards.extend(pack.get("stimulus_cards") or [])
    visual_blueprint = pack.get("visual_blueprint") or {}
    cards.extend(visual_blueprint.get("evidence_cards") or [])
    return [card for card in cards if str(card.get("content") or card.get("description") or "").strip()]


def add_evidence_material(doc: Any, pack: dict[str, Any]) -> None:
    from docx.shared import Pt

    cards = evidence_cards(pack)
    if not cards:
        return
    afrikaans = is_afrikaans_pack(pack)
    language_task = "language_integrated" in str(pack.get("blueprint") or "").lower()
    if language_task:
        heading_text = "BRONTEKSTE" if afrikaans else "SOURCE TEXTS"
    else:
        heading_text = "BEWYSMATERIAAL" if afrikaans else "EVIDENCE / STIMULUS MATERIAL"
    doc.add_heading(heading_text, level=1)
    for index, card in enumerate(cards, start=1):
        fallback_label = f"Bron {index}" if afrikaans else f"Source {index}"
        label = str(card.get("label") or card.get("id") or fallback_label)
        title = str(card.get("title") or card.get("type") or "")
        source_table = doc.add_table(rows=1, cols=1)
        source_table.style = "Table Grid"
        source_table.autofit = True
        cell = source_table.cell(0, 0)
        set_source_cell_margins(cell)
        heading = cell.paragraphs[0]
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_after = Pt(2)
        heading.add_run(f"{label}: ").bold = True
        if title:
            heading.add_run(title).bold = True
        meta_parts = []
        for key, display in [
            ("type", "Tipe" if afrikaans else "Type"),
            ("provenance", "Herkoms" if afrikaans else "Provenance"),
            ("date", "Datum" if afrikaans else "Date"),
            ("context", "Konteks" if afrikaans else "Context"),
        ]:
            if card.get(key):
                meta_parts.append(f"{display}: {card[key]}")
        if meta_parts and not language_task:
            meta = cell.add_paragraph(" | ".join(meta_parts[:4]))
            meta.paragraph_format.space_after = Pt(3)
            for run in meta.runs:
                run.font.size = Pt(9)
        body_text = str(card.get("content") or card.get("description") or "")
        body = cell.add_paragraph()
        body.paragraph_format.space_after = Pt(0)
        body.add_run(body_text)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.space_before = Pt(0)


def add_rubric_table(doc: Any, pack: dict[str, Any], design_law: dict[str, Any], title: str = "ASSESSMENT INSTRUMENT") -> None:
    if not pack.get("rubric"):
        return
    afrikaans = is_afrikaans_pack(pack)
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = design_law.get("document_style", {}).get("table_style", "Table Grid")
    headers = ["Kriterium", "Punte", "Beskrywing"] if afrikaans else ["Criterion", "Marks", "Descriptor"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for item in pack["rubric"]:
        cells = table.add_row().cells
        cells[0].text = str(item.get("criterion", ""))
        cells[1].text = str(item.get("marks", ""))
        cells[2].text = str(item.get("descriptor", ""))


def add_review_checklist(doc: Any, pack: dict[str, Any]) -> None:
    if not pack.get("teacher_review_checklist"):
        return
    is_afrikaans = (
        str(pack.get("language_of_assessment") or "").lower() == "afrikaans"
        or "afrikaans" in str(pack.get("canonical_profile_id") or "").lower()
        or "afrikaans" in str(pack.get("subject") or "").lower()
    )
    doc.add_heading("OPVOEDERKONTROLELYS" if is_afrikaans else "EDUCATOR REVIEW CHECKLIST", level=1)
    for item in pack["teacher_review_checklist"]:
        doc.add_paragraph(str(item), style="List Bullet")


def setup_doc(pack: dict[str, Any], design_law: dict[str, Any]) -> Any:
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
    afrikaans = is_afrikaans_pack(pack)
    style = doc.styles["Normal"]
    style.font.name = design_law.get("document_style", {}).get("font", "Arial")
    style.font.size = Pt(float(design_law.get("document_style", {}).get("body_size_pt", 10.5)))

    subject_title = str(pack.get("subject", "Assessering" if afrikaans else "Assessment"))
    if afrikaans and subject_title.lower() == "afrikaans language":
        subject_title = "Afrikaans"
    title = doc.add_heading(subject_title.upper(), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(str(pack.get("assessment_title", "")).upper())
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.bold = True
    doc.add_paragraph("")
    add_cover_table(doc, pack)
    doc.add_paragraph("")
    page_note = design_law.get("document_style", {}).get("page_note", "")
    if afrikaans and page_note == "Educator approval required before classroom use.":
        page_note = "Opvoedergoedkeuring word voor klasgebruik vereis."
    note = doc.add_paragraph(page_note)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    return doc


def add_task_brief(doc: Any, pack: dict[str, Any], heading: str) -> None:
    doc.add_heading(heading, level=1)
    for section in pack.get("sections", []):
        section_title = str(section.get("title", "Task")).upper()
        if section_title != heading.upper():
            doc.add_heading(section_title, level=2)
        if section.get("instructions"):
            add_paragraph(doc, str(section.get("instructions")), "Brief: ")
        if section.get("stimulus"):
            add_paragraph(doc, str(section.get("stimulus")), "Context: ")
        for question in section.get("questions", []):
            add_question_paragraph(doc, question)


def add_marking_guideline(doc: Any, pack: dict[str, Any]) -> None:
    is_afrikaans = (
        str(pack.get("language_of_assessment") or "").lower() == "afrikaans"
        or "afrikaans" in str(pack.get("canonical_profile_id") or "").lower()
        or "afrikaans" in str(pack.get("subject") or "").lower()
    )
    doc.add_page_break()
    doc.add_heading("NASIENRIGLYN" if is_afrikaans else "MARKING GUIDELINE", level=1)
    for section_index, section in enumerate(pack.get("sections", []), start=1):
        label = "TAAK" if is_afrikaans else "TASK"
        fallback = "Assesseringstaak" if is_afrikaans else "Assessment Task"
        doc.add_heading(f"{label} {section_index}: {section.get('title', fallback).upper()}", level=2)
        for question in section.get("questions", []):
            add_question_paragraph(doc, question)
            for memo_item in question.get("memo") or []:
                doc.add_paragraph(str(memo_item), style="List Bullet")


def render_question_paper(
    job_dir: Path,
    pack: dict[str, Any],
    design_law: dict[str, Any],
    assets: list[dict[str, str]],
    include_memo_sections: bool = True,
    output_name: str = "ASSESSMENT_PACK_FORMATTED.docx",
) -> Path:
    doc = setup_doc(pack, design_law)

    doc.add_heading("INSTRUCTIONS AND INFORMATION", level=1)
    for index, instruction in enumerate(default_instructions(pack), start=1):
        doc.add_paragraph(f"{index}. {instruction}")

    add_source_material(doc, pack)
    add_evidence_material(doc, pack)
    add_visual_material(doc, assets, pack)

    doc.add_heading("QUESTIONS", level=1)
    for section_index, section in enumerate(pack.get("sections", []), start=1):
        section_title = str(section.get("title") or "Assessment Task").upper()
        rendered_title = section_title if section_title.startswith("QUESTION") else f"QUESTION {section_index}: {section_title}"
        doc.add_heading(rendered_title, level=2)
        if section.get("instructions"):
            add_paragraph(doc, str(section.get("instructions")), "Instructions: ")
        if section.get("stimulus"):
            add_paragraph(doc, str(section.get("stimulus")), "Stimulus: ")
        for question in section.get("questions", []):
            add_question_paragraph(doc, question, keep_with_next=True)
            requested_lines = int(question.get("answer_lines") or 0)
            if requested_lines:
                add_answer_lines(doc, requested_lines)
            elif int(question.get("marks", 0) or 0) <= 5:
                add_answer_lines(doc, 2)
            else:
                add_answer_lines(doc, 4)

    if include_memo_sections:
        add_marking_guideline(doc, pack)
        add_rubric_table(doc, pack, design_law, "RUBRIC")
        add_review_checklist(doc, pack)

    out = job_dir / output_name
    doc.save(out)
    return out


def render_language_task(
    job_dir: Path,
    pack: dict[str, Any],
    design_law: dict[str, Any],
    assets: list[dict[str, str]],
    assessment_design: dict[str, Any],
    include_memo_sections: bool = True,
    output_name: str = "ASSESSMENT_PACK_FORMATTED.docx",
) -> Path:
    is_afrikaans = (
        str(pack.get("language_of_assessment") or "").lower() == "afrikaans"
        or "afrikaans" in str(pack.get("canonical_profile_id") or "").lower()
        or "afrikaans" in str(pack.get("subject") or "").lower()
    )
    doc = setup_doc(pack, design_law)
    doc.add_heading("INSTRUKSIES EN INLIGTING" if is_afrikaans else "INSTRUCTIONS AND INFORMATION", level=1)
    instructions = (
        [
            "Lees elke bronteks en vraag noukeurig.",
            "Beantwoord alle vrae in Afrikaans tensy anders aangedui.",
            "Gebruik bewyse uit die verskafde brontekste waar dit vereis word.",
            "Beplan skriftelike antwoorde voordat jy die finale antwoord skryf.",
            "Skryf netjies en let op spelling, leestekens, grammatika en register.",
        ]
        if is_afrikaans
        else [
            "Read each source text and question carefully.",
            "Answer in the language of assessment unless instructed otherwise.",
            "Use evidence from the supplied source texts where required.",
            "Plan written responses before writing the final answer.",
            "Write neatly and observe spelling, punctuation, grammar and register.",
        ]
    )
    for index, instruction in enumerate(instructions, start=1):
        doc.add_paragraph(f"{index}. {instruction}")

    add_source_material(doc, pack)
    add_evidence_material(doc, pack)
    add_visual_material(doc, assets, pack)

    doc.add_heading("TAALASSESSERINGSTAAK" if is_afrikaans else "LANGUAGE ASSESSMENT TASK", level=1)
    for section_index, section in enumerate(pack.get("sections", []), start=1):
        title = str(section.get("title", "Taalvaardigheid" if is_afrikaans else "Language Skill")).upper()
        heading_label = "AFDELING" if is_afrikaans else "SECTION"
        section_heading = title if title.startswith(heading_label) else f"{heading_label} {section_index}: {title}"
        section_heading_paragraph = doc.add_heading(section_heading, level=2)
        section_heading_paragraph.paragraph_format.keep_with_next = True
        if section.get("instructions"):
            instructions_paragraph = add_paragraph(doc, str(section.get("instructions")), "Instruksies: " if is_afrikaans else "Instructions: ")
            instructions_paragraph.paragraph_format.keep_with_next = True
        if section.get("stimulus"):
            stimulus_paragraph = add_paragraph(doc, str(section.get("stimulus")), "Bron: " if is_afrikaans else "Source: ")
            stimulus_paragraph.paragraph_format.keep_with_next = True
        for question in section.get("questions", []):
            add_question_paragraph(doc, question, keep_with_next=True)
            marks = int(question.get("marks", 0) or 0)
            requested_lines = int(question.get("answer_lines") or 0)
            if requested_lines:
                add_answer_lines(doc, requested_lines)
            elif marks <= 3:
                add_answer_lines(doc, 1)
            elif marks <= 8:
                add_answer_lines(doc, 3)
            else:
                add_answer_lines(doc, 9)

    if include_memo_sections and assessment_design.get("generation_constraints"):
        doc.add_heading("TAALTAAKVOORWAARDES" if is_afrikaans else "LANGUAGE TASK CONDITIONS", level=1)
        for item in assessment_design["generation_constraints"][:5]:
            text = str(item)
            if is_afrikaans:
                text = {
                    "Use language-skills integration instead of a generic content subject exam.": "Gebruik geintegreerde taalvaardighede eerder as 'n generiese inhoudsvraestel.",
                    "Scale text length, vocabulary, and writing load to the phase.": "Pas tekslengte, woordeskat en skryflading by die fase aan.",
                    "Do not import subject-case-study structure unless the language task needs it.": "Moenie 'n vak-gevallestudie-struktuur invoer tensy die taaltaak dit vereis nie.",
                }.get(text, text)
            doc.add_paragraph(text, style="List Bullet")

    if include_memo_sections:
        add_marking_guideline(doc, pack)
        add_rubric_table(doc, pack, design_law, "TAALRUBRIEK" if is_afrikaans else "LANGUAGE RUBRIC")
        add_review_checklist(doc, pack)
    out = job_dir / output_name
    doc.save(out)
    return out


def render_performance_task(
    job_dir: Path,
    pack: dict[str, Any],
    design_law: dict[str, Any],
    assets: list[dict[str, str]],
    assessment_design: dict[str, Any],
    include_memo_sections: bool = True,
    output_name: str = "ASSESSMENT_PACK_FORMATTED.docx",
) -> Path:
    doc = setup_doc(pack, design_law)
    doc.add_heading("INSTRUCTIONS AND INFORMATION", level=1)
    instructions = [
        "This is a practical performance assessment task.",
        "The learner performs the task under educator supervision.",
        "The educator records observable evidence against the assessment instrument.",
        "Written or oral reflection may be assessed only where it supports the performance evidence.",
        "Safety and performance-space readiness must be checked before the task begins.",
    ]
    for index, instruction in enumerate(instructions, start=1):
        doc.add_paragraph(f"{index}. {instruction}")

    add_source_material(doc, pack)
    add_evidence_material(doc, pack)
    add_visual_material(doc, assets, pack)
    add_task_brief(doc, pack, "PERFORMANCE ASSESSMENT TASK")

    if assessment_design.get("task_requirements"):
        doc.add_heading("TASK REQUIREMENTS", level=1)
        for requirement in assessment_design["task_requirements"][:5]:
            name = requirement.get("name", "Task")
            detail = ", ".join(f"{key}: {value}" for key, value in requirement.items() if key != "name")
            doc.add_paragraph(f"{name} - {detail}", style="List Bullet")
    if assessment_design.get("annual_weightings"):
        doc.add_heading("ANNUAL ASSESSMENT CONTEXT", level=1)
        for key, value in assessment_design["annual_weightings"].items():
            doc.add_paragraph(f"{key.replace('_', ' ').title()}: {value}%", style="List Bullet")

    add_rubric_table(doc, pack, design_law)
    if include_memo_sections:
        add_marking_guideline(doc, pack)
        add_review_checklist(doc, pack)
    out = job_dir / output_name
    doc.save(out)
    return out


def render_project_task(
    job_dir: Path,
    pack: dict[str, Any],
    design_law: dict[str, Any],
    assets: list[dict[str, str]],
    assessment_design: dict[str, Any],
    include_memo_sections: bool = True,
    output_name: str = "ASSESSMENT_PACK_FORMATTED.docx",
) -> Path:
    doc = setup_doc(pack, design_law)
    doc.add_heading("INSTRUCTIONS AND INFORMATION", level=1)
    instructions = [
        "This is a practical project/design assessment task.",
        "Complete the design, build, testing, reflection, and presentation evidence requested by the educator.",
        "Keep diagrams, test results, and changes visible for marking.",
        "Use safe classroom materials and follow educator supervision instructions.",
    ]
    for index, instruction in enumerate(instructions, start=1):
        doc.add_paragraph(f"{index}. {instruction}")
    add_source_material(doc, pack)
    add_evidence_material(doc, pack)
    add_visual_material(doc, assets, pack)
    add_task_brief(doc, pack, "PRACTICAL ASSESSMENT TASK")

    doc.add_heading("DESIGN EVIDENCE LOG", level=1)
    table = doc.add_table(rows=5, cols=2)
    table.style = design_law.get("document_style", {}).get("table_style", "Table Grid")
    for row, label in enumerate(["Problem / need", "Plan or diagram", "Materials", "Test result", "Improvement"]):
        table.rows[row].cells[0].text = label
        table.rows[row].cells[1].text = ""

    if assessment_design.get("generation_constraints"):
        doc.add_heading("ASSESSMENT CONDITIONS", level=1)
        for item in assessment_design["generation_constraints"][:5]:
            doc.add_paragraph(str(item), style="List Bullet")
    add_rubric_table(doc, pack, design_law)
    if include_memo_sections:
        add_marking_guideline(doc, pack)
        add_review_checklist(doc, pack)
    out = job_dir / output_name
    doc.save(out)
    return out


def render_activity_sheet(
    job_dir: Path,
    pack: dict[str, Any],
    design_law: dict[str, Any],
    assets: list[dict[str, str]],
    assessment_design: dict[str, Any],
    include_memo_sections: bool = True,
    output_name: str = "ASSESSMENT_PACK_FORMATTED.docx",
) -> Path:
    doc = setup_doc(pack, design_law)
    doc.add_heading("TEACHER INSTRUCTIONS", level=1)
    instructions = [
        "Use this as a teacher-led formal assessment activity.",
        "Read prompts aloud where needed and observe learner actions.",
        "Record evidence in the checklist or rubric.",
        "Keep the activity short, concrete, and age-appropriate.",
    ]
    for index, instruction in enumerate(instructions, start=1):
        doc.add_paragraph(f"{index}. {instruction}")
    add_source_material(doc, pack)
    add_visual_material(doc, assets, pack)
    add_task_brief(doc, pack, "LEARNER ACTIVITY")

    doc.add_heading("OBSERVATION CHECKLIST", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = design_law.get("document_style", {}).get("table_style", "Table Grid")
    for index, header in enumerate(["Skill / evidence", "Observed", "Notes"]):
        table.rows[0].cells[index].text = header
    for section in pack.get("sections", []):
        for question in section.get("questions", []):
            cells = table.add_row().cells
            cells[0].text = str(question.get("question", ""))[:180]
            cells[1].text = ""
            cells[2].text = ""

    add_rubric_table(doc, pack, design_law, "MARKING RUBRIC")
    if include_memo_sections:
        add_marking_guideline(doc, pack)
        add_review_checklist(doc, pack)
    out = job_dir / output_name
    doc.save(out)
    return out


def render_docx(
    job_dir: Path,
    pack: dict[str, Any],
    design_law: dict[str, Any],
    assets: list[dict[str, str]],
    assessment_design: dict[str, Any] | None = None,
    include_memo_sections: bool = True,
    output_name: str = "ASSESSMENT_PACK_FORMATTED.docx",
) -> Path:
    assessment_design = assessment_design or {}
    shell = pack.get("render_shell") or assessment_design.get("render_shell")
    blueprint = str(pack.get("blueprint") or "")
    subject = str(pack.get("subject") or "")
    if not shell:
        if blueprint == "foundation_activity_assessment":
            shell = "activity_sheet"
        elif blueprint == "practical_project_design_task":
            shell = "project_task_sheet"
        elif blueprint == "practical_performance_or_portfolio" or "Dance" in subject:
            shell = "performance_task_sheet"
        else:
            shell = "question_paper"

    if shell == "performance_task_sheet":
        return render_performance_task(job_dir, pack, design_law, assets, assessment_design, include_memo_sections, output_name)
    if shell == "project_task_sheet":
        return render_project_task(job_dir, pack, design_law, assets, assessment_design, include_memo_sections, output_name)
    if shell == "activity_sheet":
        return render_activity_sheet(job_dir, pack, design_law, assets, assessment_design, include_memo_sections, output_name)
    if shell == "language_integrated_task":
        return render_language_task(job_dir, pack, design_law, assets, assessment_design, include_memo_sections, output_name)
    if shell == "structured_case_paper":
        return render_question_paper(job_dir, pack, design_law, assets, include_memo_sections, output_name)
    return render_question_paper(job_dir, pack, design_law, assets, include_memo_sections, output_name)


def zip_outputs(job_dir: Path) -> Path:
    target = job_dir / "HOMS_FORMATTED_ASSESSMENT_PACK.zip"
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in [
            job_dir / "ASSESSMENT_PACK_FORMATTED.docx",
            job_dir / "HOMS_FORMATTER_RECEIPT.json",
            job_dir / "HOMS_VISUAL_MANIFEST.json",
        ]:
            if path.exists():
                zf.write(path, path.relative_to(job_dir))
        visual_dir = job_dir / "assessment_visuals"
        if visual_dir.exists():
            for path in visual_dir.rglob("*"):
                if path.is_file():
                    zf.write(path, path.relative_to(job_dir))
        pack_path = job_dir / "assessment_pack.json"
        if pack_path.exists():
            pack = load_json(pack_path)
            for asset in normalized_source_assets(pack):
                image_path = Path(str(asset.get("png") or ""))
                if image_path.exists():
                    zf.write(image_path, Path("source_material") / image_path.name)
    return target


def apply_design_law(job_dir: Path, design_law_path: Path, assessment_design_path: Path | None = DEFAULT_ASSESSMENT_DESIGN) -> dict[str, Any]:
    pack_path = job_dir / "assessment_pack.json"
    if not pack_path.exists():
        raise FileNotFoundError(f"assessment_pack.json not found in {job_dir}")
    design_law = load_json(design_law_path)
    assessment_design_payload = load_assessment_design(assessment_design_path)
    pack = load_json(pack_path)
    assessment_design = assessment_design_for_pack(pack, assessment_design_payload)
    blueprint = visual_family_for_pack(pack, assessment_design)
    family = (design_law.get("visual_families") or {}).get(blueprint, {})
    source_policy = (pack.get("source_embedding") or {}).get("generated_visual_policy")
    suppress_for_cards = source_policy == "suppress_generated_visuals_when_evidence_cards_present" and evidence_cards(pack)
    suppress_for_source_assets = (
        source_policy == "suppress_generated_visuals_when_source_assets_present" and normalized_source_assets(pack)
    )
    assets = [] if (suppress_for_cards or suppress_for_source_assets) else build_assets(job_dir, pack, design_law, assessment_design)
    render_shell = assessment_design.get("render_shell")
    manifest = {
        "schema": "knowedge.homs_visual_manifest.v1",
        "created_at": utc_now(),
        "assessment_title": pack.get("assessment_title"),
        "subject": pack.get("subject"),
        "grade": pack.get("grade"),
        "canonical_profile_id": pack.get("canonical_profile_id"),
        "blueprint": blueprint,
        "render_shell": render_shell,
        "visual_family": family,
        "source_assets": normalized_source_assets(pack),
        "assets": assets,
    }
    write_json(job_dir / "HOMS_VISUAL_MANIFEST.json", manifest)
    formatted = render_docx(job_dir, pack, design_law, assets, assessment_design)
    receipt = {
        "schema": "knowedge.homs_design_law_receipt.v1",
        "created_at": utc_now(),
        "status": "formatted",
        "job_dir": str(job_dir),
        "design_law": str(design_law_path),
        "assessment_design": str(assessment_design_path) if assessment_design_path else None,
        "canonical_profile_id": pack.get("canonical_profile_id"),
        "render_shell": render_shell,
        "input_pack": str(pack_path),
        "formatted_docx": str(formatted),
        "visual_manifest": str(job_dir / "HOMS_VISUAL_MANIFEST.json"),
        "source_assets": normalized_source_assets(pack),
        "assets": assets,
        "educator_approval_required": True,
    }
    write_json(job_dir / "HOMS_FORMATTER_RECEIPT.json", receipt)
    zip_path = zip_outputs(job_dir)
    receipt["formatted_zip"] = str(zip_path)
    write_json(job_dir / "HOMS_FORMATTER_RECEIPT.json", receipt)
    return receipt


def job_dirs_from_path(path: Path) -> list[Path]:
    if (path / "assessment_pack.json").exists():
        return [path]
    return sorted(p for p in path.iterdir() if p.is_dir() and (p / "assessment_pack.json").exists())


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply HOMS design law visuals and formatting to assessment pack folders.")
    parser.add_argument("path", help="Assessment job folder or run folder containing assessment job folders.")
    parser.add_argument("--design-law", default=str(DEFAULT_DESIGN_LAW))
    parser.add_argument("--assessment-design", default=str(DEFAULT_ASSESSMENT_DESIGN))
    args = parser.parse_args()

    design_law_path = Path(args.design_law).expanduser().resolve()
    assessment_design_path = Path(args.assessment_design).expanduser().resolve() if args.assessment_design else None
    jobs = job_dirs_from_path(Path(args.path).expanduser().resolve())
    receipts = [apply_design_law(job, design_law_path, assessment_design_path) for job in jobs]
    print(json.dumps({"status": "completed", "formatted": len(receipts), "receipts": receipts}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
