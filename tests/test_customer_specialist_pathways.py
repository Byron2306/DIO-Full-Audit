from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document

from products.customer_specialist_pathways import build_contract_desk_package, build_readiness_package


def _contract() -> dict:
    text = "1. Supplier must deliver 12 pumps by 2026-12-03.\n2. Customer must pay ZAR 48000 within 30 days after acceptance.\n3. Either party may terminate on 30 days written notice."
    return {"case_id": "CLINIC-104", "title": "Moya pump supply review", "sources": [{"source_id": "CONTRACT", "text": text}], "contract_source_id": "CONTRACT", "counterparties": ["Moya Clinic", "Pump Supplier"], "review_at": "2026-12-04T10:00:00+02:00"}


def test_contract_desk_runs_obligations_and_retains_unresolved_notice_terms(tmp_path: Path) -> None:
    receipt = build_contract_desk_package(_contract(), tmp_path, render_pdf=False)
    bundle = json.loads((tmp_path / "OBLIGATION_BUNDLE.json").read_text())
    assert len(bundle["obligations"]) == 2
    assert any(d["state"] == "overdue" for d in bundle["deadlines"])
    assert any("acceptance" in r["question"].casefold() for r in receipt["review_actions"])
    assert "dio.obligations.engine.build" in receipt["executed_providers"]
    document = Document(next(tmp_path.glob("customer_delivery/*.docx")))
    text = "\n".join(p.text for p in document.paragraphs) + "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "12 pumps" in text and "2026-12-03" in text
    assert "30 days written notice" in text
    assert receipt["authority_created"] is False


def test_changed_contract_alters_actual_calendar_and_report(tmp_path: Path) -> None:
    a = _contract()
    b = deepcopy(a)
    b["sources"][0]["text"] = b["sources"][0]["text"].replace("12 pumps", "20 pumps").replace("2026-12-03", "2026-12-30")
    build_contract_desk_package(a, tmp_path / "a", render_pdf=False)
    build_contract_desk_package(b, tmp_path / "b", render_pdf=False)
    first = json.loads((tmp_path / "a/OBLIGATION_BUNDLE.json").read_text())
    second = json.loads((tmp_path / "b/OBLIGATION_BUNDLE.json").read_text())
    assert first["fingerprint"] != second["fingerprint"]
    assert second["deadlines"][0]["state"] == "open"
    html = next((tmp_path / "b/customer_delivery").glob("*.html")).read_text()
    assert "20 pumps" in html and "12 pumps" not in html


def _readiness() -> dict:
    return {
        "case_id": "READINESS-104", "title": "Moya finance application", "owner_authorized": True,
        "sources": [
            {"source_id": "CHECKLIST", "text": "Company registration document. Latest management accounts."},
            {"source_id": "COMPANY", "text": "Moya Clinic registration CK-104."},
            {"source_id": "ACCOUNTS", "text": "2026 Q3 ZAR: revenue 100000; cost of sales 60000; current assets 150000; current liabilities 100000."},
        ],
        "requirements": [
            {"requirement_id": "R1", "label": "Company registration document", "source_id": "CHECKLIST", "quote": "Company registration document.", "mandatory": True},
            {"requirement_id": "R2", "label": "Latest management accounts", "source_id": "CHECKLIST", "quote": "Latest management accounts.", "mandatory": True},
        ],
        "evidence": [{"evidence_id": "E1", "source_id": "COMPANY", "quote": "Moya Clinic registration CK-104.", "supports": ["R1"]}],
        "financials": {"period": "2026 Q3", "currency": "ZAR", "source_id": "ACCOUNTS", "quote": "2026 Q3 ZAR: revenue 100000; cost of sales 60000; current assets 150000; current liabilities 100000.", "revenue": "100000", "cost_of_sales": "60000", "current_assets": "150000", "current_liabilities": "100000"},
    }


def test_readiness_joins_requirements_to_supplied_bytes_and_exports_calculations(tmp_path: Path) -> None:
    receipt = build_readiness_package("Finance Readiness", _readiness(), tmp_path, render_pdf=False)
    assert receipt["mandatory_missing_count"] == 1
    assert receipt["requirements"][0]["state"] == "RECEIVED_FOR_REVIEW"
    assert receipt["requirements"][1]["state"] == "MISSING"
    assert receipt["financial_metrics"]["gross_margin_percent"] == "40.00"
    assert receipt["financial_metrics"]["current_ratio"] == "1.50"
    assert (tmp_path / "FINANCIAL_REVIEW_WORKBOOK.xlsx").is_file()
    assert "adapters.evidex_studio.build_readiness_gap_map" in receipt["executed_providers"]
    assert receipt["professional_approval"] == "NOT_GRANTED"


def test_readiness_refuses_unmatched_quotes_and_unsupported_numbers(tmp_path: Path) -> None:
    packet = _readiness()
    packet["evidence"][0]["quote"] = "Invented signed approval"
    with pytest.raises(ValueError, match="quote"):
        build_readiness_package("Finance Readiness", packet, tmp_path / "quote", render_pdf=False)
    packet = _readiness()
    packet["financials"]["revenue"] = "999999"
    with pytest.raises(ValueError, match="source|number"):
        build_readiness_package("Finance Readiness", packet, tmp_path / "number", render_pdf=False)


def test_readiness_never_accepts_self_reported_approval_as_independent_verification(tmp_path: Path) -> None:
    packet = _readiness()
    packet["evidence"][0].update({"state": "approved", "trust_state": "trusted_for_review"})
    receipt = build_readiness_package("POPIA Readiness", packet, tmp_path, render_pdf=False)
    assert receipt["requirements"][0]["state"] == "RECEIVED_FOR_REVIEW"
    assert receipt["professional_approval"] == "NOT_GRANTED"
    assert receipt["authority_created"] is False


def test_actual_customer_evidence_profile_uses_separate_requirements_and_native_review(tmp_path: Path) -> None:
    from products.customer_specialist_pathways import build_evidence_review_package
    result = build_evidence_review_package("AuditProof", _readiness(), tmp_path, render_pdf=False)
    assert result["profile_id"] == "auditproof"
    assert result["domain_engine_invoked"] is False
    assert "products.evidence_review.run_controlled_evidence_review" in result["executed_providers"]
    native = json.loads(next((tmp_path / "native_review").glob("*_REVIEW_PACK.json")).read_text())
    assert len(native["requirement_evidence_matrix"]) == 2
    assert not any(r["review_state"] == "SUPPORTED" for r in native["requirement_evidence_matrix"])
    assert result["authority_created"] is False


def test_correspondence_is_a_specific_source_bound_draft_and_does_not_concede(tmp_path: Path) -> None:
    from products.customer_specialist_pathways import build_correspondence_package
    packet = {"case_id":"LETTER-42", "title":"Invoice INV-42 records request", "recipient":"Moya Accounts", "sender":"Nala", "requested_action":"Please send the signed scope and timesheet for reconciliation.",
              "sources":[{"source_id":"INVOICE", "text":"Invoice INV-42 records ZAR 48000. The signed scope has not been supplied."}],
              "facts":[{"source_id":"INVOICE", "quote":"Invoice INV-42 records ZAR 48000."},{"source_id":"INVOICE", "quote":"The signed scope has not been supplied."}]}
    result = build_correspondence_package(packet,tmp_path,render_pdf=False)
    text=(tmp_path / "DRAFT_CORRESPONDENCE.txt").read_text()
    assert "Dear Moya Accounts" in text and "INV-42" in text and "48000" in text and "Nala" in text
    assert packet["requested_action"] in text
    assert "Bluebird" not in text
    assert result["send_authorized"] is False


def test_new_customer_job_cannot_mix_with_a_previous_customer_output(tmp_path: Path) -> None:
    build_readiness_package("Finance Readiness", _readiness(), tmp_path, render_pdf=False)
    changed = _readiness()
    changed["case_id"] = "SECOND-CUSTOMER"
    with pytest.raises(ValueError, match="empty|existing|customer"):
        build_readiness_package("Corporate Readiness", changed, tmp_path, render_pdf=False)


def test_customer_financial_source_strings_are_literal_workbook_evidence(tmp_path: Path) -> None:
    from openpyxl import load_workbook
    packet = _readiness()
    packet['sources'][2]['text'] = '=' + packet['sources'][2]['text']
    packet['financials']['quote'] = packet['sources'][2]['text']
    build_readiness_package('Finance Readiness', packet, tmp_path, render_pdf=False)
    workbook = load_workbook(tmp_path / 'FINANCIAL_REVIEW_WORKBOOK.xlsx')
    cell = workbook['Source amounts']['F2']
    assert cell.value.startswith('=2026 Q3')
    assert cell.data_type == 's'


def test_contract_multiple_dates_do_not_create_false_overdue_delivery(tmp_path: Path):
    packet=_contract()
    packet['sources'][0]['text']='1. The agreement commences on 2026-01-01. Supplier must deliver 12 pumps by 2026-12-03.'
    packet['review_at']='2026-06-01T10:00:00+02:00'
    result=build_contract_desk_package(packet,tmp_path,render_pdf=False)
    bundle=json.loads((tmp_path/'OBLIGATION_BUNDLE.json').read_text())
    assert not bundle['deadlines']
    assert result['review_actions']


def test_contract_preamble_duty_is_not_lost_before_numbered_clauses(tmp_path: Path):
    packet=_contract()
    packet['sources'][0]['text']='The parties must keep all information confidential.\n1. Supplier must deliver 12 pumps by 2026-12-03.'
    result=build_contract_desk_package(packet,tmp_path,render_pdf=False)
    assert result['obligation_count']==2
    assert 'PREAMBLE' in (tmp_path/'OBLIGATION_BUNDLE.json').read_text()
    assert 'keep all information confidential' in next(tmp_path.glob('customer_delivery/*.html')).read_text()


def test_contract_commencement_date_is_not_a_relative_delivery_deadline(tmp_path: Path):
    packet=_contract()
    packet['sources'][0]['text']='1. The agreement commences on 2026-01-01. Supplier must deliver within 30 days after order acceptance.'
    packet['review_at']='2026-06-01T10:00:00+02:00'
    build_contract_desk_package(packet,tmp_path,render_pdf=False)
    assert not json.loads((tmp_path/'OBLIGATION_BUNDLE.json').read_text())['deadlines']
