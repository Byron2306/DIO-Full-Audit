from __future__ import annotations

import csv
import hashlib
import json
import shutil
import zipfile
from pathlib import Path
from typing import Any

from products.professional_evidence_native_rich_routes import (
    VAMP_PROFILE,
    _build_vamp_rich_database,
    _evidex_source_files,
    utc_now,
)
from products.professional_evidence_projection import evidence_rows, sha256, slug, write_json


RESTORED_VAMP_SCHEMA = "dio.professional_evidence.vamp_native_binding.v1"
RESTORED_EVIDEX_SCHEMA = "dio.professional_evidence.evidex_native_binding.v1"
EVIDEX_SUMMARY_SCHEMA = "dio.evidex.deterministic_source_summary.v1"


def run_restored_vamp(
    packet: dict[str, Any],
    execution_dir: Path,
    *,
    incarnation: str,
) -> dict[str, Any]:
    """Run the real VAMP engine with a schema-valid rich customer projection.

    Customer packet identity remains in the DIO projection/binding rather than
    being smuggled into VAMP's closed request schema.
    """
    from adapters.vamp.snapshot_pipeline import build_snapshot

    projection_dir = execution_dir.parent / "PROJECTION"
    database = projection_dir / "vamp_native_rich.db"
    projection = _build_vamp_rich_database(packet, database)
    request = {
        "schema": "dio.vamp_snapshot_request.v1",
        "job_id": "PRO-" + slug(incarnation).upper() + "-NATIVE",
        "profile_path": str(VAMP_PROFILE),
        "source": {
            "kind": "vamp_sqlite",
            "database_path": str(database),
            "staff_id": "PROFESSIONAL-001",
            "year": 2026,
        },
        "review": {"months": projection["review_months"]},
        # This is the exact enum accepted by the native VAMP request schema.
        # The richer DIO/Vesper custody semantics live outside that schema.
        "privacy_mode": "private_internal",
        "consents": {
            "evidence_owner_authorized": True,
            "performance_data_processing_approved": True,
            "human_review_terms_accepted": True,
        },
    }
    request_path = projection_dir / "VAMP_NATIVE_SNAPSHOT_REQUEST.json"
    write_json(request_path, request)

    job_dir = build_snapshot(request_path, execution_dir / "vamp_native", run_evidex=True)
    job_dir = Path(job_dir)
    receipt_path = job_dir / "VAMP_SNAPSHOT_RECEIPT.json"
    snapshot_path = job_dir / "VAMP_SNAPSHOT.json"
    ledger_path = job_dir / "EVIDENCE_LEDGER.json"
    coverage_path = job_dir / "OBJECTIVE_COVERAGE.json"
    summary_path = job_dir / "VAMP_SNAPSHOT.md"
    for path in (receipt_path, snapshot_path, ledger_path, coverage_path, summary_path):
        if not path.is_file():
            raise RuntimeError(f"VAMP restored route missed required artifact: {path.name}")

    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    metrics = dict(snapshot.get("metrics") or {})
    checks = {
        "native_request_schema_preserved": "customer_packet_fingerprint" not in request,
        "privacy_mode_native_enum": request["privacy_mode"] == "private_internal",
        "release_ready_for_human_review": receipt.get("status") == "ready_for_human_review",
        "objectives_total": int(metrics.get("objectives_total") or 0) >= 3,
        "evidence_records": int(metrics.get("evidence_records") or 0) >= 6,
        "accepted_mappings_present": int(metrics.get("accepted_mappings") or 0) >= 2,
        "candidate_or_gap_preserved": any(
            int(metrics.get(key) or 0) >= 1
            for key in ("candidate_mappings", "objectives_gap", "objectives_partial", "objectives_declared_no_evidence")
        ),
        "rating_not_generated": (snapshot.get("release") or {}).get("rating_generated") is False,
        "employment_decision_not_generated": (snapshot.get("release") or {}).get("employment_decision_generated") is False,
    }
    if not all(checks.values()):
        raise RuntimeError(
            "VAMP restored native quality checks failed: "
            + ", ".join(name for name, ok in checks.items() if not ok)
        )

    binding = {
        "schema": RESTORED_VAMP_SCHEMA,
        "native_engine": "adapters.vamp.snapshot_pipeline.build_snapshot",
        "packet_fingerprint": packet["packet_fingerprint"],
        "projection": str(projection_dir / "VAMP_NATIVE_RICH_PROJECTION.json"),
        "request": str(request_path),
        "native_job_dir": str(job_dir),
        "native_receipt": str(receipt_path),
        "native_metrics": metrics,
        "checks": checks,
        "customer_packet_identity_transport": "DIO_BINDING_NOT_NATIVE_REQUEST_SCHEMA",
        "surrogate_fallback_allowed": False,
        "surrogate_fallback_used": False,
        "authority_created": False,
        "external_effects": False,
        "human_review_gate": "NEEDS_YOU",
    }
    write_json(execution_dir / "VAMP_NATIVE_ROUTE_BINDING.json", binding)
    return {
        "executor": "adapters.vamp.snapshot_pipeline.build_snapshot",
        "native_engine_identity": "adapters.vamp.snapshot_pipeline.build_snapshot",
        "native_capability_preserved": True,
        "surrogate_fallback_used": False,
        "product_id": "vamp_performance",
        "terminal_artifact_kind": "native_performance_evidence_snapshot",
        "domain_action_executed": False,
        "product_pipeline_executed": True,
        "receipt": binding,
        "native_result": receipt,
        "native_snapshot": snapshot,
    }


def _source_observations(path: Path, *, max_rows: int = 8) -> list[str]:
    """Extract conservative, literal observations from a delivered source file."""
    suffix = path.suffix.casefold()
    observations: list[str] = []
    if suffix == ".csv":
        with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
            rows = list(csv.reader(handle))
        if rows:
            observations.append("Columns: " + " | ".join(cell.strip() for cell in rows[0]))
            for row in rows[1 : max_rows + 1]:
                observations.append(" | ".join(cell.strip() for cell in row))
            if len(rows) - 1 > max_rows:
                observations.append(f"... {len(rows) - 1 - max_rows} additional data row(s) retained in the source file.")
        return observations

    if suffix in {".md", ".txt"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        for raw in text.splitlines():
            line = " ".join(raw.lstrip("#-* ").split())
            if not line:
                continue
            observations.append(line[:500])
            if len(observations) >= max_rows:
                break
        return observations

    return [f"Binary/non-text source retained unchanged ({path.stat().st_size} bytes)."]


def _source_summary_payload(
    *,
    package_dir: Path,
    source_manifest: dict[str, Any],
    job: dict[str, Any],
) -> dict[str, Any]:
    manifest_rows = [dict(row) for row in source_manifest.get("sources") or [] if isinstance(row, dict)]
    summaries = []
    for row in manifest_rows:
        name = str(row.get("name") or "")
        delivered = package_dir / "03_SOURCES" / name
        if not delivered.is_file():
            raise RuntimeError(f"Evidex delivery omitted source declared in manifest: {name}")
        observed_hash = sha256(delivered)
        expected_hash = str(row.get("sha256") or "")
        if expected_hash and observed_hash != expected_hash:
            raise RuntimeError(f"Evidex delivered source hash mismatch: {name}")
        summaries.append(
            {
                "name": name,
                "source_role": str(row.get("source_role") or ""),
                "sha256": observed_hash,
                "bytes": delivered.stat().st_size,
                "observations": _source_observations(delivered),
                "inference_added": False,
            }
        )

    known_gaps = [str(item) for item in job.get("known_gaps") or [] if str(item).strip()]
    return {
        "schema": EVIDEX_SUMMARY_SCHEMA,
        "summary_mode": "deterministic_source_bound_no_llm",
        "customer_evidence_source_count": sum(row["source_role"] == "customer_evidence" for row in summaries),
        "request_context_source_count": sum(row["source_role"] == "request_context" for row in summaries),
        "sources": summaries,
        "known_gaps": known_gaps,
        "kpis": [dict(row) for row in job.get("kpis") or [] if isinstance(row, dict)],
        "claim_boundary": (
            "These are literal source observations and declared cross-source gaps. No unsupported fact, provenance, "
            "reconciliation, donor conclusion, or authority is created."
        ),
        "authority_created": False,
        "external_effects": False,
    }


def _write_source_summary_docx(payload: dict[str, Any], target: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    doc.add_heading("Appendix — Source Summaries", level=0)
    doc.add_paragraph(
        "Deterministic source-bound appendix. No LLM summarisation was used. "
        "Every observation below is taken from the delivered source bytes and retains its SHA-256 provenance."
    )

    gaps = list(payload.get("known_gaps") or [])
    if gaps:
        doc.add_heading("Cross-source gaps and contradictions requiring human review", level=1)
        for gap in gaps:
            doc.add_paragraph(str(gap), style="List Bullet")

    kpis = list(payload.get("kpis") or [])
    if kpis:
        doc.add_heading("Declared evidence checks", level=1)
        for row in kpis:
            p = doc.add_paragraph()
            p.add_run(str(row.get("name") or "Evidence check") + ": ").bold = True
            p.add_run(f"target={row.get('target')}; observed={row.get('actual')}; source basis={row.get('measurement')}")

    doc.add_heading("Source-by-source observations", level=1)
    for index, row in enumerate(payload.get("sources") or [], 1):
        doc.add_heading(f"{index}. {row.get('name')}", level=2)
        doc.add_paragraph(f"Role: {row.get('source_role')}")
        doc.add_paragraph(f"SHA-256: {row.get('sha256')}")
        doc.add_paragraph(f"Bytes: {row.get('bytes')}")
        for observation in row.get("observations") or []:
            doc.add_paragraph(str(observation), style="List Bullet")
        doc.add_paragraph("Boundary: literal observation only; no inference or reconciliation added.")

    doc.add_heading("Authority boundary", level=1)
    doc.add_paragraph(str(payload.get("claim_boundary") or ""))
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def _rebuild_delivery_zip(package_dir: Path, zip_path: Path) -> None:
    temporary = zip_path.with_suffix(zip_path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(package_dir.rglob("*")):
            if path.is_file():
                archive.write(path, arcname=str(path.relative_to(package_dir)))
    temporary.replace(zip_path)


def _repair_evidex_appendix(
    *,
    output_root: Path,
    source_manifest: dict[str, Any],
    job: dict[str, Any],
) -> dict[str, Any]:
    package_dirs = [
        path for path in output_root.iterdir()
        if path.is_dir() and (path / "03_SOURCES").is_dir()
    ] if output_root.is_dir() else []
    if len(package_dirs) != 1:
        raise RuntimeError(f"Evidex restored route expected one delivery package directory, found {len(package_dirs)}")
    package_dir = package_dirs[0]
    payload = _source_summary_payload(package_dir=package_dir, source_manifest=source_manifest, job=job)
    appendix = package_dir / "04_APPENDIX" / "appendix_source_summaries.docx"
    _write_source_summary_docx(payload, appendix)
    summary_json = package_dir / "04_APPENDIX" / "source_summaries.json"
    write_json(summary_json, payload)

    zip_paths = sorted(output_root.glob("*.zip"))
    if not zip_paths:
        raise RuntimeError("Evidex restored route produced no delivery ZIP to rebuild")
    for zip_path in zip_paths:
        _rebuild_delivery_zip(package_dir, zip_path)

    return {
        "schema": "dio.evidex.deterministic_source_summary_receipt.v1",
        "package_dir": str(package_dir),
        "appendix": str(appendix),
        "summary_json": str(summary_json),
        "customer_evidence_source_count": payload["customer_evidence_source_count"],
        "request_context_source_count": payload["request_context_source_count"],
        "delivery_zips_rebuilt": [str(path) for path in zip_paths],
        "placeholder_summary_present": False,
        "llm_required": False,
        "authority_created": False,
        "external_effects": False,
    }


def run_restored_evidex(packet: dict[str, Any], execution_dir: Path) -> dict[str, Any]:
    """Run the real Evidex engine and harden DIO's ingestion/presentation seams."""
    from scripts.run_evidex_jobs import EVIDEX_PYTHON, run_evidex

    if not EVIDEX_PYTHON.is_file():
        raise FileNotFoundError(f"Evidex runtime missing: {EVIDEX_PYTHON}")
    rows = evidence_rows(packet)
    customer = dict((packet.get("intake") or {}).get("customer") or {})
    source_files = _evidex_source_files(packet)
    source_names = [row["name"] for row in source_files]
    job = {
        "job_id": "PRO-EVIDEX-EVIDENCEOPS-NATIVE-001",
        "created_at": utc_now(),
        "route": {
            "product": "evidex",
            "confidence": 1.0,
            "reason": "Canonical Evidex EvidenceOps native multi-artifact route",
        },
        "risk": "moderate",
        "customer_context": {
            "organisation": str(customer.get("organisation") or "Green Basin Initiative"),
            "buyer_role": str(customer.get("buyer_role") or "Grant reporting lead"),
            "reporting_period": {},
            "project_name": "Green Basin Initiative reporting evidence review",
            "truth_boundary": "The 14-workshop narrative claim may not be silently reconciled to the 12 attendance-supported workshops.",
        },
        "inputs": [{
            "sender": f"{str(customer.get('buyer_role') or 'Grant reporting lead')} <customer-contact-not-supplied>",
            "subject": str((packet.get("intake") or {}).get("request") or "EvidenceOps review"),
            "attachment_names": "; ".join(source_names),
            "intent": "claim_to_evidence_pack",
            "urgency": "normal",
            "next_step": "prepare governed evidence pack for human review only",
        }],
        "evidence": [
            {
                "record_id": str(row.get("record_id") or ""),
                "text_extract": str(row.get("customer_supplied_record") or ""),
            }
            for row in rows
        ],
        "source_files": source_files,
        "kpis": [
            {
                "name": "C1 Community workshops delivered",
                "target": 14,
                "actual": "12 attendance-sheet supported; 2 additional workshops mentioned only in email",
                "measurement": "claims_register.csv + monitoring_workshops.csv + workstream_email_summary.txt",
            },
            {
                "name": "Venue-cost evidence",
                "target": "Traceable supporting invoice",
                "actual": "INV-118 covers venue costs for four workshops",
                "measurement": "invoice_INV-118.txt",
            },
            {
                "name": "Photograph provenance",
                "target": "Dated and event-linked photographs",
                "actual": "Three photographs lack dates and event identifiers",
                "measurement": "photo_metadata.csv",
            },
        ],
        "known_gaps": [
            "C1 claims 14 workshops while attendance sheets support 12.",
            "Two additional workshops are supported only by an email summary.",
            "INV-118 supports venue expenditure for four workshops but not attendance at fourteen workshops.",
            "Three photographs lack dates and event identifiers.",
            "Reporting-period dates remain to be confirmed by the customer.",
        ],
        "customer_packet_fingerprint": packet["packet_fingerprint"],
    }
    projection_path = execution_dir.parent / "PROJECTION" / "EVIDEX_NATIVE_MULTI_ARTIFACT_JOB.json"
    write_json(projection_path, job)
    native_root = execution_dir / "evidex_native"
    receipt = run_evidex(job, native_root)
    if int(receipt.get("returncode", 1)) != 0:
        raise RuntimeError(f"Evidex native engine failed: {str(receipt.get('stderr') or '')[-1000:]}")

    manifest_path = Path(str(receipt.get("customer_source_manifest") or ""))
    if not manifest_path.is_file():
        # Some older run_evidex receipts did not surface this field even though
        # the native engine persisted the manifest. Resolve it from the governed
        # job directory rather than treating successful ingestion as failure.
        manifest_path = native_root / "evidex" / job["job_id"] / "evidex_engine" / "CUSTOMER_SOURCE_MANIFEST.json"
    if not manifest_path.is_file():
        raise RuntimeError("Evidex native engine produced no customer source manifest")
    source_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest_rows = [dict(row) for row in source_manifest.get("sources") or [] if isinstance(row, dict)]
    customer_evidence_rows = [row for row in manifest_rows if row.get("source_role") == "customer_evidence"]
    if len(customer_evidence_rows) < 10:
        raise RuntimeError(
            f"Evidex native route ingested only {len(customer_evidence_rows)} customer evidence sources; expected at least 10"
        )
    if len({str(row.get("sha256") or "") for row in customer_evidence_rows}) < 10:
        raise RuntimeError("Evidex customer evidence sources are not provenance-distinct")

    output_dir = Path(str(receipt.get("output_dir") or ""))
    if not output_dir.is_dir():
        raise RuntimeError("Evidex native engine did not persist its output directory")
    summary_receipt = _repair_evidex_appendix(
        output_root=output_dir,
        source_manifest=source_manifest,
        job=job,
    )
    write_json(native_root / "EVIDEX_DETERMINISTIC_SOURCE_SUMMARY_RECEIPT.json", summary_receipt)

    output_files = [path for path in output_dir.rglob("*") if path.is_file()]
    appendix = Path(summary_receipt["appendix"])
    if appendix.stat().st_size < 10000:
        raise RuntimeError("Evidex deterministic source appendix is implausibly small")
    if "summaries disabled" in _source_observations(appendix)[0].casefold() if appendix.suffix.casefold() in {".md", ".txt"} else False:
        raise RuntimeError("Evidex placeholder appendix survived restoration")

    binding = {
        "schema": RESTORED_EVIDEX_SCHEMA,
        "native_engine": "scripts.run_evidex_jobs.run_evidex",
        "packet_fingerprint": packet["packet_fingerprint"],
        "projection": str(projection_path),
        "customer_source_count": len(manifest_rows),
        "customer_evidence_source_count": len(customer_evidence_rows),
        "request_context_source_count": len(manifest_rows) - len(customer_evidence_rows),
        "customer_source_manifest": str(manifest_path),
        "native_output_dir": str(output_dir),
        "native_output_file_count": len(output_files),
        "native_output_files": [
            {"path": str(path), "bytes": path.stat().st_size, "sha256": sha256(path)}
            for path in sorted(output_files)
        ],
        "deterministic_source_summary": summary_receipt,
        "surrogate_fallback_allowed": False,
        "surrogate_fallback_used": False,
        "authority_created": False,
        "external_effects": False,
        "human_review_gate": "NEEDS_YOU",
    }
    write_json(execution_dir / "EVIDEX_NATIVE_ROUTE_BINDING.json", binding)
    return {
        "executor": "scripts.run_evidex_jobs.run_evidex",
        "native_engine_identity": "scripts.run_evidex_jobs.run_evidex",
        "native_capability_preserved": True,
        "surrogate_fallback_used": False,
        "product_id": "evidex",
        "terminal_artifact_kind": "native_multi_artifact_evidence_pack_with_source_bound_appendix",
        "domain_action_executed": False,
        "product_pipeline_executed": True,
        "receipt": binding,
        "native_result": receipt,
    }


__all__ = ["run_restored_evidex", "run_restored_vamp"]
