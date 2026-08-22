from __future__ import annotations

import csv
import difflib
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageOps

from adapters.format_core import build_paragraph_semantic_content, render_semantic_asset
from adapters.sophia.review_pipeline import extract_document_text
from adapters.lingua.lifecycle import build_lingua_qa, digest_text, update_semantic_object


ROOT = Path(__file__).resolve().parents[2]
BRIDGE = ROOT / "scripts" / "document_studio_gemini_bridge.py"
NIM_BRIDGE = ROOT / "scripts" / "document_studio_nim_bridge.py"
BEAST_BRIDGE = ROOT / "scripts" / "lingua_beast_bridge.py"
DEFAULT_SOPHIA_PYTHON = Path("/home/byron/Integritas-Mechanicus/.venv/bin/python")
DEFAULT_SOPHIA_ROOT = Path("/home/byron/Integritas-Mechanicus/arda_os")
SERVICES = {"technical_edit", "translation", "edit_and_translate"}
LANGUAGE_REGISTRY_PATH = ROOT / "config" / "document_studio_languages.json"


def load_language_registry() -> dict[str, Any]:
    return json.loads(LANGUAGE_REGISTRY_PATH.read_text(encoding="utf-8"))


def canonical_language(value: Any, registry: dict[str, Any] | None = None) -> str:
    registry = registry or load_language_registry()
    raw = str(value or "").strip()
    canonical = registry.get("aliases", {}).get(raw.casefold())
    if not canonical or canonical not in registry.get("languages", {}):
        supported = ", ".join(registry.get("languages", {}))
        raise ValueError(f"Unsupported Document Studio language: {raw or '(missing)'}. Supported: {supported}")
    return str(canonical)


def normalize_language_request(request: dict[str, Any]) -> dict[str, Any]:
    normalized = dict(request)
    registry = load_language_registry()
    service = str(normalized.get("service") or "")
    source = canonical_language(normalized.get("source_language"), registry)
    normalized["source_language"] = source
    normalized["source_language_profile"] = registry["languages"][source]
    if service in {"translation", "edit_and_translate"}:
        target = canonical_language(normalized.get("target_language"), registry)
        if target == source:
            raise ValueError("Source and target language must differ for translation work.")
        if not normalized.get("human_language_review_required"):
            raise ValueError("A proficient target-language reviewer is required for translation work.")
        normalized["target_language"] = target
        normalized["target_language_profile"] = registry["languages"][target]
    else:
        normalized["target_language"] = None
        normalized["target_language_profile"] = None
    return normalized


def resolve_beast_context(request: dict[str, Any], paragraphs: list[dict[str, str]]) -> dict[str, Any]:
    if request["service"] == "technical_edit":
        return {
            "schema": "dio.lingua.beast_resolution.v1",
            "status": "not_applicable",
            "approved_units": [],
            "approved_terms": [],
            "reuse_state": "not_applicable",
            "provider_call_displaced": False,
        }
    payload = {
        "operation": "resolve",
        "target_language": request["target_language"],
        "domain": request.get("document_domain"),
        "units": [
            {
                "unit_id": row["paragraph_id"],
                "source_hash": digest_text(row["text"]),
                "source_text": row["text"],
            }
            for row in paragraphs
        ],
    }
    completed = subprocess.run(
        [sys.executable, str(BEAST_BRIDGE)],
        input=json.dumps(payload),
        text=True,
        capture_output=True,
        check=False,
        timeout=60,
        cwd=ROOT,
    )
    if completed.returncode != 0:
        raise RuntimeError(f"BEAST Lingua resolution failed: {completed.stdout[-1000:] or completed.stderr[-1000:]}")
    receipt = json.loads(completed.stdout)
    if receipt.get("status") != "resolved":
        raise RuntimeError(f"BEAST Lingua resolution unavailable: {receipt.get('error') or receipt.get('status')}")
    return receipt


def apply_beast_authority(result: dict[str, Any], beast_receipt: dict[str, Any]) -> None:
    translations = _row_map(list(result.get("translations") or []), "translated")
    applied = []
    for approved in beast_receipt.get("approved_units") or []:
        paragraph_id = str(approved.get("unit_id") or "")
        if paragraph_id not in translations or not str(approved.get("target_text") or "").strip():
            continue
        row = translations[paragraph_id]
        provider_draft = str(row.get("translated") or "")
        row["pre_beast_provider_draft"] = provider_draft
        row["translated"] = str(approved["target_text"]).strip()
        row["beast_authority_applied"] = True
        row["beast_credit_id"] = approved.get("credit_id")
        applied.append({"paragraph_id": paragraph_id, "credit_id": approved.get("credit_id"), "provider_draft_displaced": provider_draft != row["translated"]})
    result["beast_authority"] = {
        "reuse_state": beast_receipt.get("reuse_state"),
        "approved_terms": beast_receipt.get("approved_terms") or [],
        "approved_units_applied": applied,
        "provider_call_displaced": beast_receipt.get("provider_call_displaced", False),
    }


def learn_beast_patterns(request: dict[str, Any], result: dict[str, Any], validation: dict[str, Any]) -> dict[str, Any]:
    if request["service"] == "technical_edit":
        return {"status": "not_applicable", "learned": []}
    payload = {
        "operation": "learn_flags",
        "target_language": request["target_language"],
        "domain": request.get("document_domain"),
        "flags": result.get("qa_flags") or [],
        "validation": validation,
    }
    completed = subprocess.run(
        [sys.executable, str(BEAST_BRIDGE)],
        input=json.dumps(payload),
        text=True,
        capture_output=True,
        check=False,
        timeout=120,
        cwd=ROOT,
    )
    if completed.returncode != 0:
        raise RuntimeError(f"BEAST Lingua learning failed: {completed.stdout[-1000:] or completed.stderr[-1000:]}")
    receipt = json.loads(completed.stdout)
    if receipt.get("status") != "learned":
        raise RuntimeError(f"BEAST Lingua learning unavailable: {receipt.get('error') or receipt.get('status')}")
    return receipt


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def safe_job_id(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{2,100}", value):
        raise ValueError("job_id must contain only letters, numbers, dots, underscores, or hyphens")
    return value


def paragraphs_from_text(text: str) -> list[dict[str, str]]:
    values = [re.sub(r"\s+", " ", item).strip() for item in re.split(r"\n\s*\n", text) if item.strip()]
    if not values:
        raise ValueError("No readable paragraphs were found in the source document.")
    return [{"paragraph_id": f"P{index}", "text": value} for index, value in enumerate(values, 1)]


def parse_json_response(value: str) -> dict[str, Any]:
    cleaned = value.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.I)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    start, end = cleaned.find("{"), cleaned.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("Gemini did not return a JSON object.")
    return json.loads(cleaned[start : end + 1])


def build_provider_prompt(request: dict[str, Any], paragraphs: list[dict[str, str]]) -> tuple[str, str]:
    service = request["service"]
    need_edit = service in {"technical_edit", "edit_and_translate"}
    need_translation = service in {"translation", "edit_and_translate"}
    system = "\n".join([
        "You are DIO Document Studio's senior technical editor and translator.",
        "This is owner-authorised professional document work, not student assessment or ghostwriting.",
        "Preserve every operational fact, number, threshold, identifier, warning, and allocation of responsibility.",
        "Improve grammar, clarity, consistency, and technical usability without adding facts.",
        "Translate meaning rather than word order, while preserving protected tokens exactly.",
        "Before returning, audit each translation against the source for actor, action, object, negation, modality, measurement, and technical meaning.",
        "Never substitute an unrelated familiar word for an uncertain technical role or term; retain the source term and flag it when necessary.",
        "Return only valid JSON matching the requested schema. Do not use markdown fences.",
    ])
    schema = {
        "document_summary": "one sentence",
        "edits": [{
            "paragraph_id": "P1",
            "revised": "complete revised source-language paragraph",
            "category": "grammar|clarity|consistency|technical_tone|no_change",
            "rationale": "specific short explanation",
            "confidence": "high|medium|low",
        }] if need_edit else [],
        "translations": [{
            "paragraph_id": "P1",
            "translated": "complete target-language paragraph",
            "confidence": "high|medium|low",
            "review_note": "empty string unless a human language decision is needed",
        }] if need_translation else [],
        "glossary": [{
            "source_term": "term",
            "target_term": "approved equivalent",
            "note": "usage or preservation note",
        }] if need_translation else [],
        "qa_flags": [{"paragraph_id": "P1", "severity": "low|medium|high", "issue": "human review point"}],
    }
    prompt = "\n".join([
        f"Service: {service}",
        f"Source language: {request['source_language']}",
        f"Target language: {request.get('target_language') or 'not applicable'}",
        f"Source-language guidance: {request['source_language_profile']['prompt_guidance']}",
        f"Target-language guidance: {(request.get('target_language_profile') or {}).get('prompt_guidance', 'not applicable')}",
        f"Target-language support status: {(request.get('target_language_profile') or {}).get('status', 'not applicable')}",
        f"Audience: {request['audience']}",
        f"Domain: {request['document_domain']}",
        f"Style standard: {request['style_standard']}",
        f"Protected tokens: {json.dumps(request.get('protected_tokens') or [], ensure_ascii=False)}",
        f"Preferred terminology: {json.dumps(request.get('preferred_terms') or [], ensure_ascii=False)}",
        f"BEAST approved exact translation units: {json.dumps((request.get('beast_context') or {}).get('approved_units') or [], ensure_ascii=False)}",
        f"BEAST learned language-risk patterns: {json.dumps((request.get('beast_context') or {}).get('risk_patterns') or [], ensure_ascii=False)}",
        f"BEAST deterministic guards: {json.dumps((request.get('beast_context') or {}).get('deterministic_guards') or [], ensure_ascii=False)}",
        "Use each supplied paragraph ID exactly once in every required output array.",
        "For edit_and_translate, translate the revised paragraph, not the faulty original.",
        "Use low confidence and a high-severity qa_flag when any role name, safety instruction, scientific term, or administrative concept lacks a reliable equivalent.",
        "Do not silently resolve technical ambiguity. Put uncertain language choices in qa_flags.",
        "BEAST-approved terminology and exact translation units are human authority. Reuse them exactly and do not invent replacements.",
        "Required JSON shape:",
        json.dumps(schema, ensure_ascii=False),
        "SOURCE PARAGRAPHS:",
        json.dumps(paragraphs, ensure_ascii=False),
    ])
    return system, prompt


def invoke_provider(
    request: dict[str, Any],
    system: str,
    prompt: str,
    *,
    max_predict: int = 7000,
) -> tuple[dict[str, Any], dict[str, Any]]:
    provider_name = str(request.get("provider") or "gemini").strip().casefold().replace("-", "_")
    env = dict(os.environ)
    if provider_name in {"nim", "nvidia", "nvidia_nim"}:
        python = Path(sys.executable)
        bridge = NIM_BRIDGE
        payload = {
            "system_prompt": system,
            "prompt": prompt,
            "model": request.get("nim_model") or "deepseek-ai/deepseek-v4-flash-0731",
            "max_predict": max_predict,
            "temperature": 0.05,
            "secret_file": request.get("provider_secret_file") or "/home/byron/EdgeK-BEAST/.beast/provider_secrets.env",
        }
    elif provider_name in {"gemini", "google", "google_gemini"}:
        python = Path(os.environ.get("SOPHIA_PYTHON") or DEFAULT_SOPHIA_PYTHON)
        sophia_root = Path(os.environ.get("SOPHIA_ROOT") or DEFAULT_SOPHIA_ROOT)
        if not python.is_file() or not sophia_root.is_dir():
            raise FileNotFoundError("Sophia's configured Python runtime is unavailable.")
        bridge = BRIDGE
        payload = {
            "system_prompt": system,
            "prompt": prompt,
            "model": request.get("gemini_model") or "gemini-flash-lite-latest",
            "max_predict": max_predict,
            "temperature": 0.1,
        }
        env["PYTHONPATH"] = str(sophia_root)
    else:
        raise ValueError(f"Unsupported Document Studio provider: {provider_name}")
    completed = subprocess.run(
        [str(python), str(bridge)],
        input=json.dumps(payload),
        text=True,
        capture_output=True,
        check=False,
        timeout=240,
        cwd=ROOT,
        env=env,
    )
    if completed.returncode != 0:
        raise RuntimeError(f"Document Studio provider bridge failed: {completed.stderr[-1000:]}")
    provider = json.loads(completed.stdout)
    if provider.get("status") != "ok" or not provider.get("response"):
        raise RuntimeError(f"Document Studio provider unavailable: {provider.get('error') or provider.get('status')}")
    return parse_json_response(str(provider["response"])), provider


def call_provider(request: dict[str, Any], paragraphs: list[dict[str, str]]) -> tuple[dict[str, Any], dict[str, Any]]:
    system, prompt = build_provider_prompt(request, paragraphs)
    return invoke_provider(request, system, prompt)


def apply_translation_critic(
    request: dict[str, Any],
    paragraphs: list[dict[str, str]],
    result: dict[str, Any],
) -> dict[str, Any]:
    edits = _row_map(list(result.get("edits") or []), "revised") if request["service"] != "translation" else {}
    translations = _row_map(list(result.get("translations") or []), "translated")
    controlled_source = [
        {
            "paragraph_id": row["paragraph_id"],
            "source": str(edits[row["paragraph_id"]]["revised"]) if edits else row["text"],
            "draft_translation": str(translations[row["paragraph_id"]]["translated"]),
        }
        for row in paragraphs
    ]
    system = "\n".join([
        "You are DIO Document Studio's second-pass bilingual linguistic and semantic reviewer.",
        "Audit the draft independently against the source; do not defend the first translator.",
        "Correct wrong actors, actions, objects, negation, modality, measurements, spelling, concords, and unnatural literal phrasing.",
        "Do not invent technical equivalents. Retain the source term and create a review flag when no reliable equivalent is known.",
        "Preserve every protected token and operational number exactly.",
        "Return only valid JSON. This is still a machine review candidate and never replaces a proficient human reviewer.",
    ])
    schema = {
        "reviews": [{
            "paragraph_id": "P1",
            "reviewed_translation": "complete corrected target-language paragraph",
            "changed": True,
            "rationale": "specific semantic or linguistic reason, or no change required",
            "confidence": "high|medium|low",
        }],
        "qa_flags": [{"paragraph_id": "P1", "severity": "low|medium|high", "issue": "human review point"}],
    }
    prompt = "\n".join([
        f"Source language: {request['source_language']}",
        f"Target language: {request['target_language']}",
        f"Target-language guidance: {request['target_language_profile']['prompt_guidance']}",
        f"Domain: {request['document_domain']}",
        f"Audience: {request['audience']}",
        f"Protected tokens: {json.dumps(request.get('protected_tokens') or [], ensure_ascii=False)}",
        "Use every paragraph ID exactly once. Return the complete reviewed translation for every paragraph.",
        "Required JSON shape:",
        json.dumps(schema, ensure_ascii=False),
        "SOURCE AND FIRST DRAFT:",
        json.dumps(controlled_source, ensure_ascii=False),
    ])
    review_result, provider = invoke_provider(request, system, prompt, max_predict=6000)
    reviews = _row_map(list(review_result.get("reviews") or []), "reviewed_translation")
    expected_ids = {row["paragraph_id"] for row in paragraphs}
    if set(reviews) != expected_ids:
        raise ValueError("translation critic paragraph IDs do not match the source")
    changes: list[dict[str, Any]] = []
    for paragraph_id in sorted(expected_ids, key=lambda value: int(value[1:])):
        translation = translations[paragraph_id]
        review = reviews[paragraph_id]
        initial = str(translation["translated"])
        revised = str(review["reviewed_translation"]).strip()
        translation["initial_model_draft"] = initial
        translation["translated"] = revised
        translation["critic_changed"] = revised != initial
        translation["critic_rationale"] = str(review.get("rationale") or "")
        translation["critic_confidence"] = str(review.get("confidence") or "")
        if revised != initial:
            changes.append({"paragraph_id": paragraph_id, "initial": initial, "reviewed": revised, "rationale": translation["critic_rationale"]})
    result["qa_flags"] = list(result.get("qa_flags") or []) + list(review_result.get("qa_flags") or [])
    result["translation_critic"] = {
        "provider": provider.get("provider"),
        "model": provider.get("model"),
        "reviewed_count": len(reviews),
        "changed_count": len(changes),
        "changes": changes,
        "human_language_approval_still_required": True,
    }
    return provider


def _row_map(rows: list[dict[str, Any]], value_key: str) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        paragraph_id = str(row.get("paragraph_id") or "")
        if paragraph_id in result or not str(row.get(value_key) or "").strip():
            raise ValueError(f"Invalid or duplicate provider row: {paragraph_id or '(missing id)'}")
        result[paragraph_id] = row
    return result


def validate_result(request: dict[str, Any], paragraphs: list[dict[str, str]], result: dict[str, Any]) -> dict[str, Any]:
    ids = [row["paragraph_id"] for row in paragraphs]
    originals = {row["paragraph_id"]: row["text"] for row in paragraphs}
    service = request["service"]
    edits = _row_map(list(result.get("edits") or []), "revised") if service != "translation" else {}
    translations = _row_map(list(result.get("translations") or []), "translated") if service != "technical_edit" else {}
    errors: list[str] = []
    if service != "translation" and sorted(edits) != sorted(ids):
        errors.append("edit paragraph IDs do not match the source")
    if service != "technical_edit" and sorted(translations) != sorted(ids):
        errors.append("translation paragraph IDs do not match the source")
    protected = [str(item) for item in request.get("protected_tokens") or []]
    for paragraph_id in ids:
        source = originals[paragraph_id]
        outputs = []
        if edits:
            outputs.append(("edit", str(edits.get(paragraph_id, {}).get("revised") or "")))
        if translations:
            outputs.append(("translation", str(translations.get(paragraph_id, {}).get("translated") or "")))
        source_numbers = set(re.findall(r"\b\d+(?:\.\d+)?\b", source))
        for label, output in outputs:
            missing_numbers = sorted(source_numbers - set(re.findall(r"\b\d+(?:\.\d+)?\b", output)))
            if missing_numbers:
                errors.append(f"{paragraph_id} {label} dropped numbers: {', '.join(missing_numbers)}")
            for token in protected:
                if token in source and token not in output:
                    errors.append(f"{paragraph_id} {label} dropped protected token {token}")
    preferred = [row for row in request.get("preferred_terms") or [] if row.get("target")]
    translated_text = " ".join(str(row.get("translated") or "") for row in translations.values()).casefold()
    for term in preferred:
        target = str(term["target"]).casefold()
        abbreviation = str(term.get("abbreviation") or "").casefold()
        if target not in translated_text and (not abbreviation or abbreviation not in translated_text):
            errors.append(f"preferred target term not used: {term['target']}")
    return {
        "passed": not errors,
        "errors": errors,
        "paragraph_count": len(ids),
        "edit_rows": len(edits),
        "translation_rows": len(translations),
        "protected_tokens": protected,
        "preferred_terms_checked": len(preferred),
    }


def configure_document(document: Document, title: str, subtitle: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.color.rgb = RGBColor(17, 73, 75)
    paragraph = document.add_paragraph(title, style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note = document.add_paragraph(subtitle)
    note.style = styles["Subtitle"]
    note.paragraph_format.space_after = Pt(12)


def shade_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_word_redline(paragraph: Any, original: str, revised: str) -> None:
    original_words, revised_words = original.split(), revised.split()
    matcher = difflib.SequenceMatcher(a=original_words, b=revised_words)
    first = True

    def add(words: list[str], kind: str) -> None:
        nonlocal first
        if not words:
            return
        text = ("" if first else " ") + " ".join(words)
        first = False
        run = paragraph.add_run(text)
        if kind == "delete":
            run.font.strike = True
            run.font.color.rgb = RGBColor(180, 52, 45)
        elif kind == "insert":
            run.font.underline = True
            run.font.color.rgb = RGBColor(16, 91, 112)

    for opcode, a1, a2, b1, b2 in matcher.get_opcodes():
        if opcode == "equal":
            add(original_words[a1:a2], "equal")
        elif opcode == "delete":
            add(original_words[a1:a2], "delete")
        elif opcode == "insert":
            add(revised_words[b1:b2], "insert")
        else:
            add(original_words[a1:a2], "delete")
            add(revised_words[b1:b2], "insert")


def create_clean_docx(path: Path, title: str, subtitle: str, rows: list[dict[str, str]], value_key: str) -> None:
    document = Document()
    configure_document(document, title, subtitle)
    for index, row in enumerate(rows):
        value = row[value_key]
        if index == 0 and re.sub(r"\W+", "", value).casefold() == re.sub(r"\W+", "", title).casefold():
            continue
        if index == 0:
            document.add_heading(value, level=1)
        else:
            paragraph = document.add_paragraph(value)
            paragraph.paragraph_format.space_after = Pt(8)
    document.save(path)


def create_redline_docx(path: Path, title: str, paragraphs: list[dict[str, str]], edits: dict[str, dict[str, Any]]) -> None:
    document = Document()
    configure_document(document, title, "Review copy: deleted wording is red/struck; inserted wording is blue/underlined.")
    for source in paragraphs:
        paragraph_id = source["paragraph_id"]
        row = edits[paragraph_id]
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        run = heading.add_run(f"{paragraph_id}  {row.get('category', 'edit').replace('_', ' ').title()}")
        run.bold = True
        run.font.color.rgb = RGBColor(17, 73, 75)
        redline = document.add_paragraph()
        add_word_redline(redline, source["text"], str(row["revised"]))
        rationale = document.add_paragraph(f"Reason: {row.get('rationale') or 'Editorial correction.'}")
        rationale.style = document.styles["Caption"]
    document.save(path)


def create_bilingual_docx(
    path: Path,
    title: str,
    source_language: str,
    target_language: str,
    source_rows: list[dict[str, str]],
    translations: dict[str, dict[str, Any]],
) -> None:
    document = Document()
    configure_document(document, title, f"Bilingual review copy: {source_language} and {target_language}. Human language approval required.")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text, headers[1].text = source_language, target_language
    for cell in headers:
        shade_cell(cell, "DCEBE8")
        cell.paragraphs[0].runs[0].bold = True
    for source in source_rows:
        cells = table.add_row().cells
        cells[0].text = source["text"]
        cells[1].text = str(translations[source["paragraph_id"]]["translated"])
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.paragraphs[0].paragraph_format.space_after = Pt(4)
    document.save(path)



def convert_to_pdf(paths: list[Path], out_dir: Path) -> list[Path]:
    from adapters.libreoffice_low_memory import convert_many_to_pdf

    return convert_many_to_pdf(
        paths,
        out_dir,
        profile_prefix="dio-docstudio-lo-",
    )

def create_proof_image(
    redline_pdf: Path,
    bilingual_pdf: Path,
    target: Path,
    source_language: str,
    target_language: str,
) -> None:
    previews = []
    with tempfile.TemporaryDirectory(prefix="dio-docstudio-proof-") as temp:
        temp_dir = Path(temp)
        for index, pdf in enumerate((redline_pdf, bilingual_pdf), 1):
            prefix = temp_dir / f"page-{index}"
            subprocess.run(
                ["pdftoppm", "-f", "1", "-singlefile", "-png", "-r", "110", str(pdf), str(prefix)],
                check=True,
                capture_output=True,
                timeout=60,
            )
            previews.append(Image.open(prefix.with_suffix(".png")).convert("RGB"))
        canvas = Image.new("RGB", (1600, 900), "#eef3f1")
        draw = ImageDraw.Draw(canvas)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 32)
        small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 18)
        draw.text((60, 35), "DIO Document Studio", font=font, fill="#123f42")
        draw.text((60, 78), f"Technical edit and {source_language}/{target_language} bilingual review proof", font=small, fill="#52646a")
        for index, (preview, label) in enumerate(zip(previews, ("VISIBLE REDLINE", "BILINGUAL REVIEW"))):
            fitted = ImageOps.contain(preview, (680, 720))
            x = 70 + index * 760
            y = 135
            canvas.paste(fitted, (x, y))
            draw.rectangle((x - 2, y - 2, x + fitted.width + 2, y + fitted.height + 2), outline="#97aaa8", width=2)
            draw.text((x, 850), label, font=small, fill="#123f42")
        canvas.save(target, "PNG", optimize=True)


def run_document_studio(request: dict[str, Any], request_path: Path, out_root: Path) -> Path:
    job_id = safe_job_id(str(request.get("job_id") or ""))
    service = str(request.get("service") or "")
    if service not in SERVICES:
        raise ValueError(f"service must be one of: {', '.join(sorted(SERVICES))}")
    request = normalize_language_request(request)
    if not request.get("owner_authorized"):
        raise ValueError("Document-owner authority is required.")
    if not request.get("remote_processing_approved"):
        raise ValueError("Explicit remote-processing approval is required for the Gemini proof lane.")
    if request.get("certified_translation_required"):
        raise ValueError("Certified, sworn, or legally attested translation is outside this lane.")
    raw_source = Path(str(request.get("document_path") or "")).expanduser()
    source_path = raw_source if raw_source.is_absolute() else request_path.parent / raw_source
    source_path = source_path.resolve()
    if not source_path.is_file():
        raise FileNotFoundError(f"Source document not found: {source_path}")
    text, parser = extract_document_text(source_path)
    paragraphs = paragraphs_from_text(text)
    beast_receipt = resolve_beast_context(request, paragraphs)
    request["beast_context"] = beast_receipt
    existing_terms = {str(row.get("source") or row.get("source_term") or "").casefold() for row in request.get("preferred_terms") or []}
    request["preferred_terms"] = list(request.get("preferred_terms") or []) + [
        row for row in beast_receipt.get("approved_terms") or []
        if str(row.get("source") or "").casefold() not in existing_terms
    ]
    result, provider = call_provider(request, paragraphs)
    critic_provider: dict[str, Any] | None = None
    if service in {"translation", "edit_and_translate"} and request.get("automated_language_critic"):
        critic_provider = apply_translation_critic(request, paragraphs, result)
    if service in {"translation", "edit_and_translate"}:
        apply_beast_authority(result, beast_receipt)
    translation_overrides = request.get("translation_review_overrides") or {}
    applied_overrides: list[dict[str, str]] = []
    if translation_overrides:
        rows_by_id = {str(row.get("paragraph_id") or ""): row for row in result.get("translations") or []}
        valid_ids = {row["paragraph_id"] for row in paragraphs}
        for paragraph_id, reviewed_text in translation_overrides.items():
            if paragraph_id not in valid_ids or paragraph_id not in rows_by_id or not str(reviewed_text).strip():
                raise ValueError(f"Invalid translation review override: {paragraph_id}")
            row = rows_by_id[paragraph_id]
            model_draft = str(row.get("translated") or "")
            row["model_draft"] = model_draft
            row["translated"] = str(reviewed_text).strip()
            row["review_override_applied"] = True
            row["review_note"] = "Bilingual proof-review correction applied; final target-language approval remains pending."
            applied_overrides.append({"paragraph_id": paragraph_id, "model_draft": model_draft, "reviewed_text": row["translated"]})
    validation = validate_result(request, paragraphs, result)
    beast_learning_receipt = learn_beast_patterns(request, result, validation)
    if not validation["passed"]:
        raise ValueError("Provider output failed release validation: " + "; ".join(validation["errors"]))

    job_dir = (out_root / job_id).resolve()
    if job_dir.exists():
        shutil.rmtree(job_dir)
    job_dir.mkdir(parents=True)
    edits = _row_map(list(result.get("edits") or []), "revised") if service != "translation" else {}
    translations = _row_map(list(result.get("translations") or []), "translated") if service != "technical_edit" else {}
    controlled_source = [
        {"paragraph_id": row["paragraph_id"], "text": str(edits[row["paragraph_id"]]["revised"]) if edits else row["text"]}
        for row in paragraphs
    ]
    title = str(request.get("title") or source_path.stem)
    source_version = str(request.get("source_version") or "1.0.0")
    semantic_object_id = str(request.get("semantic_object_id") or f"LINGUA-{job_id}")
    semantic_content = build_paragraph_semantic_content(
        object_id=semantic_object_id,
        version=source_version,
        title=title,
        source_language=request["source_language"],
        source_rows=controlled_source,
        context={
            "product": "document_studio",
            "artifact_type": str(request.get("artifact_type") or "technical_document"),
            "audience": str(request["audience"]),
            "domain": str(request["document_domain"]),
        },
        translations=translations or None,
        target_language=request.get("target_language"),
        translation_status="human_review_required",
    )
    write_json(job_dir / "SEMANTIC_CONTENT.json", semantic_content)
    format_style = str(request.get("format_style_profile") or "dio_professional")
    format_delivery = str(request.get("format_delivery_profile") or "editable_review")
    format_channels = list(request.get("format_channels") or ["docx", "pdf", "html"])
    format_templates = dict(request.get("format_template_paths") or {})
    format_receipts = {
        request["source_language"]: render_semantic_asset(
            semantic_content,
            job_dir / "formatted" / "source",
            style_profile=format_style,
            delivery_profile=format_delivery,
            language=request["source_language"],
            channels=format_channels,
            release_mode=False,
            source_root=source_path.parent,
            template_paths=format_templates,
        )
    }
    if translations:
        format_receipts[request["target_language"]] = render_semantic_asset(
            semantic_content,
            job_dir / "formatted" / str(request["target_language"]),
            style_profile=format_style,
            delivery_profile=format_delivery,
            language=request["target_language"],
            channels=format_channels,
            release_mode=False,
            source_root=source_path.parent,
            template_paths=format_templates,
        )
    docx_files: list[Path] = []
    if edits:
        clean = job_dir / "CLEAN_EDITED_COPY.docx"
        redline = job_dir / "REDLINE_REVIEW_COPY.docx"
        create_clean_docx(clean, controlled_source[0]["text"], "Technically edited clean copy. Human approval required.", controlled_source, "text")
        create_redline_docx(redline, title, paragraphs, edits)
        docx_files.extend([clean, redline])
    if translations:
        translated_rows = [
            {"paragraph_id": row["paragraph_id"], "text": str(translations[row["paragraph_id"]]["translated"])}
            for row in controlled_source
        ]
        translated = job_dir / "TRANSLATED_CLEAN_COPY.docx"
        bilingual = job_dir / "BILINGUAL_REVIEW_COPY.docx"
        create_clean_docx(
            translated, translated_rows[0]["text"], f"{request['target_language']} translation. Human language approval required.",
            translated_rows, "text",
        )
        create_bilingual_docx(
            bilingual, title, request["source_language"], request["target_language"],
            controlled_source, translations,
        )
        docx_files.extend([translated, bilingual])

    with (job_dir / "CHANGE_LEDGER.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["paragraph_id", "category", "confidence", "original", "revised", "rationale"])
        writer.writeheader()
        for source in paragraphs:
            row = edits.get(source["paragraph_id"], {})
            writer.writerow({
                "paragraph_id": source["paragraph_id"], "category": row.get("category", "not_applicable"),
                "confidence": row.get("confidence", ""), "original": source["text"],
                "revised": row.get("revised", source["text"]), "rationale": row.get("rationale", ""),
            })
    glossary_rows = list(request.get("preferred_terms") or [])
    known = {str(row.get("source") or row.get("source_term") or "").casefold() for row in glossary_rows}
    for row in result.get("glossary") or []:
        if str(row.get("source_term") or "").casefold() not in known:
            glossary_rows.append(row)
    with (job_dir / "TERMINOLOGY_GLOSSARY.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["source_term", "target_term", "abbreviation", "note"])
        writer.writeheader()
        for row in glossary_rows:
            writer.writerow({
                "source_term": row.get("source") or row.get("source_term"),
                "target_term": row.get("target") or row.get("target_term"),
                "abbreviation": row.get("abbreviation", ""), "note": row.get("note", ""),
            })
    qa = {
        "schema": "dio.document_studio.qa.v1",
        "job_id": job_id,
        "passed": True,
        "automated_integrity_passed": True,
        "linguistic_quality_approved": False if translations else None,
        "release_readiness": "blocked_pending_human_approval",
        "automated_validation": validation,
        "provider_flags": result.get("qa_flags") or [],
        "translation_review_overrides": {
            "reviewer": request.get("translation_override_reviewer"),
            "applied_count": len(applied_overrides),
            "corrections": applied_overrides,
            "final_language_approval_still_required": True,
        },
        "language_controls": {
            "source_language": request["source_language"],
            "source_profile": request["source_language_profile"],
            "target_language": request.get("target_language"),
            "target_profile": request.get("target_language_profile"),
            "automated_critic": result.get("translation_critic"),
        },
        "format_core": {
            "style_profile": format_style,
            "delivery_profile": format_delivery,
            "channels": format_channels,
            "templates": format_templates,
            "languages": {
                language: {
                    "status": value["status"],
                    "qa": value["qa"],
                    "style_profile_hash": value["style_profile_hash"],
                    "delivery_profile_hash": value["delivery_profile_hash"],
                }
                for language, value in format_receipts.items()
            },
        },
        "human_gates": {
            "technical_editor_approval": "pending" if edits else "not_applicable",
            "target_language_reviewer_approval": "pending" if translations else "not_applicable",
            "client_final_approval": "pending",
        },
        "claim_boundaries": [
            "Typed semantic content is rebuilt through controlled style, language, and delivery profiles with post-render completeness checks.",
            "Arbitrary source files are not claimed to retain pixel-identical layout unless an explicit template profile has been supplied and approved.",
            "Translation is a review candidate, not certified, sworn, or legally attested translation.",
            "Operational facts, thresholds, identifiers, and terminology require human verification.",
        ],
    }
    write_json(job_dir / "PROVIDER_OUTPUT.json", result)
    write_json(job_dir / "BEAST_REUSE_RECEIPT.json", beast_receipt)
    write_json(job_dir / "BEAST_LEARNING_RECEIPT.json", beast_learning_receipt)
    write_json(job_dir / "DOCUMENT_STUDIO_QA.json", qa)
    (job_dir / "DOCUMENT_STUDIO_QA.md").write_text(
        "# Document Studio QA\n\n"
        f"Automated validation: **PASSED**\n\nParagraphs: {len(paragraphs)}  \n"
        f"Technical edit approval: **{qa['human_gates']['technical_editor_approval']}**  \n"
        f"Target-language approval: **{qa['human_gates']['target_language_reviewer_approval']}**\n\n"
        "## Human Flags\n\n"
        + ("\n".join(f"- {row.get('paragraph_id', 'document')}: {row.get('issue')}" for row in qa["provider_flags"]) or "- No provider flag; full human review remains required.")
        + "\n\n## Boundaries\n\n" + "\n".join(f"- {item}" for item in qa["claim_boundaries"]) + "\n",
        encoding="utf-8",
    )
    (job_dir / "HUMAN_APPROVAL.md").write_text(
        "# Human Approval Gate\n\n"
        "- [ ] Verify every number, threshold, identifier, warning, and responsibility.\n"
        "- [ ] Accept or reject each redline change.\n"
        "- [ ] Confirm terminology against the organisation's approved glossary.\n"
        "- [ ] Have a proficient target-language reviewer inspect the complete translation.\n"
        "- [ ] Confirm layout and accessibility after final formatting.\n"
        "- [ ] Record client approval before release.\n",
        encoding="utf-8",
    )
    pdfs = convert_to_pdf(docx_files, job_dir / "pdf")
    if edits and translations:
        create_proof_image(
            job_dir / "pdf" / "REDLINE_REVIEW_COPY.pdf",
            job_dir / "pdf" / "BILINGUAL_REVIEW_COPY.pdf",
            job_dir / "DOCUMENT_STUDIO_PROOF.png",
            request["source_language"],
            request["target_language"],
        )

    if translations:
        semantic_object, change_receipt = update_semantic_object(
            state_root=ROOT / "state" / "lingua",
            object_id=semantic_object_id,
            source_version=source_version,
            request=request,
            source_rows=paragraphs,
            translations=translations,
            provider=provider,
            qa_flags=list(result.get("qa_flags") or []),
        )
        write_json(job_dir / "LINGUA_SEMANTIC_OBJECT.json", semantic_object)
        write_json(job_dir / "LINGUA_CHANGE_RECEIPT.json", change_receipt)
        lingua_qa = build_lingua_qa(request, validation, result, beast_receipt)
        write_json(job_dir / "LINGUA_QA.json", lingua_qa)
        approval_template = {
            "schema": "dio.lingua.human_approval.v1",
            "approval_state": "pending",
            "semantic_object_id": semantic_object_id,
            "source_version": source_version,
            "target_language": request["target_language"],
            "domain": request.get("document_domain"),
            "reviewer": "",
            "reviewer_role": "",
            "approved_at": "",
            "units": [
                {
                    "unit_id": row["unit_id"],
                    "source_hash": row["source_hash"],
                    "target_text": next(item["target_text"] for item in semantic_object["translations"][request["target_language"]]["units"] if item["unit_id"] == row["unit_id"]),
                    "approved": False,
                }
                for row in semantic_object["source"]["units"]
            ],
            "terms": [],
            "authority_note": "A proficient target-language reviewer must correct the text, approve every unit, and identify their role before BEAST crystallization.",
        }
        write_json(job_dir / "LINGUA_APPROVAL_TEMPLATE.json", approval_template)

    receipt = {
        "schema": "dio.document_studio.receipt.v1",
        "job_id": job_id,
        "created_at": utc_now(),
        "status": "human_review_required",
        "service": service,
        "source": {"name": source_path.name, "sha256": sha256(source_path), "parser": parser, "paragraphs": len(paragraphs)},
        "processing": {
            "provider": provider.get("provider"), "model": provider.get("model"), "provider_status": provider.get("status"),
            "remote_processing_approved": True, "owner_authorized": True,
            "local_validation_passed": validation["passed"],
            "source_language": request["source_language"],
            "source_language_status": request["source_language_profile"]["status"],
            "target_language": request.get("target_language"),
            "target_language_status": (request.get("target_language_profile") or {}).get("status"),
            "target_language_reviewer_requirement": (request.get("target_language_profile") or {}).get("reviewer_requirement"),
            "automated_language_critic_provider": (critic_provider or {}).get("provider"),
            "automated_language_critic_model": (critic_provider or {}).get("model"),
            "format_core": {
                "style_profile": format_style,
                "delivery_profile": format_delivery,
                "channels": format_channels,
                "templates": format_templates,
                "rendered_languages": list(format_receipts),
                "all_output_qa_passed": all(row["qa"]["passed"] for row in format_receipts.values()),
            },
        },
        "release": {"delivery_released": False, "certified_translation": False, "human_approval_required": True},
        "outputs": [],
    }
    receipt["release"]["linguistic_quality_approved"] = False if translations else None
    receipt["release"]["release_readiness"] = "blocked_pending_human_approval"
    receipt["processing"]["beast_reuse_state"] = beast_receipt.get("reuse_state")
    receipt["processing"]["beast_learning_count"] = len(beast_learning_receipt.get("learned") or [])
    receipt["processing"]["semantic_object_id"] = str(request.get("semantic_object_id") or f"LINGUA-{job_id}") if translations else None
    write_json(job_dir / "DOCUMENT_STUDIO_RECEIPT.json", receipt)
    archive = job_dir / f"{job_id}_DOCUMENT_STUDIO_REVIEW_PACK.zip"
    generated = sorted(path for path in job_dir.rglob("*") if path.is_file() and path != archive)
    receipt["outputs"] = [{"path": str(path.relative_to(job_dir)), "sha256": sha256(path)} for path in generated]
    write_json(job_dir / "DOCUMENT_STUDIO_RECEIPT.json", receipt)
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(path for path in job_dir.rglob("*") if path.is_file() and path != archive):
            bundle.write(path, path.relative_to(job_dir))
    return job_dir
