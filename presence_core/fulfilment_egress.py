from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document

from .customer_cases import load_case, update_case
from .state import create_needs_you, write_json


SCHEMA = "dio.presence.fulfilment_egress.v1"


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _require_file(path: Path, label: str) -> Path:
    path = Path(path).resolve()
    if not path.is_file():
        raise ValueError(f"{label} not found: {path}")
    return path


def _render_commentary_pdf(
    commentary_md: Path,
    output_pdf: Path,
    *,
    title: str,
) -> Path:
    commentary_md = _require_file(commentary_md, "commentary")

    output_pdf = Path(output_pdf).resolve()
    output_pdf.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.add_heading(title, level=0)

    for raw in commentary_md.read_text(
        encoding="utf-8",
        errors="replace",
    ).splitlines():
        line = raw.strip()

        if not line:
            document.add_paragraph("")
            continue

        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("- "):
            document.add_paragraph(
                line[2:].strip(),
                style="List Bullet",
            )
        else:
            document.add_paragraph(line)

    with tempfile.TemporaryDirectory(
        prefix="dio_fulfilment_pdf_"
    ) as tmp:
        tmp = Path(tmp)
        docx_path = tmp / "customer-review.docx"
        pdf_dir = tmp / "pdf"
        pdf_dir.mkdir()

        document.save(docx_path)

        with tempfile.TemporaryDirectory(
            prefix="dio_lo_profile_"
        ) as profile:
            command = [
                "libreoffice",
                "--headless",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(pdf_dir),
                str(docx_path),
            ]

            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )

        generated = pdf_dir / "customer-review.pdf"

        if completed.returncode != 0 or not generated.is_file():
            raise RuntimeError(
                (
                    completed.stderr
                    or completed.stdout
                    or "LibreOffice PDF conversion failed"
                ).strip()
            )

        shutil.copy2(generated, output_pdf)

    return output_pdf


def build_controlled_local_sophia_request(
    case: dict[str, Any],
    *,
    source_path: Path,
    operator_id: str,
) -> dict[str, Any]:
    scope = case.get("scope") or {}
    commercial = case.get("commercial") or {}
    settlement = (
        commercial.get("controlled_test_settlement")
        or {}
    )

    if (
        case.get("stage") != "PAYMENT_VERIFIED"
        or commercial.get("payment_state")
        != "verified_controlled_test"
        or settlement.get("controlled_test") is not True
    ):
        raise ValueError(
            "controlled-test settlement is required "
            "before local Sophia processing"
        )

    case_id = str(case.get("case_id") or "")
    attachment_id = str(
        scope.get("attachment_id") or ""
    )
    source_sha256 = str(
        scope.get("scope_scan_sha256") or ""
    )

    if (
        settlement.get("case_id") != case_id
        or settlement.get("attachment_id") != attachment_id
    ):
        raise ValueError(
            "controlled-test settlement lineage mismatch"
        )

    source_path = Path(source_path)

    if not source_path.is_file():
        raise ValueError(
            "controlled Sophia source file is missing"
        )

    actual_source_sha256 = _sha256(source_path)

    if actual_source_sha256 != source_sha256:
        raise ValueError(
            "controlled Sophia source SHA-256 mismatch"
        )

    return {
        "schema": "dio.sophia_review_request.v1",
        "job_id": f"SOPHIA-{case_id}",
        "title": (
            source_path.stem.replace("-", " ").strip()
            or "Controlled Sophia Review"
        ),
        "document_path": str(source_path),
        "research_question": (
            "Perform a controlled diagnostic review of "
            "the supplied manuscript."
        ),
        "literature_queries": [],
        "citation_style": "APA 7",
        "external_retrieval": False,
        "reasoned_review_approved": True,
        "reasoned_provider": "local",
        "reasoned_model": "qwen2.5:3b",
        "gemini_review_approved": False,
        "gemini_model": "",
        "human_approval_required": True,
        "controlled_test_authority": {
            "controlled_test": True,
            "operator_id": operator_id,
            "case_id": case_id,
            "attachment_id": attachment_id,
            "source_sha256": source_sha256,
            "settlement_receipt_id": settlement.get(
                "receipt_id"
            ),
            "external_retrieval_authorized": False,
            "remote_processing_authorized": False,
            "customer_consent_claimed": False,
            "authority_created": False,
        },
    }


def prepare_review_ready_output(
    state_root: Path,
    *,
    case_id: str,
    attachment_id: str,
    source_sha256: str,
    job_id: str,
    commentary_md: Path,
    proof_pack: Path,
    review_receipt: Path,
    controlled_test: bool = False,
) -> dict[str, Any]:
    state_root = Path(state_root).resolve()

    case = load_case(state_root, case_id)
    if case is None:
        raise ValueError(f"customer case not found: {case_id}")

    commentary_md = _require_file(
        commentary_md,
        "review commentary",
    )
    proof_pack = _require_file(
        proof_pack,
        "proof pack",
    )
    review_receipt = _require_file(
        review_receipt,
        "review receipt",
    )

    source = (
        state_root
        / "quarantine"
        / attachment_id
        / "content.blob"
    ).resolve()

    source = _require_file(
        source,
        "quarantined source",
    )

    actual_source_sha = _sha256(source)

    if actual_source_sha != source_sha256:
        raise ValueError(
            "source attachment SHA-256 mismatch"
        )

    receipt = json.loads(
        review_receipt.read_text(
            encoding="utf-8",
        )
    )

    receipt_source_sha = str(
        receipt.get("source_sha256")
        or receipt.get("source_document_sha256")
        or ""
    ).strip().lower()

    if (
        receipt_source_sha
        and receipt_source_sha != actual_source_sha
    ):
        raise ValueError(
            "review receipt source SHA-256 mismatch"
        )

    production_payment_verified = (
        str(
            (case.get("commercial") or {}).get(
                "payment_state"
            )
            or ""
        ).strip().lower()
        in {
            "verified",
            "payment_verified",
            "paid_verified",
        }
        or str(case.get("stage") or "").strip().upper()
        in {
            "PAYMENT_VERIFIED",
            "WORK_QUEUED",
            "PROCESSING",
            "REVIEW_READY",
            "RELEASE_APPROVAL",
            "DELIVERED",
            "CLOSED",
        }
    )

    if not controlled_test and not production_payment_verified:
        raise ValueError(
            "production fulfilment egress requires "
            "verified payment"
        )

    if controlled_test:
        artifact_root = (
            Path("/tmp")
            / "dio-fulfilment-egress-test"
            / case_id
        )
    else:
        artifact_root = (
            state_root
            / "customer_cases"
            / "artifacts"
            / case_id
        )

    artifact_root.mkdir(
        parents=True,
        exist_ok=True,
    )

    customer_pdf = (
        artifact_root
        / f"{job_id}_CUSTOMER_REVIEW.pdf"
    )

    _render_commentary_pdf(
        commentary_md,
        customer_pdf,
        title="Sophia Integrity Review",
    )

    proof_copy = (
        artifact_root
        / proof_pack.name
    )
    receipt_copy = (
        artifact_root
        / review_receipt.name
    )

    shutil.copy2(
        proof_pack,
        proof_copy,
    )
    shutil.copy2(
        review_receipt,
        receipt_copy,
    )

    result: dict[str, Any] = {
        "schema": SCHEMA,
        "case_id": case_id,
        "product_id": case.get("product_id"),
        "job_id": job_id,
        "source": {
            "attachment_id": attachment_id,
            "sha256": actual_source_sha,
        },
        "customer_artifact": {
            "kind": "document",
            "purpose": "product_fulfilment",
            "path": str(customer_pdf),
            "file_name": customer_pdf.name,
            "mime_type": "application/pdf",
            "bytes": customer_pdf.stat().st_size,
            "sha256": _sha256(customer_pdf),
            "release_state": "HELD",
        },
        "proof_pack": {
            "path": str(proof_copy),
            "bytes": proof_copy.stat().st_size,
            "sha256": _sha256(proof_copy),
        },
        "review_receipt": {
            "path": str(receipt_copy),
            "sha256": _sha256(receipt_copy),
        },
        "authority": {
            "controlled_test": controlled_test,
            "payment_verified": (
                production_payment_verified
            ),
            "release_authority": False,
            "external_send_authority": False,
            "authority_created": False,
        },
        "egress_state": (
            "CONTROLLED_TEST_HELD"
            if controlled_test
            else "HELD_FOR_HUMAN_RELEASE"
        ),
    }

    binding_path = (
        artifact_root
        / "FULFILMENT_EGRESS.json"
    )

    write_json(
        binding_path,
        result,
    )

    result["binding_path"] = str(
        binding_path
    )

    if controlled_test:
        return result

    conversation_ids = list(
        case.get("conversation_ids") or []
    )

    conversation_id = (
        str(conversation_ids[-1])
        if conversation_ids
        else ""
    )

    needs = create_needs_you(
        state_root,
        reason="fulfilment_release_review",
        conversation_id=conversation_id,
        product=str(
            case.get("product_id") or ""
        )
        or None,
        summary=(
            f"Review completed fulfilment for "
            f"{case_id} / {job_id}. "
            f"Customer artifact SHA-256 "
            f"{result['customer_artifact']['sha256']}."
        ),
        priority="normal",
    )

    needs_ids = list(
        case.get("needs_you_ids") or []
    )

    if needs["needs_you_id"] not in needs_ids:
        needs_ids.append(
            needs["needs_you_id"]
        )

    updated = update_case(
        state_root,
        case,
        stage="REVIEW_READY",
        patch={
            "needs_you_ids": needs_ids,
            "fulfilment": {
                "job_id": job_id,
                "state": "review_ready",
                "customer_artifact": (
                    result["customer_artifact"]
                ),
                "proof_pack": (
                    result["proof_pack"]
                ),
                "review_receipt": (
                    result["review_receipt"]
                ),
                "release_authority": False,
                "external_send_authority": False,
            },
        },
        evidence_ref=(
            f"fulfilment:{job_id};"
            f"sha256:"
            f"{result['customer_artifact']['sha256']}"
        ),
    )

    result["needs_you_id"] = (
        needs["needs_you_id"]
    )
    result["case_stage"] = (
        updated["stage"]
    )

    write_json(
        binding_path,
        result,
    )

    return result
