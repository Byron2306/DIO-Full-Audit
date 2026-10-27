"""Source-driven contract and readiness products over existing DIO engines."""
from __future__ import annotations

import hashlib
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Mapping

from adapters.evidex_studio import build_readiness_gap_map
from adapters.format_core.renderer import render_semantic_asset
from adapters.lingua.lifecycle import register_product_source
from dio.obligations.engine import build as build_obligations
from dio.obligations.extractor import extract as extract_obligations


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")


def _new_output(path: str | Path) -> Path:
    out = Path(path).resolve()
    if out.exists() and (not out.is_dir() or any(out.iterdir())):
        raise ValueError("customer output directory must be empty; retain previous jobs separately")
    out.mkdir(parents=True, exist_ok=True, mode=0o700)
    return out


def _sources(packet: Mapping[str, Any]) -> dict[str, dict[str, str]]:
    result = {}
    for row in packet.get("sources") or []:
        key, text = str(row.get("source_id") or ""), str(row.get("text") or "")
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,79}", key) or key in result or not text.strip():
            raise ValueError("each source requires unique source_id and actual text")
        digest = _sha(text.encode())
        if row.get("sha256") and row["sha256"] != digest:
            raise ValueError(f"source hash mismatch: {key}")
        result[key] = {"source_id": key, "text": text, "sha256": digest, "title": str(row.get("title") or key)}
    if not result:
        raise ValueError("actual customer sources are required")
    return result


def _quote(row: Mapping[str, Any], sources: Mapping[str, Any]) -> dict[str, Any]:
    source = sources.get(str(row.get("source_id") or ""))
    quote = str(row.get("quote") or "")
    if source is None or not quote.strip() or quote not in source["text"]:
        raise ValueError("source quote must match supplied source bytes")
    start = source["text"].index(quote)
    return {"source_id": source["source_id"], "source_sha256": source["sha256"], "quote": quote, "start": start, "end": start + len(quote)}


def _paragraph(text: str) -> dict[str, Any]:
    return {"type": "paragraph", "text": text}


def _heading(text: str) -> dict[str, Any]:
    return {"type": "heading", "level": 1, "text": text}


def _display_state(value: Any) -> str:
    return str(value).replace('_', ' ').capitalize()


def _table(headers: list[str], rows: list[list[Any]]) -> dict[str, Any]:
    return {"type": "table", "headers": headers, "rows": [[str(v) for v in row] for row in rows]}


def _render(product: str, packet: Mapping[str, Any], sources: dict, blocks: list[dict], out: Path, *, render_pdf: bool, details: dict) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    case_id = str(packet.get("case_id") or "").strip()
    title = str(packet.get("title") or "").strip()
    if not case_id or not title:
        raise ValueError("customer case_id and title are required")
    identity = json.dumps({"product": product, "packet": packet}, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    packet_sha = _sha(identity.encode())
    object_id = "CUSTOMER-" + packet_sha[:24]
    semantic, change = register_product_source(
        state_root=out / "lingua", object_id=object_id, source_version="1.0.0", source_language="English",
        source_rows=[{"paragraph_id": row["source_id"], "text": row["text"]} for row in sources.values()],
        origin={"product": product, "artifact_type": "customer_specialist_review", "artifact_id": case_id, "audience": "customer and professional reviewer", "channel": "document"},
        domain="customer supplied business records",
    )
    content = {"schema": "dio.semantic_content.v1", "object_id": object_id, "version": "1.0.0", "title": title,
               "source_language": "en", "context": {"product": product, "artifact_type": "customer_review_package", "audience": "customer and professional reviewer"},
               "translations": {}, "blocks": [{**block, "block_id": f"B{i:04d}"} for i, block in enumerate([{"type":"title","text":title}, _paragraph(product + " · review copy"), *blocks], 1)]}
    _json(out / "SOURCE_PACKET.json", packet)
    _json(out / "SOURCE_INVENTORY.json", [{k:v for k,v in row.items() if k!='text'} for row in sources.values()])
    _json(out / "LINGUA_CHANGE_RECEIPT.json", change)
    _json(out / "REVIEW_CONTENT.json", content)
    rendered = render_semantic_asset(content, out / "customer_delivery", style_profile="minimal_print", delivery_profile="editable_review", channels=["docx", "html"] + (["pdf"] if render_pdf else []), release_mode=False)
    outputs = [{**a, "path": "customer_delivery/" + a["path"]} for a in rendered["outputs"]]
    receipt = {"schema": "dio.customer_specialist_package.v1", "product_name": product, "case_id": case_id,
               "status": "HELD_REVIEW_PACKAGE_RENDERED", "customer_packet_sha256": packet_sha,
               "source_inventory": [{k:v for k,v in row.items() if k!='text'} for row in sources.values()],
               "lingua_source_document_hash": semantic["source"]["document_hash"], "format_core_receipt": rendered,
               "outputs": outputs, "professional_approval": "NOT_GRANTED", "human_review_required": True,
               "authority_created": False, "external_effects": False, "commercial_validation": "UNPROVED", **details}
    receipt["executed_providers"] += ["adapters.lingua.lifecycle.register_product_source", "adapters.format_core.renderer.render_semantic_asset"]
    _json(out / "CUSTOMER_PACKAGE_RECEIPT.json", receipt)
    return receipt


def build_contract_desk_package(packet: Mapping[str, Any], output_dir: str | Path, *, render_pdf: bool = True) -> dict[str, Any]:
    from datetime import datetime, timezone
    review_as_of = str(packet.get("review_at") or datetime.now(timezone.utc).replace(microsecond=0).isoformat())
    sources = _sources(packet)
    key = str(packet.get("contract_source_id") or "")
    if key not in sources:
        raise ValueError("contract_source_id must name an actual supplied contract")
    source = sources[key]
    clauses, actions, terms = [], [], []
    chunks = re.split(r"(?m)^(\d+(?:\.\d+)*)(?:[.)]\s+|\s+)", source["text"])
    if len(chunks) == 1:
        raw = [(f"C{i}", line) for i, line in enumerate(re.split(r"\n\s*\n", source["text"]), 1) if line.strip()]
    else:
        raw = [(chunks[i], chunks[i+1].strip()) for i in range(1, len(chunks), 2)]
        if chunks[0].strip():
            raw.insert(0, ("PREAMBLE", chunks[0].strip()))
    for locator, text in raw:
        if re.search(r"ignore (?:all |previous |prior )?instructions|system prompt|bypass (?:human|approval)", text, re.I):
            actions.append({"clause_id": locator, "question": "Review source instruction as untrusted content; it was excluded from obligation extraction.", "source_quote": text})
            continue
        dates = re.findall(r"\b\d{4}-\d{2}-\d{2}\b", text)
        row = {"clause_id": locator, "text": text}
        # Only an explicitly stated 'by/on YYYY-MM-DD' binds a due date. Other
        # dates and relative windows remain owner questions, not inferred law.
        due_candidates = []
        for sentence in re.split(r"(?<=[.!?;])\s+", text):
            due = re.search(r"\b(?:by|on)\s+(\d{4}-\d{2}-\d{2})\b", sentence, re.I)
            if due and re.search(r"\b(?:shall|must|is required to|are required to|will be required to)\b", sentence, re.I) and not re.search(r"\b(?:within|after|before|commences|effective)\b", sentence, re.I):
                due_candidates.append(due.group(1))
        if len(due_candidates) == 1 and len(dates) == 1:
            row["due_at"] = due_candidates[0] + "T23:59:59+02:00"
        if re.search(r"\b(?:within|after|before)\b.*?\b\d+\s+(?:business\s+)?days\b|\b\d+\s+days?\b.*?\b(?:after|before|notice)\b", text, re.I):
            actions.append({"clause_id": locator, "question": "Confirm the trigger, acceptance date and calendar/business-day convention before calculating: " + text, "source_quote": text})
        if len(dates) > 1:
            actions.append({"clause_id": locator, "question": "Confirm which explicit date applies to each obligation in this clause.", "source_quote": text})
        clauses.append(row)
        for topic in ["payment", "renewal", "termination", "liability", "confidentiality", "notice"]:
            needle = {"payment": r"\bpay\w*", "renewal": r"\brenew\w*", "termination": r"\bterminat\w*", "confidentiality": r"\bconfiden\w*"}.get(topic, rf"\b{topic}\b")
            if re.search(needle, text, re.I):terms.append([topic, locator, text])
    if not clauses:
        raise ValueError("no reviewable contract clauses remain")
    engine_source = {"source_id": key, "source_ref": f"customer-packet://{key}", "source_type": "contract", "sha256": source["sha256"], "clauses": clauses}
    candidates = extract_obligations(engine_source)
    by_locator = {c["source_locator"]: c["obligation_id"] for c in candidates}
    known_clauses = {c["clause_id"] for c in clauses}
    evidence = []
    for index, item in enumerate(packet.get("evidence") or [], 1):
        anchor = _quote(item, sources)
        target = [str(t) for t in item.get("supports_clause_ids") or []]
        if not target or not set(target) <= known_clauses:
            raise ValueError("contract evidence requires known clause IDs")
        ids = [by_locator[t] for t in target if t in by_locator]
        if ids:
            evidence.append({"evidence_id": str(item.get("evidence_id") or f"E{index}"), "obligation_ids": ids,
                             "source_ref": f"customer-packet://{anchor['source_id']}#{anchor['start']}:{anchor['end']}", "relation": "supports",
                             "trust_state": "captured_untrusted", "freshness_state": "unknown"})
    bundle = build_obligations(engine_source, evidence_records=evidence, now=review_as_of)
    out = _new_output(output_dir)
    _json(out / "OBLIGATION_BUNDLE.json", bundle)
    _json(out / "REVIEW_ACTIONS.json", actions)
    blocks = [_heading("Your contract review"), _paragraph("Review as of: " + review_as_of + ". Calendar states use this case review time."), _paragraph(f"{len(clauses)} supplied clauses; {len(bundle['obligations'])} obligation candidates; {len(actions)} questions for the contract owner. This review organises the supplied agreement and records; it does not decide enforceability or settle a dispute."),
              _heading("Clause register"), _table(["Clause", "Supplied wording"], [[c["clause_id"],c["text"]] for c in clauses])]
    if terms: blocks += [_heading("Commercial terms to review"), _table(["Topic","Clause","Exact wording"],terms)]
    if bundle['obligations']: blocks += [_heading("Obligations and evidence"), _table(["Clause","Obligation candidate","Due date","Review state"], [[r['source']['locator'],r['statement'],str(r.get('due_at') or '')[:10] or 'Trigger/date to confirm',_display_state(r['status'])] for r in bundle['obligations']])]
    if bundle['deadlines']: blocks += [_heading("Calendar"),_table(["Obligation","Date","Calendar state"],[[r['obligation_id'],str(r['due_at'])[:10],_display_state(r['state'])] for r in bundle['deadlines']])]
    if actions: blocks += [_heading("Questions for your reviewer"),_table(["Clause","Action"],[[r['clause_id'],r['question']] for r in actions])]
    blocks += [_heading("Source record"),_paragraph(f"{source['title']} ({key}). Counterparties supplied: {', '.join(str(p) for p in packet.get('counterparties') or []) or 'not supplied'}. Original source bytes are retained in the source packet.")]
    return _render("Contract Desk",packet,sources,blocks,out,render_pdf=render_pdf,details={"executed_providers":["dio.obligations.extractor.extract","dio.obligations.engine.build"],"review_as_of":review_as_of,"review_actions":actions,"obligation_count":len(bundle['obligations']),"clause_count":len(clauses)})


_FINANCIAL_FIELDS = {"revenue":"revenue", "cost_of_sales":"cost of sales", "current_assets":"current assets", "current_liabilities":"current liabilities", "cash":"cash", "monthly_cash_outflow":"monthly cash outflow"}


def _financial_metrics(financials: Mapping[str, Any], sources: dict) -> tuple[dict[str, str], dict[str, Decimal]]:
    anchor = _quote(financials,sources)
    quote = anchor['quote']
    if not financials.get('period') or str(financials['period']) not in quote or not financials.get('currency') or str(financials['currency']) not in quote:
        raise ValueError('financial source must explicitly support period and currency')
    numbers = {}
    for field,label in _FINANCIAL_FIELDS.items():
        if field not in financials:continue
        try:value=Decimal(str(financials[field]).replace(',',''))
        except InvalidOperation as exc:raise ValueError('invalid financial number') from exc
        if not value.is_finite():raise ValueError('financial numbers must be finite')
        found=re.search(rf'\b{re.escape(label)}\b\s*[:=]?\s*(?:ZAR|R)?\s*([-+]?\d[\d,]*(?:\.\d+)?)',quote,re.I)
        if not found or Decimal(found.group(1).replace(',','')) != value:
            raise ValueError(f'financial number is not supported by source label: {field}')
        numbers[field]=value
    metrics={}
    if {'revenue','cost_of_sales'} <= numbers.keys() and numbers['revenue']>0:
        metrics['gross_margin_percent']=f"{(numbers['revenue']-numbers['cost_of_sales'])/numbers['revenue']*100:.2f}"
    if {'current_assets','current_liabilities'} <= numbers.keys() and numbers['current_liabilities']>0:
        metrics['current_ratio']=f"{numbers['current_assets']/numbers['current_liabilities']:.2f}"
    if {'cash','monthly_cash_outflow'} <= numbers.keys() and numbers['monthly_cash_outflow']>0:
        metrics['cash_runway_months']=f"{numbers['cash']/numbers['monthly_cash_outflow']:.2f}"
    return metrics,numbers


def build_readiness_package(product_name: str, packet: Mapping[str, Any], output_dir: str | Path, *, render_pdf: bool = True) -> dict[str, Any]:
    if product_name not in {'Corporate Readiness','Finance Readiness','Finance Readiness Studio','POPIA Readiness','InvestorProof','EntrepreneurProof'}:
        raise ValueError('unsupported readiness product')
    sources=_sources(packet); requirements=[]; seen=set()
    for row in packet.get('requirements') or []:
        key=str(row.get('requirement_id') or '')
        if not key or key in seen:raise ValueError('requirements need unique IDs')
        seen.add(key);anchor=_quote(row,sources)
        requirements.append({'requirement_id':key,'label':str(row.get('label') or anchor['quote']),'mandatory':bool(row.get('mandatory',True)),'source_anchor':anchor})
    if not requirements:raise ValueError('separate source-bound requirements are required')
    evidence=[];evidence_ids=set()
    for row in packet.get('evidence') or []:
        key=str(row.get('evidence_id') or '')
        if not key or key in evidence_ids:raise ValueError('evidence requires unique IDs')
        evidence_ids.add(key);anchor=_quote(row,sources);targets=list(row.get('supports') or [])
        if not targets or not set(targets)<=seen:raise ValueError('evidence must map to known requirements')
        evidence.append({'evidence_id':key,'label':anchor['quote'],'supports':targets,'state':'supplied','source_anchor':anchor})
    native=build_readiness_gap_map({'studio_id':str(packet.get('case_id') or ''),'artifact_contract':{'kind':'finance_readiness','published_requirements_fixture':requirements,'supplied_evidence_fixture':evidence,'decision_boundary':'Supplied evidence presence only; applicability, factual accuracy and professional approval remain for the responsible reviewer.'}})
    matrix=[{**r,'state':'RECEIVED_FOR_REVIEW' if r['supporting_evidence_ids'] else 'MISSING'} for r in native['requirements']]
    actions=[{'requirement_id':r['requirement_id'],'requested_document':r['label'],'owner':'Customer evidence owner','next_action':'Supply the missing record or confirm that this requirement is inapplicable.'} for r in matrix if r['state']=='MISSING']
    out=_new_output(output_dir)
    _json(out/'READINESS_EVIDENCE_MATRIX.json',matrix);_json(out/'DOCUMENT_REQUESTS.json',actions);_json(out/'EVIDEX_READINESS_RECEIPT.json',native)
    blocks=[_heading('Readiness overview'),_paragraph(f"{len(matrix)} supplied requirements reviewed. {len(matrix)-len(actions)} have records received; {len(actions)} need documents. Receipt of a record does not independently verify its validity or meet a legal, lender or investor approval condition."),_heading('Requirement-to-evidence review'),_table(['Requirement','Required record','Received evidence','Review state'],[[r['requirement_id'],r['label'],', '.join(r['supporting_evidence_ids']) or 'None supplied',_display_state(r['state'])] for r in matrix])]
    if actions:blocks += [_heading('Documents to request'),_table(['Requirement','Requested document','Owner','Next action'],[[a['requirement_id'],a['requested_document'],a['owner'],a['next_action']] for a in actions])]
    metrics={}
    if packet.get('financials') and product_name in {'Finance Readiness','Finance Readiness Studio','InvestorProof','EntrepreneurProof'}:
        metrics,numbers=_financial_metrics(packet['financials'],sources)
        blocks += [_heading('Source-based financial calculations'),_paragraph(f"Period: {packet['financials']['period']}; currency: {packet['financials']['currency']}. Calculations use only the explicitly supplied amounts; they are not affordability, lending or valuation decisions."),_table(['Input','Supplied amount'],[[_display_state(k),f'{v:,.2f}'] for k,v in numbers.items()])]
        metric_labels={'gross_margin_percent':'Gross margin (%)','current_ratio':'Current ratio','cash_runway_months':'Cash runway (months)'}
        if metrics:blocks += [_table(['Metric','Calculated value'],[[metric_labels.get(k,_display_state(k)),v] for k,v in metrics.items()])]
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.title='Source amounts';sheet.append(['Field','Amount','Period','Currency','Source','Exact source quote'])
        for key,value in numbers.items():sheet.append([key,float(value),packet['financials']['period'],packet['financials']['currency'],packet['financials']['source_id'],packet['financials']['quote']])
        # Customer strings are evidence, including a leading '='; never turn a
        # supplied period or quotation into an executable spreadsheet formula.
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        calculations=book.create_sheet('Calculated metrics');calculations.append(['Metric','Value','Basis']);
        for key,value in metrics.items():calculations.append([key,float(value),'Arithmetic from supplied amounts; human review required'])
        book.save(out/'FINANCIAL_REVIEW_WORKBOOK.xlsx')
    blocks += [_heading('Source and reviewer handoff'),_table(['Source','Description'],[[k,row['title']] for k,row in sources.items()]),_paragraph('Review applicability, freshness, signatures and factual accuracy with the appropriate evidence owner before using this pack for a consequential decision.')]
    result=_render(product_name,packet,sources,blocks,out,render_pdf=render_pdf,details={'executed_providers':['adapters.evidex_studio.build_readiness_gap_map'],'requirements':matrix,'mandatory_missing_count':native['mandatory_gap_count'],'review_actions':actions,'financial_metrics':metrics})
    if (out/'FINANCIAL_REVIEW_WORKBOOK.xlsx').is_file():
        path=out/'FINANCIAL_REVIEW_WORKBOOK.xlsx';result['outputs'].append({'path':path.name,'sha256':_sha(path.read_bytes()),'channel':'xlsx','bytes':path.stat().st_size});_json(out/'CUSTOMER_PACKAGE_RECEIPT.json',result)
    return result


def build_evidence_review_package(product_name: str, packet: Mapping[str, Any], output_dir: str | Path, *, render_pdf: bool = True) -> dict[str, Any]:
    """Apply an existing governed review profile to actual separate source inputs."""
    from products.evidence_review import run_controlled_evidence_review, validate_profile
    from products.governed_case import new_case
    from datetime import datetime, timezone

    if packet.get('owner_authorized') is not True:
        raise ValueError('explicit source-owner authorization is required')

    root=Path(__file__).resolve().parents[1]
    routes=json.loads((root/'config/professional_evidence_portfolio/v1/routes.json').read_text())['routes']
    route=routes.get(product_name) or {}
    if route.get('route') not in {'evidence_profile','high_risk_profile','education_research_profile'}:
        raise ValueError('product has no generic evidence-review provider contract')
    profile_id=route['profile_id']
    profile_path=root/'config/products/profiles'/f'{profile_id}.json'
    profile=json.loads(profile_path.read_text());validate_profile(profile)
    sources=_sources(packet);requirements=[];keys=set()
    for row in packet.get('requirements') or []:
        key=str(row.get('requirement_id') or '')
        if not key or key in keys:raise ValueError('separate unique requirement IDs are required')
        keys.add(key);anchor=_quote(row,sources)
        # A caller's display label is not evidence of what the source says.
        # Keep the exact matched quotation as the native review statement.
        requirements.append({'requirement_key':key,'statement':anchor['quote'],
                             'source_ref':f"customer-packet://{anchor['source_id']}#{anchor['start']}:{anchor['end']}",
                             'mandatory':bool(row.get('mandatory',True)),'dependency_requirement_keys':[]})
    if not requirements:raise ValueError('actual source-bound review requirements are required')
    evidence=[]
    for row in packet.get('evidence') or []:
        anchor=_quote(row,sources);targets=list(row.get('supports') or [])
        if not targets or not set(targets)<=keys:raise ValueError('evidence must name actual requirement IDs')
        evidence.append({'evidence_kind':'customer_record','source_ref':f"customer-packet://{anchor['source_id']}#{anchor['start']}:{anchor['end']}",
                         'sha256':anchor['source_sha256'],'target_requirement_keys':targets,'relation':'supports',
                         'authority_grade':'self_asserted','trust_state':'captured_untrusted','freshness_state':'unknown'})
    issues=[]
    for row in packet.get('issues') or []:
        if row.get('requirement_id') not in keys:raise ValueError('issue must name a known requirement')
        anchor=_quote(row,sources)
        issues.append({'requirement_key':row['requirement_id'],'challenge_type':'alternative_hypothesis','severity':'material',
                       'hypothesis':'Source-reported issue for verification: '+anchor['quote']})
    out=_new_output(output_dir);_json(out/'SOURCE_PACKET.json',packet)
    now=str(packet.get('review_at') or datetime.now(timezone.utc).replace(microsecond=0).isoformat())
    case=new_case(product=profile['product_id'],job_id=str(packet.get('case_id') or ''),
                  source={'kind':'customer_source_packet','source_sha256':_sha((out/'SOURCE_PACKET.json').read_bytes())},source_path=out/'SOURCE_PACKET.json',
                  evidence_inputs=['source-bound customer requirements','source-bound customer evidence'],
                  expected_outputs=['controlled evidence review pack','proof manifest'],required_authorities=['customer_source_owner','domain_reviewer'],intake_state='approved',now=now)
    native=run_controlled_evidence_review(case,profile=profile,requirements=requirements,evidence_inputs=evidence,issues=issues,
                                         output_dir=out/'native_review',operator_id='customer-specialist-workbench',now=now)
    matrix=native['review_pack']['requirement_evidence_matrix'];issue_register=native['review_pack']['issue_register']
    blocks=[_heading('Your evidence review'),_paragraph(f"{len(matrix)} requirements and {len(evidence)} supplied records were processed using the {product_name} evidence-review profile. Source presence and quoted statements are recorded; independent truth, freshness and domain conclusions remain unverified."),
            _heading('Requirement and evidence matrix'),_table(['Requirement','Source requirement','Evidence references','Review state'],[[r['requirement_key'],r['statement'],', '.join(r['evidence_ids']) or 'No record supplied',r['review_state']] for r in matrix])]
    if issue_register:blocks += [_heading('Issues requiring a decision'),_table(['Requirement','Source-reported issue','Review priority'],[[r['requirement_key'],r['hypothesis'],r['severity']] for r in issue_register])]
    blocks += [_heading('Next review steps'),_paragraph('Verify the supplied records with their owners, check currency and applicability of the requirements, and resolve the listed documentary gaps before making a domain decision.'),
               _heading('Scope and evidence handoff'),_paragraph(str(profile.get('scope_limit') or profile.get('human_gate_reason') or 'Evidence preparation and review support.'))]
    return _render(product_name,packet,sources,blocks,out,render_pdf=render_pdf,
                   details={'executed_providers':['products.evidence_review.run_controlled_evidence_review'],'profile_id':profile_id,
                            'profile_source_sha256':_sha(profile_path.read_bytes()),'native_processing_receipt':native['receipt'],
                            'domain_engine_invoked':False,'requirements':matrix,'review_actions':issue_register})


def build_correspondence_package(packet: Mapping[str, Any], output_dir: str | Path, *, render_pdf: bool = True) -> dict[str, Any]:
    """Draft a specific factual letter; the sender decides tone, position and send."""
    sources=_sources(packet)
    recipient,sender,action=(str(packet.get(k) or '').strip() for k in ['recipient','sender','requested_action'])
    if not recipient or not sender or not action:raise ValueError('recipient, sender and requested_action are required')
    anchors=[_quote(row,sources) for row in packet.get('facts') or []]
    if not anchors:raise ValueError('correspondence requires exact source-bound facts')
    retained=[];quarantined=[]
    for anchor in anchors:
        if re.search(r'ignore (?:previous |prior )?instructions|declare .*approved|bypass .*approval',anchor['quote'],re.I):quarantined.append(anchor)
        else:retained.append(anchor)
    if not retained:raise ValueError('no factual source statements remain for correspondence')
    lines=[f"Subject: {packet.get('title')}",'',f'Dear {recipient},','',
           'Thank you for your correspondence. The records supplied for this matter state:','']
    lines += [f"- {a['quote']}" for a in retained]
    lines += ['',action,'','Please let us know if any of these records need correction or if further supporting documents are available.','',
              'This draft records a documentary review request; it makes no settlement, payment, waiver or liability commitment.','', 'Kind regards,',sender]
    text='\n'.join(lines)+'\n'
    out=_new_output(output_dir);(out/'DRAFT_CORRESPONDENCE.txt').write_text(text)
    _json(out/'CORRESPONDENCE_SOURCE_MAP.json',{'facts':retained,'quarantined':quarantined,'send_authorized':False})
    blocks=[_paragraph(line) for line in lines if line.strip()]
    result=_render('Professional Correspondence',packet,sources,blocks,out,render_pdf=render_pdf,
                   details={'executed_providers':['source_bound_correspondence_composition'],'send_authorized':False,'source_facts':retained,'quarantined_source_instructions':quarantined})
    path=out/'DRAFT_CORRESPONDENCE.txt';result['outputs'].append({'path':path.name,'sha256':_sha(path.read_bytes()),'channel':'txt','bytes':path.stat().st_size})
    _json(out/'CUSTOMER_PACKAGE_RECEIPT.json',result)
    return result


def build_evidex_customer_package(packet: Mapping[str, Any], output_dir: str | Path, *, render_pdf: bool = True) -> dict[str, Any]:
    """Run the full Evidex engine on actual records, then finish the review pack."""
    from datetime import datetime, timezone
    from scripts.run_evidex_jobs import run_evidex
    from products.evidex_customer_review import build_record_review
    if packet.get('owner_authorized') is not True:
        raise ValueError('explicit source-owner authorization is required')
    sources = _sources(packet)
    records = []
    for row in packet.get('records') or [{'record_id': key, 'source_id': key, 'quote': source['text']} for key, source in sources.items()]:
        anchor = _quote(row, sources)
        records.append({'record_id': str(row.get('record_id') or ''), 'text': anchor['quote']})
    review = build_record_review(records)
    out = _new_output(output_dir)
    job = {'job_id': 'CUSTOMER-' + _sha(json.dumps(packet,sort_keys=True).encode())[:20],
           'created_at': datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
           'route': {'product': 'evidex', 'confidence': 1, 'reason': 'Explicit source-owner-authorized customer record review'},
           'customer_context': {'organisation': str(packet.get('organisation') or 'Customer organisation not supplied'), 'reporting_period': dict(packet.get('reporting_period') or {})},
           'inputs': [{'sender': str(packet.get('contact_name') or 'Customer contact not supplied'), 'subject': str(packet.get('title') or 'Customer evidence review')}],
           'evidence': [{'text_extract': '\n'.join(row['text'] for row in records)}], 'source_records': records}
    _json(out/'EVIDEX_CUSTOMER_JOB.json',job)
    native = run_evidex(job,out/'native_evidex')
    if native.get('returncode') != 0:
        raise RuntimeError('full Evidex engine failed: '+str(native.get('stderr') or native.get('studio_finish'))[-700:])
    _json(out/'CUSTOMER_EVIDENCE_REVIEW.json',review)
    blocks = [_heading('Claims and documentary leads'), _paragraph(review['scope']),
              _table(['Record','Exact supplied statement'], [[r['record_id'],r['text']] for r in review['source_records'] if r['record_id'] not in review['quarantined_record_ids']])]
    if review['quantity_comparisons']:
        blocks += [_heading('Reported quantities to reconcile'),_table(['Claim','Declared quantity','Reported record quantity','Difference','Boundary'],
                    [[r['claim_id'],r['declared_value'],r['reported_record_value'],r['difference'],r['comparability']] for r in review['quantity_comparisons']])]
    blocks += [_heading('Questions and records to obtain'),_table(['Record','Reviewer action'],[[r['record_id'],r['question']] for r in review['review_actions']]),
               _paragraph('The underlying source records, engine output and supporting-document candidates remain available for review. No source assertion has been independently verified.')]
    return _render('Evidex EvidenceOps',packet,sources,blocks,out,render_pdf=render_pdf,
                   details={'executed_providers':['scripts.run_evidex_jobs.run_evidex','products.evidex_customer_review.build_record_review'],
                            'full_evidex_pack_engine_invoked':True,'native_processing_receipt':native,'record_review':review})


def build_vamp_customer_package(packet: Mapping[str, Any], output_dir: str | Path, *, render_pdf: bool = True) -> dict[str, Any]:
    """Build a coverage review from actual customer objective statements."""
    import csv
    from datetime import datetime, timezone
    from products.professional_evidence_customer_surface_compat import _build_vamp_database_v2
    from products.vamp_customer_projection import review_window
    from adapters.vamp.snapshot_pipeline import build_snapshot
    if packet.get('owner_authorized') is not True:
        raise ValueError('explicit source-owner authorization is required')
    consents=packet.get('consents') or {}
    if any(consents.get(k) is not True for k in ['evidence_owner_authorized','performance_data_processing_approved','human_review_terms_accepted']):
        raise ValueError('the native VAMP local processing consents must be explicit')
    sources=_sources(packet)
    rows=[]
    for raw in packet.get('records') or [{'record_id':k,'source_id':k,'quote':v['text']} for k,v in sources.items()]:
        anchor=_quote(raw,sources)
        rows.append({'record_id':str(raw.get('record_id') or ''),'customer_supplied_record':anchor['quote']})
    window=review_window(rows)
    out=_new_output(output_dir)
    source_dir=out/'source_capture/SOURCES';source_dir.mkdir(parents=True)
    with (source_dir/'02_evidence_register.csv').open('w',newline='',encoding='utf-8') as handle:
        writer=csv.DictWriter(handle,fieldnames=['record_id','customer_supplied_record']);writer.writeheader();writer.writerows(rows)
    source_packet={'packet_dir':source_dir.parent,'packet_fingerprint':'sha256:'+_sha(json.dumps(packet,sort_keys=True).encode())}
    database=_build_vamp_database_v2(source_packet,out/'PROJECTION/progress.db',source_rows=rows)
    request={'schema':'dio.vamp_snapshot_request.v1','job_id':'CUSTOMER-'+_sha(json.dumps(packet,sort_keys=True).encode())[:20],
             'profile_path':str(Path(__file__).resolve().parents[1]/'config/vamp_profiles/university_generic_v1.json'),
             'source':{'kind':'vamp_sqlite','database_path':str(database),'staff_id':'PROFESSIONAL-001','year':window['year']},
             'review':{'months':window['months']},'privacy_mode':'private_internal','consents':consents}
    request_path=out/'VAMP_CUSTOMER_REQUEST.json';_json(request_path,request)
    native=build_snapshot(request_path,out/'native_vamp',run_evidex=True)
    snapshot=json.loads((native/'VAMP_SNAPSHOT.json').read_text())
    native_receipt=json.loads((native/'VAMP_SNAPSHOT_RECEIPT.json').read_text())
    if native_receipt.get('status')!='ready_for_human_review':raise RuntimeError('VAMP snapshot failed review preparation')
    objective_label=lambda row: (re.search(r'\bKPA\d+\b',row['title']) or [row['objective_id']])[0]
    blocks=[_heading('Performance evidence coverage'),_paragraph(f"Customer-declared review window: {window['start']} to {window['end']}. {len(snapshot['objectives'])} objectives. These are evidence candidates for a performance conversation; no employee rating or promotion/CPD decision is generated."),
            _table(['Objective','Supplied requirement'],[[objective_label(r),r['title']] for r in snapshot['objectives']]),
            _heading('Coverage to review'),_table(['Objective','Required','Accepted','Candidates','State'],
                   [[objective_label(r),r['minimum_required'] if r.get('minimum_required_basis')=='explicit_source_quantity' else 'To confirm',r['accepted_evidence'],r['candidate_evidence'],_display_state(r['coverage_status'])] for r in snapshot['objectives']]),
            _heading('Documents and decisions to obtain'),_paragraph('Obtain the underlying dated outputs and authorisations. A statement that an article, timetable or planning email exists is not independent inspection of that artifact. Confirm objective mapping and evidence sufficiency with the responsible reviewer.'),
            _paragraph('The supplied notes identify a review period, but do not independently establish dates for the underlying work. Obtain dated source records before treating work as occurring within this period.'),
            _heading('Exact supplied records'),_table(['Record','Source quote'],[[r['record_id'],r['customer_supplied_record']] for r in rows])]
    return _render('VAMP Performance',packet,sources,blocks,out,render_pdf=render_pdf,
                   details={'executed_providers':['adapters.vamp.snapshot_pipeline.build_snapshot'],'native_processing_receipt':native_receipt,
                            'review_window':window,'coverage_metrics':snapshot['metrics'],'rating_generated':False,'evidence_event_dates_verified':False})
