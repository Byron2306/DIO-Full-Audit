"""Behavior checks on real engine code; source and template fixtures are explicit."""
import csv
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parents[1]
ROOT_REL=Path('cross_folder_variants/Evidex/A_CODE')
SOURCE=HERE/'verification-source'
if not SOURCE.exists():SOURCE=HERE

class QualityConsumptionTests(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
  self.root=Path(self.temp.name)/'DIO';self.root.mkdir()
  self.engine=self.root/ROOT_REL
  package=self.engine/'src/evidence_pack_engine';package.mkdir(parents=True)
  for source in (SOURCE/ROOT_REL/'src/evidence_pack_engine').glob('*.py'):shutil.copy2(source,package/source.name)
  if (HERE/'payload').exists():shutil.copytree(HERE/'payload'/ROOT_REL,self.engine,dirs_exist_ok=True)
  (self.root/'scripts').mkdir();shutil.copy2(SOURCE/'scripts/dio_epistemic_spine.py',self.root/'scripts/dio_epistemic_spine.py')
  self.uploads=self.root/'uploads';self.uploads.mkdir()
  self.intake=self.root/'intake.json';self.intake.write_text(json.dumps({'client':{'organization':'SYNTHETIC FIXTURE'},'pack':{'purpose':'Internal review','project_name':'Test','reporting_period':{'start':'2026-01-01','end':'2026-08-31'}},'kpis':[{'name':'Participants trained','target':20,'actual':12,'measurement':'Attendance register'}]}))
  templates=self.engine/'templates/grant_report';templates.mkdir(parents=True,exist_ok=True)
  (templates/'executive_summary.md.j2').write_text('# Synthetic executive\n')
  (templates/'narrative.md.j2').write_text('# Synthetic narrative\n')
  ops_templates=self.engine/'templates/ops';ops_templates.mkdir(parents=True,exist_ok=True)
  for name in ('invoice.md.j2','delivery_email.txt.j2'):
   shutil.copy2(SOURCE/ROOT_REL/'templates/ops'/name,ops_templates/name)
 def run_code(self,body):
  # Only replace unavailable local Jinja rendering at its external boundary.
  prelude="""import sys,os,json,hashlib,types,importlib.util
from pathlib import Path
if __import__('importlib').util.find_spec('jinja2') is None:
 m=types.ModuleType('evidence_pack_engine.render')
 m.render_template=lambda **kw:'# Explicit template-boundary fixture\\n'
 sys.modules['evidence_pack_engine.render']=m
from evidence_pack_engine.pack import build_evidence_rows,load_intake,write_pack
from evidence_pack_engine.ingest import ingest_uploads
from evidence_pack_engine.llm import LlmClient,load_llm_config
root=Path(os.environ['DIO_ROOT']);uploads=root/'uploads';intake=load_intake(root/'intake.json');llm=LlmClient(load_llm_config())
"""
  env={k:v for k,v in os.environ.items() if not k.startswith(('OPENAI_','OLLAMA_','LLM_'))}
  env.update(DIO_ROOT=str(self.root),PYTHONPATH=os.pathsep.join([str(self.engine/'src'),str(self.root)]),LLM_DISABLED='1',SUMMARY_USE_LLM='0',NARRATIVE_USE_LLM='0')
  r=subprocess.run([sys.executable,'-c',prelude+body],cwd=self.root,env=env,text=True,capture_output=True,timeout=25)
  self.assertEqual(r.returncode,0,r.stderr);return json.loads(r.stdout)
 def test_word_overlap_does_not_become_verified_support(self):
  (self.uploads/'register.txt').write_text('Participants trained: planned only. No training was delivered.')
  value=self.run_code("rows,flags=build_evidence_rows(intake,ingest_uploads(uploads),llm)\nprint(json.dumps({'state':getattr(rows[0],'epistemic_state',None),'summary':rows[0].evidence_summary}))")
  self.assertEqual(value['state'],'UNVERIFIED');self.assertIn('No training was delivered.',value['summary']);self.assertNotIn('supporting measurement',value['summary'])
 def test_same_filename_changed_fact_changes_customer_visible_summary(self):
  value=self.run_code("p=uploads/'register.txt';p.write_text('Participants trained: 12')\na=build_evidence_rows(intake,ingest_uploads(uploads),llm)[0][0].evidence_summary\np.write_text('Participants trained: 13')\nb=build_evidence_rows(intake,ingest_uploads(uploads),llm)[0][0].evidence_summary\nprint(json.dumps({'a':a,'b':b}))")
  self.assertNotEqual(value['a'],value['b']);self.assertIn('Source extract: Participants trained: 13',value['b'])
 def test_source_modified_after_ingestion_is_refused(self):
  value=self.run_code("p=uploads/'register.txt';p.write_text('Participants trained: 12');docs=ingest_uploads(uploads);p.write_text('Participants trained: 99')\ntry:\n build_evidence_rows(intake,docs,llm);held=False\nexcept (RuntimeError,ValueError):held=True\nprint(json.dumps({'held':held}))")
  self.assertTrue(value['held'])
 def test_missing_sources_remain_unverified(self):
  value=self.run_code("rows,flags=build_evidence_rows(intake,[],llm)\nprint(json.dumps({'state':getattr(rows[0],'epistemic_state',None),'flags':flags}))")
  self.assertEqual(value['state'],'UNVERIFIED');self.assertTrue(value['flags'])
 def test_actual_pack_contains_review_state_bound_excerpts_and_declared_actual(self):
  (self.uploads/'register.txt').write_text('Participants trained: 12; independent attendance audit not supplied.')
  value=self.run_code("archive,pack,flags=write_pack(intake_path=root/'intake.json',uploads_dir=uploads,out_dir=root/'out',job_name='SYNTHETIC')\nprint(json.dumps({'archive':str(archive),'pack':str(pack)}))")
  pack=Path(value['pack']);self.assertTrue((pack/'SOURCE_REVIEW.json').is_file());review=json.loads((pack/'SOURCE_REVIEW.json').read_text())
  self.assertFalse(review['release_authorized']);self.assertTrue(review['human_review_required']);self.assertEqual(review['claims'][0]['epistemic_state'],'UNVERIFIED')
  binding=review['claims'][0]['source_bindings'][0];self.assertEqual(binding['sha256'],hashlib.sha256((self.uploads/'register.txt').read_bytes()).hexdigest())
  self.assertIn('independent attendance audit not supplied',binding['excerpt'])
  with (pack/'01_EVIDENCE_TABLE/evidence_table.csv').open() as f: row=next(csv.DictReader(f))
  self.assertEqual(row['Declared actual'],'12');self.assertEqual(row['Epistemic state'],'UNVERIFIED')
  with zipfile.ZipFile(value['archive']) as z:self.assertIn('SOURCE_REVIEW.json',z.namelist())
  from docx import Document
  for rel in ['00_EXEC_SUMMARY/executive_summary.docx','02_NARRATIVE/narrative_justification.docx']:
   text='\n'.join(p.text for p in Document(pack/rel).paragraphs);self.assertIn('UNVERIFIED',text);self.assertIn('HUMAN REVIEW REQUIRED',text)
 def test_bytes_changed_during_copy_do_not_produce_a_sealed_archive(self):
  (self.uploads/'register.txt').write_text('Participants trained: 12')
  value=self.run_code("import evidence_pack_engine.pack as module\noriginal=module.shutil.copy2\ndef altered(src,dst):\n if Path(src).name=='register.txt':Path(src).write_text('Participants trained: 99')\n return original(src,dst)\nmodule.shutil.copy2=altered\ntry:\n write_pack(intake_path=root/'intake.json',uploads_dir=uploads,out_dir=root/'out',job_name='SYNTHETIC');held=False\nexcept (RuntimeError,ValueError):held=True\nprint(json.dumps({'held':held,'archives':len(list((root/'out').glob('*.zip')))}))")
  self.assertTrue(value['held']);self.assertEqual(value['archives'],0)
 def test_claim_classifier_is_consumed_and_cannot_promote_lexical_match(self):
  (self.uploads/'register.txt').write_text('Participants trained: 12')
  (self.root/'scripts/dio_epistemic_spine.py').write_text("def epistemic_tokens():return ('SUPPORTED',)\ndef claim_epistemic_state(**kwargs):return 'SUPPORTED'\n")
  value=self.run_code("try:\n build_evidence_rows(intake,ingest_uploads(uploads),llm);held=False\nexcept (RuntimeError,ValueError):held=True\nprint(json.dumps({'held':held}))")
  self.assertTrue(value['held'])

 def _generated_pack(self, body_prefix=""):
  (self.uploads/'register.txt').write_text('Participants trained: 12; independent verification absent.')
  value=self.run_code(body_prefix+"\narchive,pack,flags=write_pack(intake_path=root/'intake.json',uploads_dir=uploads,out_dir=root/'out',job_name='SYNTHETIC')\nprint(json.dumps({'archive':str(archive),'pack':str(pack)}))")
  return Path(value['pack'])

 def test_customer_docx_renders_emphasis_without_markdown_markers(self):
  pack=self._generated_pack()
  for rel in ('00_EXEC_SUMMARY/executive_summary.docx','02_NARRATIVE/narrative_justification.docx','INVOICE.docx'):
   doc=Document(pack/rel)
   text='\n'.join(p.text for p in doc.paragraphs)
   self.assertNotIn('**',text,rel)
  invoice=Document(pack/'INVOICE.docx')
  self.assertTrue(any(run.bold for p in invoice.paragraphs for run in p.runs))

 def test_evidence_workbook_has_readable_column_layout(self):
  pack=self._generated_pack()
  sheet=load_workbook(pack/'01_EVIDENCE_TABLE/evidence_table.xlsx').active
  self.assertEqual(sheet.freeze_panes,'A2')
  self.assertEqual(sheet.auto_filter.ref,sheet.dimensions)
  self.assertGreaterEqual(sheet.column_dimensions['A'].width,24)
  self.assertGreaterEqual(sheet.column_dimensions['F'].width,60)
  self.assertGreaterEqual(sheet.column_dimensions['G'].width,24)
  self.assertTrue(sheet['F2'].alignment.wrap_text)

 def test_missing_payment_configuration_marks_commercial_artifacts_draft(self):
  pack=self._generated_pack("\nfor key in ('PAYMENT_LINK','PAYPAL_LINK','PAYMENT_INSTRUCTIONS','PAYMENT_SHORT','PAYPAL_SHORT'):\n os.environ.pop(key,None)")
  invoice_doc=Document(pack/'INVOICE.docx')
  invoice='\n'.join(p.text for p in invoice_doc.paragraphs)
  email=(pack/'DELIVERY_EMAIL.txt').read_text()
  self.assertIn('DRAFT',invoice)
  self.assertIn('PAYMENT DETAILS REQUIRED',invoice)
  self.assertNotIn('(add payment link)',invoice)
  self.assertIn('DRAFT',email)
  self.assertIn('DO NOT SEND',email)
  self.assertFalse(any(not paragraph.text.strip() for paragraph in invoice_doc.paragraphs))

 def test_payment_instructions_without_url_are_configured(self):
  prefix="""
for key in ('PAYMENT_LINK','PAYPAL_LINK','PAYMENT_SHORT','PAYPAL_SHORT'):
 os.environ.pop(key,None)
os.environ['PAYMENT_INSTRUCTIONS']='Bank transfer after operator review.'
"""
  pack=self._generated_pack(prefix)
  invoice='\n'.join(p.text for p in Document(pack/'INVOICE.docx').paragraphs)
  email=(pack/'DELIVERY_EMAIL.txt').read_text()
  self.assertNotIn('DRAFT',invoice)
  self.assertIn('Bank transfer after operator review.',invoice)
  self.assertNotIn('Payment details are not configured',email)
  self.assertIn('See invoice payment instructions',email)

if __name__=='__main__':unittest.main()
