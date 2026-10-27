"""Read native deliverable content for review rendering, without inferring facts."""
from __future__ import annotations

import csv
import io
import json
import shutil
import subprocess
from html.parser import HTMLParser
from pathlib import Path
from typing import Any


class _HTMLContent(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[dict[str, Any]] = []
        self.text: list[str] = []
        self.hidden = 0
        self.table: list[list[str]] | None = None
        self.row: list[str] = []
        self.cell: list[str] | None = None
        self.heading = False
        self.caption: list[str] | None = None
        self.links: list[str] = []

    def flush(self) -> None:
        value = " ".join(" ".join(self.text).split())
        if value:
            self.blocks.append({"type": "heading" if self.heading else "paragraph", "text": value, **({"level": 1} if self.heading else {})})
        self.text = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"head", "script", "style"}:
            self.hidden += 1
        if self.hidden:
            return
        if tag == "a":
            self.links.append(next((value or "" for key, value in attrs if key == "href"), ""))
        if tag == "table":
            self.flush()
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in {"td", "th"} and self.table is not None:
            self.cell = []
        elif tag == "caption" and self.table is not None:
            self.caption = []
        elif self.table is None and tag in {"p", "div", "li", "br", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.flush()
            self.heading = tag.startswith("h") and len(tag) == 2

    def handle_endtag(self, tag: str) -> None:
        if tag in {"head", "script", "style"}:
            self.hidden = max(0, self.hidden - 1)
            return
        if self.hidden:
            return
        if tag == "a" and self.links:
            link = self.links.pop()
            if link:
                self.handle_data(f" ({link})")
        if tag == "caption" and self.caption is not None:
            value = " ".join(" ".join(self.caption).split())
            if value:
                self.blocks.append({"type": "paragraph", "text": value})
            self.caption = None
        elif tag in {"td", "th"} and self.cell is not None:
            self.row.append(" ".join(" ".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.table is not None:
            if self.row:
                self.table.append(self.row)
            self.row = []
        elif tag == "table" and self.table is not None:
            rows = self.table
            self.table = None
            if rows:
                width = max(map(len, rows))
                rows = [r + [""] * (width - len(r)) for r in rows]
                if len(rows) > 1:
                    self.blocks.append({"type": "table", "headers": rows[0], "rows": rows[1:]})
                else:
                    self.blocks.append({"type": "paragraph", "text": " | ".join(rows[0])})
        elif self.table is None and tag in {"p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.flush()
            self.heading = False

    def handle_data(self, data: str) -> None:
        if self.hidden:
            return
        if self.caption is not None:
            self.caption.append(data)
        elif self.cell is not None:
            self.cell.append(data)
        elif self.table is None:
            self.text.append(data)


def _paragraphs(text: str) -> list[dict[str, Any]]:
    blocks = []
    for part in text.splitlines():
        value = part.strip()
        if value:
            heading = value.startswith("#")
            blocks.append({"type": "heading" if heading else "paragraph", "text": value.lstrip("# ") if heading else value, **({"level": 1} if heading else {})})
    return blocks


def extract_artifact_blocks(path: Path) -> list[dict[str, Any]]:
    """Extract visible text/tables. Receipts and binary media are not domain prose.

    Unsupported types return no blocks and remain in the source inventory. A
    caller must refuse a review pack consisting solely of metadata/binary media.
    """
    suffix = path.suffix.casefold()
    if suffix in {".html", ".htm"}:
        parser = _HTMLContent()
        parser.feed(path.read_text(encoding="utf-8"))
        parser.flush()
        return parser.blocks
    if suffix in {".txt", ".md"}:
        return _paragraphs(path.read_text(encoding="utf-8"))
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"))))
        if len(rows) < 2:
            return _paragraphs("\n".join(", ".join(r) for r in rows))
        width = max(map(len, rows))
        rows = [r + [""] * (width - len(r)) for r in rows]
        return [{"type": "table", "headers": rows[0], "rows": rows[1:]}]
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(value, dict) and value.get("schema") in {"dio.phase8.market_intelligence_pack.v1", "dio.vesper_journey_view.v1"}:
            # These versioned native products deliver structured domain data,
            # not an execution receipt. Preserve their original keys, values
            # and evidence/authority qualifications without inferring prose.
            rows = []
            def flatten(item: Any, prefix: str) -> None:
                if isinstance(item, dict):
                    for key, child in item.items():
                        flatten(child, f"{prefix}.{key}" if prefix else str(key))
                elif isinstance(item, list):
                    for index, child in enumerate(item):
                        flatten(child, f"{prefix}[{index}]")
                    if not item:
                        rows.append([prefix, "[]"])
                else:
                    rows.append([prefix, json.dumps(item, ensure_ascii=False)])
            flatten(value, "")
            return [{"type": "table", "headers": ["Native field", "Recorded value"], "rows": rows}]
        if isinstance(value, dict) and value.get("schema") == "dio.semantic_content.v1":
            from adapters.format_core.renderer import validate_semantic_content
            validation = validate_semantic_content(value)
            if not validation["passed"]:
                raise ValueError("Invalid native semantic content: " + "; ".join(validation["errors"]))
            # Do not carry figure paths into a different source-root context.
            return [dict(b) for b in value["blocks"] if b["type"] not in {"figure", "diagram", "page_break"}]
        return []
    if suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        blocks = []
        for element in document.element.body:
            if element.tag.endswith("}p"):
                blocks.extend(_paragraphs(Paragraph(element, document).text))
            elif element.tag.endswith("}tbl"):
                table = Table(element, document)
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                if len(rows) > 1:
                    blocks.append({"type": "table", "headers": rows[0], "rows": rows[1:]})
                elif rows:
                    blocks.extend(_paragraphs(" | ".join(rows[0])))
        return blocks
    if suffix == ".pptx":
        from pptx import Presentation
        blocks = []
        for slide in Presentation(path).slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    blocks.extend(_paragraphs(shape.text))
                if getattr(shape, "has_table", False):
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    if len(rows) > 1:
                        blocks.append({"type": "table", "headers": rows[0], "rows": rows[1:]})
        return blocks
    if suffix == ".pdf":
        binary = shutil.which("pdftotext")
        if not binary:
            raise ValueError("pdftotext is required to finish native PDF content")
        completed = subprocess.run([binary, "-layout", str(path), "-"], capture_output=True, text=True, timeout=30, check=True)
        return _paragraphs(completed.stdout)
    return []
