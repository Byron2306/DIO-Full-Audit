#!/usr/bin/env python3
from __future__ import annotations

import argparse
import asyncio
import importlib.util
import json
import os
import re
import shutil
import sys
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from build_caps_assessment_ontology import build_profiles, canonical_subject  # noqa: E402
from apply_homs_design_law import (  # noqa: E402
    DEFAULT_ASSESSMENT_DESIGN,
    DEFAULT_DESIGN_LAW,
    apply_design_law,
    assessment_design_for_pack,
    load_json as load_design_json,
    render_docx,
)
from resolve_caps_source import PHASE_BY_GRADE, extract_text, load_manifest, resolve as resolve_caps  # noqa: E402

DEFAULT_SECRET_FILE = Path("/home/byron/EdgeK-BEAST/.beast/provider_secrets.env")
DEFAULT_HYMARK_BACKEND = Path("/home/byron/Downloads/NoEdge-Multi-Hymark-main/backend/server.py")
DEFAULT_SUBJECT_PROFILE_DIR = ROOT / "config" / "homs_subject_profiles"
DEFAULT_GRADE_LADDER = ROOT / "config" / "homs_grade_ladder.json"
DEFAULT_CAPS_MANIFEST = ROOT / "corpora" / "caps" / "caps_corpus_manifest.json"
DEFAULT_CAPS_MATRIX = ROOT / "deliverables" / "caps_matrix_analysis" / "caps_matrix_analysis.json"
DEFAULT_CAPS_ONTOLOGY = ROOT / "deliverables" / "caps_assessment_ontology" / "caps_assessment_ontology.json"
DEFAULT_CAPS_ASSESSMENT_DESIGN = DEFAULT_ASSESSMENT_DESIGN
DEFAULT_EXAM_SOURCE_CONTRACTS = ROOT / "deliverables" / "exam_source_matrix" / "exam_source_contracts.json"
DEFAULT_EXAM_SOURCE_CATALOGUE = ROOT / "deliverables" / "exam_source_catalogue_fet_2024_november" / "FET_EXAM_SOURCE_CATALOGUE.json"
DEFAULT_EVIDENCE_GRAMMAR = ROOT / "config" / "homs_source_construction_grammar.json"
DEFAULT_NIM_MODEL = "deepseek-ai/deepseek-v4-flash-0731"

AFRIKAANS_TERM_FOCI = {
    "1": {
        "focus": "begripstoets, taal in konteks en teksgebaseerde interpretasie",
        "secondary": "toon, teikengroep, doel, afleiding en bewys uit 'n gegewe teks",
        "methodology": "opsomming en noukeurige taalredigering in konteks",
    },
    "2": {
        "focus": "literere en nieliterere teksbegrip, karakter, tema en taalontleding",
        "secondary": "standpunt, register, figuurlike taal en teksbewys",
        "methodology": "aanhalingskeuse, bewysinvoeging en paragraafstruktuur",
    },
    "3": {
        "focus": "geintegreerde begrip, opsomming en taalstrukture in konteks",
        "secondary": "toon, doel, afleiding, woordkeuse en kritiese taalbewustheid",
        "methodology": "redigering vir grammatika, woordkeuse, register en leestekens",
    },
    "4": {
        "focus": "finale geintegreerde taaleksamenoefening met onvoorbereide tekste",
        "secondary": "sintese, toon, doel, standpunt en bewysgebaseerde interpretasie",
        "methodology": "tydbestuurde opsomming, redigering en presiese teksbewys",
    },
}

ENGLISH_TERM_FOCI = {
    "1": {
        "focus": "comprehension, language in context and text-based interpretation",
        "secondary": "tone, audience, purpose, inference and textual evidence",
        "methodology": "summary writing and precise language editing in context",
    },
    "2": {
        "focus": "literary and non-literary comprehension, viewpoint and language analysis",
        "secondary": "register, figurative language, argument and textual support",
        "methodology": "selecting evidence and structuring concise analytical responses",
    },
    "3": {
        "focus": "integrated comprehension, summary and language structures in context",
        "secondary": "tone, purpose, inference, word choice and critical language awareness",
        "methodology": "editing for grammar, word choice, register and punctuation",
    },
    "4": {
        "focus": "integrated language examination practice using unseen texts",
        "secondary": "synthesis, tone, purpose, viewpoint and evidence-based interpretation",
        "methodology": "timed summary, editing and precise textual evidence",
    },
}


DEFAULT_REQUEST = {
    "module_code": "HISE411",
    "module_name": "HISTORY SNR & FET 4A",
    "topics": [
        "The Cuban Missile Crisis (1962)",
        "The Division of Germany (1945-1949)",
    ],
    "methodology_topic": "The Berlin Airlift",
    "essay_topic": "The role of media in shaping public opinion during the Vietnam War",
    "total_marks": 125,
    "duration_hours": 3,
    "additional_instructions": "Controlled demo generation. Human subject expert must verify all historical sources before classroom use.",
}


def default_request_for_subject(subject_profile: dict[str, Any], grade_profile: dict[str, Any]) -> dict[str, Any]:
    if str(subject_profile.get("subject_id") or "").lower() == "history":
        return dict(DEFAULT_REQUEST)
    subject = subject_profile["display_name"]
    grade = grade_profile["grade"]
    return {
        "module_code": re.sub(r"[^A-Z0-9]+", "", subject.upper())[:6] or "CAPS",
        "module_name": f"{subject} Grade {grade}",
        "topics": [
            f"{subject} CAPS core knowledge and skills",
            f"{subject} applied task or assessment evidence",
        ],
        "methodology_topic": f"{subject} applied assessment task",
        "essay_topic": f"{subject} extended or practical synthesis task",
        "total_marks": 50,
        "duration_hours": 2,
        "additional_instructions": "Controlled CAPS-aligned demo generation. A qualified educator must verify topic scope, mark allocation, and assessment form before classroom use.",
    }


def apply_subject_term_focus(request: dict[str, Any]) -> None:
    subject_id = subject_id_from_request(request)
    term = str(request.get("term") or "auto")
    term_maps = {
        "afrikaans_language": AFRIKAANS_TERM_FOCI,
        "english_language": ENGLISH_TERM_FOCI,
    }
    term_map = term_maps.get(subject_id)
    if not term_map or term not in term_map:
        return
    focus = term_map[term]
    request["topics"] = [focus["focus"], focus["secondary"]]
    request["methodology_topic"] = focus["methodology"]
    request["essay_topic"] = ""
    request["assessment_scope"] = {
        "term": term,
        "language_variant": (
            "Afrikaans Eerste Addisionele Taal"
            if subject_id == "afrikaans_language"
            else "English First Additional Language"
        ),
        "paper_form": (
            "teksgebaseerde gekontroleerde toets"
            if subject_id == "afrikaans_language"
            else "text-based controlled test"
        ),
        "focus": focus,
    }


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def parse_env_file(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values
    for raw in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.replace("export ", "").strip()] = value.strip().strip('"').strip("'")
    return values


def configure_provider(secret_file: Path, provider_name: str, model: str) -> dict[str, Any]:
    values = parse_env_file(secret_file)
    for key, value in values.items():
        os.environ.setdefault(key, value)

    if provider_name in {"nim", "nvidia", "nvidia_nim"}:
        key_name = "NVIDIA_API_KEY"
        api_key = values.get(key_name) or os.environ.get(key_name)
        if not api_key:
            raise RuntimeError("NVIDIA_API_KEY was not found in the provider secret file or environment.")
        os.environ["OPENAI_API_KEY"] = api_key
        os.environ["OPENAI_BASE_URL"] = os.environ.get("NVIDIA_BASE_URL") or "https://integrate.api.nvidia.com/v1"
        os.environ["HOMS_AI_MODEL"] = model or os.environ.get("HOMS_NIM_MODEL") or DEFAULT_NIM_MODEL
        selected = "nvidia_nim"
    else:
        key_name = "OPENAI_API_KEY"
        api_key = values.get(key_name) or os.environ.get(key_name)
        if not api_key:
            raise RuntimeError("OPENAI_API_KEY was not found in the provider secret file or environment.")
        os.environ["OPENAI_API_KEY"] = api_key
        os.environ.pop("OPENAI_BASE_URL", None)
        os.environ["HOMS_AI_MODEL"] = model or os.environ.get("OPENAI_MODEL") or "gpt-4o"
        selected = "openai"

    return {
        "selected_provider": selected,
        "selected_model": os.environ["HOMS_AI_MODEL"],
        "selected_key_name": key_name,
        "base_url": os.environ.get("OPENAI_BASE_URL") or None,
        "available_secret_names": sorted(k for k, v in values.items() if v and ("KEY" in k or "TOKEN" in k or "SECRET" in k)),
        "values_redacted": True,
    }


def load_hymark_backend(path: Path) -> Any:
    spec = importlib.util.spec_from_file_location("hymark_backend_server", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load HyMark backend module from {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def load_subject_profile(profile_name: str, profile_path: Path | None = None) -> dict[str, Any]:
    if profile_path:
        path = profile_path
    else:
        safe_name = profile_name.strip().lower().replace(" ", "_") or "history"
        path = DEFAULT_SUBJECT_PROFILE_DIR / f"{safe_name}.json"
    if not path.exists():
        return synthesize_subject_profile(profile_name)
    profile = json.loads(path.read_text(encoding="utf-8"))
    if profile.get("schema") != "knowedge.homs_subject_profile.v1":
        raise ValueError(f"Unsupported subject profile schema in {path}")
    required = [
        "subject_id",
        "display_name",
        "assessment_modes",
        "source_types",
        "question_families",
        "cognitive_targets",
        "rubric_dimensions",
        "memo_expectations",
        "human_review_checks",
    ]
    missing = [field for field in required if not profile.get(field)]
    if missing:
        raise ValueError(f"Subject profile {path} is missing: {', '.join(missing)}")
    profile["_profile_path"] = str(path)
    return profile


def title_from_id(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("_", " ")).strip().title()


def synthesize_subject_profile(profile_name: str) -> dict[str, Any]:
    subject_id = re.sub(r"[^a-z0-9_]+", "_", profile_name.strip().lower().replace(" ", "_")).strip("_") or "general_subject"
    display_name = title_from_id(subject_id)
    return {
        "schema": "knowedge.homs_subject_profile.v1",
        "subject_id": subject_id,
        "display_name": display_name,
        "default_level": "caps_school",
        "assessment_modes": ["caps_shell_assessment", "structured_task", "memo", "marking_relief"],
        "source_types": ["caps_stimulus", "diagram", "table", "scenario", "task_brief"],
        "question_families": ["recall", "apply", "explain", "investigate", "design_or_perform", "justify"],
        "cognitive_targets": ["remember", "understand", "apply", "analyze", "evaluate", "create"],
        "rubric_dimensions": ["content_accuracy", "skill_demonstration", "reasoning", "communication", "caps_alignment"],
        "memo_expectations": [
            "reconcile total marks with section marks and rubric marks",
            "credit grade-appropriate equivalent responses",
            "separate learner-facing instructions from educator-facing memo",
            "tie every mark to observable evidence or answer points",
        ],
        "retrieval_queries": [f"{display_name} CAPS assessment requirements"],
        "human_review_checks": [
            "verify CAPS topic scope",
            "check mark allocation",
            "confirm assessment form matches subject and phase",
            "approve final pack before classroom use",
        ],
        "blocked_claims_or_modes": [
            "unverified official quotation",
            "internal uncertainty in learner-facing pack",
            "unreconciled marks",
        ],
        "_profile_path": "synthesized_from_subject_name",
    }


def load_grade_profile(grade: int, ladder_path: Path = DEFAULT_GRADE_LADDER) -> dict[str, Any]:
    if grade < 1 or grade > 12:
        raise ValueError("Grade must be between 1 and 12.")
    if not ladder_path.exists():
        raise FileNotFoundError(f"Grade ladder not found: {ladder_path}")
    ladder = json.loads(ladder_path.read_text(encoding="utf-8"))
    if ladder.get("schema") != "knowedge.homs_grade_ladder.v1":
        raise ValueError(f"Unsupported grade ladder schema in {ladder_path}")
    for row in ladder.get("grades", []):
        if int(row.get("grade") or 0) == grade:
            profile = dict(row)
            profile["_ladder_path"] = str(ladder_path)
            return profile
    raise ValueError(f"Grade {grade} was not found in {ladder_path}")


def normalized_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def load_caps_ontology(ontology_path: Path | None, matrix: dict[str, Any] | None) -> list[dict[str, Any]]:
    if ontology_path and ontology_path.exists():
        return json.loads(ontology_path.read_text(encoding="utf-8")).get("profiles", [])
    if matrix:
        return build_profiles(matrix.get("rows", []))
    return []


def load_caps_assessment_design(design_path: Path | None) -> dict[str, Any]:
    if design_path and design_path.exists():
        return json.loads(design_path.read_text(encoding="utf-8"))
    return {}


def find_assessment_design_profile(design_payload: dict[str, Any], ontology_profile: dict[str, Any] | None) -> dict[str, Any] | None:
    if not ontology_profile:
        return None
    profile_id = ontology_profile.get("profile_id")
    profiles_by_id = design_payload.get("profiles_by_id") or {}
    if profile_id and profile_id in profiles_by_id:
        return profiles_by_id[profile_id]
    for profile in design_payload.get("profiles") or []:
        if profile.get("profile_id") == profile_id:
            return profile
    return None


def find_caps_ontology_profile(
    profiles: list[dict[str, Any]],
    matrix_row: dict[str, Any] | None,
    subject_profile: dict[str, Any],
    grade_profile: dict[str, Any],
) -> dict[str, Any] | None:
    phase = PHASE_BY_GRADE[int(grade_profile["grade"])]
    candidate_names = [
        str(subject_profile.get("display_name") or ""),
        str(subject_profile.get("subject_id") or "").replace("_", " "),
    ]
    if matrix_row:
        subject, _family, _language = canonical_subject(matrix_row)
        candidate_names.insert(0, subject)
    wanted = {normalized_name(name) for name in candidate_names if name}
    for profile in profiles:
        if profile.get("phase") == phase and normalized_name(profile.get("subject", "")) in wanted:
            return profile
    for profile in profiles:
        if profile.get("phase") == phase and any(name and name in normalized_name(profile.get("subject", "")) for name in wanted):
            return profile
    return None


def select_caps_excerpt(text: str, subject_profile: dict[str, Any], grade_profile: dict[str, Any], max_chars: int) -> str:
    if not text.strip() or max_chars <= 0:
        return ""
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]+", " ", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    search_terms = [
        "assessment",
        "formal assessment",
        "curriculum",
        "content",
        "cognitive",
        "grade " + str(grade_profile["grade"]),
        "examination",
        subject_profile.get("display_name", ""),
    ]
    lowered = text.lower()
    for term in search_terms:
        term = str(term or "").lower().strip()
        if not term:
            continue
        index = lowered.find(term)
        if index >= 0:
            start = max(0, index - 900)
            excerpt = text[start : start + max_chars].strip()
            return excerpt
    return text[:max_chars].strip()


def compact_lines(value: str, limit: int = 12) -> list[str]:
    lines: list[str] = []
    for raw in re.split(r"[\r\n]+|•|--|;", str(value or "")):
        line = re.sub(r"\s+", " ", raw).strip(" -\t")
        if len(line) < 8:
            continue
        if "...." in line or re.search(r"\.{4,}\s*\d+$", line):
            continue
        if re.fullmatch(r"[\d\s.]+", line):
            continue
        if line not in lines:
            lines.append(line[:220])
        if len(lines) >= limit:
            break
    return lines


def extract_term_theme_lines(text: str, grade: int, term: str | int | None, max_lines: int = 14) -> list[str]:
    if not text.strip():
        return []
    cleaned = re.sub(r"[ \t]{2,}", " ", text)
    lower = cleaned.lower()
    term_value = str(term or "").strip()
    if term_value and term_value.lower() not in {"auto", "0", "none"}:
        explicit_matches: list[tuple[int, str]] = []
        pattern = re.compile(rf"(?im)^\s*term\s+{re.escape(term_value)}\s*$")
        for match in pattern.finditer(cleaned):
            window = cleaned[match.start() : match.start() + 4200]
            if re.search(rf"(?i)term\s+{re.escape(term_value)}\s+to\s+term", window[:80]):
                continue
            if "topic:" in window.lower():
                before = cleaned[max(0, match.start() - 1600) : match.start()].lower()
                grade_score = 0
                if f"grade {grade}" in before or f"graad {grade}" in before:
                    grade_score = 2
                elif "grade " not in before[-500:] and "graad " not in before[-500:]:
                    grade_score = 1
                explicit_matches.append((grade_score, window))
        if explicit_matches:
            best_score = max(score for score, _window in explicit_matches)
            if best_score > 0:
                explicit_matches = [item for item in explicit_matches if item[0] == best_score]
            lines = []
            for _score, window in explicit_matches:
                for line in compact_lines(window, max_lines * 2):
                    line_lower = line.lower()
                    if line_lower.startswith("term ") or line_lower in {"topic:", "time:", "additional resources:"}:
                        continue
                    if any(token in line_lower for token in ["topic", "map", "graph", "table", "source", "investigation", "project", "practical", "atmosphere", "settlement", "climate", "geomorphology", "population", "water", "gis", "structure", "earth", "plate", "weather"]):
                        if line not in lines:
                            lines.append(line)
                    if len(lines) >= max_lines:
                        return lines
            if lines:
                return lines[:max_lines]
    grade_terms = [f"grade {grade}", f"grade {grade} term", f"graad {grade}"]
    term_terms = [f"term {term_value}", f"kwartaal {term_value}"] if term_value and term_value.lower() not in {"auto", "0", "none"} else []
    search_terms = [
        f"overview of geography content" if "geography" in lower else "",
        f"grade {grade} fet band",
        *grade_terms,
        *term_terms,
        "annual teaching plans",
        "content",
        "topic",
        "assessment",
        "programme of assessment",
    ]
    windows: list[str] = []
    for needle in search_terms:
        if not needle:
            continue
        start = lower.find(needle.lower())
        if start >= 0:
            windows.append(cleaned[max(0, start - 700) : start + 1800])
    if not windows:
        windows.append(cleaned[:2600])
    lines = []
    for window in windows:
        for line in compact_lines(window, max_lines):
            line_lower = line.lower()
            if any(token in line_lower for token in ["term", "grade", "topic", "content", "assessment", "map", "graph", "table", "source", "investigation", "project", "practical", "atmosphere", "settlement", "climate", "geomorphology", "gis"]):
                if line not in lines:
                    lines.append(line)
            if len(lines) >= max_lines:
                return lines
    return lines[:max_lines]


def load_exam_source_contracts() -> dict[str, Any]:
    cache = getattr(load_exam_source_contracts, "_cache", None)
    if cache is not None:
        return cache
    if not DEFAULT_EXAM_SOURCE_CONTRACTS.exists():
        setattr(load_exam_source_contracts, "_cache", {})
        return {}
    try:
        data = json.loads(DEFAULT_EXAM_SOURCE_CONTRACTS.read_text(encoding="utf-8"))
    except Exception:
        data = {}
    setattr(load_exam_source_contracts, "_cache", data)
    return data


def load_exam_source_catalogue() -> dict[str, Any]:
    cache = getattr(load_exam_source_catalogue, "_cache", None)
    if cache is not None:
        return cache
    if not DEFAULT_EXAM_SOURCE_CATALOGUE.exists():
        setattr(load_exam_source_catalogue, "_cache", {})
        return {}
    try:
        data = json.loads(DEFAULT_EXAM_SOURCE_CATALOGUE.read_text(encoding="utf-8"))
    except Exception:
        data = {}
    setattr(load_exam_source_catalogue, "_cache", data)
    return data


def map_exam_source_type(source_type: str) -> str:
    return {
        "topographic_map_and_orthophoto": "map_extract",
        "synoptic_weather_map": "map_extract",
        "map_extract": "map_extract",
        "graph_or_chart": "data_table",
        "data_table": "data_table",
        "text_extract": "source_extract",
        "photograph_or_image": "visual_source",
        "diagram_or_model": "diagram",
        "cartoon": "visual_source",
        "image_rendered_or_scanned_pages": "diagram",
    }.get(source_type, "structured_prompt")


def exam_source_contract_for_subject(subject: str, theme_blob: str) -> dict[str, Any]:
    data = load_exam_source_contracts()
    if not data:
        return {}
    subject_l = subject.lower()
    contract_subject = str(data.get("subject") or "").lower()
    if contract_subject and contract_subject not in subject_l and subject_l not in contract_subject:
        return {}
    contracts = data.get("contracts") or {}
    theme_l = theme_blob.lower()
    p2_terms = ["settlement", "urban", "rural", "economic geography", "industry", "trade", "agriculture", "population"]
    p1_terms = ["climate", "weather", "synoptic", "cyclone", "geomorphology", "river", "slope", "atmosphere"]
    if any(token in theme_l for token in p2_terms):
        preferred = "P2:settlement_economic_geography_and_mapwork"
    elif any(token in theme_l for token in p1_terms):
        preferred = "P1:climate_weather_geomorphology_and_mapwork"
    else:
        preferred = ""
    if preferred and preferred in contracts:
        return contracts[preferred]
    if contracts:
        return max(contracts.values(), key=lambda item: int(item.get("paper_count") or 0))
    return {}


def exam_catalogue_contract_for_subject(subject: str) -> dict[str, Any]:
    data = load_exam_source_catalogue()
    subjects = data.get("subjects") or {}
    subject_l = subject.lower()
    for name, row in subjects.items():
        name_l = str(name).lower()
        if name_l == subject_l or name_l in subject_l or subject_l in name_l:
            return {**row, "subject": name}
    return {}


def source_type_contract(subject: str, family: str, theme_blob: str, matrix: dict[str, Any] | None = None) -> dict[str, Any]:
    matrix = matrix or {}
    subject_l = subject.lower()
    family_l = family.lower()
    theme_l = theme_blob.lower()
    blob = f"{subject_l} {family_l} {theme_l}"
    source_types: list[str] = []
    reasons: list[str] = []
    if "geography" in subject_l:
        allowed = {"map_extract", "data_table", "source_extract", "evidence_table", "visual_source", "diagram"}
    elif "history" in subject_l:
        allowed = {"source_extract", "evidence_table", "visual_source", "map_extract", "data_table"}
    elif "mathematics" in subject_l or "calculation" in family_l:
        allowed = {"calculation_workspace", "data_table", "diagram"}
    elif any(token in subject_l for token in ["physical sciences", "life sciences", "natural sciences", "agricultural sciences", "technical sciences", "marine sciences"]):
        allowed = {"data_table", "diagram", "investigation_sheet", "calculation_workspace"}
    elif any(token in subject_l for token in ["business studies", "economics", "accounting", "consumer studies", "tourism", "hospitality"]):
        allowed = {"case_file", "data_table", "calculation_workspace", "source_extract", "evidence_table"}
    elif "language" in family_l or "language" in subject_l:
        allowed = {"source_extract", "visual_source", "language_text", "evidence_table"}
    elif any(token in subject_l for token in ["dance", "dramatic", "visual arts", "music", "design studies"]):
        allowed = {"performance_observation", "portfolio_evidence", "visual_source", "source_extract", "evidence_table"}
    elif any(token in subject_l for token in ["technology", "engineering graphics", "coding", "robotics"]):
        allowed = {"design_brief", "diagram", "data_table", "investigation_sheet", "calculation_workspace"}
    else:
        allowed = {"structured_prompt", "source_extract", "evidence_table", "data_table"}

    def add(source_type: str, reason: str) -> None:
        if source_type not in allowed:
            return
        if source_type not in source_types:
            source_types.append(source_type)
            reasons.append(reason)

    old_paper_contract = exam_source_contract_for_subject(subject, theme_blob)
    old_paper_catalogue = exam_catalogue_contract_for_subject(subject)
    for item in old_paper_contract.get("required_repeated_source_types", []):
        mapped = map_exam_source_type(str(item.get("source_type") or ""))
        add(
            mapped,
            (
                f"Official old-paper contract: {item.get('source_type')} recurs in "
                f"{old_paper_contract.get('paper')} {old_paper_contract.get('paper_theme')} "
                f"({item.get('hits')} object hit(s) across {old_paper_contract.get('paper_count')} paper(s))."
            ),
        )
    for source_type in old_paper_catalogue.get("observed_contract", [])[:8]:
        mapped = map_exam_source_type(str(source_type or ""))
        add(
            mapped,
            (
                f"FET old-paper catalogue: {source_type} observed in "
                f"{old_paper_catalogue.get('subject')} "
                f"({old_paper_catalogue.get('unique_source_object_count')} source/stimulus object(s) across "
                f"{old_paper_catalogue.get('paper_count')} paper(s))."
            ),
        )

    if (matrix.get("signals_source_based") or "source_based" in family_l or "history" in subject_l) and "geography" not in subject_l:
        add("source_extract", "CAPS/source matrix or subject profile supports source-based assessment.")
        add("evidence_table", "Learners must cite, compare, or reason from source evidence.")
    if matrix.get("signals_data_graph_diagram"):
        add("data_table", "CAPS/source matrix signals data, graph, or diagram assessment.")
        add("diagram", "CAPS/source matrix signals data, graph, or diagram assessment.")
    if matrix.get("signals_practical_investigation"):
        add("investigation_sheet", "CAPS/source matrix signals practical investigation or experiment evidence.")
    if matrix.get("signals_case_study"):
        add("case_file", "CAPS/source matrix signals case/scenario assessment.")
    if matrix.get("signals_calculation_problem"):
        add("calculation_workspace", "CAPS/source matrix signals calculation/problem-solving evidence.")
    if matrix.get("signals_oral_or_performance"):
        add("performance_observation", "CAPS/source matrix signals oral/performance evidence.")
    if matrix.get("signals_project_or_sba"):
        add("design_brief", "CAPS/source matrix signals project or SBA evidence.")
        add("portfolio_evidence", "CAPS/source matrix signals project or portfolio evidence.")

    if "geography" in blob:
        geo_map_terms = ["map", "settlement", "spatial", "river", "urban", "rural", "climate", "weather", "synoptic", "gis", "isobar", "plate", "tectonic"]
        geo_data_terms = ["population", "rainfall", "temperature", "economic", "data", "table", "graph", "distribution", "density"]
        if any(token in blob for token in geo_map_terms):
            add("map_extract", "Geography term/theme content requires spatial evidence.")
        if any(token in blob for token in geo_data_terms):
            add("data_table", "Geography term/theme content requires numeric, tabular, or graph evidence.")
    if "history" in subject_l:
        add("source_extract", "History assessment requires historical source material.")
        add("evidence_table", "History questions need provenance and claim-evidence-reasoning support.")
    if any(token in blob for token in ["photograph", "cartoon", "advertisement", "poster", "visual source"]):
        add("visual_source", "Term/theme content references visual source material.")
    if any(token in blob for token in ["case study", "case studies", "scenario", "stakeholder"]):
        add("case_file", "Term/theme content references a scenario or case study.")
    if any(token in blob for token in ["experiment", "investigation", "practical", "variables", "observations"]):
        add("investigation_sheet", "Term/theme content references practical investigation evidence.")
    if any(token in blob for token in ["calculate", "calculation", "formula", "equation", "function", "geometry", "measurement"]):
        add("calculation_workspace", "Term/theme content requires calculation or represented working.")
    if any(token in blob for token in ["performance", "dance", "dramatic", "music", "portfolio"]):
        add("performance_observation", "Term/theme content requires observable performance or portfolio evidence.")
        add("portfolio_evidence", "Term/theme content requires portfolio evidence.")

    if not source_types:
        source_types.append("structured_prompt")
        reasons.append("No specific source type resolved; use structured prompt only after educator approval.")

    return {
        "schema": "knowedge.homs_source_type_contract.v1",
        "source_types": source_types,
        "source_type_reasons": reasons,
        "matrix_signals_used": {
            key: matrix.get(key)
            for key in [
                "signals_source_based",
                "signals_essay_or_extended_response",
                "signals_case_study",
                "signals_data_graph_diagram",
                "signals_practical_investigation",
                "signals_project_or_sba",
                "signals_oral_or_performance",
                "signals_calculation_problem",
                "signals_formal_test_or_exam",
            ]
            if key in matrix
        },
        "old_paper_contract_used": {
            "path": str(DEFAULT_EXAM_SOURCE_CONTRACTS) if old_paper_contract else None,
            "paper": old_paper_contract.get("paper"),
            "paper_theme": old_paper_contract.get("paper_theme"),
            "paper_count": old_paper_contract.get("paper_count"),
        } if old_paper_contract else None,
        "exam_source_catalogue_used": {
            "path": str(DEFAULT_EXAM_SOURCE_CATALOGUE) if old_paper_catalogue else None,
            "subject": old_paper_catalogue.get("subject"),
            "paper_count": old_paper_catalogue.get("paper_count"),
            "observed_contract": old_paper_catalogue.get("observed_contract"),
        } if old_paper_catalogue else None,
    }


def visual_kind_rules(subject: str, family: str, theme_blob: str, source_contract: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    subject_l = subject.lower()
    family_l = family.lower()
    theme_l = theme_blob.lower()
    blob = f"{subject_l} {family_l} {theme_l}"
    source_types = set((source_contract or {}).get("source_types") or [])
    visuals: list[dict[str, Any]] = []

    def add(visual_id: str, title: str, role: str, kind: str, must_show: list[str], use: str) -> None:
        if any(item["id"] == visual_id for item in visuals):
            return
        visuals.append(
            {
                "id": visual_id,
                "title": title,
                "role": role,
                "visual_kind": kind,
                "must_show": must_show,
                "assessment_use": use,
            }
        )

    if "geography" in blob:
        if "map_extract" in source_types or any(token in blob for token in ["map", "settlement", "spatial", "river", "urban", "rural", "climate", "weather"]):
            add(
                "geo_map_extract",
                "Map extract with legend, scale, north arrow and labelled features",
                "stimulus",
                "map_extract",
                ["map frame", "legend", "scale", "north arrow", "labelled physical or human features"],
                "Learners identify spatial patterns and cite map evidence.",
            )
        if "data_table" in source_types or any(token in blob for token in ["population", "rainfall", "temperature", "climate", "economic", "data", "table", "graph"]):
            add(
                "geo_data_table_or_graph",
                "Geography data table or graph",
                "stimulus",
                "data_table",
                ["units", "years/categories", "clear trend", "source label"],
                "Learners compare data with the map/source before explaining.",
            )
        if "diagram" in source_types:
            add(
                "geo_process_diagram",
                "Geography process diagram",
                "stimulus",
                "axis_diagram",
                ["clear process labels", "direction arrows where needed", "feature labels", "caption"],
                "Learners interpret a geographic process, structure or relationship before explaining.",
            )
        if "visual_source" in source_types:
            add(
                "geo_photo_or_infographic_source",
                "Geography photograph or infographic source",
                "stimulus",
                "source_extract",
                ["caption", "visible evidence cues", "source note", "labelled features"],
                "Learners cite visible evidence from the image/source in structured answers.",
            )
    if "history" in subject_l or "source_extract" in source_types or ("source_based" in family_l and "geography" not in subject_l):
        add(
            "source_provenance_panel",
            "Source provenance panel",
            "stimulus",
            "source_extract",
            ["author/origin", "date/context", "caption", "extract or visual description"],
            "Learners evaluate usefulness, reliability and perspective.",
        )
        add(
            "claim_evidence_reasoning_table",
            "Claim-evidence-reasoning table",
            "learner_workspace",
            "evidence_table",
            ["claim column", "evidence column", "reasoning/inference column"],
            "Learners connect source evidence to answers.",
        )
    if "mathematics" in subject_l or "calculation_workspace" in source_types or "calculation" in family_l:
        add(
            "working_grid",
            "Working grid with method mark space",
            "learner_workspace",
            "working_grid",
            ["question number", "marks", "working area", "final answer cue"],
            "Learners show method before final answer.",
        )
        if any(token in blob for token in ["graph", "data", "function", "pattern", "geometry", "measurement"]):
            add(
                "axis_or_diagram",
                "Axis, diagram or table required by the topic",
                "stimulus_or_workspace",
                "axis_diagram",
                ["labelled axes or diagram labels", "units where needed", "plot/table space"],
                "Learners interpret or construct visual mathematical evidence.",
            )
    science_subject = any(token in subject_l for token in ["life sciences", "natural sciences", "physical sciences", "agricultural sciences"])
    if science_subject or "investigation_sheet" in source_types or "data_diagram_practical_investigation" in family_l:
        if "diagram" in source_types:
            add(
                "science_diagram",
                "Science labelled diagram stimulus",
                "stimulus",
                "axis_diagram",
                ["clear structure/process labels", "visible arrows or stages where needed", "caption", "label key"],
                "Learners identify, label, compare, and explain using the diagram.",
            )
        if "data_table" in source_types:
            add(
                "science_data_or_graph",
                "Science data table or graph stimulus",
                "stimulus",
                "data_table",
                ["variables", "units", "readable values", "trend or comparison"],
                "Learners calculate, compare, interpret trends, and support conclusions from data.",
            )
        add(
            "investigation_method_data",
            "Investigation method and data sheet",
            "stimulus_and_workspace",
            "investigation_sheet",
            ["aim", "variables", "method", "data table", "conclusion line"],
            "Learners record observations and draw a supported conclusion.",
        )
    commerce_subject = any(token in subject_l for token in ["business studies", "economics", "accounting"])
    if "case_study" in family_l or "case_file" in source_types or commerce_subject:
        add(
            "case_decision_table",
            "Case facts and decision table",
            "stimulus_and_workspace",
            "case_table",
            ["scenario facts", "stakeholder", "problem", "evidence", "recommendation"],
            "Learners justify decisions from case evidence.",
        )
    if "language" in family_l or "language" in subject_l:
        add(
            "language_response_planner",
            "Text response and language planner",
            "learner_workspace",
            "language_planner",
            ["text panel", "main idea", "evidence", "plan", "draft", "edit"],
            "Learners plan reading/viewing and writing responses.",
        )
    project_subject = any(token in subject_l for token in ["coding", "robotics", "technology", "engineering graphics", "design studies"])
    if "practical_project_design_task" in family_l or project_subject:
        add(
            "design_test_log",
            "Design, build, test and improve log",
            "learner_workspace",
            "design_log",
            ["problem", "constraints", "diagram", "test result", "improvement"],
            "Learners retain process evidence for marking.",
        )
    performance_subject = any(token in subject_l for token in ["dance", "dramatic", "visual arts", "music"])
    if "practical_performance_or_portfolio" in family_l or "performance_observation" in source_types or performance_subject:
        add(
            "performance_space_map",
            "Performance space or portfolio evidence map",
            "stimulus_and_workspace",
            "performance_map",
            ["space/stage orientation", "sequence/process labels", "observable evidence fields"],
            "Learners plan performance evidence; educators retain it for moderation.",
        )
        add(
            "observation_instrument",
            "Teacher observation instrument",
            "marking_evidence",
            "rubric_observation_table",
            ["criterion", "marks", "evidence seen", "concern", "score"],
            "Educator records live observable evidence against the rubric.",
        )
    if not visuals:
        add(
            "structured_response_workspace",
            "Structured response workspace",
            "learner_workspace",
            "structured_table",
            ["question", "evidence/working", "answer"],
            "Learners make reasoning visible.",
        )
    return visuals


def content_alignment_notes(focus_items: list[str], term_lines: list[str]) -> list[str]:
    focus_blob = " ".join(focus_items).lower()
    term_blob = " ".join(term_lines).lower()
    stop = {
        "grade",
        "using",
        "assessment",
        "explain",
        "change",
        "source",
        "data",
        "evidence",
        "patterns",
    }
    tokens = {
        token
        for token in re.findall(r"[a-z][a-z-]{4,}", focus_blob)
        if token not in stop
    }
    if not tokens or not term_blob:
        return ["CAPS term/content alignment requires educator confirmation."]
    matched = sorted(token for token in tokens if token in term_blob)
    if matched:
        return [f"Requested focus shares CAPS term/theme signals: {', '.join(matched[:6])}."]
    return [
        "Requested focus was not clearly found in extracted CAPS term/theme lines. Confirm term fit before generation.",
        "Use this as a visual planning warning, not an automatic blocker; PDF extraction can miss table structure.",
    ]


def build_visual_blueprint(request: dict[str, Any]) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    ontology = caps.get("assessment_ontology_profile") or {}
    design = caps.get("assessment_design_profile") or {}
    selected = caps.get("selected") or {}
    term = request.get("term") or "auto"
    raw_text = str(caps.get("full_text") or caps.get("extracted_excerpt") or "")
    source_path = selected.get("path")
    if source_path and not caps.get("full_text"):
        try:
            raw_text = extract_text(Path(source_path), 120000)
        except Exception:
            pass
    term_lines = extract_term_theme_lines(raw_text, int(grade.get("grade") or request.get("grade") or 0), term)
    theme_blob = " ".join(
        [
            " ".join(request.get("topics") or []),
            str(request.get("methodology_topic") or ""),
            str(request.get("essay_topic") or ""),
            " ".join(term_lines),
            json.dumps(design, ensure_ascii=True),
        ]
    )
    family = str(ontology.get("assessment_family") or design.get("assessment_family") or "structured_test_or_task")
    matrix = caps.get("assessment_matrix_row") or {}
    source_contract = source_type_contract(str(subject.get("display_name") or ""), family, theme_blob, matrix)
    visuals = visual_kind_rules(str(subject.get("display_name") or ""), family, theme_blob, source_contract)
    focus_items = [
        *[str(item) for item in request.get("topics") or []],
        str(request.get("methodology_topic") or ""),
        str(request.get("essay_topic") or ""),
    ]
    blueprint = {
        "schema": "knowedge.homs_visual_blueprint.v1",
        "subject": subject.get("display_name"),
        "grade": grade.get("grade"),
        "phase": caps.get("phase") or grade.get("phase"),
        "term": term,
        "caps_source": {
            "title": selected.get("title"),
            "path": selected.get("path"),
            "sha256": selected.get("sha256"),
        },
        "canonical_profile_id": ontology.get("profile_id"),
        "assessment_family": family,
        "render_shell": design.get("render_shell") or render_shell_from_request(request),
        "source_type_contract": source_contract,
        "caps_theme_candidates": term_lines,
        "assessment_content_focus": focus_items,
        "content_alignment_notes": content_alignment_notes(focus_items, term_lines),
        "required_visuals": visuals,
        "generation_constraints": [
            "Plan visual evidence before writing questions.",
            "If a question references a map, graph, table, diagram, source, performance space, or practical setup, the matching visual must be present in required_visuals.",
            "Do not create blank source placeholders unless the task explicitly requires the educator to insert an external source.",
            "Use the visual as assessable evidence: questions must ask learners to read, complete, compare, cite, label, calculate from, or reflect on it.",
        ],
        "validation_rules": [
            "No learner-facing question may refer to a missing visual.",
            "Every required visual must have a clear assessment role: stimulus, learner workspace, marking evidence, or moderation evidence.",
            "Visuals must match subject, phase, term/content focus, and assessment family.",
        ],
    }
    return blueprint


def resolve_caps_context(
    subject_profile: dict[str, Any],
    grade_profile: dict[str, Any],
    manifest_path: Path,
    preferred_language: str,
    include_policy: bool,
    max_chars: int,
    matrix_path: Path | None = None,
    ontology_path: Path | None = None,
    assessment_design_path: Path | None = None,
) -> dict[str, Any]:
    manifest = load_manifest(manifest_path)
    candidates = resolve_caps(
        manifest,
        str(subject_profile["subject_id"]),
        int(grade_profile["grade"]),
        preferred_language,
        include_policy,
    )
    context: dict[str, Any] = {
        "schema": "knowedge.homs_caps_context.v1",
        "created_at": utc_now(),
        "enabled": True,
        "manifest": str(manifest_path),
        "subject_id": subject_profile["subject_id"],
        "subject": subject_profile["display_name"],
        "grade": grade_profile["grade"],
        "phase": PHASE_BY_GRADE[int(grade_profile["grade"])],
        "preferred_language": preferred_language,
        "candidate_count": len(candidates),
        "candidates": candidates[:5],
        "selected": candidates[0] if candidates else None,
        "assessment_matrix_row": None,
        "assessment_ontology_profile": None,
        "assessment_design_profile": None,
        "extracted_excerpt": "",
        "extraction_error": None,
    }
    matrix_payload = None
    if candidates and matrix_path and matrix_path.exists():
        try:
            matrix_payload = json.loads(matrix_path.read_text(encoding="utf-8"))
            selected_sha = str(candidates[0].get("sha256") or "")
            matrix_row = next((row for row in matrix_payload.get("rows", []) if row.get("sha256") == selected_sha), None)
            if matrix_row:
                context["assessment_matrix_row"] = {
                    "recommended_blueprint": matrix_row.get("recommended_blueprint"),
                    "subject_family": matrix_row.get("subject_family"),
                    "language_guess": matrix_row.get("language_guess"),
                    "assessment_flags": matrix_row.get("assessment_flags"),
                    "assessment_keyword_counts": matrix_row.get("assessment_keyword_counts"),
                    "signals_source_based": matrix_row.get("signals_source_based"),
                    "signals_essay_or_extended_response": matrix_row.get("signals_essay_or_extended_response"),
                    "signals_case_study": matrix_row.get("signals_case_study"),
                    "signals_data_graph_diagram": matrix_row.get("signals_data_graph_diagram"),
                    "signals_practical_investigation": matrix_row.get("signals_practical_investigation"),
                    "signals_project_or_sba": matrix_row.get("signals_project_or_sba"),
                    "signals_oral_or_performance": matrix_row.get("signals_oral_or_performance"),
                    "signals_calculation_problem": matrix_row.get("signals_calculation_problem"),
                    "signals_formal_test_or_exam": matrix_row.get("signals_formal_test_or_exam"),
                }
                profiles = load_caps_ontology(ontology_path, matrix_payload)
                profile = find_caps_ontology_profile(profiles, matrix_row, subject_profile, grade_profile)
                if profile:
                    context["assessment_ontology_profile"] = {
                        "profile_id": profile.get("profile_id"),
                        "subject": profile.get("subject"),
                        "phase": profile.get("phase"),
                        "phase_group": profile.get("phase_group"),
                        "grades": profile.get("grades"),
                        "subject_family": profile.get("subject_family"),
                        "assessment_family": profile.get("assessment_family"),
                        "allowed": profile.get("allowed"),
                        "default_off": profile.get("default_off"),
                        "required_distribution": profile.get("required_distribution"),
                        "generation_constraints": profile.get("generation_constraints"),
                        "validation_rules": profile.get("validation_rules"),
                        "evidence_summary": profile.get("evidence_summary"),
                    }
                    design_payload = load_caps_assessment_design(assessment_design_path)
                    design_profile = find_assessment_design_profile(design_payload, profile)
                    if design_profile:
                        context["assessment_design_profile"] = {
                            "profile_id": design_profile.get("profile_id"),
                            "subject": design_profile.get("subject"),
                            "assessment_family": design_profile.get("assessment_family"),
                            "render_shell": design_profile.get("render_shell"),
                            "confidence": design_profile.get("confidence"),
                            "assessment_forms": design_profile.get("assessment_forms"),
                            "annual_weightings": design_profile.get("annual_weightings"),
                            "task_requirements": design_profile.get("task_requirements"),
                            "paper_structure": design_profile.get("paper_structure"),
                            "cognitive_distribution": design_profile.get("cognitive_distribution"),
                            "marking_instruments": design_profile.get("marking_instruments"),
                            "generation_constraints": design_profile.get("generation_constraints"),
                            "validation_rules": design_profile.get("validation_rules"),
                        }
        except Exception as exc:
            context["assessment_matrix_error"] = str(exc)
    if candidates and max_chars > 0:
        try:
            raw_text = extract_text(Path(candidates[0]["path"]), max(max_chars * 6, max_chars))
            context["extracted_excerpt"] = select_caps_excerpt(raw_text, subject_profile, grade_profile, max_chars)
        except Exception as exc:
            context["extraction_error"] = str(exc)
    return context


def caps_instruction(caps_context: dict[str, Any], subject_profile: dict[str, Any], grade_profile: dict[str, Any]) -> str:
    selected = caps_context.get("selected") or {}
    if not selected:
        return (
            "CAPS resolver found no matching official DBE CAPS document. "
            "Treat this as a human-review blocker for school-level use."
        )
    excerpt = str(caps_context.get("extracted_excerpt") or "").strip()
    matrix = caps_context.get("assessment_matrix_row") or {}
    ontology = caps_context.get("assessment_ontology_profile") or {}
    assessment_design = caps_context.get("assessment_design_profile") or {}
    blueprint = ontology.get("assessment_family") or matrix.get("recommended_blueprint") or "unresolved"
    signals = []
    for label, key in [
        ("source-based", "signals_source_based"),
        ("essay/extended response", "signals_essay_or_extended_response"),
        ("case study", "signals_case_study"),
        ("data/diagram", "signals_data_graph_diagram"),
        ("practical investigation", "signals_practical_investigation"),
        ("project/SBA", "signals_project_or_sba"),
        ("oral/performance", "signals_oral_or_performance"),
        ("calculation/problem", "signals_calculation_problem"),
    ]:
        if matrix.get(key):
            signals.append(label)
    excerpt_note = f" CAPS excerpt: {excerpt[:1200]}" if excerpt else ""
    ontology_note = ""
    if ontology:
        distribution = ontology.get("required_distribution") or {}
        ontology_note = (
            f" Canonical ontology profile: {ontology.get('profile_id')} / {ontology.get('subject')}. "
            f"Allowed modes: {', '.join(ontology.get('allowed') or [])}. "
            f"Default-off modes: {', '.join(ontology.get('default_off') or [])}. "
            f"Distribution guide: {', '.join(f'{k}={v}%' for k, v in distribution.items())}. "
            f"Generation constraints: {'; '.join(ontology.get('generation_constraints') or [])}. "
            "Treat document keyword signals as supporting evidence, not the final assessment authority."
        )
    design_note = ""
    if assessment_design:
        annual_weightings = assessment_design.get("annual_weightings") or {}
        cognitive = assessment_design.get("cognitive_distribution") or {}
        design_note = (
            f" CAPS assessment design profile: render_shell={assessment_design.get('render_shell')}; "
            f"assessment_forms={', '.join(assessment_design.get('assessment_forms') or [])}; "
            f"annual_weightings={', '.join(f'{k}={v}%' for k, v in annual_weightings.items()) or 'not extracted'}; "
            f"cognitive_or_skill_distribution={', '.join(f'{k}={v}%' for k, v in cognitive.items()) or 'not extracted'}; "
            f"marking_instruments={', '.join(assessment_design.get('marking_instruments') or [])}; "
            f"design_constraints={' ; '.join(assessment_design.get('generation_constraints') or [])}."
        )
    return (
        "CAPS CONTROL: Align this assessment draft to the official DBE CAPS source "
        f"'{selected.get('title')}' for {subject_profile['display_name']}, "
        f"Grade {grade_profile['grade']} / {grade_profile['phase']}. "
        f"CAPS SHA-256: {selected.get('sha256')}. "
        f"CAPS assessment blueprint: {blueprint}. "
        f"Assessment format signals: {', '.join(signals) if signals else 'none detected; use structured grade-appropriate tasks'}. "
        f"{ontology_note} "
        f"{design_note} "
        "Do not force source-based sections or essay matrices unless the canonical ontology profile and requested brief justify them. "
        f"Bloom targets: {', '.join(grade_profile.get('bloom_targets', []))}. "
        f"Question styles: {', '.join(grade_profile.get('question_style', [])[:5])}. "
        f"Memo style: {grade_profile.get('memo_style')}."
        f"{excerpt_note}"
    )


def generation_topic(topic: str, request: dict[str, Any]) -> str:
    note = str(request.get("caps_generation_note") or "").strip()
    if not note:
        return topic
    compact_note = " ".join(note.split())
    return f"{topic}\n\n{compact_note[:1800]}"


def subject_id_from_request(request: dict[str, Any]) -> str:
    return str((request.get("subject_profile") or {}).get("subject_id") or "").strip().lower()


def normalized_request(
    path: Path | None,
    subject_profile: dict[str, Any],
    grade_profile: dict[str, Any],
    caps_context: dict[str, Any] | None = None,
    term: str | int | None = None,
) -> dict[str, Any]:
    request = default_request_for_subject(subject_profile, grade_profile)
    if path:
        user_request = json.loads(path.read_text(encoding="utf-8"))
        request.update(user_request)
    if term and str(term).strip():
        request["term"] = str(term).strip()
    else:
        request["term"] = str(request.get("term") or "auto")
    request["grade"] = int(grade_profile["grade"])
    request["subject_profile"] = {
        key: value for key, value in subject_profile.items() if not key.startswith("_")
    }
    request["subject_profile_path"] = subject_profile.get("_profile_path")
    request["grade_profile"] = {
        key: value for key, value in grade_profile.items() if not key.startswith("_")
    }
    request["grade_ladder_path"] = grade_profile.get("_ladder_path")
    if caps_context:
        request["caps_context"] = caps_context
        request["caps_generation_note"] = caps_instruction(caps_context, subject_profile, grade_profile)
    review_checks = "; ".join(subject_profile.get("human_review_checks", [])[:4])
    blocked_modes = "; ".join(subject_profile.get("blocked_claims_or_modes", [])[:3])
    profile_instruction = (
        f"Subject profile: {subject_profile['display_name']}. "
        f"Grade profile: Grade {grade_profile['grade']} ({grade_profile['phase']}); "
        f"learner level {grade_profile['learner_level']}; "
        f"Bloom targets {', '.join(grade_profile.get('bloom_targets', []))}; "
        f"ZPD {grade_profile['zpd_level']}; "
        f"question style {', '.join(grade_profile.get('question_style', [])[:5])}; "
        f"reading load {grade_profile['reading_load']}; writing load {grade_profile['writing_load']}. "
        f"Question families: {', '.join(subject_profile.get('question_families', [])[:6])}. "
        f"Memo expectations: {', '.join(subject_profile.get('memo_expectations', [])[:4])}. "
        f"Grade memo style: {grade_profile['memo_style']}. "
        f"Human review checks: {review_checks}. "
        f"Do not use blocked modes: {blocked_modes}."
    )
    if caps_context:
        profile_instruction += " " + request["caps_generation_note"]
    if request.get("visual_blueprint"):
        visual_titles = [
            f"{item.get('title')} ({item.get('role')})"
            for item in (request.get("visual_blueprint") or {}).get("required_visuals", [])
        ]
        profile_instruction += (
            " Visual blueprint: "
            + "; ".join(visual_titles)
            + ". Questions must use these visuals as assessable evidence, not decorative inserts."
        )
    existing_instructions = str(request.get("additional_instructions") or "").strip()
    request["additional_instructions"] = (
        f"{existing_instructions}\n\n{profile_instruction}".strip()
        if existing_instructions
        else profile_instruction
    )
    request["topics"] = [str(t).strip() for t in request.get("topics", []) if str(t).strip()][:2]
    if len(request["topics"]) < 2:
        raise ValueError("Exam builder request must include two source-question topics.")
    for field in ["module_code", "module_name", "methodology_topic", "essay_topic"]:
        if not str(request.get(field, "")).strip():
            raise ValueError(f"Exam builder request is missing {field}.")
    request["total_marks"] = int(request.get("total_marks") or 125)
    request["duration_hours"] = int(request.get("duration_hours") or 3)
    return request


def ensure_source_shape(source: dict[str, Any], index: int) -> dict[str, Any]:
    return {
        "label": str(source.get("label") or f"Source {chr(65 + index)}"),
        "type": str(source.get("type") or "Historical source"),
        "title": str(source.get("title") or "Untitled source"),
        "author": str(source.get("author") or source.get("origin") or "Unknown"),
        "date": str(source.get("date") or "Date not specified"),
        "publication": str(source.get("publication") or ""),
        "content": str(source.get("content") or source.get("description") or ""),
        "context": str(source.get("context") or ""),
    }


def ensure_question_shape(question: dict[str, Any], index: int, source_label: str) -> dict[str, Any]:
    marks = question.get("marks", 2)
    try:
        marks = int(marks)
    except Exception:
        marks = 2
    return {
        "number": str(question.get("number") or f"1.{index + 1}"),
        "source_reference": str(question.get("source_reference") or source_label),
        "question": str(question.get("question") or "Answer the question using the source."),
        "marks": marks,
        "cognitive_level": question.get("cognitive_level", 1),
        "expected_answer_points": question.get("expected_answer_points") or [],
    }


def model_json(hymark: Any, system: str, prompt: str, max_tokens: int = 2500, temperature: float = 0.55) -> dict[str, Any]:
    response = hymark.openai_client.chat.completions.create(
        model=hymark.AI_MODEL,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": prompt},
        ],
        temperature=temperature,
        max_tokens=max_tokens,
        response_format={"type": "json_object"},
        timeout=300.0,
    )
    content = response.choices[0].message.content
    if not content:
        raise ValueError("provider returned an empty JSON response")
    cleaned = content.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE | re.DOTALL).strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start >= 0 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise


def fallback_caps_source(topic: str, index: int, request: dict[str, Any]) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    selected = caps.get("selected") or {}
    excerpt = str(caps.get("extracted_excerpt") or "").strip()
    return {
        "label": f"Source {chr(65 + index)}",
        "type": "CAPS-aligned fallback stimulus",
        "title": f"{subject.get('display_name', 'Subject')} stimulus: {topic[:80]}",
        "author": "HOMS Exam Studio fallback generator",
        "date": utc_now()[:10],
        "publication": selected.get("title") or "Official DBE CAPS source",
        "content": (
            f"Controlled fallback stimulus for Grade {grade.get('grade')} {subject.get('display_name')} on {topic}. "
            f"Use this as a review placeholder grounded in the selected CAPS document. "
            f"CAPS excerpt: {excerpt[:900] if excerpt else 'No extracted excerpt available.'}"
        ),
        "context": (
            "Generated because the provider did not return usable source JSON. "
            "Educator review must replace or approve this stimulus before classroom use."
        ),
    }


def ensure_sources_available(sources: list[dict[str, Any]], topic: str, request: dict[str, Any]) -> list[dict[str, Any]]:
    if sources:
        return sources
    return [fallback_caps_source(topic, 0, request), fallback_caps_source(topic, 1, request)]


async def generate_caps_sources(hymark: Any, topic: str, question_number: int, opportunity: int, request: dict[str, Any]) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    variation = ""
    if opportunity == 2:
        variation = "This is a second-opportunity paper. Use equivalent difficulty but different contexts, sources, data, or case facts."
    prompt = f"""Create CAPS-aligned assessment stimulus material.

Subject: {subject.get('display_name')}
Grade: {grade.get('grade')} / {grade.get('phase')}
Topic: {topic}
Source/stimulus types allowed: {', '.join(subject.get('source_types', []))}
Question families: {', '.join(subject.get('question_families', []))}
Bloom targets: {', '.join(grade.get('bloom_targets', []))}
Question style: {', '.join(grade.get('question_style', []))}
CAPS context: {request.get('caps_generation_note', '')}
{variation}

Generate 2-3 legitimate educational stimuli appropriate to the subject and grade. These may be source extracts, case studies, tables, graph descriptions, diagram descriptions, scenarios, or short data sets depending on the subject profile.

Respond in JSON:
{{
  "topic": "{topic}",
  "sources": [
    {{
      "label": "Source A",
      "type": "Case study / table / diagram / extract",
      "title": "Short title",
      "author": "DBE/CAPS-aligned classroom stimulus or supplied context",
      "date": "Not applicable",
      "publication": "",
      "content": "Full stimulus text or detailed description",
      "context": "Why this is appropriate for the grade and subject"
    }}
  ]
}}"""
    try:
        return model_json(
            hymark,
            "You are a South African CAPS assessment designer. Create grade-appropriate, review-ready classroom assessment stimuli without fabricating official CAPS wording.",
            prompt,
            max_tokens=3000,
            temperature=0.55,
        )
    except Exception as exc:
        print(f"[CAPS Source Generation Error] {exc}")
        return {"topic": topic, "sources": [], "error": str(exc)}


async def generate_caps_source_questions(hymark: Any, sources: list[dict[str, Any]], topic: str, total_marks: int, request: dict[str, Any]) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    sources_text = "\n\n".join(
        [
            f"{s['label']}: {s['type']} - {s.get('title', '')}\n{s.get('content', '')}"
            for s in sources
        ]
    )
    prompt = f"""Create CAPS-aligned questions totaling {total_marks} marks.

Subject: {subject.get('display_name')}
Grade: {grade.get('grade')} / {grade.get('phase')}
Topic: {topic}
Bloom targets: {', '.join(grade.get('bloom_targets', []))}
Question families: {', '.join(subject.get('question_families', []))}
Rubric dimensions: {', '.join(subject.get('rubric_dimensions', []))}
Memo expectations: {', '.join(subject.get('memo_expectations', []))}
Grade memo style: {grade.get('memo_style')}
CAPS context: {request.get('caps_generation_note', '')}

Stimuli:
{sources_text}

Return questions with varied cognitive demand appropriate to the grade. Avoid undergraduate assumptions for school grades.

Respond in JSON:
{{
  "questions": [
    {{
      "number": "1.1",
      "source_reference": "Source A",
      "question": "Question text",
      "marks": 2,
      "cognitive_level": "remember/understand/apply/analyze/evaluate/create",
      "expected_answer_points": ["Point 1", "Point 2"]
    }}
  ],
  "total_marks": {total_marks}
}}"""
    try:
        return model_json(
            hymark,
            "You are a CAPS-aligned question writer. Write clear, grade-appropriate assessment questions and memo points.",
            prompt,
            max_tokens=2500,
            temperature=0.5,
        )
    except Exception as exc:
        print(f"[CAPS Question Generation Error] {exc}")
        return {"questions": [], "error": str(exc)}


async def generate_caps_applied_question(hymark: Any, topic: str, marks: int, request: dict[str, Any]) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    prompt = f"""Create a CAPS-aligned applied/structured question.

Subject: {subject.get('display_name')}
Grade: {grade.get('grade')} / {grade.get('phase')}
Topic: {topic}
Marks: {marks}
Bloom targets: {', '.join(grade.get('bloom_targets', []))}
Question style: {', '.join(grade.get('question_style', []))}
CAPS context: {request.get('caps_generation_note', '')}

The question should test application, interpretation, or structured reasoning appropriate to the subject and grade. Include requirements and a compact rubric.

Respond in JSON with keys: question_number, title, topic, task, requirements, marks, rubric.
The rubric must be an object whose criteria each contain exactly these keys: "3 marks", "2 marks", "1 mark"."""
    try:
        return model_json(
            hymark,
            "You are a CAPS assessment specialist creating structured school assessment tasks.",
            prompt,
            max_tokens=2000,
            temperature=0.5,
        )
    except Exception as exc:
        print(f"[CAPS Applied Question Error] {exc}")
        return {"error": str(exc), "marks": marks}


async def generate_caps_extended_question(hymark: Any, topic: str, marks: int, opportunity: int, request: dict[str, Any]) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    variation = ""
    if opportunity == 2:
        variation = "Make this an equivalent but different second-opportunity question."
    prompt = f"""Create a CAPS-aligned extended response or synthesis question.

Subject: {subject.get('display_name')}
Grade: {grade.get('grade')} / {grade.get('phase')}
Topic: {topic}
Marks: {marks}
Bloom targets: {', '.join(grade.get('bloom_targets', []))}
Rubric dimensions: {', '.join(subject.get('rubric_dimensions', []))}
Memo expectations: {', '.join(subject.get('memo_expectations', []))}
CAPS context: {request.get('caps_generation_note', '')}
{variation}

For lower grades, keep the extended response shorter and scaffolded. For FET, allow more independent argument or justified recommendation.

Respond in JSON with keys: question_number, title, topic, question, context, marks, matrix.
The matrix must be an object. Each criterion must contain:
- "weight": number
- "levels": an object with keys containing "Excellent", "Good", "Satisfactory", and "Needs Work"."""
    try:
        return model_json(
            hymark,
            "You are a CAPS-aligned assessment designer creating extended response questions and marking matrices.",
            prompt,
            max_tokens=2500,
            temperature=0.55,
        )
    except Exception as exc:
        print(f"[CAPS Extended Question Error] {exc}")
        return {"error": str(exc), "marks": marks}


def assessment_design_from_request(request: dict[str, Any]) -> dict[str, Any]:
    return ((request.get("caps_context") or {}).get("assessment_design_profile") or {})


def render_shell_from_request(request: dict[str, Any]) -> str:
    design = assessment_design_from_request(request)
    if design.get("render_shell"):
        return str(design["render_shell"])
    ontology = ((request.get("caps_context") or {}).get("assessment_ontology_profile") or {})
    family = ontology.get("assessment_family") or ((request.get("caps_context") or {}).get("assessment_matrix_row") or {}).get("recommended_blueprint")
    return {
        "foundation_activity_assessment": "activity_sheet",
        "practical_project_design_task": "project_task_sheet",
        "practical_performance_or_portfolio": "performance_task_sheet",
        "case_study_structured_questions": "structured_case_paper",
        "calculation_problem_solving": "question_paper",
        "data_diagram_practical_investigation": "investigation_task_sheet",
        "source_based_plus_extended_response": "source_response_paper",
        "source_based_plus_essay": "source_essay_paper",
        "language_integrated_assessment": "language_integrated_task",
    }.get(str(family), "question_paper")


def shell_guidance(shell: str) -> str:
    return {
        "performance_task_sheet": (
            "Create a practical performance assessment instrument. Use task brief, assessment conditions, "
            "observable criteria, and rubric. Do not create answer-line written questions for physical performance."
        ),
        "project_task_sheet": (
            "Create a practical project/design task. Use challenge, constraints, materials, deliverables, "
            "design log, testing evidence, and rubric. Written questions may only support design/theory/reflection."
        ),
        "activity_sheet": (
            "Create a teacher-led activity sheet. Use oral prompts, concrete learner actions, observation evidence, "
            "minimal reading load, and a simple memo or checklist."
        ),
        "structured_case_paper": (
            "Create a structured case/scenario paper. Use compact case facts, applied questions, justified alternatives, "
            "and a memo that ties marks to case evidence."
        ),
        "question_paper": (
            "Create a formal question paper. Use grade-appropriate items, method marks, and a memorandum. "
            "Do not use essays unless the assessment design permits them."
        ),
        "investigation_task_sheet": (
            "Create an investigation task. Separate method, observation, data, interpretation, conclusion, and safety."
        ),
    }.get(shell, "Create a subject-aware CAPS assessment pack that follows the assessment design profile.")


def assessment_language_for_request(request: dict[str, Any]) -> str:
    subject = request.get("subject_profile") or {}
    explicit = subject.get("language_of_assessment") or subject.get("output_language")
    subject_id = str(subject.get("subject_id") or "").lower()
    display = str(subject.get("display_name") or "").lower()
    if explicit:
        return str(explicit)
    if "afrikaans" in subject_id or "afrikaans" in display:
        return "Afrikaans"
    return "English"


def language_contract_text(request: dict[str, Any]) -> str:
    language = assessment_language_for_request(request)
    if language.lower() == "afrikaans":
        return (
            "LANGUAGE CONTRACT: The assessment language is Afrikaans. "
            "All learner-facing and educator-facing document text must be in Afrikaans: "
            "assessment_title, section titles, instructions, stimuli, questions, memo bullets, "
            "rubric criteria/descriptors, and teacher_review_checklist. JSON keys may remain English. "
            "Do not write English instructions such as 'Answer all questions', 'Use the supplied source', "
            "'Credit accurate subject knowledge', or 'Teacher review checklist'."
        )
    return (
        f"LANGUAGE CONTRACT: The assessment language is {language}. "
        "Write learner-facing and educator-facing document text in that language. JSON keys may remain English."
    )


def model_shell_pack(hymark: Any, request: dict[str, Any], opportunity: int) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    ontology = caps.get("assessment_ontology_profile") or {}
    design = assessment_design_from_request(request)
    shell = render_shell_from_request(request)
    language = assessment_language_for_request(request)
    visual_blueprint = request.get("visual_blueprint") or {}
    variation = "First opportunity." if opportunity == 1 else "Second opportunity: equivalent difficulty, different context, different item wording, and no copy of first opportunity."
    prompt = f"""Create one complete CAPS-aligned assessment pack for HyMark Exam Studio.

Subject: {subject.get('display_name')}
Grade: {grade.get('grade')} / {grade.get('phase')}
Language of assessment: {language}
Opportunity: {variation}
Module: {request.get('module_code')} - {request.get('module_name')}
Required render shell: {shell}
Shell guidance: {shell_guidance(shell)}
Topics / focus areas: {'; '.join(request.get('topics') or [])}
Methodology / applied focus: {request.get('methodology_topic', '')}
Extended / synthesis focus: {request.get('essay_topic', '')}
Total marks target: {request.get('total_marks')}

Subject profile:
- Assessment modes: {', '.join(subject.get('assessment_modes') or [])}
- Source/stimulus types: {', '.join(subject.get('source_types') or [])}
- Question families: {', '.join(subject.get('question_families') or [])}
- Rubric dimensions: {', '.join(subject.get('rubric_dimensions') or [])}
- Memo expectations: {', '.join(subject.get('memo_expectations') or [])}
- Blocked modes: {', '.join(subject.get('blocked_claims_or_modes') or [])}

CAPS ontology:
- Profile: {ontology.get('profile_id')}
- Assessment family: {ontology.get('assessment_family')}
- Allowed modes: {', '.join(ontology.get('allowed') or [])}
- Default-off modes: {', '.join(ontology.get('default_off') or [])}
- Distribution: {json.dumps(ontology.get('required_distribution') or {}, ensure_ascii=True)}

CAPS assessment design:
{json.dumps(design, indent=2, ensure_ascii=True)[:3500]}

Visual requirements blueprint:
{json.dumps(visual_blueprint, indent=2, ensure_ascii=True)[:3200]}

CAPS excerpt:
{str(caps.get('extracted_excerpt') or '')[:1800]}

Hard rules:
- {language_contract_text(request)}
- Reconcile marks exactly: total_marks must equal the sum of all question marks and the sum of all rubric marks.
- Build questions around the visual requirements blueprint. Required visuals are not decoration; each must be used as stimulus, learner workspace, marking evidence, or moderation evidence.
- If the visual blueprint requires a map/graph/table/diagram/source, the section stimulus must describe concrete data/features that the formatter can render and learners can answer from.
- Do not include self-corrections, apologies, uncertainty language, "demo purposes", "adjust if", or internal pipeline terms in learner/educator text.
- Do not invent official CAPS quotations. Use CAPS as design authority, not as quoted source text unless supplied verbatim.
- Keep learner-facing instructions separate from memo/rubric guidance.
- Include teacher_review_checklist items that are final checks, not excuses.

Return strict JSON:
{{
  "assessment_title": "...",
  "subject": "{subject.get('display_name')}",
  "grade": {grade.get('grade')},
  "phase": "{caps.get('phase') or grade.get('phase')}",
  "canonical_profile_id": "{ontology.get('profile_id', '')}",
  "blueprint": "{ontology.get('assessment_family', '')}",
  "render_shell": "{shell}",
  "visual_blueprint": {json.dumps(visual_blueprint, ensure_ascii=True)},
  "duration": "...",
  "total_marks": {request.get('total_marks')},
  "sections": [
    {{
      "title": "...",
      "mode": "{shell}",
      "instructions": "...",
      "stimulus": "...",
      "questions": [
        {{"number": "1.1", "question": "...", "marks": 0, "memo": ["..."]}}
      ]
    }}
  ],
  "rubric": [
    {{"criterion": "...", "marks": 0, "descriptor": "..."}}
  ],
  "teacher_review_checklist": ["..."]
}}"""
    return model_json(
        hymark,
        "You are the subject-aware HyMark Exam Builder. Return only valid JSON that passes mark reconciliation.",
        prompt,
        max_tokens=4200,
        temperature=0.45,
    )


def load_evidence_grammar(path: Path = DEFAULT_EVIDENCE_GRAMMAR) -> dict[str, Any]:
    if not path.exists():
        return {"subjects": {}, "global_rules": []}
    return json.loads(path.read_text(encoding="utf-8"))


def evidence_profile_for_request(request: dict[str, Any], grammar_path: Path = DEFAULT_EVIDENCE_GRAMMAR) -> dict[str, Any]:
    grammar = load_evidence_grammar(grammar_path)
    subject_id = subject_id_from_request(request)
    profile = dict((grammar.get("subjects") or {}).get(subject_id) or {})
    profile["global_rules"] = grammar.get("global_rules") or []
    profile.setdefault("evidence_policy", "no_external_source_required")
    profile.setdefault("minimum_sources", 0)
    profile.setdefault("source_families", [])
    profile.setdefault("allowed_question_moves", [])
    return profile


def should_use_evidence_first_route(request: dict[str, Any]) -> bool:
    subject_id = subject_id_from_request(request)
    if subject_id == "history":
        return False
    profile = evidence_profile_for_request(request)
    policy = str(profile.get("evidence_policy") or "")
    return subject_id in {"english_language", "afrikaans_language", "life_orientation"} and policy in {"source_required", "stimulus_required"}


def apply_evidence_route_overrides(request: dict[str, Any]) -> None:
    subject_id = subject_id_from_request(request)
    caps = request.setdefault("caps_context", {})
    if not isinstance(caps.get("assessment_ontology_profile"), dict):
        caps["assessment_ontology_profile"] = {}
    if not isinstance(caps.get("assessment_design_profile"), dict):
        caps["assessment_design_profile"] = {}
    ontology = caps["assessment_ontology_profile"]
    design = caps["assessment_design_profile"]
    if not isinstance(request.get("visual_blueprint"), dict):
        request["visual_blueprint"] = {}
    visual_blueprint = request["visual_blueprint"]
    if subject_id == "life_orientation":
        ontology["assessment_family"] = "case_study_structured_questions"
        design["assessment_family"] = "case_study_structured_questions"
        design["render_shell"] = "structured_case_paper"
        visual_blueprint["assessment_family"] = "case_study_structured_questions"
        visual_blueprint["render_shell"] = "structured_case_paper"
        visual_blueprint["required_visuals"] = [
            {
                "id": "life_orientation_case_facts",
                "title": "Scenario facts and decision table",
                "role": "stimulus_and_workspace",
                "visual_kind": "case_table",
                "must_show": ["scenario facts", "choices", "consequences", "support options"],
                "assessment_use": "Learners analyse scenario evidence and justify safe, responsible decisions.",
            }
        ]
    elif subject_id in {"english_language", "afrikaans_language"}:
        ontology["assessment_family"] = "language_integrated_assessment"
        design["assessment_family"] = "language_integrated_assessment"
        design["render_shell"] = "language_integrated_task"
        visual_blueprint["assessment_family"] = "language_integrated_assessment"
        visual_blueprint["render_shell"] = "language_integrated_task"
        visual_blueprint["required_visuals"] = []
        visual_blueprint["generation_constraints"] = [
            "Render complete generated learner-facing texts before the questions.",
            "Do not insert a generic planner, source placeholder, or described-but-missing visual.",
        ]


def evidence_labels(cards: list[dict[str, Any]]) -> list[str]:
    labels = []
    for index, card in enumerate(cards, start=1):
        label = str(card.get("label") or card.get("id") or f"Source {index}").strip()
        if label:
            labels.append(label)
    return labels


def afrikaans_language_blueprint(total_marks: int) -> dict[str, Any]:
    if total_marks < 30:
        raise ValueError("Afrikaans evidence-first assessments require at least 30 marks.")
    summary_marks = 10 if total_marks >= 40 else 5
    remaining = total_marks - summary_marks
    comprehension_marks = int(round(remaining * 0.6))
    language_marks = remaining - comprehension_marks
    return {
        "paper_form": "teksgebaseerde gekontroleerde toets",
        "language_variant": "Afrikaans Eerste Addisionele Taal",
        "sections": [
            {
                "title": "AFDELING A: BEGRIPSTOETS",
                "marks": comprehension_marks,
                "question_count": "10-20",
                "moves": ["letterlike begrip", "afleiding", "toon", "doel", "woordkeuse", "kritiese taalbewustheid"],
            },
            {
                "title": "AFDELING B: OPSOMMING",
                "marks": summary_marks,
                "question_count": "1",
                "moves": ["identifiseer kernfeite", "skryf in eie woorde", "gehoorsaam woordbeperking"],
            },
            {
                "title": "AFDELING C: TAALSTRUKTURE EN -KONVENSIES",
                "marks": language_marks,
                "question_count": "10",
                "mark_pattern": [1, 2, 1, 2, 2, 1, 2, 2, 1, 2] if language_marks == 16 else [],
                "moves": ["redigering", "woordvorming", "sinsbou", "leestekens", "register", "betekenis in konteks"],
            },
        ],
        "rubric_policy": "Gebruik memorandum-antwoordpunte; geen volpapier-rubriek nie.",
        "question_mark_range": "1-5 punte per subvraag; die opsomming mag die volle afdelingspunt dra.",
    }


def english_language_blueprint(total_marks: int) -> dict[str, Any]:
    if total_marks < 30:
        raise ValueError("English evidence-first assessments require at least 30 marks.")
    summary_marks = 10 if total_marks >= 40 else 5
    remaining = total_marks - summary_marks
    comprehension_marks = int(round(remaining * 0.6))
    language_marks = remaining - comprehension_marks
    return {
        "paper_form": "text-based controlled test",
        "language_variant": "English First Additional Language",
        "sections": [
            {
                "title": "SECTION A: COMPREHENSION",
                "marks": comprehension_marks,
                "question_count": "10-20",
                "moves": ["literal comprehension", "inference", "tone", "purpose", "word choice", "critical language awareness"],
            },
            {
                "title": "SECTION B: SUMMARY",
                "marks": summary_marks,
                "question_count": "1",
                "moves": ["identify key facts", "write in own words", "observe the word limit"],
            },
            {
                "title": "SECTION C: LANGUAGE STRUCTURES AND CONVENTIONS",
                "marks": language_marks,
                "question_count": "10",
                "mark_pattern": [1, 2, 1, 2, 2, 1, 2, 2, 1, 2] if language_marks == 16 else [],
                "moves": ["editing", "word formation", "sentence structure", "punctuation", "register", "meaning in context"],
            },
        ],
        "rubric_policy": "Use question-specific memorandum points; do not attach a duplicate whole-paper rubric.",
        "question_mark_range": "1-5 marks per subquestion; the summary may carry the full section mark.",
    }


def normalize_evidence_cards_for_language(cards: list[dict[str, Any]], request: dict[str, Any]) -> list[dict[str, Any]]:
    language = assessment_language_for_request(request).lower()
    if language != "afrikaans":
        return cards
    replacements = {
        "Text A": "Teks A",
        "Text B": "Teks B",
        "Visual Text A": "Visuele Teks A",
        "Visual Text B": "Teks B",
        "Visuele Teks B": "Teks B",
        "Source A": "Bron A",
        "Source B": "Bron B",
        "constructed classroom evidence": "gekonstrueerde klaskamerbewys",
        "Constructed CAPS-aligned classroom stimulus for educator review": "Gekonstrueerde CAPS-belynde klaskamerstimulus vir opvoederhersiening",
        "Not applicable": "Nie van toepassing nie",
    }
    normalized = []
    for card in cards:
        item = dict(card)
        for key in ["label", "type", "title", "provenance", "date", "content", "context"]:
            value = item.get(key)
            if isinstance(value, str):
                for source, target in replacements.items():
                    value = value.replace(source, target)
                item[key] = value
        normalized.append(item)
    return normalized


def validate_evidence_grounding(pack: dict[str, Any], request: dict[str, Any], evidence_profile: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    policy = str(evidence_profile.get("evidence_policy") or "no_external_source_required")
    if policy == "no_external_source_required":
        return errors
    cards = list(pack.get("evidence_cards") or pack.get("stimulus_cards") or [])
    minimum = int(evidence_profile.get("minimum_sources") or 1)
    if len(cards) < minimum:
        errors.append(f"evidence-first route produced {len(cards)} evidence card(s), expected at least {minimum}")
    labels = evidence_labels(cards)
    if not labels:
        errors.append("evidence-first route has no usable evidence labels")
        return errors
    sections = pack.get("sections") or []
    questions = [question for section in sections for question in section.get("questions") or []]
    if not questions:
        return errors
    grounded = 0
    for section in sections:
        for question in section.get("questions") or []:
            q_blob = " ".join(
                [
                    str(question.get("source_reference") or ""),
                    str(question.get("question") or ""),
                    " ".join(question.get("memo") or []),
                ]
            ).lower()
            if any(label.lower() in q_blob for label in labels):
                grounded += 1
    required = max(1, len(questions) // 2)
    if grounded < required:
        errors.append(f"too few evidence-grounded questions: {grounded}/{len(questions)} cite supplied evidence labels")
    return errors


def generate_evidence_cards(hymark: Any, request: dict[str, Any], opportunity: int, evidence_profile: dict[str, Any]) -> list[dict[str, Any]]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    language = assessment_language_for_request(request)
    variation = "First opportunity." if opportunity == 1 else "Second opportunity: use different text/case/stimulus details with equivalent difficulty."
    label_guidance = (
        "For Life Orientation, label cards Scenario A and Decision Tool B."
        if subject_id_from_request(request) == "life_orientation"
        else (
            "For Afrikaans language tasks, label the complete text cards Teks A and Teks B."
            if assessment_language_for_request(request).lower() == "afrikaans"
            else "For English language tasks, label the complete text cards Text A and Text B."
        )
    )
    subject_id = subject_id_from_request(request)
    afrikaans_source_contract = ""
    english_source_contract = ""
    if subject_id == "afrikaans_language":
        afrikaans_source_contract = """
Afrikaans source contract:
- Return exactly TWO fully written text sources, labelled Teks A and Teks B.
- Teks A must be an original 450-650 word Afrikaans comprehension text with numbered paragraphs [1], [2], etc.
- Teks B must be an original 180-260 word Afrikaans informational text suitable for a summary task.
- These are complete learner-facing texts, not descriptions of an article, graph, photograph, infographic or missing visual.
- Include concrete names, actions, claims, contrasts, figures and language choices that support specific questions.
- Use idiomatic South African Afrikaans appropriate to Grade 12 Eerste Addisionele Taal.
- Do not call either text a source exemplar, classroom evidence, route test or assessment artefact in learner-facing content.
- Add a language_targets array to Teks A with exactly ten verified language-in-context operations totalling 16 marks in this pattern: [1,2,1,2,2,1,2,2,1,2].
- Each language target must contain kind, question, marks and answers. Every question must quote or adapt a real sentence from Teks A, and every answer must be grammatically correct and non-identical to a supplied transformation word.
"""
    elif subject_id == "english_language":
        english_source_contract = """
English source contract:
- Return exactly TWO fully written text sources, labelled Text A and Text B.
- Text A must be an original 450-650 word South African English comprehension text with numbered paragraphs [1], [2], etc.
- Text B must be an original 180-260 word informational text suitable for a summary task.
- These are complete learner-facing texts, not descriptions of an article, graph, photograph, infographic or missing visual.
- Include concrete names, actions, claims, contrasts, figures and language choices that support specific questions.
- Use idiomatic English appropriate to Grade 12 First Additional Language learners.
- Do not call either text a source exemplar, classroom evidence, route test or assessment artefact in learner-facing content.
- Add a language_targets array to Text A with exactly ten verified language-in-context operations totalling 16 marks in this pattern: [1,2,1,2,2,1,2,2,1,2].
- Each language target must contain kind, question, marks and answers. Every question must quote or adapt a real sentence from Text A, and every answer must be grammatically correct and non-identical to a supplied transformation word.
- Use short, unambiguous source sentences for transformations. Avoid passive-voice targets containing stacked auxiliary verbs such as 'has been able to'.
"""
    prompt = f"""Construct CAPS-aligned assessment evidence before questions are written.

Subject: {subject.get('display_name')}
Subject ID: {subject.get('subject_id')}
Grade: {grade.get('grade')} / {grade.get('phase')}
Language of assessment: {language}
Opportunity: {variation}
Topics / focus: {'; '.join(request.get('topics') or [])}
Applied focus: {request.get('methodology_topic', '')}
Extended focus: {request.get('essay_topic', '')}

Evidence policy:
{json.dumps(evidence_profile, indent=2, ensure_ascii=False)[:4500]}

CAPS authority excerpt:
{str(caps.get('extracted_excerpt') or '')[:2200]}

Rules:
- {language_contract_text(request)}
- CAPS content and term focus are the authority.
- Build the evidence/stimulus cards first; do not write questions yet.
- Do not quote exact official CAPS or exam-paper text unless supplied verbatim.
- For English/Afrikaans, create comprehension-style text/visual-text material.
- For Life Orientation, create realistic, school-safe scenario/case evidence.
- {label_guidance}
- Every card must contain assessable details that later questions can cite.
- Label cards clearly, e.g. Text A, Visual Text B, Scenario A, Source A.
{afrikaans_source_contract}
{english_source_contract}

Return strict JSON only:
{{
  "evidence_cards": [
    {{
      "label": "Teks A",
      "type": "constructed classroom evidence",
      "title": "...",
      "provenance": "Constructed CAPS-aligned classroom stimulus for educator review",
      "date": "Not applicable",
      "content": "...",
      "context": "...",
      "usable_evidence": ["specific assessable fact, phrase, value, visible feature or scenario detail"],
      "allowed_question_moves": ["..."]
    }}
  ]
}}"""
    raw = model_json(
        hymark,
        "You are HyMark's evidence construction layer. Return valid JSON only. Create evidence, not questions.",
        prompt,
        max_tokens=3800,
        temperature=0.45,
    )
    cards = raw.get("evidence_cards") or raw.get("stimulus_cards") or raw.get("sources") or []
    cards = [card for card in cards if str(card.get("content") or card.get("description") or "").strip()]
    return normalize_evidence_cards_for_language(cards, request)


def model_evidence_first_pack(
    hymark: Any,
    request: dict[str, Any],
    opportunity: int,
    evidence_profile: dict[str, Any],
    cards: list[dict[str, Any]],
) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    ontology = caps.get("assessment_ontology_profile") or {}
    design = assessment_design_from_request(request)
    shell = render_shell_from_request(request)
    language = assessment_language_for_request(request)
    variation = "First opportunity." if opportunity == 1 else "Second opportunity: equivalent difficulty, different wording, and no copy of first opportunity."
    afrikaans_contract = ""
    if subject_id_from_request(request) == "afrikaans_language":
        afrikaans_contract = f"""
MANDATORY AFRIKAANS PAPER BLUEPRINT:
{json.dumps(afrikaans_language_blueprint(int(request.get('total_marks') or 50)), indent=2, ensure_ascii=False)}

- Use Teks A for comprehension and language-in-context questions.
- Use Teks B for the summary task. State a concrete summary instruction and word limit.
- Write at least 12 substantive subquestions across the three required sections.
- Every question must be answerable from an exact phrase, fact, contrast or language feature in Teks A or Teks B.
- Ask about the CONTENT and LANGUAGE of the texts. Never ask how a source could be used as evidence or which source would make a good assessment question.
- Questions must use realistic short-question marks. No ordinary comprehension question may exceed 5 marks.
- Memo entries must contain the expected answer or acceptable alternatives. Never write only 'Ken punte toe vir...' or another marking instruction.
- Return an empty rubric array because this controlled test is marked with question-specific memo points.
"""
    prompt = f"""Create one complete CAPS-aligned assessment pack from locked evidence cards.

Subject: {subject.get('display_name')}
Grade: {grade.get('grade')} / {grade.get('phase')}
Language of assessment: {language}
Opportunity: {variation}
Render shell: {shell}
Shell guidance: {shell_guidance(shell)}
Topics / focus: {'; '.join(request.get('topics') or [])}
Total marks target: {request.get('total_marks')}

Evidence construction policy:
{json.dumps(evidence_profile, indent=2, ensure_ascii=False)[:4200]}

LOCKED EVIDENCE CARDS. These are the only source/stimulus facts you may assess:
{json.dumps(cards, indent=2, ensure_ascii=False)[:6500]}

CAPS assessment design:
{json.dumps(design, indent=2, ensure_ascii=False)[:2500]}

Hard rules:
- {language_contract_text(request)}
- Reconcile question marks exactly to total_marks. A rubric is independent and is required only for an extended writing/performance task.
- Put the locked evidence cards in the returned JSON as evidence_cards.
- Every evidence-dependent question must explicitly reference a card label such as Text A, Visual Text B, Scenario A or Source A.
- Do not add new evidence facts that are not in the locked evidence cards.
- Do not include pipeline, route, smoke-test, product-QA, apology, uncertainty or demo language.
- Memo bullets must be specific enough for marking and must refer to card evidence where relevant.
{afrikaans_contract}

Return strict JSON:
{{
  "assessment_title": "...",
  "subject": "{subject.get('display_name')}",
  "grade": {grade.get('grade')},
  "phase": "{caps.get('phase') or grade.get('phase')}",
  "canonical_profile_id": "{ontology.get('profile_id', '')}",
  "blueprint": "{ontology.get('assessment_family', '')}",
  "render_shell": "{shell}",
  "language_of_assessment": "{language}",
  "duration": "...",
  "total_marks": {request.get('total_marks')},
  "evidence_cards": {json.dumps(cards, ensure_ascii=False)},
  "sections": [
    {{
      "title": "...",
      "mode": "{shell}",
      "instructions": "...",
      "stimulus": "Refer to Text A / Scenario A / Visual Text B as appropriate.",
      "questions": [
        {{"number": "1.1", "question": "...", "marks": 0, "memo": ["..."]}}
      ]
    }}
  ],
  "rubric": [
    {{"criterion": "...", "marks": 0, "descriptor": "..."}}
  ],
  "teacher_review_checklist": ["..."]
}}"""
    return model_json(
        hymark,
        "You are HyMark's old source-first exam builder adapted for CAPS evidence/stimulus tasks. Return valid JSON only.",
        prompt,
        max_tokens=5200,
        temperature=0.35,
    )


def repair_evidence_first_pack(
    hymark: Any,
    request: dict[str, Any],
    opportunity: int,
    evidence_profile: dict[str, Any],
    cards: list[dict[str, Any]],
    failed_pack: dict[str, Any],
    errors: list[str],
) -> dict[str, Any]:
    language = assessment_language_for_request(request)
    subject_id = subject_id_from_request(request)
    blueprint = afrikaans_language_blueprint(int(request.get("total_marks") or 50)) if subject_id == "afrikaans_language" else {}
    prompt = f"""Repair a failed evidence-first assessment pack. Return a complete replacement JSON object, not commentary.

Subject ID: {subject_id}
Language: {language}
Opportunity: {opportunity}
Total marks: {request.get('total_marks')}

Validation failures:
{json.dumps(errors, indent=2, ensure_ascii=False)}

Locked evidence cards:
{json.dumps(cards, indent=2, ensure_ascii=False)[:9000]}

Required blueprint:
{json.dumps(blueprint, indent=2, ensure_ascii=False)}

Failed pack to replace:
{json.dumps(failed_pack, indent=2, ensure_ascii=False)[:9000]}

Hard repair rules:
- Question marks must sum exactly to {request.get('total_marks')}.
- Preserve the locked evidence cards verbatim and use their exact labels.
- For Afrikaans: provide Begripstoets, Opsomming, and Taalstrukture en -konvensies sections with at least 12 substantive subquestions.
- Every question must test actual text content or an exact language feature in context.
- Include specific expected answers in each question's memo array.
- Do not create a full-paper rubric for a memorandum-based language test; return "rubric": [].
- Never ask about source quality, assessment design, how evidence could be used, or which source makes a good question.
- No ordinary Afrikaans comprehension or language subquestion may exceed 5 marks; the summary may carry up to 10.

Use the same assessment-pack JSON schema as the failed pack. Return JSON only.
"""
    return model_json(
        hymark,
        "You repair CAPS assessment JSON. Produce specific learner questions and exact memorandum answers. Return JSON only.",
        prompt,
        max_tokens=7200,
        temperature=0.2,
    )


def afrikaans_section_errors(section: dict[str, Any], spec: dict[str, Any], source_label: str) -> list[str]:
    errors: list[str] = []
    questions = list(section.get("questions") or [])
    target = int(spec["marks"])
    actual = sum(int(question.get("marks") or 0) for question in questions)
    if actual != target:
        errors.append(f"section marks sum to {actual}, expected {target}")
    count_contract = str(spec.get("question_count") or "1")
    if "-" in count_contract:
        low, high = (int(value) for value in count_contract.split("-", 1))
    else:
        low = high = int(count_contract)
    if not low <= len(questions) <= high:
        errors.append(f"section has {len(questions)} questions, expected {low}-{high}")
    is_summary = "opsom" in str(spec.get("title") or "").lower()
    is_comprehension = "begrip" in str(spec.get("title") or "").lower()
    is_language = "taalstrukture" in str(spec.get("title") or "").lower()
    for question in questions:
        number = str(question.get("number") or "?")
        marks = int(question.get("marks") or 0)
        if marks < 1 or marks > (10 if is_summary else 5):
            errors.append(f"question {number} has invalid mark value {marks}")
        question_blob = " ".join(
            [str(question.get("source_reference") or ""), str(question.get("question") or "")]
        ).lower()
        memo_blob = " ".join(str(item) for item in question.get("memo") or []).lower()
        if source_label.lower() not in question_blob and source_label.lower() not in memo_blob:
            errors.append(f"question {number} does not cite {source_label}")
        if is_comprehension and any(
            phrase in question_blob
            for phrase in ["intensiewe vorm", "lydende vorm", "bedrywende vorm", "trap van vergelyking", "meervoud", "verkleining"]
        ):
            errors.append(f"question {number} is a language-structure item misplaced in the comprehension section")
        if is_language and "korrekte vorm" in question_blob and "_____" not in question_blob:
            errors.append(f"question {number} asks for a correct form but provides no visible blank in the sentence")
        if is_language and "korrekte vorm" in question_blob:
            bracketed = re.findall(r"\(([^()]*)\)", str(question.get("question") or ""))
            memo_answers = [re.sub(r"[^a-zA-Z\u00c0-\u024f-]+", "", item).lower() for item in question.get("memo") or []]
            if bracketed:
                base = re.sub(r"[^a-zA-Z\u00c0-\u024f-]+", "", bracketed[-1]).lower()
                if base and base in memo_answers:
                    errors.append(f"question {number} is a no-op language transformation: {base} -> {base}")
        if any(phrase in question_blob for phrase in ["verduidelik", "motiveer", "twee redes", "twee dinge", "twee eienskappe"]) and marks < 2:
            errors.append(f"question {number} demands explanation or two points but allocates only {marks} mark")
        if not memo_blob:
            errors.append(f"question {number} has no memorandum answer")
        elif any(phrase in memo_blob for phrase in GENERIC_MEMO_PHRASES):
            errors.append(f"question {number} has a generic memorandum instruction")
        if is_summary and ("inhoud" not in memo_blob or "taal" not in memo_blob):
            errors.append(f"question {number} summary memo must state the 7 content + 3 language allocation")
    return errors


def normalize_afrikaans_section(section: dict[str, Any], spec: dict[str, Any], source_label: str) -> dict[str, Any]:
    normalized = dict(section)
    questions = [dict(question) for question in section.get("questions") or []]
    for index, question in enumerate(questions, start=1):
        question["number"] = str(question.get("number") or f"1.{index}")
        question["source_reference"] = source_label
        question["marks"] = int(question.get("marks") or 0)
        question["memo"] = [str(item) for item in question.get("memo") or []]
        question_blob = str(question.get("question") or "").lower()
        minimum = 2 if any(
            phrase in question_blob
            for phrase in ["verduidelik", "motiveer", "twee redes", "twee dinge", "twee eienskappe", "noem twee"]
        ) else 1
        if question["marks"] < minimum:
            question["marks"] = minimum
        if "opsom" in str(spec.get("title") or "").lower():
            points = [re.sub(r"\s*\(1\)\s*$", "", item).strip() for item in question["memo"]]
            question["memo"] = [
                *points,
                "Inhoud: Ken 1 punt per korrekte kernfeit toe, tot 'n maksimum van 7 punte.",
                "Taal: Ken tot 3 punte toe vir eie woorde, volsinne, samehang en nakoming van die woordlimiet.",
            ]
    target = int(spec["marks"])
    actual = sum(question["marks"] for question in questions)
    delta = target - actual
    # Permit only a small reconciliation where an existing explanatory answer
    # clearly supports one additional mark. Larger gaps return to the provider.
    if 0 < delta <= 3:
        ranked = sorted(
            questions,
            key=lambda question: (
                any(verb in str(question.get("question") or "").lower() for verb in ["verduidelik", "motiveer", "vergelyk", "bespreek", "evalueer"]),
                len(question.get("memo") or []),
                -int(question.get("marks") or 0),
            ),
            reverse=True,
        )
        for question in ranked:
            if delta <= 0:
                break
            if int(question.get("marks") or 0) < 5:
                question["marks"] += 1
                delta -= 1
    elif -3 <= delta < 0:
        for question in reversed(questions):
            if delta >= 0:
                break
            question_blob = str(question.get("question") or "").lower()
            minimum = 2 if any(
                phrase in question_blob
                for phrase in ["verduidelik", "motiveer", "twee redes", "twee dinge", "twee eienskappe", "noem twee"]
            ) else 1
            if int(question.get("marks") or 0) > minimum:
                question["marks"] -= 1
                delta += 1
    normalized["questions"] = questions
    return normalized


def generate_afrikaans_section(
    hymark: Any,
    request: dict[str, Any],
    opportunity: int,
    spec: dict[str, Any],
    card: dict[str, Any],
    section_index: int,
    opportunity_dir: Path,
) -> dict[str, Any]:
    label = str(card.get("label") or "Teks A")
    content = str(card.get("content") or "")
    is_summary = "opsom" in str(spec.get("title") or "").lower()
    question_count = str(spec.get("question_count") or "1")
    prompt = f"""Skryf EEN afdeling van 'n Graad 12 Afrikaans Eerste Addisionele Taal-toets.

Geleentheid: {opportunity}
Kwartaal: {request.get('term')}
Afdeling: {spec.get('title')}
Presiese afdelingstotaal: {spec.get('marks')} punte
Aantal vrae/subvrae: {question_count}
Presiese puntpatroon: {json.dumps(spec.get('mark_pattern') or [], ensure_ascii=False)}
Toegelate vraagbewegings: {', '.join(spec.get('moves') or [])}
Bronetiket: {label}
Brontitel: {card.get('title')}

VOLLEDIGE GESLOTE TEKS:
{content}

Reels:
- Skryf alle leerderteks en memorandumantwoorde in idiomatiese Afrikaans.
- Elke vraag moet na {label} verwys en toets 'n presiese feit, frase, paragraaf, afleiding, toonkeuse of taalvorm uit die teks.
- Memorandumlyste moet die werklike verwagte antwoord(e) bevat, met geldige alternatiewe waar nodig.
- Moenie vra hoe 'n bron as bewys gebruik kan word, hoe sterk 'n assesseringsvraag is, of hoe die vraestel ontwerp is nie.
- Gewone subvrae dra 1-5 punte. {'Die enkele opsommingstaak dra presies 10 punte: 7 vir SEWE korrekte inhoudspunte en 3 vir taal, eie woorde, volsinne, samehang en woordlimiet. Die memorandum moet hierdie 7+3 verdeling uitdruklik wys.' if is_summary else 'Versprei die punte realisties oor die vereiste aantal subvrae.'}
- In AFDELING A, toets slegs begrip, afleiding, toon, doel, woordkeuse en kritiese taalbewustheid. Plaas geen lydende vorm, intensiewe vorm, trappe van vergelyking of ander grammatika-oefeninge daar nie.
- In AFDELING C, skryf presies 10 vrae volgens die gegewe puntpatroon. Gebruik 'n gebalanseerde mengsel van taal-in-konteks, redigering, woordvorming, sinsbou, direkte/indirekte rede, ontkenning, leestekens en register.
- As 'n leerder die korrekte vorm van 'n woord moet gee, plaas die grondwoord tussen hakies EN 'n sigbare _____ in die sin. Die antwoord mag nooit identies aan die grondwoord wees nie.
- Kontroleer elke taalantwoord self. Moenie 'n gebroke direkte aanhaling of 'n sin sonder 'n werklike transformasie gebruik nie.
- 'n Vraag wat TWEE feite vra, moet minstens 2 punte dra. 'n Vraag wat verduideliking plus motivering vereis, moet genoeg punte vir albei dele dra.
- Die som van vraagpunte moet presies {spec.get('marks')} wees.

Return strict JSON only:
{{
  "title": "{spec.get('title')}",
  "mode": "language_integrated_task",
  "instructions": "...",
  "stimulus": "Gebruik {label}.",
  "questions": [
    {{"number": "{section_index}.1", "source_reference": "{label}", "question": "...", "marks": 2, "memo": ["spesifieke antwoord"]}}
  ]
}}
"""
    system = "Jy is 'n ervare Suid-Afrikaanse Afrikaans EAT-vraestelopsteller. Skryf spesifieke vrae en presiese memorandumantwoorde. Return JSON only."
    raw = model_json(hymark, system, prompt, max_tokens=2600, temperature=0.25)
    opportunity_dir.joinpath(f"HOMS_PROVIDER_SECTION_{section_index}_ATTEMPT_1.json").write_text(json.dumps(raw, indent=2), encoding="utf-8")
    section = dict(raw.get("section") or raw)
    section["title"] = str(section.get("title") or spec.get("title"))
    section["mode"] = "language_integrated_task"
    section = normalize_afrikaans_section(section, spec, label)
    opportunity_dir.joinpath(f"HOMS_PROVIDER_SECTION_{section_index}_NORMALIZED_1.json").write_text(json.dumps(section, indent=2), encoding="utf-8")
    errors = afrikaans_section_errors(section, spec, label)
    if errors:
        repair_prompt = prompt + "\n\nDie vorige poging het misluk:\n" + "\n".join(f"- {error}" for error in errors) + "\n\nHerstel die hele afdeling en return JSON only."
        repaired = model_json(hymark, system, repair_prompt, max_tokens=2800, temperature=0.15)
        opportunity_dir.joinpath(f"HOMS_PROVIDER_SECTION_{section_index}_ATTEMPT_2.json").write_text(json.dumps(repaired, indent=2), encoding="utf-8")
        section = dict(repaired.get("section") or repaired)
        section["title"] = str(section.get("title") or spec.get("title"))
        section["mode"] = "language_integrated_task"
        section = normalize_afrikaans_section(section, spec, label)
        opportunity_dir.joinpath(f"HOMS_PROVIDER_SECTION_{section_index}_NORMALIZED_2.json").write_text(json.dumps(section, indent=2), encoding="utf-8")
        errors = afrikaans_section_errors(section, spec, label)
    if errors:
        raise RuntimeError(f"{spec.get('title')} failed validation: " + "; ".join(errors))
    return section


def build_afrikaans_language_section_from_targets(card: dict[str, Any], spec: dict[str, Any]) -> dict[str, Any]:
    targets = list(card.get("language_targets") or [])
    if len(targets) != 10:
        raise RuntimeError(f"Teks A requires exactly 10 reviewed language targets, found {len(targets)}")
    questions = []
    for index, target in enumerate(targets, start=1):
        questions.append(
            {
                "number": f"3.{index}",
                "source_reference": str(card.get("label") or "Teks A"),
                "question": str(target.get("question") or ""),
                "marks": int(target.get("marks") or 0),
                "memo": [str(answer) for answer in target.get("answers") or []],
                "language_target_kind": str(target.get("kind") or ""),
            }
        )
    section = {
        "title": str(spec.get("title") or "AFDELING C: TAALSTRUKTURE EN -KONVENSIES"),
        "mode": "language_integrated_task",
        "instructions": "Beantwoord die taalvrae. Gebruik Teks A en volg elke opdrag noukeurig.",
        "stimulus": "Gebruik Teks A.",
        "questions": questions,
    }
    errors = afrikaans_section_errors(section, spec, str(card.get("label") or "Teks A"))
    if errors:
        raise RuntimeError("reviewed Teks A language targets failed validation: " + "; ".join(errors))
    return section


def model_afrikaans_evidence_pack(
    hymark: Any,
    request: dict[str, Any],
    opportunity: int,
    cards: list[dict[str, Any]],
    opportunity_dir: Path,
) -> dict[str, Any]:
    if len(cards) < 2:
        raise RuntimeError("Afrikaans split builder requires Teks A and Teks B.")
    total = int(request.get("total_marks") or 50)
    blueprint = afrikaans_language_blueprint(total)
    sections = []
    for index, spec in enumerate(blueprint["sections"], start=1):
        card = cards[1] if "opsom" in str(spec.get("title") or "").lower() else cards[0]
        locked_key = "locked_afrikaans_section_a_path" if index == 1 else ("locked_afrikaans_section_b_path" if index == 2 else "")
        locked_section_path = Path(str(request.get(locked_key) or "")).expanduser() if locked_key else Path()
        if index == 3 and card.get("language_targets"):
            language_section = build_afrikaans_language_section_from_targets(card, spec)
            opportunity_dir.joinpath("HOMS_PROVIDER_SECTION_3_FROM_LOCKED_TARGETS.json").write_text(
                json.dumps(language_section, indent=2), encoding="utf-8"
            )
            sections.append(language_section)
        elif locked_key and request.get(locked_key) and locked_section_path.exists():
            locked_section = json.loads(locked_section_path.read_text(encoding="utf-8"))
            locked_section = normalize_afrikaans_section(locked_section.get("section") or locked_section, spec, str(card.get("label") or "Teks A"))
            locked_errors = afrikaans_section_errors(locked_section, spec, str(card.get("label") or "Teks A"))
            if locked_errors:
                raise RuntimeError("reviewed Afrikaans Section A failed current validation: " + "; ".join(locked_errors))
            opportunity_dir.joinpath(f"HOMS_PROVIDER_SECTION_{index}_REUSED.json").write_text(json.dumps(locked_section, indent=2), encoding="utf-8")
            sections.append(locked_section)
        else:
            sections.append(generate_afrikaans_section(hymark, request, opportunity, spec, card, index, opportunity_dir))
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    ontology = caps.get("assessment_ontology_profile") or {}
    return {
        "assessment_title": f"Graad {grade.get('grade', 12)} Afrikaans Eerste Addisionele Taal: Kwartaal {request.get('term')} Gekontroleerde Toets",
        "subject": "Afrikaans Eerste Addisionele Taal",
        "grade": grade.get("grade", 12),
        "phase": caps.get("phase") or grade.get("phase"),
        "canonical_profile_id": ontology.get("profile_id") or "fet.afrikaans_language",
        "blueprint": "language_integrated_assessment",
        "render_shell": "language_integrated_task",
        "language_of_assessment": "Afrikaans",
        "duration": f"{request.get('duration_hours', 2)} uur",
        "total_marks": total,
        "evidence_cards": cards,
        "sections": sections,
        "rubric": [],
        "teacher_review_checklist": [
            "Bevestig dat Teks A en Teks B se taalvlak by Graad 12 EAT pas.",
            "Bevestig dat elke memorandumantwoord uit die betrokke teks afgelei kan word.",
            "Bevestig die woordlimiet, puntetoekenning en aanvaarbare alternatiewe.",
            "Keur die finale vraestel en memorandum voor klasgebruik goed.",
        ],
        "generation_backend": "hymark_afrikaans_split_source_first",
    }


def english_section_errors(section: dict[str, Any], spec: dict[str, Any], source_label: str) -> list[str]:
    errors: list[str] = []
    questions = list(section.get("questions") or [])
    target = int(spec["marks"])
    actual = sum(int(question.get("marks") or 0) for question in questions)
    if actual != target:
        errors.append(f"section marks sum to {actual}, expected {target}")
    count_contract = str(spec.get("question_count") or "1")
    if "-" in count_contract:
        low, high = (int(value) for value in count_contract.split("-", 1))
    else:
        low = high = int(count_contract)
    if not low <= len(questions) <= high:
        errors.append(f"section has {len(questions)} questions, expected {low}-{high}")
    is_summary = "summary" in str(spec.get("title") or "").lower()
    is_comprehension = "comprehension" in str(spec.get("title") or "").lower()
    is_language = "language structures" in str(spec.get("title") or "").lower()
    for question in questions:
        number = str(question.get("number") or "?")
        marks = int(question.get("marks") or 0)
        if marks < 1 or marks > (10 if is_summary else 5):
            errors.append(f"question {number} has invalid mark value {marks}")
        question_blob = " ".join(
            [str(question.get("source_reference") or ""), str(question.get("question") or "")]
        ).lower()
        memo_blob = " ".join(str(item) for item in question.get("memo") or []).lower()
        if source_label.lower() not in question_blob and source_label.lower() not in memo_blob:
            errors.append(f"question {number} does not cite {source_label}")
        if is_comprehension and any(
            phrase in question_blob
            for phrase in ["passive voice", "active voice", "reported speech", "plural form", "punctuate the sentence"]
        ):
            errors.append(f"question {number} is a language-structure item misplaced in comprehension")
        if is_language and "correct form" in question_blob and "_____" not in question_blob:
            errors.append(f"question {number} asks for a correct form but provides no visible blank")
        if any(phrase in question_blob for phrase in ["explain", "justify", "two reasons", "two ways", "two qualities"]) and marks < 2:
            errors.append(f"question {number} demands explanation or two points but allocates only {marks} mark")
        if not memo_blob:
            errors.append(f"question {number} has no memorandum answer")
        elif any(phrase in memo_blob for phrase in GENERIC_MEMO_PHRASES):
            errors.append(f"question {number} has a generic memorandum instruction")
        if is_summary and ("content" not in memo_blob or "language" not in memo_blob):
            errors.append(f"question {number} summary memo must state the 7 content + 3 language allocation")
    return errors


def normalize_english_section(section: dict[str, Any], spec: dict[str, Any], source_label: str) -> dict[str, Any]:
    normalized = dict(section)
    questions = [dict(question) for question in section.get("questions") or []]
    for index, question in enumerate(questions, start=1):
        question["number"] = str(question.get("number") or f"1.{index}")
        question["source_reference"] = source_label
        question["marks"] = int(question.get("marks") or 0)
        question["memo"] = [str(item) for item in question.get("memo") or []]
        if "summary" in str(spec.get("title") or "").lower():
            points = [re.sub(r"\s*\(1\)\s*$", "", item).strip() for item in question["memo"]]
            question["memo"] = [
                *points,
                "Content: Award 1 mark per correct key fact, to a maximum of 7 marks.",
                "Language: Award up to 3 marks for own words, complete sentences, coherence and adherence to the word limit.",
            ]
    normalized["questions"] = questions
    return normalized


def generate_english_section(
    hymark: Any,
    request: dict[str, Any],
    opportunity: int,
    spec: dict[str, Any],
    card: dict[str, Any],
    section_index: int,
    opportunity_dir: Path,
) -> dict[str, Any]:
    label = str(card.get("label") or "Text A")
    content = str(card.get("content") or "")
    is_summary = "summary" in str(spec.get("title") or "").lower()
    prompt = f"""Write ONE section of a Grade 12 English First Additional Language controlled test.

Opportunity: {opportunity}
Term: {request.get('term')}
Section: {spec.get('title')}
Exact section total: {spec.get('marks')} marks
Required questions/subquestions: {spec.get('question_count')}
Permitted assessment moves: {', '.join(spec.get('moves') or [])}
Source label: {label}
Source title: {card.get('title')}

COMPLETE LOCKED TEXT:
{content}

Rules:
- Write all learner text and memorandum answers in idiomatic South African English.
- Every question must cite {label} and assess an exact fact, phrase, paragraph, inference, tone choice or language feature in the text.
- Memorandum arrays must contain the actual expected answer and acceptable alternatives where appropriate.
- Never ask how a source could be used as evidence, how strong an assessment question would be, or how the paper was designed.
- Ordinary subquestions carry 1-5 marks. {'The single summary task carries exactly 10 marks: 7 for SEVEN correct content points and 3 for language, own words, complete sentences, coherence and a strict 70-word limit. State this 7+3 allocation explicitly in the memorandum.' if is_summary else 'Distribute the marks realistically across the required number of subquestions.'}
- In SECTION A, assess comprehension, inference, tone, purpose, word choice and critical language awareness only. Do not insert passive voice, reported speech, plurals or punctuation exercises.
- A question asking for TWO facts must carry at least 2 marks. An explanation or justification must have enough marks for its required parts.
- The sum of question marks must be exactly {spec.get('marks')}.

Return strict JSON only:
{{
  "title": "{spec.get('title')}",
  "mode": "language_integrated_task",
  "instructions": "...",
  "stimulus": "Use {label}.",
  "questions": [
    {{"number": "{section_index}.1", "source_reference": "{label}", "question": "...", "marks": 2, "memo": ["specific answer"]}}
  ]
}}
"""
    system = "You are an experienced South African English FAL paper setter. Write specific questions and exact memorandum answers. Return JSON only."
    raw = model_json(hymark, system, prompt, max_tokens=2800, temperature=0.25)
    opportunity_dir.joinpath(f"HOMS_PROVIDER_SECTION_{section_index}_ATTEMPT_1.json").write_text(json.dumps(raw, indent=2), encoding="utf-8")
    section = normalize_english_section(dict(raw.get("section") or raw), spec, label)
    errors = english_section_errors(section, spec, label)
    if errors:
        repair_prompt = prompt + "\n\nThe previous attempt failed:\n" + "\n".join(f"- {error}" for error in errors) + "\nRepair the whole section and return JSON only."
        repaired = model_json(hymark, system, repair_prompt, max_tokens=3000, temperature=0.15)
        opportunity_dir.joinpath(f"HOMS_PROVIDER_SECTION_{section_index}_ATTEMPT_2.json").write_text(json.dumps(repaired, indent=2), encoding="utf-8")
        section = normalize_english_section(dict(repaired.get("section") or repaired), spec, label)
        errors = english_section_errors(section, spec, label)
    if errors:
        raise RuntimeError(f"{spec.get('title')} failed validation: " + "; ".join(errors))
    return section


def build_english_language_section_from_targets(card: dict[str, Any], spec: dict[str, Any]) -> dict[str, Any]:
    targets = list(card.get("language_targets") or [])
    if len(targets) != 10:
        raise RuntimeError(f"Text A requires exactly 10 reviewed language targets, found {len(targets)}")
    questions = [
        {
            "number": f"3.{index}",
            "source_reference": str(card.get("label") or "Text A"),
            "question": str(target.get("question") or ""),
            "marks": int(target.get("marks") or 0),
            "memo": [str(answer) for answer in target.get("answers") or []],
            "language_target_kind": str(target.get("kind") or ""),
        }
        for index, target in enumerate(targets, start=1)
    ]
    section = {
        "title": str(spec.get("title") or "SECTION C: LANGUAGE STRUCTURES AND CONVENTIONS"),
        "mode": "language_integrated_task",
        "instructions": "Answer the language questions. Use Text A and follow each instruction carefully.",
        "stimulus": "Use Text A.",
        "questions": questions,
    }
    errors = english_section_errors(section, spec, str(card.get("label") or "Text A"))
    if errors:
        raise RuntimeError("reviewed Text A language targets failed validation: " + "; ".join(errors))
    return section


def model_english_evidence_pack(
    hymark: Any,
    request: dict[str, Any],
    opportunity: int,
    cards: list[dict[str, Any]],
    opportunity_dir: Path,
) -> dict[str, Any]:
    if len(cards) < 2:
        raise RuntimeError("English split builder requires Text A and Text B.")
    total = int(request.get("total_marks") or 50)
    blueprint = english_language_blueprint(total)
    sections = []
    for index, spec in enumerate(blueprint["sections"], start=1):
        card = cards[1] if "summary" in str(spec.get("title") or "").lower() else cards[0]
        if index == 3 and card.get("language_targets"):
            section = build_english_language_section_from_targets(card, spec)
            opportunity_dir.joinpath("HOMS_PROVIDER_SECTION_3_FROM_LOCKED_TARGETS.json").write_text(
                json.dumps(section, indent=2), encoding="utf-8"
            )
            sections.append(section)
        else:
            sections.append(generate_english_section(hymark, request, opportunity, spec, card, index, opportunity_dir))
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    ontology = caps.get("assessment_ontology_profile") or {}
    return {
        "assessment_title": f"Grade {grade.get('grade', 12)} English First Additional Language: Term {request.get('term')} Controlled Test",
        "subject": "English First Additional Language",
        "grade": grade.get("grade", 12),
        "phase": caps.get("phase") or grade.get("phase"),
        "canonical_profile_id": ontology.get("profile_id") or "fet.english_language",
        "blueprint": "language_integrated_assessment",
        "render_shell": "language_integrated_task",
        "language_of_assessment": "English",
        "duration": f"{request.get('duration_hours', 2)} hours",
        "total_marks": total,
        "evidence_cards": cards,
        "sections": sections,
        "rubric": [],
        "teacher_review_checklist": [
            "Confirm that Text A and Text B are appropriate for Grade 12 FAL learners.",
            "Confirm that every memorandum answer is traceable to the relevant text.",
            "Confirm the summary word limit, mark allocation and acceptable alternatives.",
            "Approve the final paper and memorandum before classroom use.",
        ],
        "generation_backend": "hymark_english_split_source_first",
    }


def fallback_evidence_first_pack(
    request: dict[str, Any],
    opportunity: int,
    cards: list[dict[str, Any]],
    evidence_profile: dict[str, Any],
    error: str = "",
) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    ontology = caps.get("assessment_ontology_profile") or {}
    shell = render_shell_from_request(request)
    language = assessment_language_for_request(request)
    subject_name = subject.get("display_name", "Subject")
    total = int(request.get("total_marks") or 50)
    labels = evidence_labels(cards)
    primary = labels[0] if labels else ("Teks A" if language.lower() == "afrikaans" else "Text A")
    secondary = labels[1] if len(labels) > 1 else primary
    third = labels[2] if len(labels) > 2 else secondary
    first = max(10, total // 2)
    second = max(8, total // 4)
    third_marks = total - first - second
    if third_marks < 1:
        third_marks = max(1, total - first)
        second = total - first - third_marks

    if language.lower() == "afrikaans":
        sections = [
            {
                "title": "Afdeling A: Begrip En Bewysgebruik",
                "mode": shell,
                "instructions": f"Lees {primary} en {secondary} noukeurig en beantwoord die vrae in Afrikaans.",
                "stimulus": f"Gebruik slegs die besonderhede in {primary} en {secondary} as bewys.",
                "questions": [
                    {"number": "1.1", "question": f"Identifiseer die hoofgedagte van {primary} en motiveer dit met EEN bewys uit die teks.", "marks": first, "memo": [f"Ken punte toe vir 'n geldige hoofgedagte en relevante bewys uit {primary}."]},
                    {"number": "1.2", "question": f"Verduidelik hoe taal, toon of visuele besonderhede in {secondary} die boodskap versterk.", "marks": second, "memo": [f"Ken punte toe vir akkurate verwysing na {secondary} en duidelike verduideliking."]},
                    {"number": "1.3", "question": f"Skryf 'n kort, gestruktureerde respons wat {primary} en {third} verbind.", "marks": third_marks, "memo": [f"Ken punte toe vir samehangende respons en korrekte gebruik van {primary} en {third}."]},
                ],
            }
        ]
        rubric = [
            {"criterion": "Begrip en bewysgebruik", "marks": first, "descriptor": "Gebruik toepaslike bewyse uit die gegewe teks of stimulus."},
            {"criterion": "Taal- en visuele interpretasie", "marks": second, "descriptor": "Verduidelik toon, register, beeld of taalkeuse akkuraat."},
            {"criterion": "Gestruktureerde respons", "marks": third_marks, "descriptor": "Skryf duidelik, samehangend en taakgerig."},
        ]
        checklist = [
            "Bevestig CAPS-onderwerpomvang en graadvlak.",
            "Bevestig dat die teks- of visuele stimulus geskik is.",
            "Bevestig dat alle vrae na die bewysetikette verwys.",
            "Keur memorandum en punteverdeling goed voor gebruik.",
        ]
    else:
        scenario_word = "Scenario" if "life_orientation" in subject_id_from_request(request) else "Text"
        sections = [
            {
                "title": "Section A: Evidence-Based Response",
                "mode": shell,
                "instructions": f"Read {primary} and {secondary} carefully before answering.",
                "stimulus": f"Use only the details in {primary}, {secondary} and {third} as evidence.",
                "questions": [
                    {"number": "1.1", "question": f"Identify the main issue, idea or message in {primary} and support your answer with one detail from the evidence.", "marks": first, "memo": [f"Credit a valid point supported by evidence from {primary}."]},
                    {"number": "1.2", "question": f"Explain how a detail in {secondary} affects the meaning, decision or viewpoint.", "marks": second, "memo": [f"Credit accurate interpretation of {secondary} and a clear explanation."]},
                    {"number": "1.3", "question": f"Write a structured response that uses {primary} and {third} to justify a conclusion or recommendation.", "marks": third_marks, "memo": [f"Credit a coherent response that uses evidence from {primary} and {third}."]},
                ],
            }
        ]
        rubric = [
            {"criterion": f"{scenario_word} comprehension and evidence use", "marks": first, "descriptor": "Uses supplied evidence accurately."},
            {"criterion": "Interpretation and explanation", "marks": second, "descriptor": "Explains meaning, consequence or viewpoint clearly."},
            {"criterion": "Structured response", "marks": third_marks, "descriptor": "Uses evidence to support a coherent conclusion or recommendation."},
        ]
        checklist = [
            "Confirm CAPS topic scope and grade level.",
            "Confirm the evidence/stimulus cards are suitable and school-safe.",
            "Confirm every evidence-dependent question cites a supplied label.",
            "Approve memo and mark allocation before classroom use.",
        ]

    return {
        "assessment_title": f"Grade {grade.get('grade', request.get('grade', ''))} {subject_name}: Evidence-First CAPS Assessment",
        "subject": subject_name,
        "grade": grade.get("grade", request.get("grade", "")),
        "phase": caps.get("phase") or grade.get("phase"),
        "canonical_profile_id": ontology.get("profile_id"),
        "blueprint": ontology.get("assessment_family") or "evidence_first_assessment",
        "render_shell": shell,
        "language_of_assessment": language,
        "visual_blueprint": request.get("visual_blueprint") or {},
        "duration": f"{request.get('duration_hours', 2)} hours",
        "total_marks": total,
        "evidence_cards": cards,
        "evidence_policy": evidence_profile.get("evidence_policy"),
        "sections": sections,
        "rubric": rubric,
        "teacher_review_checklist": checklist,
        "fallback_reason": error,
        "generation_backend": "hymark_evidence_first_deterministic_fallback",
    }


def fallback_shell_pack(request: dict[str, Any], opportunity: int, error: str = "") -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    ontology = caps.get("assessment_ontology_profile") or {}
    design = assessment_design_from_request(request)
    shell = render_shell_from_request(request)
    subject_name = subject.get("display_name", "Subject")
    grade_value = grade.get("grade", request.get("grade", ""))
    opp_note = "First Opportunity" if opportunity == 1 else "Second Opportunity"
    blueprint = ontology.get("assessment_family") or design.get("assessment_family") or "structured_test_or_task"
    language = assessment_language_for_request(request)

    if language.lower() == "afrikaans":
        total = int(request.get("total_marks") or 50)
        first = max(10, total // 2)
        second = total - first
        sections = [
            {
                "title": "Leesbegrip En Taalstrukture",
                "mode": shell,
                "instructions": "Beantwoord al die vrae in Afrikaans. Gebruik bewyse uit die verskafde teks of visuele bron waar dit vereis word.",
                "stimulus": "CAPS-belynde Afrikaanse lees- of visuele stimulus wat deur die opvoeder verskaf word.",
                "questions": [
                    {"number": "1.1", "question": "Identifiseer die hoofgedagte van die stimulus en motiveer jou antwoord met EEN teksbewys.", "marks": first, "memo": ["Ken punte toe vir 'n geldige hoofgedagte en gepaste teksbewys."]},
                    {"number": "1.2", "question": "Voltooi die toegepaste taal- of skryftaak met gepaste register, struktuur en taalgebruik.", "marks": second, "memo": ["Ken punte toe vir akkurate taalgebruik, gepaste register en duidelike struktuur."]},
                ],
            }
        ]
        rubric = [
            {"criterion": "Teksbegrip en bewysgebruik", "marks": first, "descriptor": "Antwoorde toon akkurate begrip en gebruik relevante teksbewyse."},
            {"criterion": "Taalgebruik en toepassing", "marks": second, "descriptor": "Taal, register, spelling, leestekens en struktuur pas by die taak."},
        ]
    elif shell == "performance_task_sheet":
        total = 50
        sections = [
            {
                "title": "Performance Assessment Task",
                "mode": "performance_task_sheet",
                "instructions": "Prepare and perform the required sequence under educator supervision. The educator records observable performance evidence against the rubric.",
                "stimulus": "Use classroom-taught technique, movement vocabulary, composition principles, and safe dance practice.",
                "questions": [
                    {"number": "1.1", "question": "Perform the prepared practical sequence according to the task brief and safety conditions.", "marks": 30, "memo": ["Assess observable technique, control, accuracy, spatial use, and performance readiness."]},
                    {"number": "1.2", "question": "Present a short oral or written reflection on movement choices, style, and improvement.", "marks": 20, "memo": ["Credit correct terminology, relevant reflection, and links to the studied dance context."]},
                ],
            }
        ]
        rubric = [
            {"criterion": "Technique and safe body use", "marks": 15, "descriptor": "Accurate, controlled, safe execution of required movement vocabulary."},
            {"criterion": "Composition, space, dynamics and musicality", "marks": 15, "descriptor": "Clear movement structure with effective use of space, timing, dynamics and transitions."},
            {"criterion": "Performance quality and expression", "marks": 10, "descriptor": "Committed, focused performance with appropriate expressive intent."},
            {"criterion": "Reflection and dance terminology", "marks": 10, "descriptor": "Clear reflection using relevant subject vocabulary and evidence from the performed work."},
        ]
    elif shell == "project_task_sheet":
        total = 50
        sections = [
            {
                "title": "Practical Project Task",
                "mode": "project_task_sheet",
                "instructions": "Complete the design, build, testing and reflection evidence required by the task. Keep all design notes and test results available for marking.",
                "stimulus": "Use the supplied classroom problem, available materials, and subject-specific constraints.",
                "questions": [
                    {"number": "1.1", "question": "Define the problem and list the design constraints.", "marks": 10, "memo": ["Credit a clear problem statement, realistic constraints, and subject vocabulary."]},
                    {"number": "1.2", "question": "Produce the planned solution, diagram, artefact or procedure required by the task.", "marks": 20, "memo": ["Assess feasibility, correctness, execution quality, and alignment with the brief."]},
                    {"number": "1.3", "question": "Test the solution and record evidence of what worked and what did not.", "marks": 10, "memo": ["Credit honest test evidence and relevant observations."]},
                    {"number": "1.4", "question": "Explain one improvement using evidence from the test.", "marks": 10, "memo": ["Credit justified improvement linked to evidence."]},
                ],
            }
        ]
        rubric = [
            {"criterion": "Problem and constraints", "marks": 10, "descriptor": "Clear problem definition and realistic constraints."},
            {"criterion": "Practical execution", "marks": 20, "descriptor": "Accurate, feasible, complete solution or artefact."},
            {"criterion": "Testing evidence", "marks": 10, "descriptor": "Relevant observations and test records."},
            {"criterion": "Reflection and improvement", "marks": 10, "descriptor": "Evidence-based improvement using subject vocabulary."},
        ]
    elif shell == "activity_sheet":
        total = 20
        sections = [
            {
                "title": "Teacher-Led Learner Activity",
                "mode": "activity_sheet",
                "instructions": "The educator reads prompts aloud, observes learner actions, and records evidence. Learners use concrete materials or simple drawings.",
                "stimulus": "Use classroom manipulatives, picture prompts, number lines, matching cards, or teacher-provided examples.",
                "questions": [
                    {"number": "1.1", "question": "Complete the concrete activity as instructed by the educator.", "marks": 8, "memo": ["Credit accurate observable actions and correct use of materials."]},
                    {"number": "1.2", "question": "Explain or show your thinking using words, numbers, drawings, or actions.", "marks": 6, "memo": ["Credit grade-appropriate reasoning and communication."]},
                    {"number": "1.3", "question": "Complete the short recorded response.", "marks": 6, "memo": ["Credit correct answer and visible working where appropriate."]},
                ],
            }
        ]
        rubric = [
            {"criterion": "Observable skill", "marks": 8, "descriptor": "Learner demonstrates the required skill using concrete materials or actions."},
            {"criterion": "Reasoning communication", "marks": 6, "descriptor": "Learner explains or shows thinking at grade level."},
            {"criterion": "Recorded response", "marks": 6, "descriptor": "Learner records a correct or partially correct response."},
        ]
    else:
        total = int(request.get("total_marks") or 50)
        first = max(10, total // 2)
        second = total - first
        sections = [
            {
                "title": "Structured Assessment Task",
                "mode": shell,
                "instructions": "Answer all questions. Use the provided stimulus and show working or reasoning where required.",
                "stimulus": "CAPS-aligned controlled classroom stimulus supplied by the educator.",
                "questions": [
                    {"number": "1.1", "question": f"Answer the first {subject_name} structured task using the stimulus.", "marks": first, "memo": ["Credit accurate subject knowledge and evidence-based reasoning."]},
                    {"number": "1.2", "question": f"Complete the applied {subject_name} task with justified reasoning.", "marks": second, "memo": ["Credit application, explanation, and grade-appropriate communication."]},
                ],
            }
        ]
        rubric = [
            {"criterion": "Subject knowledge", "marks": first, "descriptor": "Accurate CAPS-aligned knowledge."},
            {"criterion": "Application and reasoning", "marks": second, "descriptor": "Clear application and justified response."},
        ]

    return {
        "assessment_title": f"Grade {grade_value} {subject_name}: {opp_note} CAPS-Aligned Assessment",
        "subject": subject_name,
        "grade": grade_value,
        "phase": caps.get("phase") or grade.get("phase"),
        "canonical_profile_id": ontology.get("profile_id"),
        "blueprint": blueprint,
        "render_shell": shell,
        "language_of_assessment": language,
        "visual_blueprint": request.get("visual_blueprint") or {},
        "duration": f"{request.get('duration_hours', 2)} hours",
        "total_marks": total,
        "sections": sections,
        "rubric": rubric,
        "teacher_review_checklist": (
            [
                "Bevestig CAPS-onderwerpomvang en fasegeskiktheid.",
                "Bevestig dat alle punte korrek optel voor gebruik.",
                "Bevestig dat die assesseringsvorm by die vak en fase pas.",
                "Keur leerdergerigte bewoording en memorandum goed voor klasgebruik.",
            ]
            if language.lower() == "afrikaans"
            else [
                "Confirm CAPS topic scope and phase appropriateness.",
                "Confirm all marks reconcile before use.",
                "Confirm the assessment form matches the subject and phase.",
                "Approve learner-facing wording and memo before classroom use.",
            ]
        ),
        "generation_backend": "deterministic_shell_fallback",
    }


def sum_question_marks(pack: dict[str, Any]) -> int:
    return sum(int(q.get("marks") or 0) for section in pack.get("sections") or [] for q in section.get("questions") or [])


def sum_rubric_marks(pack: dict[str, Any]) -> int:
    return sum(int(item.get("marks") or 0) for item in pack.get("rubric") or [])


def normalize_assessment_pack(pack: dict[str, Any], request: dict[str, Any], opportunity: int) -> dict[str, Any]:
    subject = request.get("subject_profile") or {}
    grade = request.get("grade_profile") or {}
    caps = request.get("caps_context") or {}
    ontology = caps.get("assessment_ontology_profile") or {}
    shell = render_shell_from_request(request)
    normalized = dict(pack or {})
    normalized.setdefault("assessment_title", f"Grade {grade.get('grade')} {subject.get('display_name')}: CAPS-Aligned Assessment")
    normalized["subject"] = str(normalized.get("subject") or subject.get("display_name") or "")
    normalized["grade"] = int(normalized.get("grade") or grade.get("grade") or request.get("grade") or 0)
    normalized["phase"] = str(normalized.get("phase") or caps.get("phase") or grade.get("phase") or "")
    normalized["canonical_profile_id"] = normalized.get("canonical_profile_id") or ontology.get("profile_id")
    normalized["blueprint"] = normalized.get("blueprint") or ontology.get("assessment_family") or "structured_test_or_task"
    normalized["render_shell"] = normalized.get("render_shell") or shell
    normalized["language_of_assessment"] = normalized.get("language_of_assessment") or assessment_language_for_request(request)
    normalized["visual_blueprint"] = normalized.get("visual_blueprint") or request.get("visual_blueprint") or {}
    normalized.setdefault("duration", f"{request.get('duration_hours', 2)} hours")
    normalized["sections"] = list(normalized.get("sections") or [])
    normalized["rubric"] = list(normalized.get("rubric") or [])
    normalized["teacher_review_checklist"] = list(normalized.get("teacher_review_checklist") or [])
    normalized["opportunity"] = opportunity
    return normalized


BANNED_PACK_PHRASES = [
    "demo purposes",
    "adjust if",
    "if strict",
    "implicit",
    "corrected:",
    "scaled to",
    "as per opportunity",
    "pipeline",
    "design law",
    "assessment family",
    "default-off",
    "strongest assessment question",
    "sterkste assesseringsvraag",
    "identify one important feature of the",
    "identifiseer een belangrike kenmerk van die",
    "how it could be used as evidence",
    "hoe dit as bewys in afrikaans gebruik kan word",
    "official-paper source exemplars",
    "amptelike vraestel-bronvoorbeelde",
    "generic generated visuals",
    "generiese gegenereerde visuele materiaal",
    "kontroles wat 'n opvoeder moet uitvoer voordat hierdie assessering",
]

NO_CONTEXT_PACK_PHRASES = [
    "answer the first",
    "complete the applied",
    "using the stimulus",
    "use the provided stimulus",
    "caps-aligned controlled classroom stimulus supplied by the educator",
    "structured task using the stimulus",
    "credit accurate subject knowledge",
    "evidence-based reasoning",
]

AFRIKAANS_REQUIRED_HINTS = [
    "beantwoord",
    "vrae",
    "teks",
    "bron",
    "punte",
    "memo",
    "taal",
    "leerder",
    "opvoeder",
]

AFRIKAANS_BANNED_ENGLISH_PHRASES = [
    "answer all questions",
    "use the provided",
    "use the supplied",
    "credit accurate",
    "subject knowledge",
    "teacher review",
    "assessment task",
    "source interpretation",
    "method and review",
    "text a",
    "visual text b",
]

GENERIC_MEMO_PHRASES = [
    "credit a valid",
    "credit accurate",
    "ken punte toe vir",
    "relevante kenmerk word geidentifiseer",
    "antwoord toon begrip",
]


def is_memorandum_language_paper(pack: dict[str, Any], request: dict[str, Any]) -> bool:
    subject_id = subject_id_from_request(request)
    shell = str(pack.get("render_shell") or render_shell_from_request(request))
    return subject_id in {"english_language", "afrikaans_language"} and shell == "language_integrated_task"


def validate_afrikaans_language_pack(pack: dict[str, Any], request: dict[str, Any]) -> list[str]:
    if subject_id_from_request(request) != "afrikaans_language":
        return []
    errors: list[str] = []
    cards = list(pack.get("evidence_cards") or [])
    labels = [str(card.get("label") or "") for card in cards]
    if labels[:2] != ["Teks A", "Teks B"]:
        errors.append(f"Afrikaans source contract requires Teks A and Teks B, got {labels[:2]}")
    minimum_words = [350, 140]
    for index, minimum in enumerate(minimum_words):
        if index >= len(cards):
            break
        content = str(cards[index].get("content") or "")
        words = re.findall(r"\b[\w'-]+\b", content, flags=re.UNICODE)
        if len(words) < minimum:
            errors.append(f"{labels[index] or f'Teks {index + 1}'} is too short: {len(words)} words, minimum {minimum}")
        source_blob = " ".join(str(cards[index].get(key) or "") for key in ["type", "title", "content"]).lower()
        if any(term in source_blob for term in ["description of", "beskrywing van 'n grafiek", "beskrywing van die infografiek"]):
            errors.append(f"{labels[index] or f'Teks {index + 1}'} is a description rather than a complete learner-facing text")

    sections = list(pack.get("sections") or [])
    section_titles = " ".join(str(section.get("title") or "") for section in sections).lower()
    for required in ["begrip", "opsom", "taal"]:
        if required not in section_titles:
            errors.append(f"Afrikaans paper is missing the required {required} section")
    questions = [question for section in sections for question in section.get("questions") or []]
    if len(questions) < 12:
        errors.append(f"Afrikaans paper has only {len(questions)} questions; at least 12 substantive subquestions are required")
    for section in sections:
        is_summary = "opsom" in str(section.get("title") or "").lower()
        for question in section.get("questions") or []:
            marks = int(question.get("marks") or 0)
            if marks > (10 if is_summary else 5):
                errors.append(f"question {question.get('number')} has implausible short-question marks: {marks}")
            memo_blob = " ".join(str(item) for item in question.get("memo") or []).lower()
            if not memo_blob:
                errors.append(f"question {question.get('number')} has no specific memorandum answer")
            elif any(phrase in memo_blob for phrase in GENERIC_MEMO_PHRASES):
                errors.append(f"question {question.get('number')} uses a generic marking instruction instead of an answer")
    if pack.get("rubric"):
        errors.append("memorandum-based Afrikaans controlled test must not carry a duplicate full-paper rubric")
    return errors


def validate_english_language_pack(pack: dict[str, Any], request: dict[str, Any]) -> list[str]:
    if subject_id_from_request(request) != "english_language":
        return []
    errors: list[str] = []
    cards = list(pack.get("evidence_cards") or [])
    labels = [str(card.get("label") or "") for card in cards]
    if labels[:2] != ["Text A", "Text B"]:
        errors.append(f"English source contract requires Text A and Text B, got {labels[:2]}")
    for index, minimum in enumerate([350, 140]):
        if index >= len(cards):
            break
        content = str(cards[index].get("content") or "")
        words = re.findall(r"\b[\w'-]+\b", content, flags=re.UNICODE)
        if len(words) < minimum:
            errors.append(f"{labels[index] or f'Text {index + 1}'} is too short: {len(words)} words, minimum {minimum}")
        source_blob = " ".join(str(cards[index].get(key) or "") for key in ["type", "title", "content"]).lower()
        if any(term in source_blob for term in ["description of a graph", "description of the infographic", "image showing"]):
            errors.append(f"{labels[index] or f'Text {index + 1}'} is a description rather than a complete learner-facing text")
    sections = list(pack.get("sections") or [])
    section_titles = " ".join(str(section.get("title") or "") for section in sections).lower()
    for required in ["comprehension", "summary", "language"]:
        if required not in section_titles:
            errors.append(f"English paper is missing the required {required} section")
    questions = [question for section in sections for question in section.get("questions") or []]
    if len(questions) < 12:
        errors.append(f"English paper has only {len(questions)} questions; at least 12 substantive subquestions are required")
    for section in sections:
        is_summary = "summary" in str(section.get("title") or "").lower()
        for question in section.get("questions") or []:
            marks = int(question.get("marks") or 0)
            if marks > (10 if is_summary else 5):
                errors.append(f"question {question.get('number')} has implausible short-question marks: {marks}")
            memo_blob = " ".join(str(item) for item in question.get("memo") or []).lower()
            if not memo_blob:
                errors.append(f"question {question.get('number')} has no specific memorandum answer")
            elif any(phrase in memo_blob for phrase in GENERIC_MEMO_PHRASES):
                errors.append(f"question {question.get('number')} uses a generic marking instruction instead of an answer")
    if pack.get("rubric"):
        errors.append("memorandum-based English controlled test must not carry a duplicate full-paper rubric")
    return errors


def validate_assessment_pack(pack: dict[str, Any], request: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    total = int(pack.get("total_marks") or 0)
    question_total = sum_question_marks(pack)
    rubric_total = sum_rubric_marks(pack)
    if not total:
        errors.append("total_marks is missing or zero")
    if question_total != total:
        errors.append(f"question marks sum to {question_total}, expected {total}")
    memorandum_language_paper = is_memorandum_language_paper(pack, request)
    if not memorandum_language_paper and rubric_total != total:
        errors.append(f"rubric marks sum to {rubric_total}, expected {total}")
    if not pack.get("sections"):
        errors.append("sections are missing")
    if not memorandum_language_paper and not pack.get("rubric"):
        errors.append("rubric is missing")
    learner_facing = {
        "assessment_title": pack.get("assessment_title"),
        "subject": pack.get("subject"),
        "grade": pack.get("grade"),
        "duration": pack.get("duration"),
        "instructions": pack.get("instructions"),
        "sections": pack.get("sections"),
        "rubric": pack.get("rubric"),
        "teacher_review_checklist": pack.get("teacher_review_checklist"),
    }
    blob = json.dumps(learner_facing, ensure_ascii=True).lower()
    for phrase in BANNED_PACK_PHRASES:
        if phrase in blob:
            errors.append(f"banned internal/uncertain phrase found: {phrase}")
    for phrase in NO_CONTEXT_PACK_PHRASES:
        if phrase in blob:
            errors.append(f"no-context placeholder phrase found: {phrase}")
    if assessment_language_for_request(request).lower() == "afrikaans":
        visible_blob = json.dumps(
            {
                "assessment_title": pack.get("assessment_title"),
                "sections": pack.get("sections"),
                "rubric": pack.get("rubric"),
                "teacher_review_checklist": pack.get("teacher_review_checklist"),
            },
            ensure_ascii=False,
        ).lower()
        hint_count = sum(1 for hint in AFRIKAANS_REQUIRED_HINTS if hint in visible_blob)
        if hint_count < 4:
            errors.append("Afrikaans language contract failed: not enough Afrikaans learner-facing text")
        for phrase in AFRIKAANS_BANNED_ENGLISH_PHRASES:
            if phrase in visible_blob:
                errors.append(f"Afrikaans language contract failed: English phrase found: {phrase}")
    shell = str(pack.get("render_shell") or render_shell_from_request(request))
    if shell == "performance_task_sheet" and "answer the following questions" in blob:
        errors.append("performance task contains written-paper question phrasing")
    if shell == "performance_task_sheet" and "essay" in blob:
        errors.append("performance task contains essay language")
    errors.extend(validate_afrikaans_language_pack(pack, request))
    errors.extend(validate_english_language_pack(pack, request))
    visual_blueprint = pack.get("visual_blueprint") or request.get("visual_blueprint") or {}
    visual_kinds = {str(item.get("visual_kind") or "") for item in visual_blueprint.get("required_visuals") or []}
    language_pack = (
        "language" in str(pack.get("canonical_profile_id") or "").lower()
        or "language_integrated" in str(pack.get("blueprint") or "").lower()
        or "taal" in str(pack.get("blueprint") or "").lower()
    )
    references_map = bool(re.search(r"\bmap(?:s|work)?\b", blob))
    references_graph = bool(re.search(r"\bgraph(?:s)?\b", blob))
    if not language_pack and references_map and not any("map" in kind for kind in visual_kinds):
        errors.append("pack references a map but the visual blueprint has no map visual")
    if not language_pack and references_graph and not any(kind in visual_kinds for kind in ["data_table", "axis_diagram", "investigation_sheet"]):
        errors.append("pack references a graph but the visual blueprint has no graph/data visual")
    return errors


def write_validation_report(path: Path, errors: list[str], pack: dict[str, Any]) -> Path:
    lines = [
        "# HyMark Assessment Validation",
        "",
        f"Status: {'passed' if not errors else 'failed'}",
        f"Total marks: {pack.get('total_marks')}",
        f"Question marks: {sum_question_marks(pack)}",
        f"Rubric marks: {sum_rubric_marks(pack)}",
        "",
    ]
    if errors:
        lines.append("## Errors")
        lines.extend(f"- {error}" for error in errors)
    else:
        lines.append("No validation errors.")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def create_shell_memo_docx(pack: dict[str, Any], output_path: Path) -> Path:
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(22)
    section.right_margin = Mm(22)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    is_afrikaans = "afrikaans" in str(pack.get("subject", "")).lower()
    labels = {
        "memo_title": "MEMORANDUM" if is_afrikaans else "MEMORANDUM / MARKING GUIDE",
        "assessment": "Assessering" if is_afrikaans else "Assessment",
        "total": "Totaal" if is_afrikaans else "Total",
        "marks": "punte" if is_afrikaans else "marks",
        "section": "AFDELING" if is_afrikaans else "SECTION",
        "rubric": "NASIENRUBRIEK" if is_afrikaans else "RUBRIC",
        "criterion": "Kriterium" if is_afrikaans else "Criterion",
        "descriptor": "Beskrywing" if is_afrikaans else "Descriptor",
        "educator_review": "OPVOEDERHERSIENING" if is_afrikaans else "EDUCATOR REVIEW",
    }
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{pack.get('subject', 'Assessment')} - {labels['memo_title']}")
    run.bold = True
    run.font.size = Pt(15)
    doc.add_paragraph(f"{labels['assessment']}: {pack.get('assessment_title', '')}")
    doc.add_paragraph(f"{labels['total']}: {pack.get('total_marks', '')} {labels['marks']}")
    for section_index, section in enumerate(pack.get("sections") or [], start=1):
        section_title = str(section.get("title") or f"{labels['section']} {section_index}")
        doc.add_heading(section_title, level=1)
        for question in section.get("questions") or []:
            p = doc.add_paragraph()
            p.add_run(f"{question.get('number', '')}. ").bold = True
            p.add_run(str(question.get("question", "")))
            p.add_run(f" ({question.get('marks', 0)})").bold = True
            for memo_item in question.get("memo") or []:
                doc.add_paragraph(str(memo_item), style="List Bullet")
    if pack.get("rubric"):
        doc.add_heading(labels["rubric"], level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for index, header in enumerate([labels["criterion"], labels["marks"].capitalize(), labels["descriptor"]]):
            table.rows[0].cells[index].text = header
        for item in pack["rubric"]:
            cells = table.add_row().cells
            cells[0].text = str(item.get("criterion", ""))
            cells[1].text = str(item.get("marks", ""))
            cells[2].text = str(item.get("descriptor", ""))
    doc.add_heading(labels["educator_review"], level=1)
    for item in pack.get("teacher_review_checklist") or []:
        doc.add_paragraph(str(item), style="List Bullet")
    doc.save(output_path)
    return output_path


async def build_evidence_first_opportunity(
    hymark: Any,
    request: dict[str, Any],
    opp: int,
    job_dir: Path,
    design_law_path: Path,
    assessment_design_path: Path,
) -> dict[str, Any]:
    apply_evidence_route_overrides(request)
    opp_label = "1st" if opp == 1 else "2nd"
    opp_suffix = "1stOpp" if opp == 1 else "2ndOpp"
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    opportunity_dir = job_dir / opp_suffix
    opportunity_dir.mkdir(parents=True, exist_ok=True)
    evidence_profile = evidence_profile_for_request(request)
    status = "generated"
    fallback_reason = ""
    cards: list[dict[str, Any]] = []
    try:
        locked_cards_path = Path(str(request.get("locked_evidence_cards_path") or "")).expanduser()
        if request.get("locked_evidence_cards_path") and locked_cards_path.exists():
            cards = json.loads(locked_cards_path.read_text(encoding="utf-8"))
            cards = normalize_evidence_cards_for_language(cards, request)
        else:
            cards = generate_evidence_cards(hymark, request, opp, evidence_profile)
        (opportunity_dir / "HOMS_EVIDENCE_CARDS.json").write_text(json.dumps(cards, indent=2), encoding="utf-8")
        subject_id = subject_id_from_request(request)
        split_language = subject_id in {"afrikaans_language", "english_language"}
        if subject_id == "afrikaans_language":
            raw_pack = model_afrikaans_evidence_pack(hymark, request, opp, cards, opportunity_dir)
        elif subject_id == "english_language":
            raw_pack = model_english_evidence_pack(hymark, request, opp, cards, opportunity_dir)
        else:
            raw_pack = model_evidence_first_pack(hymark, request, opp, evidence_profile, cards)
        if not split_language:
            (opportunity_dir / "HOMS_PROVIDER_PACK_ATTEMPT_1.json").write_text(json.dumps(raw_pack, indent=2), encoding="utf-8")
        pack = normalize_assessment_pack(raw_pack, request, opp)
        pack["evidence_cards"] = cards or list(pack.get("evidence_cards") or [])
        pack["evidence_policy"] = evidence_profile.get("evidence_policy")
        pack["generation_backend"] = str(raw_pack.get("generation_backend") or "hymark_evidence_first_caps")
        errors = validate_assessment_pack(pack, request)
        errors.extend(validate_evidence_grounding(pack, request, evidence_profile))
        if errors and not split_language:
            status = "provider_repair"
            repaired = repair_evidence_first_pack(hymark, request, opp, evidence_profile, cards, pack, errors)
            (opportunity_dir / "HOMS_PROVIDER_PACK_ATTEMPT_2.json").write_text(json.dumps(repaired, indent=2), encoding="utf-8")
            pack = normalize_assessment_pack(repaired, request, opp)
            pack["evidence_cards"] = cards
            pack["evidence_policy"] = evidence_profile.get("evidence_policy")
            pack["generation_backend"] = "hymark_evidence_first_caps_repaired"
            errors = validate_assessment_pack(pack, request)
            errors.extend(validate_evidence_grounding(pack, request, evidence_profile))
            if errors:
                failure = {
                    "status": "blocked_invalid_provider_pack",
                    "provider_model": os.environ.get("HOMS_AI_MODEL"),
                    "errors": sorted(set(errors)),
                    "policy": "No deterministic learner-paper fallback is permitted.",
                }
                (opportunity_dir / "HOMS_EVIDENCE_GENERATION_BLOCKED.json").write_text(json.dumps(failure, indent=2), encoding="utf-8")
                raise RuntimeError("provider assessment failed educational validation after repair: " + "; ".join(sorted(set(errors))))
        elif errors:
            failure = {
                "status": "blocked_invalid_split_pack",
                "provider_model": os.environ.get("HOMS_AI_MODEL"),
                "errors": sorted(set(errors)),
                "policy": "No deterministic learner-paper fallback is permitted.",
            }
            (opportunity_dir / "HOMS_EVIDENCE_GENERATION_BLOCKED.json").write_text(json.dumps(failure, indent=2), encoding="utf-8")
            raise RuntimeError("split language assessment failed educational validation: " + "; ".join(sorted(set(errors))))
    except Exception as exc:
        failure_path = opportunity_dir / "HOMS_EVIDENCE_GENERATION_BLOCKED.json"
        if not failure_path.exists():
            failure_path.write_text(
                json.dumps(
                    {
                        "status": "blocked_generation_error",
                        "provider_model": os.environ.get("HOMS_AI_MODEL"),
                        "error": str(exc),
                        "policy": "No deterministic learner-paper fallback is permitted.",
                    },
                    indent=2,
                ),
                encoding="utf-8",
            )
        raise

    pack["validation_errors"] = sorted(set(errors))
    (opportunity_dir / "assessment_pack.json").write_text(json.dumps(pack, indent=2), encoding="utf-8")
    (opportunity_dir / "HOMS_EVIDENCE_CONSTRUCTION_PROFILE.json").write_text(json.dumps(evidence_profile, indent=2), encoding="utf-8")
    write_validation_report(opportunity_dir / "HYMARK_ASSESSMENT_VALIDATION.md", pack["validation_errors"], pack)
    design_receipt = apply_design_law(opportunity_dir, design_law_path, assessment_design_path)
    design_law = load_design_json(design_law_path)
    design_payload = load_design_json(assessment_design_path) if assessment_design_path.exists() else {}
    assessment_design = assessment_design_for_pack(pack, design_payload)
    learner_docx = render_docx(
        opportunity_dir,
        pack,
        design_law,
        design_receipt.get("assets") or [],
        assessment_design,
        include_memo_sections=False,
        output_name=f"ASSESSMENT_LEARNER_{opp_suffix}.docx",
    )

    exam_filename = f"{request['module_code']}_Assessment_{opp_suffix}_{timestamp}.docx"
    memo_filename = f"{request['module_code']}_Memo_{opp_suffix}_{timestamp}.docx"
    formatted_source = Path(design_receipt["formatted_docx"])
    shutil.copy2(learner_docx, job_dir / exam_filename)
    create_shell_memo_docx(pack, job_dir / memo_filename)

    return {
        "module_code": request["module_code"],
        "module_name": request["module_name"],
        "total_marks": pack["total_marks"],
        "duration_hours": request["duration_hours"],
        "opportunity": opp,
        "opportunity_label": f"{opp_label} Opportunity",
        "generation_backend": pack.get("generation_backend"),
        "status": status if not pack["validation_errors"] else "failed_validation",
        "render_shell": pack.get("render_shell"),
        "fallback_reason": fallback_reason,
        "evidence_policy": evidence_profile.get("evidence_policy"),
        "evidence_cards": len(cards),
        "assessment_pack": pack,
        "calculated_total": sum_question_marks(pack),
        "rubric_total": sum_rubric_marks(pack),
        "validation_errors": pack["validation_errors"],
        "opportunity_dir": str(opportunity_dir),
        "filename": exam_filename,
        "memo_filename": memo_filename,
        "learner_docx": str(learner_docx),
        "formatted_review_docx": str(formatted_source),
        "formatted_review_zip": design_receipt.get("formatted_zip"),
        "created_at": utc_now(),
    }


async def build_caps_shell_opportunity(
    hymark: Any,
    request: dict[str, Any],
    opp: int,
    job_dir: Path,
    design_law_path: Path,
    assessment_design_path: Path,
) -> dict[str, Any]:
    opp_label = "1st" if opp == 1 else "2nd"
    opp_suffix = "1stOpp" if opp == 1 else "2ndOpp"
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    opportunity_dir = job_dir / opp_suffix
    opportunity_dir.mkdir(parents=True, exist_ok=True)
    status = "generated"
    fallback_reason = ""
    try:
        raw_pack = model_shell_pack(hymark, request, opp)
        pack = normalize_assessment_pack(raw_pack, request, opp)
        errors = validate_assessment_pack(pack, request)
        if errors:
            status = "fallback_after_validation"
            fallback_reason = "; ".join(errors)
            pack = fallback_shell_pack(request, opp, fallback_reason)
            errors = validate_assessment_pack(pack, request)
    except Exception as exc:
        status = "fallback_after_generation_error"
        fallback_reason = str(exc)
        pack = fallback_shell_pack(request, opp, fallback_reason)
        errors = validate_assessment_pack(pack, request)

    pack["validation_errors"] = errors
    (opportunity_dir / "assessment_pack.json").write_text(json.dumps(pack, indent=2), encoding="utf-8")
    write_validation_report(opportunity_dir / "HYMARK_ASSESSMENT_VALIDATION.md", errors, pack)
    design_receipt = apply_design_law(opportunity_dir, design_law_path, assessment_design_path)
    design_law = load_design_json(design_law_path)
    design_payload = load_design_json(assessment_design_path) if assessment_design_path.exists() else {}
    assessment_design = assessment_design_for_pack(pack, design_payload)
    learner_docx = render_docx(
        opportunity_dir,
        pack,
        design_law,
        design_receipt.get("assets") or [],
        assessment_design,
        include_memo_sections=False,
        output_name=f"ASSESSMENT_LEARNER_{opp_suffix}.docx",
    )

    exam_filename = f"{request['module_code']}_Assessment_{opp_suffix}_{timestamp}.docx"
    memo_filename = f"{request['module_code']}_Memo_{opp_suffix}_{timestamp}.docx"
    formatted_source = Path(design_receipt["formatted_docx"])
    shutil.copy2(learner_docx, job_dir / exam_filename)
    create_shell_memo_docx(pack, job_dir / memo_filename)

    return {
        "module_code": request["module_code"],
        "module_name": request["module_name"],
        "total_marks": pack["total_marks"],
        "duration_hours": request["duration_hours"],
        "opportunity": opp,
        "opportunity_label": f"{opp_label} Opportunity",
        "generation_backend": "caps_shell",
        "status": status if not errors else "failed_validation",
        "render_shell": pack.get("render_shell"),
        "fallback_reason": fallback_reason,
        "assessment_pack": pack,
        "calculated_total": sum_question_marks(pack),
        "rubric_total": sum_rubric_marks(pack),
        "validation_errors": errors,
        "opportunity_dir": str(opportunity_dir),
        "filename": exam_filename,
        "memo_filename": memo_filename,
        "learner_docx": str(learner_docx),
        "formatted_review_docx": str(formatted_source),
        "formatted_review_zip": design_receipt.get("formatted_zip"),
        "created_at": utc_now(),
    }


async def build_opportunity(hymark: Any, request: dict[str, Any], opp: int, job_dir: Path) -> dict[str, Any]:
    opp_label = "1st" if opp == 1 else "2nd"
    exam_data: dict[str, Any] = {
        "module_code": request["module_code"],
        "module_name": request["module_name"],
        "total_marks": request["total_marks"],
        "duration_hours": request["duration_hours"],
        "opportunity": opp,
        "opportunity_label": f"{opp_label} Opportunity",
        "source_questions": [],
        "methodology_question": None,
        "essay_question": None,
        "created_at": utc_now(),
    }
    use_generic_caps = bool(request.get("caps_context")) and subject_id_from_request(request) != "history"
    exam_data["generation_backend"] = "caps_generic" if use_generic_caps else "hymark_history"

    for index, topic in enumerate(request["topics"]):
        prompt_topic = generation_topic(topic, request)
        if use_generic_caps:
            sources_data = await generate_caps_sources(hymark, prompt_topic, index + 1, opp, request)
        else:
            sources_data = await hymark.generate_exam_sources(prompt_topic, index + 1, opportunity=opp)
        sources = [ensure_source_shape(source, i) for i, source in enumerate(sources_data.get("sources", []))]
        sources = ensure_sources_available(sources, topic, request)
        if use_generic_caps:
            questions_data = await generate_caps_source_questions(hymark, sources, prompt_topic, 25, request)
        else:
            questions_data = await hymark.generate_source_questions(sources, prompt_topic, total_marks=25)
        questions = [
            ensure_question_shape(question, q_index, sources[0]["label"] if sources else "Source A")
            for q_index, question in enumerate(questions_data.get("questions", []))
        ]
        exam_data["source_questions"].append(
            {
                "question_number": index + 1,
                "topic": topic,
                "generation_topic": prompt_topic,
                "sources": sources,
                "questions": questions,
                "total_marks": 25,
            }
        )

    exam_data["caps_context"] = request.get("caps_context")
    if use_generic_caps:
        exam_data["methodology_question"] = await generate_caps_applied_question(
            hymark,
            generation_topic(request["methodology_topic"], request),
            25,
            request,
        )
        exam_data["essay_question"] = await generate_caps_extended_question(
            hymark,
            generation_topic(request["essay_topic"], request),
            50,
            opp,
            request,
        )
    else:
        exam_data["methodology_question"] = await hymark.generate_methodology_question(
            generation_topic(request["methodology_topic"], request),
            marks=25,
        )
        exam_data["essay_question"] = await hymark.generate_essay_question(
            generation_topic(request["essay_topic"], request),
            marks=50,
            opportunity=opp,
        )
    actual_total = sum(sq.get("total_marks", 0) for sq in exam_data["source_questions"])
    actual_total += (exam_data.get("methodology_question") or {}).get("marks", 0)
    actual_total += (exam_data.get("essay_question") or {}).get("marks", 0)
    exam_data["calculated_total"] = actual_total

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    opp_suffix = "1stOpp" if opp == 1 else "2ndOpp"
    exam_filename = f"{request['module_code']}_Exam_{opp_suffix}_{timestamp}.docx"
    memo_filename = f"{request['module_code']}_Memo_{opp_suffix}_{timestamp}.docx"
    hymark.create_exam_docx(exam_data, job_dir / exam_filename)
    hymark.create_memorandum_docx(exam_data, job_dir / memo_filename)
    exam_data["filename"] = exam_filename
    exam_data["memo_filename"] = memo_filename
    return exam_data


def zip_dir(source: Path, target: Path) -> None:
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in source.rglob("*"):
            if path.is_file() and path.resolve() != target.resolve():
                zf.write(path, path.relative_to(source))


def write_summary(job_dir: Path, request: dict[str, Any], exam_set: dict[str, Any], provider: dict[str, Any]) -> Path:
    profile = request.get("subject_profile") or {}
    grade_profile = request.get("grade_profile") or {}
    caps_context = request.get("caps_context") or {}
    visual_blueprint = request.get("visual_blueprint") or {}
    selected_caps = caps_context.get("selected") or {}
    matrix = caps_context.get("assessment_matrix_row") or {}
    ontology = caps_context.get("assessment_ontology_profile") or {}
    summary = [
        "# HyMark Exam Builder Review Summary",
        "",
        f"Created: {utc_now()}",
        f"Module: {request['module_code']} - {request['module_name']}",
        f"Subject profile: {profile.get('display_name', 'Unspecified')}",
        f"Grade profile: Grade {grade_profile.get('grade', request.get('grade', 'unspecified'))} / {grade_profile.get('phase', 'unspecified')}",
        f"Provider: {provider['selected_provider']} / {provider['selected_model']}",
        "",
        "## Outputs",
    ]
    opportunity_labels = [label for label in ["first_opportunity", "second_opportunity"] if label in exam_set]
    for label in opportunity_labels:
        data = exam_set[label]
        summary.append(f"- {data['opportunity_label']}: `{data['filename']}` and `{data['memo_filename']}`")
    backend = exam_set["first_opportunity"].get("generation_backend", "hymark_history")
    summary.extend(["", "## Structure"])
    if str(backend).startswith("hymark_evidence_first") or str(backend).startswith("hymark_afrikaans_split"):
        for label in opportunity_labels:
            data = exam_set[label]
            errors = data.get("validation_errors") or []
            summary.extend(
                [
                    f"- {data['opportunity_label']}: evidence-first route / {data.get('total_marks')} marks / validation {'passed' if not errors else 'failed'}",
                    f"- Evidence cards: {data.get('evidence_cards', 0)} / policy `{data.get('evidence_policy', '')}`",
                    f"- Review pack: `{Path(data.get('opportunity_dir', '')).name}/ASSESSMENT_PACK_FORMATTED.docx`",
                ]
            )
            if errors:
                summary.extend(f"  - {error}" for error in errors)
        summary.append("- Evidence/stimulus cards are generated before questions, then rendered into the learner paper.")
        summary.append("- Separate memorandum/marking guide per opportunity.")
    elif backend == "caps_shell":
        for label in opportunity_labels:
            data = exam_set[label]
            errors = data.get("validation_errors") or []
            summary.extend(
                [
                    f"- {data['opportunity_label']}: `{data.get('render_shell', 'unknown')}` / {data.get('total_marks')} marks / validation {'passed' if not errors else 'failed'}",
                    f"- Review pack: `{Path(data.get('opportunity_dir', '')).name}/ASSESSMENT_PACK_FORMATTED.docx`",
                ]
            )
            if errors:
                summary.extend(f"  - {error}" for error in errors)
        summary.append("- Separate memorandum/marking guide per opportunity.")
        summary.append("- CAPS design-law formatted review pack per opportunity.")
    else:
        summary.extend(
            [
                "- Two source-based questions at 25 marks each.",
                "- One methodology question at 25 marks.",
                "- One essay question at 50 marks.",
                "- Separate memorandum/marking guide per opportunity.",
            ]
        )
    assessment_design = caps_context.get("assessment_design_profile") or {}
    summary.extend(
        [
            "",
            "## Subject Profile Controls",
            f"- Source types: {', '.join(profile.get('source_types', [])) or 'not specified'}",
            f"- Question families: {', '.join(profile.get('question_families', [])) or 'not specified'}",
            f"- Rubric dimensions: {', '.join(profile.get('rubric_dimensions', [])) or 'not specified'}",
            f"- Bloom targets: {', '.join(grade_profile.get('bloom_targets', [])) or 'not specified'}",
            f"- ZPD level: {grade_profile.get('zpd_level', 'not specified')}",
            f"- Question style: {', '.join(grade_profile.get('question_style', [])) or 'not specified'}",
            "",
            "## CAPS Source",
            f"- Selected: {selected_caps.get('title', 'not resolved')}",
            f"- Phase: {selected_caps.get('phase', caps_context.get('phase', 'not resolved'))}",
            f"- Path: {selected_caps.get('path', 'not resolved')}",
            f"- SHA-256: {selected_caps.get('sha256', 'not resolved')}",
            f"- Candidate count: {caps_context.get('candidate_count', 0)}",
            f"- Recommended blueprint: {matrix.get('recommended_blueprint', 'not resolved')}",
            f"- Format signals: {', '.join(k.replace('signals_', '') for k, v in matrix.items() if k.startswith('signals_') and v) or 'none detected'}",
            "",
            "## CAPS Ontology",
            f"- Profile: {ontology.get('profile_id', 'not resolved')}",
            f"- Canonical subject: {ontology.get('subject', 'not resolved')}",
            f"- Assessment family: {ontology.get('assessment_family', 'not resolved')}",
            f"- Allowed modes: {', '.join(ontology.get('allowed') or []) or 'not resolved'}",
            f"- Default-off modes: {', '.join(ontology.get('default_off') or []) or 'not resolved'}",
            "",
            "## CAPS Assessment Design",
            f"- Render shell: {assessment_design.get('render_shell', 'not resolved')}",
            f"- Assessment forms: {', '.join(assessment_design.get('assessment_forms') or []) or 'not resolved'}",
            f"- Marking instruments: {', '.join(assessment_design.get('marking_instruments') or []) or 'not resolved'}",
            "",
            "## Visual Blueprint",
            f"- Term: {visual_blueprint.get('term', request.get('term', 'auto'))}",
            f"- Required visuals: {', '.join(item.get('title', '') for item in visual_blueprint.get('required_visuals', [])) or 'not resolved'}",
            f"- CAPS theme hints: {'; '.join((visual_blueprint.get('caps_theme_candidates') or [])[:4]) or 'not extracted'}",
            "",
            "## Review Boundary",
            "Human subject expert verification is required before classroom or assessment use, especially source authenticity, mark allocation, and institutional format compliance.",
        ]
    )
    path = job_dir / "HYMARK_EXAM_REVIEW_SUMMARY.md"
    path.write_text("\n".join(summary).strip() + "\n", encoding="utf-8")
    return path


def run_builder(
    request_path: Path | None,
    out_root: Path,
    secret_file: Path,
    backend_path: Path,
    provider_name: str,
    model: str,
    subject_profile_name: str,
    subject_profile_path: Path | None,
    grade: int,
    grade_ladder_path: Path,
    caps_manifest_path: Path,
    caps_matrix_path: Path,
    caps_ontology_path: Path,
    caps_assessment_design_path: Path,
    design_law_path: Path,
    preferred_language: str,
    caps_extract_chars: int,
    skip_caps: bool,
    term: str,
    locked_evidence_cards_path: Path | None = None,
    opportunities: str = "both",
    locked_afrikaans_section_a_path: Path | None = None,
    locked_afrikaans_section_b_path: Path | None = None,
) -> dict[str, Any]:
    started = time.perf_counter()
    provider = configure_provider(secret_file, provider_name, model)
    hymark = load_hymark_backend(backend_path)
    subject_profile = load_subject_profile(subject_profile_name, subject_profile_path)
    grade_profile = load_grade_profile(grade, grade_ladder_path)
    effective_language = preferred_language
    if str(subject_profile.get("language_of_assessment") or "").lower() == "afrikaans":
        effective_language = "Afrikaans"
    caps_context = None
    if not skip_caps:
        caps_context = resolve_caps_context(
            subject_profile,
            grade_profile,
            caps_manifest_path,
            effective_language,
            include_policy=False,
            max_chars=caps_extract_chars,
            matrix_path=caps_matrix_path,
            ontology_path=caps_ontology_path,
            assessment_design_path=caps_assessment_design_path,
        )
        if subject_profile.get("subject_id") == "afrikaans_language" and not caps_context.get("selected"):
            raise RuntimeError("Afrikaans generation blocked: no authoritative Afrikaans CAPS document resolved.")
    request = normalized_request(request_path, subject_profile, grade_profile, caps_context, term)
    apply_subject_term_focus(request)
    if locked_evidence_cards_path:
        request["locked_evidence_cards_path"] = str(locked_evidence_cards_path)
    if locked_afrikaans_section_a_path:
        request["locked_afrikaans_section_a_path"] = str(locked_afrikaans_section_a_path)
    if locked_afrikaans_section_b_path:
        request["locked_afrikaans_section_b_path"] = str(locked_afrikaans_section_b_path)
    request["visual_blueprint"] = build_visual_blueprint(request)
    if should_use_evidence_first_route(request):
        apply_evidence_route_overrides(request)
    visual_titles = [
        f"{item.get('title')} ({item.get('role')})"
        for item in request["visual_blueprint"].get("required_visuals", [])
    ]
    visual_instruction = (
        "VISUAL BLUEPRINT CONTROL: "
        f"Term={request['visual_blueprint'].get('term')}. "
        f"Required visuals: {'; '.join(visual_titles)}. "
        "Questions must use these visuals as assessable evidence, not as decorative inserts."
    )
    request["additional_instructions"] = (
        str(request.get("additional_instructions") or "").strip() + "\n\n" + visual_instruction
    ).strip()

    out_root.mkdir(parents=True, exist_ok=True)
    job_id = (
        "hymark-exam-"
        + subject_profile["subject_id"].lower().replace(" ", "_")
        + "-grade"
        + str(grade_profile["grade"])
        + "-"
        + request["module_code"].lower()
        + "-"
        + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    )
    job_dir = out_root / job_id
    if job_dir.exists():
        shutil.rmtree(job_dir)
    job_dir.mkdir(parents=True)
    (job_dir / "exam_builder_request.json").write_text(json.dumps(request, indent=2), encoding="utf-8")
    (job_dir / "HOMS_VISUAL_BLUEPRINT.json").write_text(json.dumps(request["visual_blueprint"], indent=2), encoding="utf-8")

    use_evidence_first = bool(request.get("caps_context")) and should_use_evidence_first_route(request)
    use_caps_shell = bool(request.get("caps_context")) and subject_id_from_request(request) != "history" and not use_evidence_first
    opportunity_numbers = [1] if opportunities == "first" else ([2] if opportunities == "second" else [1, 2])
    if use_evidence_first:
        exam_set = {
            "first_opportunity" if number == 1 else "second_opportunity": asyncio.run(
                build_evidence_first_opportunity(hymark, request, number, job_dir, design_law_path, caps_assessment_design_path)
            )
            for number in opportunity_numbers
        }
    elif use_caps_shell:
        exam_set = {
            "first_opportunity" if number == 1 else "second_opportunity": asyncio.run(
                build_caps_shell_opportunity(hymark, request, number, job_dir, design_law_path, caps_assessment_design_path)
            )
            for number in opportunity_numbers
        }
    else:
        exam_set = {
            "first_opportunity" if number == 1 else "second_opportunity": asyncio.run(build_opportunity(hymark, request, number, job_dir))
            for number in opportunity_numbers
        }
    if "first_opportunity" not in exam_set:
        exam_set["first_opportunity"] = exam_set.pop("second_opportunity")
    (job_dir / "exam_set.json").write_text(json.dumps(exam_set, indent=2), encoding="utf-8")
    summary_path = write_summary(job_dir, request, exam_set, provider)

    zip_path = job_dir / "HYMARK_EXAM_BUILDER_PACK.zip"
    receipt = {
        "schema": "knowedge.hymark_exam_builder_receipt.v1",
        "created_at": utc_now(),
        "job_id": job_id,
        "status": "completed",
        "request": request,
        "subject_profile": {
            "subject_id": subject_profile["subject_id"],
            "display_name": subject_profile["display_name"],
            "profile_path": subject_profile.get("_profile_path"),
        },
        "grade_profile": {
            "grade": grade_profile["grade"],
            "phase": grade_profile["phase"],
            "learner_level": grade_profile["learner_level"],
            "bloom_targets": grade_profile["bloom_targets"],
            "ladder_path": grade_profile.get("_ladder_path"),
        },
        "caps_context": caps_context,
        "visual_blueprint": request.get("visual_blueprint"),
        "assessor": {
            "name": "HyMark Exam Builder",
            "backend_path": str(backend_path),
            "generation_backend": exam_set["first_opportunity"].get("generation_backend", "hymark_history"),
            "function_sources": (
                [
                    "generate_evidence_cards",
                    "model_evidence_first_pack",
                    "repair_evidence_first_pack",
                    "validate_afrikaans_language_pack",
                    "apply_design_law",
                ]
                if use_evidence_first
                else [
                    "model_shell_pack",
                    "validate_assessment_pack",
                    "apply_design_law",
                    "create_shell_memo_docx",
                ]
                if use_caps_shell
                else [
                    "generate_exam_sources",
                    "generate_source_questions",
                    "generate_methodology_question",
                    "generate_essay_question",
                    "create_exam_docx",
                    "create_memorandum_docx",
                ]
            ),
        },
        "provider": provider,
        "duration_seconds": round(time.perf_counter() - started, 3),
        "outputs": {
            "job_dir": str(job_dir),
            "request": str(job_dir / "exam_builder_request.json"),
            "visual_blueprint": str(job_dir / "HOMS_VISUAL_BLUEPRINT.json"),
            "exam_set": str(job_dir / "exam_set.json"),
            "summary": str(summary_path),
            "review_zip": str(zip_path),
            "first_exam": str(job_dir / exam_set["first_opportunity"]["filename"]),
            "first_memo": str(job_dir / exam_set["first_opportunity"]["memo_filename"]),
            "second_exam": str(job_dir / exam_set["second_opportunity"]["filename"]) if "second_opportunity" in exam_set else None,
            "second_memo": str(job_dir / exam_set["second_opportunity"]["memo_filename"]) if "second_opportunity" in exam_set else None,
            "first_formatted_review_pack": exam_set["first_opportunity"].get("formatted_review_docx"),
            "first_formatted_review_zip": exam_set["first_opportunity"].get("formatted_review_zip"),
            "second_formatted_review_pack": exam_set["second_opportunity"].get("formatted_review_docx") if "second_opportunity" in exam_set else None,
            "second_formatted_review_zip": exam_set["second_opportunity"].get("formatted_review_zip") if "second_opportunity" in exam_set else None,
        },
    }
    (job_dir / "HYMARK_EXAM_BUILDER_RECEIPT.json").write_text(json.dumps(receipt, indent=2), encoding="utf-8")
    zip_dir(job_dir, zip_path)
    return receipt


def main() -> int:
    parser = argparse.ArgumentParser(description="Run HyMark Exam Builder without the web UI.")
    parser.add_argument("--request", default="", help="Optional JSON request file.")
    parser.add_argument("--out", default=str(ROOT / "campaigns" / "phase3" / "homs" / "exam_builder" / "proof"))
    parser.add_argument("--secret-file", default=str(DEFAULT_SECRET_FILE))
    parser.add_argument("--backend", default=str(DEFAULT_HYMARK_BACKEND))
    parser.add_argument("--provider", default="nim", choices=["nim", "nvidia", "nvidia_nim", "openai"])
    parser.add_argument("--model", default="")
    parser.add_argument("--subject-profile", default="history", help="Subject profile name from config/homs_subject_profiles.")
    parser.add_argument("--subject-profile-path", default="", help="Optional explicit subject profile JSON path.")
    parser.add_argument("--grade", type=int, default=12, help="School grade level, 1 through 12.")
    parser.add_argument("--term", default="auto", help="CAPS term focus, e.g. 1, 2, 3, 4, or auto.")
    parser.add_argument("--grade-ladder", default=str(DEFAULT_GRADE_LADDER), help="Grade ladder JSON path.")
    parser.add_argument("--caps-manifest", default=str(DEFAULT_CAPS_MANIFEST), help="Official DBE CAPS corpus manifest path.")
    parser.add_argument("--caps-matrix", default=str(DEFAULT_CAPS_MATRIX), help="CAPS assessment matrix analysis JSON path.")
    parser.add_argument("--caps-ontology", default=str(DEFAULT_CAPS_ONTOLOGY), help="Canonical CAPS assessment ontology JSON path.")
    parser.add_argument("--caps-assessment-design", default=str(DEFAULT_CAPS_ASSESSMENT_DESIGN), help="CAPS assessment-design profile JSON path.")
    parser.add_argument("--design-law", default=str(DEFAULT_DESIGN_LAW), help="HOMS design-law formatting JSON path.")
    parser.add_argument("--preferred-language", default="English", help="Preferred CAPS document language.")
    parser.add_argument("--caps-extract-chars", type=int, default=3500, help="Maximum CAPS excerpt characters injected into the generation note.")
    parser.add_argument("--skip-caps", action="store_true", help="Do not resolve CAPS context for this run.")
    parser.add_argument("--locked-evidence-cards", default="", help="Reuse a reviewed HOMS_EVIDENCE_CARDS.json file.")
    parser.add_argument("--opportunities", default="both", choices=["first", "second", "both"], help="Generate one opportunity or both.")
    parser.add_argument("--locked-afrikaans-section-a", default="", help="Reuse a reviewed Afrikaans Section A JSON response.")
    parser.add_argument("--locked-afrikaans-section-b", default="", help="Reuse a reviewed Afrikaans Section B JSON response.")
    args = parser.parse_args()

    receipt = run_builder(
        Path(args.request).expanduser().resolve() if args.request else None,
        Path(args.out).expanduser().resolve(),
        Path(args.secret_file).expanduser().resolve(),
        Path(args.backend).expanduser().resolve(),
        args.provider,
        args.model,
        args.subject_profile,
        Path(args.subject_profile_path).expanduser().resolve() if args.subject_profile_path else None,
        args.grade,
        Path(args.grade_ladder).expanduser().resolve(),
        Path(args.caps_manifest).expanduser().resolve(),
        Path(args.caps_matrix).expanduser().resolve(),
        Path(args.caps_ontology).expanduser().resolve(),
        Path(args.caps_assessment_design).expanduser().resolve(),
        Path(args.design_law).expanduser().resolve(),
        args.preferred_language,
        args.caps_extract_chars,
        args.skip_caps,
        args.term,
        Path(args.locked_evidence_cards).expanduser().resolve() if args.locked_evidence_cards else None,
        args.opportunities,
        Path(args.locked_afrikaans_section_a).expanduser().resolve() if args.locked_afrikaans_section_a else None,
        Path(args.locked_afrikaans_section_b).expanduser().resolve() if args.locked_afrikaans_section_b else None,
    )
    print(json.dumps({"status": receipt["status"], "assessor": receipt["assessor"]["name"], "provider": receipt["provider"]["selected_provider"], "outputs": receipt["outputs"]}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
