from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from presence_core import shared_production_finish as finish


def _source(root: Path, text: str, name: str = "assessment.md") -> dict:
    path = root / name
    path.write_text(text, encoding="utf-8")
    return {
        "artifact_id": "NATIVE-REPORT",
        "kind": "assessment_report",
        "path": str(path),
        "file_name": name,
        "mime_type": "text/markdown",
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "release_state": "HELD",
    }


def _render(root: Path, artifacts: list[dict]) -> dict:
    return finish.render_held_artifacts_for_review(
        state_root=root / "finished",
        case_id="CASE-QUALITY-001",
        source_sha256="1" * 64,
        product_name="Finance Readiness",
        artifacts=artifacts,
    )


def test_finishing_renders_native_findings_through_real_format_core(tmp_path: Path) -> None:
    artifact = _source(tmp_path, "# Finance findings\n\nFR-03: Bank statement missing.\nThe supplied total is R 42,750; accountant review remains required.")
    result = _render(tmp_path, [artifact])
    by_type = {row["mime_type"]: row for row in result["artifacts"]}
    html = Path(by_type["text/html"]["path"]).read_text()
    docx = Document(by_type["application/vnd.openxmlformats-officedocument.wordprocessingml.document"]["path"])
    text = "\n".join(p.text for p in docx.paragraphs)
    for expected in ("FR-03: Bank statement missing.", "R 42,750", "accountant review remains required"):
        assert expected in html
        assert expected in text
    manifest = json.loads(Path(by_type["application/json"]["path"]).read_text())
    assert manifest["format_core"]["receipt"]["schema"] == "dio.format_core.receipt.v1"
    assert manifest["format_core"]["receipt"]["qa"]["passed"] is True
    assert manifest["source_artifacts"][0]["sha256"] == artifact["sha256"]
    assert all(row["release_state"] == "HELD" for row in result["artifacts"])
    assert manifest["authority_created"] is False


def test_finishing_preserves_html_table_content_and_source_qualifications(tmp_path: Path) -> None:
    artifact = _source(tmp_path, '<html><head><style>DO NOT INCLUDE CSS</style></head><body><h1>Readiness</h1><table><tr><th>Evidence</th><th>Status</th></tr><tr><td>FR-03</td><td>Missing</td></tr></table><p>Not a lending decision.</p></body></html>', "report.html")
    result = _render(tmp_path, [artifact])
    html = Path(next(a["path"] for a in result["artifacts"] if a["mime_type"] == "text/html")).read_text()
    assert "FR-03" in html and "Missing" in html and "Not a lending decision." in html
    assert "<table" in html
    assert "DO NOT INCLUDE CSS" not in html


def test_finishing_refuses_missing_or_changed_source_bytes(tmp_path: Path) -> None:
    artifact = _source(tmp_path, "Real source findings")
    Path(artifact["path"]).write_text("Different findings")
    with pytest.raises(ValueError, match="hash"):
        _render(tmp_path, [artifact])
    Path(artifact["path"]).unlink()
    with pytest.raises(ValueError, match="missing"):
        _render(tmp_path, [artifact])


def test_inventory_receipt_alone_cannot_pass_semantic_finishing(tmp_path: Path) -> None:
    artifact = _source(tmp_path, '{"schema":"dio.some.receipt.v1","status":"PASS"}', "RECEIPT.json")
    with pytest.raises(ValueError, match="semantic|substantive"):
        _render(tmp_path, [artifact])


@pytest.mark.parametrize("schema, content", [
    ("dio.phase8.market_intelligence_pack.v1", {"snapshot": {"impressions": 120, "evidence_grade": "controlled_test"}, "outreach_authority": False}),
    ("dio.vesper_journey_view.v1", {"case_id": "CASE-42", "stage": "PROCESSING", "authority_created": False}),
])
def test_known_structured_domain_outputs_retain_data_and_qualifications(tmp_path: Path, schema: str, content: dict) -> None:
    artifact = _source(tmp_path, json.dumps({"schema": schema, **content}), "domain.json")
    result = _render(tmp_path, [artifact])
    rendered = Path(next(a["path"] for a in result["artifacts"] if a["mime_type"] == "text/html")).read_text()
    for value in ("120", "controlled_test") if "snapshot" in content else ("CASE-42", "PROCESSING"):
        assert value in rendered
    assert "false" in rendered


def test_repeated_finish_keeps_hashes_and_refuses_tampered_finished_artifacts(tmp_path: Path) -> None:
    artifact = _source(tmp_path, "# Findings\n\nCustomer evidence remains incomplete.")
    first = _render(tmp_path, [artifact])
    second = _render(tmp_path, [artifact])
    assert first == second
    Path(first["artifacts"][0]["path"]).write_text("Tampered review")
    with pytest.raises(ValueError, match="hash|tamper"):
        _render(tmp_path, [artifact])


def test_actual_vesper_product_journey_retains_substantive_shared_finish(tmp_path: Path) -> None:
    from products.professional_evidence_vesper_gate import execute_customer_case_via_vesper

    execute_customer_case_via_vesper(
        "AuditProof", tmp_path, operator_id="quality-regression", now="2026-10-27T16:00:00+00:00"
    )
    receipt_path = next(tmp_path.glob("*/PROFESSIONAL_EVIDENCE_RECEIPT.json"))
    receipt = json.loads(receipt_path.read_text())
    assert receipt["status"] == "PASS_FULL_PIPELINE", receipt.get("error")
    shared = receipt["shared_production_finish"]
    assert shared["organ_steps"][0]["provider"] == "adapters.format_core.renderer.render_semantic_asset"
    html = Path(next(a["path"] for a in shared["artifacts"] if a["mime_type"] == "text/html")).read_text()
    assert "AuditProof" in html
    assert "sign-off" in html.casefold() or "signoff" in html.casefold()
    assert receipt["authority_created"] is False


def test_customer_statements_alone_cannot_become_supported_audit_requirements(tmp_path: Path) -> None:
    from products.professional_evidence_vesper_gate import execute_customer_case_via_vesper

    execute_customer_case_via_vesper("AuditProof", tmp_path, operator_id="quality-regression", now="2026-10-27T16:00:00+00:00")
    pack = json.loads(next(tmp_path.glob("*/EXECUTION/AUDITPROOF_REVIEW_PACK.json")).read_text())
    assert all(row["review_state"] != "SUPPORTED" for row in pack["requirement_evidence_matrix"])
    assert pack["human_review"]["final_domain_decision"] == "NEEDS_YOU"


def test_html_finishing_preserves_caption_qualifications_and_source_urls(tmp_path: Path) -> None:
    artifact = _source(tmp_path, '<table><caption>Unaudited estimates, not approved lending values.</caption><tr><th>Amount</th></tr><tr><td>R999999</td></tr></table><p>Evidence: <a href="https://example.org/source-42">source record</a>.</p>', "qualified.html")
    result = _render(tmp_path, [artifact])
    html = Path(next(a["path"] for a in result["artifacts"] if a["mime_type"] == "text/html")).read_text()
    assert "Unaudited estimates, not approved lending values." in html
    assert "https://example.org/source-42" in html
    assert html.index("Unaudited estimates") < html.index("R999999")


def test_docx_finishing_keeps_paragraph_and_table_associations(tmp_path: Path) -> None:
    path = tmp_path / "ordered.docx"
    document = Document()
    document.add_paragraph("Approved figures follow.")
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Amount"
    table.cell(1, 0).text = "R100"
    document.add_paragraph("Unverified estimates follow.")
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Amount"
    table.cell(1, 0).text = "R999999"
    document.save(path)
    artifact = {"artifact_id": "ORDERED", "path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "release_state": "HELD"}
    result = _render(tmp_path, [artifact])
    html = Path(next(a["path"] for a in result["artifacts"] if a["mime_type"] == "text/html")).read_text()
    assert html.index("Approved figures") < html.index("R100") < html.index("Unverified estimates") < html.index("R999999")


@pytest.mark.parametrize("corruption", ["delete", "modify"])
def test_prepared_journey_checks_provider_hashes_before_shared_finishing(tmp_path: Path, monkeypatch, corruption: str) -> None:
    from products import professional_evidence_prepared_executor as prepared
    from products.professional_evidence_vesper_gate import execute_customer_case_via_vesper

    original = prepared._route_execute

    def corrupt_after_provider(packet, execution_dir, **kwargs):
        result = original(packet, execution_dir, **kwargs)
        if corruption == "delete":
            (execution_dir / "BOUNDED_OFFER.json").unlink()
        else:
            (execution_dir / "BOUNDED_OFFER.md").write_text("Corrupted customer offer. Price is R999999.")
        return result

    monkeypatch.setattr(prepared, "_route_execute", corrupt_after_provider)
    execute_customer_case_via_vesper("Offer Lab", tmp_path, operator_id="quality-regression", now="2026-10-27T16:00:00Z")
    receipt = json.loads(next(tmp_path.glob("*/PROFESSIONAL_EVIDENCE_RECEIPT.json")).read_text())
    assert receipt["status"] == "FAIL_EXECUTION"
    assert "native artifact" in receipt["error"].casefold()


def test_nested_document_provider_receipts_keep_their_relative_paths(tmp_path: Path) -> None:
    from products.professional_evidence_prepared_executor import _verify_provider_artifacts

    job = tmp_path / "document_studio" / "LOCALIZE"
    path = job / "formatted" / "Afrikaans" / "FORMAT_CORE_RECEIPT.json"
    path.parent.mkdir(parents=True)
    path.write_text('{"status":"rendered_review_candidate"}')
    receipt = {"outputs": [{"path": "formatted/Afrikaans/FORMAT_CORE_RECEIPT.json", "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}]}
    verified = _verify_provider_artifacts({"receipt": receipt}, tmp_path)
    assert str(path.resolve()) in verified
    path.write_text("tampered")
    with pytest.raises(ValueError, match="hash mismatch"):
        _verify_provider_artifacts({"receipt": receipt}, tmp_path)
