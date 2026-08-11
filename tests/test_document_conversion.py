from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from adapters.document_studio.conversion import execute_conversion, plan_conversion


def test_text_to_docx_preserves_source_and_emits_receipt(tmp_path: Path) -> None:
    source = tmp_path / "source.txt"
    source.write_text("First paragraph.\n\nSecond paragraph.\n", encoding="utf-8")
    output = tmp_path / "converted.docx"
    result = execute_conversion(source, output)
    assert source.read_text(encoding="utf-8") == "First paragraph.\n\nSecond paragraph.\n"
    assert output.is_file()
    receipt_path = Path(result["receipt_path"])
    assert receipt_path.is_file()
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    assert receipt["source_preserved"] is True
    assert receipt["semantic_editing_performed"] is False
    assert receipt["source"]["sha256"] != receipt["output"]["sha256"]
    assert receipt["strategy"] == "direct_text_to_docx"
    document = Document(output)
    assert [p.text for p in document.paragraphs if p.text] == ["First paragraph.", "Second paragraph."]


def test_docx_to_text_is_marked_lossy(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Alpha")
    document.add_paragraph("Beta")
    document.save(source)
    output = tmp_path / "plain.txt"
    result = execute_conversion(source, output)
    assert output.read_text(encoding="utf-8") == "Alpha\n\nBeta\n"
    assert result["receipt"]["lossy"] is True
    assert result["receipt"]["semantic_editing_performed"] is False


def test_dry_run_plans_without_writing_output(tmp_path: Path) -> None:
    source = tmp_path / "source.md"
    source.write_text("# Heading\n\nBody", encoding="utf-8")
    output = tmp_path / "source.docx"
    result = execute_conversion(source, output, dry_run=True)
    assert result["status"] == "planned"
    assert not output.exists()
    assert result["plan"]["plan_hash"]


def test_same_path_and_unsupported_format_are_refused(tmp_path: Path) -> None:
    source = tmp_path / "source.txt"
    source.write_text("x", encoding="utf-8")
    try:
        plan_conversion(source, source)
        raise AssertionError("same-path conversion should fail")
    except ValueError as exc:
        assert "must differ" in str(exc)
    bad = tmp_path / "source.bin"
    bad.write_bytes(b"x")
    try:
        plan_conversion(bad, tmp_path / "source.pdf")
        raise AssertionError("unsupported conversion should fail")
    except ValueError as exc:
        assert "Unsupported Document Studio source format" in str(exc)


def test_existing_output_is_not_overwritten_without_force(tmp_path: Path) -> None:
    source = tmp_path / "source.txt"
    source.write_text("source", encoding="utf-8")
    output = tmp_path / "converted.docx"
    output.write_bytes(b"existing")
    try:
        execute_conversion(source, output)
        raise AssertionError("existing output should fail")
    except FileExistsError as exc:
        assert "Refusing to overwrite" in str(exc)
    assert output.read_bytes() == b"existing"
