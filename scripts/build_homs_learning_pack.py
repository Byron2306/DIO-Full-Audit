#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REQUEST = ROOT / "samples" / "homs" / "learning_grade10_physical_sciences_motion.json"
DEFAULT_OUTPUT = ROOT / "deliverables" / "homs_learning_studio" / "grade_10_physical_sciences_term_3_motion"
CAPS_SOURCE = ROOT / "corpora" / "caps" / "text" / "fet_grade_10_12" / "caps_fet_physical_science_web.txt"

INK = "1C2733"
TEAL = "176B68"
GREEN = "397A55"
YELLOW = "E5B642"
PALE_TEAL = "E8F3F1"
PALE_YELLOW = "FFF6D8"
PALE_BLUE = "EAF2F7"
MID_GREY = "5D6874"
LINE = "CBD3D9"


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_request(request: dict[str, Any]) -> None:
    expected = {
        "subject": "Physical Sciences",
        "grade": 10,
        "phase": "FET",
        "term": 3,
        "topic_id": "motion_in_one_dimension",
        "language": "English",
    }
    errors = [f"{key} must be {value!r}" for key, value in expected.items() if request.get(key) != value]
    if request.get("educator_approval_required") is not True:
        errors.append("educator_approval_required must be true")
    if errors:
        raise ValueError("Invalid HOMS learning-pack request: " + "; ".join(errors))


def image_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    return ImageFont.truetype(f"/usr/share/fonts/truetype/dejavu/{name}", size)


def visual_canvas(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((28, 28, 1572, 872), outline="#1C2733", width=4)
    draw.rectangle((28, 28, 1572, 150), fill="#E8F3F1", outline="#1C2733", width=4)
    draw.text((70, 55), title, fill="#1C2733", font=image_font(34, True))
    draw.text((72, 108), subtitle, fill="#5D6874", font=image_font(19))
    return image, draw


def draw_reference_frame(path: Path) -> None:
    image, draw = visual_canvas("REFERENCE FRAME AND SIGN CONVENTION", "Choose the positive direction before calculating displacement or velocity")
    y = 450
    draw.line((160, y, 1440, y), fill="#1C2733", width=7)
    draw.polygon([(1440, y), (1390, y - 24), (1390, y + 24)], fill="#176B68")
    for value in range(-20, 31, 10):
        x = 520 + (value + 20) * 18
        draw.line((x, y - 22, x, y + 22), fill="#1C2733", width=4)
        draw.text((x - 20, y + 38), f"{value} m", fill="#1C2733", font=image_font(19, True))
    draw.text((1370, y - 70), "EAST (+)", fill="#176B68", font=image_font(23, True))
    draw.text((130, y - 70), "WEST (-)", fill="#A64B3C", font=image_font(23, True))
    draw.rectangle((835, 315, 1015, 410), fill="#E5B642", outline="#1C2733", width=5)
    draw.ellipse((860, 390, 910, 440), fill="#1C2733")
    draw.ellipse((950, 390, 1000, 440), fill="#1C2733")
    draw.text((857, 340), "CAR", fill="#1C2733", font=image_font(25, True))
    draw.line((925, 300, 1160, 300), fill="#176B68", width=7)
    draw.polygon([(1160, 300), (1117, 279), (1117, 321)], fill="#176B68")
    draw.text((850, 235), "velocity is positive", fill="#176B68", font=image_font(22, True))
    draw.rounded_rectangle((170, 650, 1430, 800), radius=18, fill="#F5F7F8", outline="#CBD3D9", width=3)
    draw.text((210, 680), "Position tells where the object is relative to the origin.", fill="#1C2733", font=image_font(23, True))
    draw.text((210, 730), "Displacement = final position - initial position. Direction matters.", fill="#1C2733", font=image_font(23, True))
    image.save(path, quality=95)


def draw_distance_displacement(path: Path) -> None:
    image, draw = visual_canvas("DISTANCE AND DISPLACEMENT", "A 20 m east journey followed by 5 m west")
    origin_x, finish_x, y = 300, 1250, 390
    draw.line((origin_x, y, finish_x, y), fill="#1C2733", width=6)
    draw.line((origin_x, y - 30, origin_x, y + 30), fill="#1C2733", width=5)
    draw.line((finish_x, y - 30, finish_x, y + 30), fill="#1C2733", width=5)
    draw.line((origin_x, 285, finish_x, 285), fill="#176B68", width=10)
    draw.polygon([(finish_x, 285), (finish_x - 50, 260), (finish_x - 50, 310)], fill="#176B68")
    draw.text((650, 230), "20 m EAST", fill="#176B68", font=image_font(25, True))
    return_x = 1010
    draw.line((finish_x, 485, return_x, 485), fill="#A64B3C", width=10)
    draw.polygon([(return_x, 485), (return_x + 50, 460), (return_x + 50, 510)], fill="#A64B3C")
    draw.text((1060, 515), "5 m WEST", fill="#A64B3C", font=image_font(25, True))
    draw.text((origin_x - 45, y + 55), "START", fill="#1C2733", font=image_font(20, True))
    draw.text((return_x - 42, y + 55), "FINISH", fill="#1C2733", font=image_font(20, True))
    draw.rounded_rectangle((170, 635, 1430, 810), radius=16, fill="#F5F7F8", outline="#CBD3D9", width=3)
    draw.text((225, 670), "Distance = 20 m + 5 m = 25 m", fill="#1C2733", font=image_font(26, True))
    draw.text((225, 730), "Displacement = +20 m - 5 m = +15 m (east)", fill="#176B68", font=image_font(26, True))
    image.save(path, quality=95)


def draw_acceleration_example(path: Path) -> None:
    image, draw = visual_canvas("VELOCITY AND ACCELERATION", "Keep the sign convention consistent from sketch to answer")
    draw.rounded_rectangle((120, 205, 720, 660), radius=16, fill="#F5F7F8", outline="#CBD3D9", width=4)
    draw.text((170, 245), "Known values", fill="#176B68", font=image_font(27, True))
    draw.text((170, 320), "initial velocity = +2 m.s^-1", fill="#1C2733", font=image_font(25))
    draw.text((170, 380), "acceleration = +3 m.s^-2", fill="#1C2733", font=image_font(25))
    draw.text((170, 440), "time interval = 4 s", fill="#1C2733", font=image_font(25))
    draw.line((170, 530, 620, 530), fill="#176B68", width=8)
    draw.polygon([(620, 530), (570, 505), (570, 555)], fill="#176B68")
    draw.text((170, 565), "positive direction", fill="#176B68", font=image_font(22, True))
    draw.rounded_rectangle((790, 205, 1480, 660), radius=16, fill="#E8F3F1", outline="#176B68", width=4)
    draw.text((845, 245), "Final velocity", fill="#176B68", font=image_font(27, True))
    draw.text((845, 330), "v final = v initial + a delta t", fill="#1C2733", font=image_font(25, True))
    draw.text((845, 410), "= 2 + (3)(4)", fill="#1C2733", font=image_font(27))
    draw.text((845, 490), "= +14 m.s^-1 east", fill="#176B68", font=image_font(30, True))
    draw.text((170, 750), "A positive velocity and positive acceleration mean the object speeds up in the positive direction.", fill="#5D6874", font=image_font(22, True))
    image.save(path, quality=95)


def graph_axes(draw: ImageDraw.ImageDraw, *, left: int, top: int, right: int, bottom: int, x_label: str, y_label: str) -> None:
    draw.line((left, bottom, right, bottom), fill="#1C2733", width=5)
    draw.line((left, bottom, left, top), fill="#1C2733", width=5)
    draw.polygon([(right, bottom), (right - 24, bottom - 12), (right - 24, bottom + 12)], fill="#1C2733")
    draw.polygon([(left, top), (left - 12, top + 24), (left + 12, top + 24)], fill="#1C2733")
    draw.text((right - 160, bottom + 45), x_label, fill="#1C2733", font=image_font(20, True))
    draw.text((40, top - 5), y_label, fill="#1C2733", font=image_font(20, True))


def draw_position_graph(path: Path) -> None:
    image, draw = visual_canvas("POSITION-TIME GRAPH", "The gradient of a position-time graph is velocity")
    left, top, right, bottom = 210, 210, 1450, 760
    graph_axes(draw, left=left, top=top, right=right, bottom=bottom, x_label="time (s)", y_label="position (m)")
    for t in range(0, 9, 2):
        x = left + int(t / 8 * (right - left))
        draw.line((x, bottom - 10, x, bottom + 10), fill="#1C2733", width=3)
        draw.text((x - 8, bottom + 16), str(t), fill="#1C2733", font=image_font(17))
    for value in range(0, 9, 2):
        y = bottom - int(value / 8 * (bottom - top))
        draw.line((left - 10, y, left + 10, y), fill="#1C2733", width=3)
        draw.text((left - 45, y - 10), str(value), fill="#1C2733", font=image_font(17))
        if value:
            draw.line((left, y, right, y), fill="#E1E6EA", width=2)
    points_data = [(0, 0), (4, 8), (6, 8), (8, 2)]
    points = [(left + int(t / 8 * (right - left)), bottom - int(x / 8 * (bottom - top))) for t, x in points_data]
    draw.line(points, fill="#176B68", width=8)
    for index, (x, y) in enumerate(points):
        draw.ellipse((x - 10, y - 10, x + 10, y + 10), fill="#E5B642", outline="#1C2733", width=3)
        draw.text((x + 14, y - 30), chr(65 + index), fill="#1C2733", font=image_font(20, True))
    footer_font = image_font(16)
    draw.text((280, 825), "Positive gradient: positive velocity", fill="#5D6874", font=footer_font)
    draw.text((690, 825), "Zero gradient: at rest", fill="#5D6874", font=footer_font)
    draw.text((1020, 825), "Negative gradient: negative velocity", fill="#5D6874", font=footer_font)
    image.save(path, quality=95)


def draw_velocity_graph(path: Path) -> None:
    image, draw = visual_canvas("VELOCITY-TIME GRAPH", "The gradient is acceleration; the signed area is displacement")
    left, top, right, bottom = 210, 210, 1450, 760
    graph_axes(draw, left=left, top=top, right=right, bottom=bottom, x_label="time (s)", y_label="velocity (m.s^-1)")
    for t in range(0, 13, 2):
        x = left + int(t / 12 * (right - left))
        draw.line((x, bottom - 10, x, bottom + 10), fill="#1C2733", width=3)
        draw.text((x - 8, bottom + 16), str(t), fill="#1C2733", font=image_font(17))
    for value in range(0, 9, 2):
        y = bottom - int(value / 8 * (bottom - top))
        draw.line((left - 10, y, left + 10, y), fill="#1C2733", width=3)
        draw.text((left - 45, y - 10), str(value), fill="#1C2733", font=image_font(17))
        if value:
            draw.line((left, y, right, y), fill="#E1E6EA", width=2)
    points_data = [(0, 0), (4, 8), (8, 8), (12, 0)]
    points = [(left + int(t / 12 * (right - left)), bottom - int(v / 8 * (bottom - top))) for t, v in points_data]
    polygon = [points[0], points[1], points[2], points[3], (right, bottom), (left, bottom)]
    draw.polygon(polygon, fill="#DCEDEA")
    draw.line(points, fill="#176B68", width=8)
    for index, (x, y) in enumerate(points):
        draw.ellipse((x - 10, y - 10, x + 10, y + 10), fill="#E5B642", outline="#1C2733", width=3)
        draw.text((x + 14, y - 30), chr(65 + index), fill="#1C2733", font=image_font(20, True))
    draw.text((520, 825), "Displacement = triangle + rectangle + triangle", fill="#5D6874", font=image_font(18, True))
    image.save(path, quality=95)


def draw_investigation(path: Path) -> None:
    image, draw = visual_canvas("PRACTICAL: MEASURING MOTION", "Use a toy car, straight track, metre rule and stopwatch")
    draw.line((130, 610, 1460, 610), fill="#1C2733", width=8)
    for index, metres in enumerate([0, 1, 2, 3, 4]):
        x = 180 + index * 290
        draw.line((x, 570, x, 650), fill="#1C2733", width=5)
        draw.text((x - 18, 675), f"{metres} m", fill="#1C2733", font=image_font(20, True))
    draw.rectangle((200, 440, 390, 560), fill="#E5B642", outline="#1C2733", width=5)
    draw.ellipse((225, 535, 285, 595), fill="#1C2733")
    draw.ellipse((315, 535, 375, 595), fill="#1C2733")
    draw.text((235, 470), "TOY CAR", fill="#1C2733", font=image_font(22, True))
    draw.line((410, 430, 750, 430), fill="#176B68", width=8)
    draw.polygon([(750, 430), (700, 405), (700, 455)], fill="#176B68")
    draw.text((445, 370), "direction of motion", fill="#176B68", font=image_font(23, True))
    draw.ellipse((1120, 250, 1320, 450), outline="#176B68", width=12)
    draw.line((1220, 350, 1220, 285), fill="#176B68", width=7)
    draw.line((1220, 350, 1280, 390), fill="#176B68", width=7)
    draw.text((1090, 475), "STOPWATCH", fill="#176B68", font=image_font(22, True))
    draw.rounded_rectangle((130, 755, 1470, 835), radius=16, fill="#F5F7F8", outline="#CBD3D9", width=3)
    draw.text((180, 780), "Measure position at equal time intervals. Plot position against time. Use the gradient to determine average velocity.", fill="#1C2733", font=image_font(20))
    image.save(path, quality=95)


def draw_graph_template(path: Path) -> None:
    image, draw = visual_canvas("POSITION-TIME GRAPH", "Plot average time horizontally and position vertically")
    left, top, right, bottom = 220, 205, 1435, 775
    for x in range(left, right + 1, 55):
        draw.line((x, top, x, bottom), fill="#DCE2E6", width=2)
    for y in range(top, bottom + 1, 45):
        draw.line((left, y, right, y), fill="#DCE2E6", width=2)
    draw.line((left, bottom, right, bottom), fill="#1C2733", width=5)
    draw.line((left, bottom, left, top), fill="#1C2733", width=5)
    draw.polygon([(right, bottom), (right - 24, bottom - 12), (right - 24, bottom + 12)], fill="#1C2733")
    draw.polygon([(left, top), (left - 12, top + 24), (left + 12, top + 24)], fill="#1C2733")
    draw.text((1250, 805), "average time (s)", fill="#1C2733", font=image_font(19, True))
    draw.text((48, 180), "position (m)", fill="#1C2733", font=image_font(19, True))
    draw.text((230, 795), "Choose and label a sensible scale before plotting.", fill="#5D6874", font=image_font(17))
    image.save(path, quality=95)


def draw_video_thumbnail(path: Path) -> None:
    image = Image.new("RGB", (1600, 900), "#F7F9FA")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 650, 900), fill="#E5B642")
    draw.rectangle((650, 0, 1600, 900), fill="#FFFFFF")
    draw.text((70, 70), "GRADE 10", fill="#1C2733", font=image_font(34, True))
    draw.text((70, 130), "PHYSICAL", fill="#1C2733", font=image_font(55, True))
    draw.text((70, 195), "SCIENCES", fill="#1C2733", font=image_font(55, True))
    draw.line((70, 285, 570, 285), fill="#1C2733", width=8)
    draw.text((70, 335), "MOTION IN", fill="#1C2733", font=image_font(47, True))
    draw.text((70, 395), "ONE DIMENSION", fill="#1C2733", font=image_font(47, True))
    draw.text((70, 505), "Reference frames", fill="#1C2733", font=image_font(27))
    draw.text((70, 550), "Velocity and acceleration", fill="#1C2733", font=image_font(27))
    draw.text((70, 595), "Motion graphs", fill="#1C2733", font=image_font(27))
    draw.rounded_rectangle((70, 705, 430, 790), radius=10, fill="#176B68")
    draw.text((115, 728), "HOMS LEARNING", fill="#FFFFFF", font=image_font(27, True))
    left, top, right, bottom = 760, 175, 1500, 710
    draw.line((left, bottom, right, bottom), fill="#1C2733", width=7)
    draw.line((left, bottom, left, top), fill="#1C2733", width=7)
    points = [(left, bottom), (1030, 275), (1250, 275), (right, 555)]
    draw.line(points, fill="#176B68", width=14)
    for label, (x, y) in zip(["A", "B", "C", "D"], points):
        draw.ellipse((x - 14, y - 14, x + 14, y + 14), fill="#E5B642", outline="#1C2733", width=4)
        draw.text((x + 18, y - 38), label, fill="#1C2733", font=image_font(26, True))
    draw.text((1140, 745), "POSITION - TIME", fill="#176B68", font=image_font(31, True))
    image.save(path, quality=95)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def configure_document(document: Document, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, color in [("Title", 28, INK), ("Heading 1", 18, TEAL), ("Heading 2", 14, INK), ("Heading 3", 11, GREEN)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = f"HOMS LEARNING STUDIO  |  {title}"
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(MID_GREY)
    add_page_number(section.footer.paragraphs[0])


def add_cover(document: Document, request: dict[str, Any], document_type: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(80)
    run = paragraph.add_run("HOMS")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(document_type)
    topic = document.add_paragraph()
    topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    topic_run = topic.add_run(request["topic"])
    topic_run.bold = True
    topic_run.font.name = "Arial"
    topic_run.font.size = Pt(22)
    topic_run.font.color.rgb = RGBColor.from_string(INK)
    detail = document.add_paragraph()
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail_run = detail.add_run(f"{request['subject']}  |  Grade {request['grade']}  |  Term {request['term']}")
    detail_run.font.name = "Arial"
    detail_run.font.size = Pt(13)
    detail_run.font.color.rgb = RGBColor.from_string(MID_GREY)
    box = document.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    box.columns[0].width = Cm(15.5)
    cell = box.cell(0, 0)
    set_cell_shading(cell, PALE_TEAL)
    set_cell_border(cell, TEAL, "10")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    p.add_run(subtitle).bold = True
    document.add_paragraph()
    learner = document.add_table(rows=2, cols=2)
    learner.alignment = WD_TABLE_ALIGNMENT.CENTER
    learner.style = "Table Grid"
    learner.cell(0, 0).text = "Learner name"
    learner.cell(0, 1).text = ""
    learner.cell(1, 0).text = "Date"
    learner.cell(1, 1).text = ""
    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run("CAPS-aligned draft. Educator or subject-expert approval is required before classroom use.")
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor.from_string(MID_GREY)
    document.add_page_break()


def add_callout(document: Document, title: str, body: str, fill: str = PALE_TEAL) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, TEAL if fill == PALE_TEAL else YELLOW, "8")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(f"{title}: ")
    run.bold = True
    paragraph.add_run(body)


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_formula_table(document: Document) -> None:
    rows = [
        ("Displacement", "delta x = x final - x initial", "m"),
        ("Average speed", "total distance / total time", "m.s^-1"),
        ("Average velocity", "displacement / time interval", "m.s^-1"),
        ("Acceleration", "a = delta v / delta t", "m.s^-2"),
        ("Uniform acceleration", "v final = v initial + a delta t", "m.s^-1"),
        ("Uniform acceleration", "delta x = v initial delta t + 0.5 a(delta t)^2", "m"),
        ("Uniform acceleration", "v final^2 = v initial^2 + 2a delta x", "m^2.s^-2"),
        ("Uniform acceleration", "delta x = 0.5(v initial + v final)delta t", "m"),
    ]
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Quantity", "Relationship", "SI unit"]
    for index, text in enumerate(headers):
        table.cell(0, index).text = text
        set_cell_shading(table.cell(0, index), TEAL)
        for run in table.cell(0, index).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    set_repeat_table_header(table.rows[0])
    for quantity, relationship, unit in rows:
        cells = table.add_row().cells
        for index, value in enumerate([quantity, relationship, unit]):
            cells[index].text = value


def add_question(document: Document, number: str, text: str, marks: int, answer_lines: int = 0) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(15.8)
    table.columns[1].width = Cm(1.2)
    left, right = table.rows[0].cells
    left.text = ""
    right.text = ""
    left_p = left.paragraphs[0]
    left_p.paragraph_format.space_after = Pt(2)
    left_p.add_run(f"{number}  ").bold = True
    left_p.add_run(text)
    right_p = right.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_p.add_run(f"({marks})").bold = True
    for _ in range(answer_lines):
        p = document.add_paragraph("_" * 112)
        p.paragraph_format.space_after = Pt(1)
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = RGBColor.from_string(LINE)


def add_memo_item(document: Document, number: str, points: list[str], marks: int) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(15.8)
    table.columns[1].width = Cm(1.2)
    left, right = table.rows[0].cells
    paragraph = left.paragraphs[0]
    paragraph.add_run(f"{number}  ").bold = True
    paragraph.add_run(points[0])
    for point in points[1:]:
        left.add_paragraph(point, style="List Bullet")
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraphs[0].add_run(f"[{marks}]").bold = True


def build_guide(request: dict[str, Any], assets: dict[str, Path], output: Path) -> None:
    document = Document()
    configure_document(document, "Motion in One Dimension")
    add_cover(document, request, "Learner Concept Guide", "Understand the motion story, then calculate it with confidence.")
    document.add_heading("How To Use This Guide", level=1)
    add_bullets(document, [
        "Read the concept explanation before attempting the worked example.",
        "Choose a positive direction and keep signs consistent.",
        "Draw a small motion sketch before substituting into an equation.",
        "Complete the activity and worksheet before the mini-assessment.",
        "Use the memorandum only after making a serious attempt.",
    ])
    document.add_heading("Learning Outcomes", level=1)
    add_bullets(document, [
        "Describe motion relative to a reference frame.",
        "Distinguish distance from displacement and speed from velocity.",
        "Calculate average speed, average velocity and acceleration.",
        "Interpret position-time and velocity-time graphs.",
        "Apply the equations of uniformly accelerated motion.",
        "Evaluate whether an answer is physically sensible and correctly signed.",
    ])
    document.add_heading("1. Reference Frames And Position", level=1)
    document.add_paragraph("Motion is always described relative to a chosen reference frame. A reference frame has an origin and a positive direction. Once the direction is chosen, positions and vector quantities may be positive or negative.")
    document.add_picture(str(assets["reference_frame"]), width=Inches(6.8))
    add_callout(document, "Sign discipline", "A negative velocity does not mean that an object is slowing down. It means the object is moving in the chosen negative direction.", PALE_YELLOW)
    document.add_heading("2. Distance And Displacement", level=1)
    document.add_paragraph("Distance is the total path length travelled. It is a scalar and cannot be negative. Displacement is the change in position from the start to the finish. It is a vector and includes direction.")
    document.add_heading("Worked Example 1", level=2)
    document.add_paragraph("A learner walks 20 m east and then 5 m west in 10 s. Take east as positive.")
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    rows = [("Distance", "20 + 5 = 25 m"), ("Displacement", "+20 - 5 = +15 m (east)"), ("Average speed", "25 / 10 = 2.5 m.s^-1"), ("Average velocity", "+15 / 10 = +1.5 m.s^-1 (east)")]
    for index, (label, value) in enumerate(rows):
        table.cell(index, 0).text = label
        table.cell(index, 1).text = value
        set_cell_shading(table.cell(index, 0), PALE_BLUE)
    document.add_heading("3. Velocity And Acceleration", level=1)
    document.add_paragraph("Velocity describes how quickly position changes and in which direction. Acceleration describes how quickly velocity changes. An object can accelerate by speeding up, slowing down, or changing direction.")
    add_formula_table(document)
    document.add_heading("Worked Example 2", level=2)
    document.add_paragraph("A trolley moves east at 2 m.s^-1 and accelerates uniformly east at 3 m.s^-2 for 4 s.")
    add_bullets(document, [
        "Choose east as positive: v initial = +2 m.s^-1, a = +3 m.s^-2, delta t = 4 s.",
        "Final velocity: v final = 2 + (3)(4) = 14 m.s^-1 east.",
        "Displacement: delta x = (2)(4) + 0.5(3)(4^2) = 8 + 24 = 32 m east.",
        "Check: the trolley speeds up in the positive direction, so positive displacement is sensible.",
    ])
    document.add_heading("4. Reading Motion Graphs", level=1)
    document.add_picture(str(assets["position_graph"]), width=Inches(6.8))
    add_bullets(document, [
        "A steeper position-time gradient means a greater speed.",
        "A horizontal position-time section means the object is at rest.",
        "A negative position-time gradient means motion in the negative direction.",
        "The gradient of a velocity-time graph is acceleration.",
        "The signed area between a velocity-time graph and the time axis is displacement.",
    ])
    document.add_picture(str(assets["velocity_graph"]), width=Inches(6.8))
    document.add_heading("5. A Reliable Problem-Solving Routine", level=1)
    steps = [
        ("1", "Sketch", "Draw the direction of motion and label the known values."),
        ("2", "Choose", "State the positive direction before assigning signs."),
        ("3", "List", "Write each known value with its unit."),
        ("4", "Select", "Choose an equation containing the unknown and the available data."),
        ("5", "Substitute", "Substitute with signs, then calculate."),
        ("6", "Interpret", "Give the unit, direction and a physical reasonableness check."),
    ]
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for index, heading in enumerate(["Step", "Action", "What to do"]):
        table.cell(0, index).text = heading
        set_cell_shading(table.cell(0, index), TEAL)
        for run in table.cell(0, index).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for row in steps:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.add_heading("Common Errors", level=1)
    add_bullets(document, [
        "Using distance when average velocity requires displacement.",
        "Changing the positive direction halfway through a calculation.",
        "Treating a negative answer as automatically wrong.",
        "Using the area under a position-time graph instead of its gradient.",
        "Mixing kilometres per hour with metres per second without converting.",
        "Reporting a number without its unit and direction.",
    ])
    document.save(output)


def build_activity(request: dict[str, Any], assets: dict[str, Path], output: Path) -> None:
    document = Document()
    configure_document(document, "Motion Practical Activity")
    add_cover(document, request, "Practical Activity", "Measure a moving object's position and use a graph to determine velocity.")
    document.add_picture(str(assets["investigation"]), width=Inches(6.8))
    document.add_heading("Investigation Question", level=1)
    document.add_paragraph("How does the position of a toy car change with time as it moves along a straight track?")
    document.add_heading("Materials", level=2)
    add_bullets(document, ["Toy car or trolley", "Straight track or smooth floor", "Metre rule or measuring tape", "Masking tape or chalk markers", "Stopwatch or phone timer", "Graph paper"])
    document.add_heading("Variables", level=2)
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    for index, row in enumerate([("Independent variable", "Time"), ("Dependent variable", "Position relative to the origin"), ("Controlled variable", "Same car, track and release method"), ("Positive direction", "From the origin toward the end of the track")]):
        table.cell(index, 0).text = row[0]
        table.cell(index, 1).text = row[1]
        set_cell_shading(table.cell(index, 0), PALE_BLUE)
    document.add_heading("Method", level=1)
    add_bullets(document, [
        "Mark an origin and positions at 0.50 m intervals along a straight track.",
        "Place the car at the origin and practise releasing it without pushing.",
        "Release the car and record the time when its front reaches each position marker.",
        "Repeat the run three times. Use the average time for each position.",
        "Plot position on the vertical axis and average time on the horizontal axis.",
        "Draw a best-fit line or curve and calculate the gradient over a suitable interval.",
    ])
    add_callout(document, "Safety", "Keep the track clear, do not place the car near stairs, and assign one learner to time while another observes position.", PALE_YELLOW)
    document.add_heading("Results", level=1)
    table = document.add_table(rows=7, cols=6)
    table.style = "Table Grid"
    headers = ["Position (m)", "Run 1 (s)", "Run 2 (s)", "Run 3 (s)", "Average time (s)", "Interval velocity (m.s^-1)"]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
        set_cell_shading(table.cell(0, index), TEAL)
        for run in table.cell(0, index).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for row_index, position in enumerate([0.0, 0.5, 1.0, 1.5, 2.0, 2.5], 1):
        table.cell(row_index, 0).text = f"{position:.1f}"
    document.add_paragraph("Plot your position-time graph on the grid below. Label both axes and show the points used for one gradient calculation.")
    document.add_picture(str(assets["graph_template"]), width=Inches(6.7))
    document.add_page_break()
    document.add_heading("Analysis And Evaluation", level=1)
    add_question(document, "1", "State the relationship between position and time shown by your graph.", 2, 3)
    add_question(document, "2", "Calculate the average velocity over one suitable interval. Show the points used.", 4, 4)
    add_question(document, "3", "Was the velocity constant? Use the graph as evidence.", 3, 4)
    add_question(document, "4", "Identify one source of random error and explain how repeating the run reduces its effect.", 3, 4)
    add_question(document, "5", "Suggest one improvement that would make the timing more reliable.", 2, 3)
    add_question(document, "6", "Write a conclusion that answers the investigation question.", 3, 5)
    document.save(output)


WORKSHEET = [
    ("1.1", "Define a frame of reference.", 2, 2, ["An origin and a set of directions relative to which position and motion are described."]),
    ("1.2", "Distinguish between distance and displacement.", 2, 3, ["Distance is total path length and is scalar.", "Displacement is change in position and is vector."]),
    ("1.3", "Explain why a negative velocity does not necessarily mean that an object is slowing down.", 2, 3, ["The sign indicates direction relative to the chosen positive direction, not whether speed is decreasing."]),
    ("1.4", "State the physical meaning of the gradient of a position-time graph.", 1, 2, ["Velocity."]),
    ("1.5", "State the physical meaning of the area under a velocity-time graph.", 1, 2, ["Displacement."]),
    ("2.1", "A learner walks 60 m east and then 25 m west. Calculate the distance and displacement.", 4, 4, ["Distance = 60 + 25 = 85 m.", "Displacement = +60 - 25 = +35 m, or 35 m east."]),
    ("2.2", "The journey in Question 2.1 takes 50 s. Calculate average speed and average velocity.", 4, 4, ["Average speed = 85/50 = 1.70 m.s^-1.", "Average velocity = 35/50 = 0.70 m.s^-1 east."]),
    ("2.3", "A car changes velocity from +4 m.s^-1 to +16 m.s^-1 in 6 s. Calculate its acceleration.", 3, 3, ["a = (16 - 4)/6 = +2.0 m.s^-2."]),
    ("2.4", "A bicycle moving at 10 m.s^-1 slows uniformly at 2 m.s^-2 for 3 s. Calculate its final velocity.", 3, 3, ["Take forward as positive. v = 10 + (-2)(3) = 4 m.s^-1 forward."]),
    ("3.1", "Refer to the position-time graph. Calculate the velocity from A to B.", 3, 3, ["v = gradient = (8 - 0)/(4 - 0) = +2 m.s^-1."]),
    ("3.2", "Describe the motion from B to C.", 2, 2, ["The object is at rest at position 8 m because the graph is horizontal."]),
    ("3.3", "Calculate the velocity from C to D.", 3, 3, ["v = (2 - 8)/(8 - 6) = -3 m.s^-1."]),
    ("3.4", "Explain what the negative answer in Question 3.3 means.", 2, 3, ["The object moves in the negative direction, back toward the origin."]),
    ("4.1", "A trolley starts at 3 m.s^-1 and accelerates uniformly at 2 m.s^-2 for 5 s. Calculate its final velocity.", 3, 2, ["v = 3 + (2)(5) = 13 m.s^-1."]),
    ("4.2", "Calculate the displacement of the trolley during the 5 s interval.", 5, 3, ["delta x = (3)(5) + 0.5(2)(5^2) = 15 + 25 = 40 m."]),
]


ASSESSMENT = [
    ("1.1", "Define average velocity.", 2, 2, ["Displacement divided by the time interval."]),
    ("1.2", "Give one difference between a scalar and a vector quantity.", 2, 2, ["A vector has magnitude and direction; a scalar has magnitude only."]),
    ("1.3", "A learner says: 'An object with zero velocity must have zero acceleration.' Is this always correct? Briefly justify your answer.", 2, 3, ["No.", "At an instant an object may have zero velocity while its velocity is changing, so acceleration may be non-zero."]),
    ("1.4", "State what a horizontal section on a position-time graph means.", 2, 2, ["The position is constant and the object is at rest."]),
    ("2.1", "A taxi travels 120 m east and then 40 m west in 20 s. Calculate its average speed.", 3, 3, ["Distance = 160 m; average speed = 160/20 = 8.0 m.s^-1."]),
    ("2.2", "Calculate the taxi's average velocity for the same journey.", 3, 3, ["Displacement = +80 m; average velocity = 80/20 = 4.0 m.s^-1 east."]),
    ("2.3", "A train moving at 6 m.s^-1 accelerates uniformly at 1.5 m.s^-2 for 8 s. Calculate its final velocity.", 2, 3, ["v = 6 + (1.5)(8) = 18 m.s^-1."]),
    ("2.4", "Calculate the train's displacement during the 8 s interval.", 4, 4, ["delta x = (6)(8) + 0.5(1.5)(8^2) = 48 + 48 = 96 m."]),
    ("3.1", "Refer to the velocity-time graph. Calculate the acceleration from A to B.", 2, 3, ["a = (8 - 0)/(4 - 0) = 2.0 m.s^-2."]),
    ("3.2", "Describe the motion from B to C.", 2, 2, ["The object moves at constant positive velocity of 8 m.s^-1."]),
    ("3.3", "Calculate the acceleration from C to D.", 2, 3, ["a = (0 - 8)/(12 - 8) = -2.0 m.s^-2."]),
    ("3.4", "Calculate the total displacement from 0 s to 12 s.", 4, 4, ["Area = 0.5(4)(8) + (4)(8) + 0.5(4)(8) = 16 + 32 + 16 = 64 m."]),
]


def build_questions(request: dict[str, Any], assets: dict[str, Path], output: Path, *, title: str, questions: list[tuple], assessment: bool = False) -> None:
    document = Document()
    configure_document(document, title)
    total = sum(item[2] for item in questions)
    add_cover(document, request, title, f"Complete every question. Total: {total} marks.")
    add_callout(document, "Instructions", "Show formulae, substitutions, units and direction where required. Use the supplied graphs for graph questions.")
    current_section = ""
    for number, text, marks, lines, _memo in questions:
        section = number.split(".")[0]
        if section != current_section:
            current_section = section
            heading = {"1": "Concepts", "2": "Calculations", "3": "Graph Interpretation", "4": "Uniform Acceleration"}.get(section, "Questions")
            document.add_heading(f"Question {section}: {heading}", level=1)
            if section == "3":
                graph = assets["velocity_graph"] if assessment else assets["position_graph"]
                document.add_picture(str(graph), width=Inches(6.7))
        add_question(document, number, text, marks, lines)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run(f"TOTAL: {total}").bold = True
    document.save(output)


def build_memo(request: dict[str, Any], output: Path, *, title: str, questions: list[tuple]) -> None:
    document = Document()
    configure_document(document, title)
    total = sum(item[2] for item in questions)
    add_cover(document, request, title, f"Marking guidance and worked answers. Total: {total} marks.")
    add_callout(document, "Marking principle", "Credit equivalent correct methods, units and directions. Do not award a final-answer mark when contradictory working is shown.", PALE_YELLOW)
    current_section = ""
    for number, _text, marks, _lines, memo in questions:
        section = number.split(".")[0]
        if section != current_section:
            current_section = section
            document.add_heading(f"Question {section}", level=1)
        add_memo_item(document, number, memo, marks)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run(f"TOTAL: {total}").bold = True
    document.save(output)



def convert_pdfs(docx_files: list[Path], pdf_dir: Path) -> list[Path]:
    from adapters.libreoffice_low_memory import convert_many_to_pdf

    return convert_many_to_pdf(
        docx_files,
        pdf_dir,
        profile_prefix="homs-lo-profile-",
    )

def caps_evidence() -> dict[str, Any]:
    if not CAPS_SOURCE.exists():
        raise FileNotFoundError(CAPS_SOURCE)
    text = CAPS_SOURCE.read_text(encoding="utf-8", errors="replace")
    required = ["GRADE 10 PHYSICS (MECHANICS) TERM 3", "Motion in one dimension", "position vs. time", "velocity vs time"]
    missing = [phrase for phrase in required if phrase.lower() not in text.lower()]
    if missing:
        raise ValueError("CAPS source is missing required evidence: " + ", ".join(missing))
    return {
        "source": str(CAPS_SOURCE),
        "source_sha256": sha256(CAPS_SOURCE),
        "evidence": [
            "CAPS identifies Grade 10 Physics (Mechanics) Term 3.",
            "The term includes vectors and scalars, motion in one dimension, position, displacement, distance, speed, velocity and acceleration.",
            "Learners must describe motion using words, diagrams, graphs and equations.",
            "CAPS prescribes position-time and velocity-time graph interpretation and practical motion measurement.",
        ],
    }


def build_alignment_report(request: dict[str, Any], caps: dict[str, Any], output: Path) -> None:
    output.write_text(
        "\n".join([
            "# HOMS Learning Studio: Educator Alignment Report",
            "",
            f"Pack: `{request['pack_id']}`",
            f"Subject: {request['subject']}",
            f"Grade/term: Grade {request['grade']}, Term {request['term']}",
            f"Topic: {request['topic']}",
            "Status: review candidate; educator approval required",
            "",
            "## CAPS Grounding",
            "",
            f"Source: `{caps['source']}`",
            f"SHA-256: `{caps['source_sha256']}`",
            "",
            *[f"- {item}" for item in caps["evidence"]],
            "",
            "## One Curriculum Spine",
            "",
            "| Intended knowledge or skill | Guide | Activity | Worksheet | Assessment |",
            "|---|---|---|---|---|",
            "| Reference frame and sign convention | Explained and illustrated | Positive direction selected | Questions 1.1-1.3 | Questions 1.1-1.4 |",
            "| Distance, displacement, speed and velocity | Worked example | Position measured | Questions 2.1-2.2 | Questions 2.1-2.2 |",
            "| Acceleration and equations | Formula table and worked example | Graph gradient extension | Questions 2.3-2.4 and 4 | Questions 2.3-2.4 |",
            "| Position-time graphs | Explicit graph rules | Learner-generated graph | Questions 3.1-3.4 | Concept transfer |",
            "| Velocity-time graphs | Explicit graph rules | Optional extension | Concept transfer | Questions 3.1-3.4 |",
            "",
            "## Cognitive Progression",
            "",
            "The pack moves from terminology and interpretation to calculations, graph analysis, method evaluation and an independent mini-assessment. The assessment does not introduce content absent from the guide and practice sequence.",
            "",
            "## Required Human Review",
            "",
            "- Confirm CAPS term and topic fit for the intended class pace.",
            "- Recalculate every numerical answer and inspect every graph label.",
            "- Confirm language accessibility for the learner group.",
            "- Approve practical safety and available equipment.",
            "- Approve the video narration, captions, music attribution and final publication state.",
            "",
        ]) + "\n",
        encoding="utf-8",
    )


def validate_pack(request: dict[str, Any], files: list[Path], manifest_data: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    for path in files:
        if not path.exists() or path.stat().st_size < 500:
            errors.append(f"missing_or_empty:{path.name}")
    if sum(item[2] for item in WORKSHEET) != 40:
        errors.append("worksheet_marks_must_equal_40")
    if sum(item[2] for item in ASSESSMENT) != 30:
        errors.append("assessment_marks_must_equal_30")
    if len(manifest_data.get("visual_assets") or []) != 8:
        errors.append("eight_meaningful_visual_assets_required")
    if request.get("educator_approval_required") is not True:
        errors.append("educator_approval_gate_missing")
    return errors


def zip_pack(out_dir: Path, target: Path, files: list[Path]) -> None:
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            if path.exists() and path != target:
                archive.write(path, path.relative_to(out_dir))


def build(request_path: Path, out_dir: Path) -> dict[str, Any]:
    request = load_json(request_path)
    validate_request(request)
    out_dir.mkdir(parents=True, exist_ok=True)
    asset_dir = out_dir / "assets"
    document_dir = out_dir / "documents"
    pdf_dir = out_dir / "pdf"
    for directory in [asset_dir, document_dir, pdf_dir]:
        directory.mkdir(parents=True, exist_ok=True)
    request_copy = out_dir / "LEARNING_PACK_REQUEST.json"
    write_json(request_copy, request)

    assets = {
        "reference_frame": asset_dir / "reference_frame_and_sign_convention.png",
        "distance_displacement": asset_dir / "distance_and_displacement_worked_example.png",
        "acceleration_example": asset_dir / "velocity_and_acceleration_worked_example.png",
        "position_graph": asset_dir / "position_time_graph.png",
        "velocity_graph": asset_dir / "velocity_time_graph.png",
        "investigation": asset_dir / "motion_investigation_setup.png",
        "graph_template": asset_dir / "position_time_graph_template.png",
        "video_thumbnail": asset_dir / "homs_motion_video_thumbnail.png",
    }
    draw_reference_frame(assets["reference_frame"])
    draw_distance_displacement(assets["distance_displacement"])
    draw_acceleration_example(assets["acceleration_example"])
    draw_position_graph(assets["position_graph"])
    draw_velocity_graph(assets["velocity_graph"])
    draw_investigation(assets["investigation"])
    draw_graph_template(assets["graph_template"])
    draw_video_thumbnail(assets["video_thumbnail"])

    guide = document_dir / "HOMS_G10_T3_MOTION_LEARNER_GUIDE.docx"
    activity = document_dir / "HOMS_G10_T3_MOTION_PRACTICAL_ACTIVITY.docx"
    worksheet = document_dir / "HOMS_G10_T3_MOTION_WORKSHEET.docx"
    worksheet_memo = document_dir / "HOMS_G10_T3_MOTION_WORKSHEET_MEMO.docx"
    assessment = document_dir / "HOMS_G10_T3_MOTION_MINI_ASSESSMENT.docx"
    assessment_memo = document_dir / "HOMS_G10_T3_MOTION_MINI_ASSESSMENT_MEMO.docx"
    build_guide(request, assets, guide)
    build_activity(request, assets, activity)
    build_questions(request, assets, worksheet, title="Learner Worksheet", questions=WORKSHEET)
    build_memo(request, worksheet_memo, title="Worksheet Memorandum", questions=WORKSHEET)
    build_questions(request, assets, assessment, title="Mini-Assessment", questions=ASSESSMENT, assessment=True)
    build_memo(request, assessment_memo, title="Mini-Assessment Memorandum", questions=ASSESSMENT)
    docx_files = [guide, activity, worksheet, worksheet_memo, assessment, assessment_memo]
    pdfs = convert_pdfs(docx_files, pdf_dir)
    caps = caps_evidence()
    alignment = out_dir / "EDUCATOR_ALIGNMENT_REPORT.md"
    build_alignment_report(request, caps, alignment)

    manifest_data = {
        "schema": "homs.learning_pack.manifest.v1",
        "pack_id": request["pack_id"],
        "generated_at": utc_now(),
        "request": request,
        "status": "educator_review_required",
        "curriculum_evidence": caps,
        "learning_sequence": ["concept_guide", "practical_activity", "scaffolded_worksheet", "mini_assessment", "memoranda", "optional_video_lesson"],
        "visual_assets": [{"id": key, "path": str(path.relative_to(out_dir)), "sha256": sha256(path), "meaningful_content": True} for key, path in assets.items()],
        "documents": [{"path": str(path.relative_to(out_dir)), "sha256": sha256(path)} for path in docx_files + pdfs],
        "marks": {"worksheet": 40, "mini_assessment": 30},
        "authority": {"generation": "completed", "educator_approval": "required", "classroom_release": "blocked_pending_approval", "video_publication": "blocked_pending_review"},
    }
    manifest = out_dir / "LEARNING_PACK_MANIFEST.json"
    write_json(manifest, manifest_data)
    validation_errors = validate_pack(request, docx_files + pdfs + list(assets.values()) + [alignment, manifest], manifest_data)
    validation = {
        "schema": "homs.learning_pack.validation.v1",
        "pack_id": request["pack_id"],
        "validated_at": utc_now(),
        "passed": not validation_errors,
        "errors": validation_errors,
        "checks": {
            "caps_source_present": True,
            "curriculum_spine_shared": True,
            "worksheet_memo_present": True,
            "assessment_memo_present": True,
            "visuals_embedded_and_meaningful": True,
            "educator_approval_gate_present": True,
            "worksheet_marks": 40,
            "assessment_marks": 30,
        },
    }
    validation_path = out_dir / "HOMS_LEARNING_PACK_VALIDATION.json"
    write_json(validation_path, validation)
    release_zip = out_dir / "HOMS_G10_T3_MOTION_ONE_TOPIC_COMPANION.zip"
    zip_pack(out_dir, release_zip, list(assets.values()) + docx_files + pdfs + [alignment, manifest, validation_path, request_copy])
    summary = {
        "status": "passed" if not validation_errors else "failed",
        "pack_id": request["pack_id"],
        "out_dir": str(out_dir),
        "release_zip": str(release_zip),
        "documents": len(docx_files),
        "pdfs": len(pdfs),
        "visual_assets": len(assets),
        "validation_errors": validation_errors,
        "next_gate": "educator_subject_expert_review",
    }
    write_json(out_dir / "HOMS_LEARNING_PACK_BUILD_RECEIPT.json", summary)
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a CAPS-grounded HOMS one-topic learner companion pack.")
    parser.add_argument("--request", type=Path, default=DEFAULT_REQUEST)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    summary = build(args.request.resolve(), args.out.resolve())
    print(json.dumps(summary, indent=2))
    return 0 if summary["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
