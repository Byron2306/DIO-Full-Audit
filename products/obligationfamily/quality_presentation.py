"""Human-readable obligation-family reports.

The evidence pack remains the canonical machine-readable source.
These renderers do not make fulfilment, award, permit, or release decisions.
"""

from __future__ import annotations

import html
from io import BytesIO
from typing import Any


def _value(value: Any, fallback: str = "Not supplied") -> str:
    if value is None or value == "":
        return fallback
    if isinstance(value, (dict, list)):
        return fallback
    return str(value)



def report_sections(title: str, pack: dict[str, Any]):
    ledger = pack.get("requirement_or_obligation_ledger") or []
    evidence = pack.get("evidence_map") or []
    contested = pack.get("contested_state_register") or []
    missing = pack.get("missing_evidence_register") or []
    deadlines = pack.get("deadline_register") or []
    gate = pack.get("human_review_register") or {}

    sections = [
        ("Executive overview", [
            f"Product: {title}",
            f"Case: {_value(pack.get('case_id'))}",
            f"Requirements or obligations identified: {len(ledger)}",
            f"Evidence mappings recorded: {len(evidence)}",
            f"Items requiring review: {len(contested)}",
            "This document is an internal evidence-review candidate, "
            "not a final approval or determination.",
        ]),
    ]

    obligations = []
    for index, row in enumerate(ledger, 1):
        evaluation = row.get("evaluation") or {}
        statement = (
            row.get("statement")
            or row.get("text")
            or row.get("requirement_text")
            or row.get("description")
            or row.get("raw_text")
        )
        source_record = row.get("source") or {}
        locator = (
            row.get("source_locator")
            or source_record.get("locator")
            or row.get("locator")
            or row.get("obligation_id")
        )
        status = evaluation.get("status") or "NEEDS_REVIEW"

        obligations.append(
            f"{index}. [{_value(status)}] "
            f"{_value(locator, 'Unlocated')}: "
            f"{_value(statement, 'Statement unavailable; inspect source evidence.')}"
        )

        basis = evaluation.get("status_basis") or row.get("status_basis") or []
        if isinstance(basis, list) and basis:
            obligations.append("Review basis: " + "; ".join(map(str, basis)))

        required = evaluation.get("missing_evidence_requirements") or []
        if required:
            obligations.append(
                "Evidence requiring confirmation: "
                + "; ".join(map(str, required))
            )

        obligations.append(
            "Responsible party: "
            + _value(row.get("responsible_party"), "Not assigned")
        )

        if row.get("due_at"):
            obligations.append("Recorded deadline: " + str(row["due_at"]))

    sections.append((
        "Requirements and obligation review",
        obligations or ["No requirements were extracted. Human review required."],
    ))

    evidence_lines = []
    for row in evidence:
        evidence_lines.append(
            f"Evidence {_value(row.get('evidence_id'))}: "
            f"{_value(row.get('source_ref'))}; "
            f"relation={_value(row.get('relation'), 'unresolved')}; "
            f"trust={_value(row.get('trust_state'), 'unknown')}; "
            f"freshness={_value(row.get('freshness_state'), 'unknown')}; "
            f"obligation={_value(row.get('obligation_id') or ', '.join(map(str, row.get('obligation_ids') or [])), 'not mapped')}"
        )

    sections.append((
        "Source and evidence register",
        evidence_lines or ["No evidence mappings supplied."],
    ))

    review_lines = []
    for row in contested:
        review_lines.append(
            f"{_value(row.get('obligation_id'))}: "
            f"{_value(row.get('status'), 'NEEDS_REVIEW')}; "
            f"contradicting evidence: "
            f"{', '.join(map(str, row.get('contradicting_evidence_ids') or [])) or 'none identified'}"
        )

    sections.append((
        "Contradictions and outstanding review",
        review_lines or ["No contested items recorded by this assessment."],
    ))

    gap_lines = []
    for row in missing:
        if isinstance(row, dict):
            gap_lines.append(
                f"{_value(row.get('requirement_id') or row.get('obligation_id'))}: "
                f"{_value(row.get('reason') or row.get('description') or row.get('status'), 'Evidence gap requires review')}"
            )
        else:
            gap_lines.append(str(row))

    sections.append((
        "Missing evidence",
        gap_lines or [
            "The separate gap register is empty. This does not establish "
            "evidence sufficiency; review each obligation's assessment."
        ],
    ))

    deadline_lines = [
        f"{_value(row.get('due_at') or row.get('date'))}: "
        f"{_value(row.get('kind'), 'Deadline')} "
        f"({_value(row.get('state'), 'unverified')})"
        for row in deadlines
    ]
    sections.append((
        "Deadline register",
        deadline_lines or ["No explicit deadline entries recorded."],
    ))

    sections.append((
        "Recommended human review sequence",
        [
            "1. Confirm each extracted requirement against the original source.",
            "2. Investigate contested statements and evidence gaps.",
            "3. Verify any dates, conditions and source mappings before reliance.",
            "4. Obtain the required authorised human decision before disclosure.",
        ],
    ))

    sections.append((
        "Authority and release boundary",
        [
            f"Fulfilment: {_value(gate.get('fulfilment'), 'NEEDS_YOU')}",
            f"Disclosure: {_value(gate.get('disclosure'), 'NEEDS_YOU')}",
            f"External release: {_value(gate.get('external_release'), 'REFUSE')}",
            "No award, permit, legal determination, waiver or external "
            "release is authorised by this report.",
        ],
    ))

    return sections


def render_html(title: str, pack: dict[str, Any]) -> bytes:
    sections = report_sections(title, pack)

    body = []
    for heading, lines in sections:
        body.append(f"<section><h2>{html.escape(heading)}</h2>")
        for line in lines:
            body.append(f"<p>{html.escape(line)}</p>")
        body.append("</section>")

    content = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)} | Evidence Review</title>
<style>
body {{ margin:0; background:#eef2f7; color:#1a2638;
       font:16px/1.65 Arial,sans-serif; }}
main {{ max-width:920px; margin:30px auto; padding:48px;
        background:white; }}
header {{ border-bottom:4px solid #203b64; padding-bottom:22px; }}
h1 {{ color:#183458; font-size:34px; }}
h2 {{ color:#24527c; margin-top:34px; }}
p {{ overflow-wrap:anywhere; }}
.notice {{ background:#fff6e5; border-left:5px solid #b87916;
           padding:16px; }}
section {{ border-bottom:1px solid #e3e8ef; padding-bottom:14px; }}
@media print {{
 body {{ background:white; }}
 main {{ margin:0; padding:20px; }}
}}
</style>
</head>
<body>
<main>
<header>
<p>DIO | Controlled Professional Evidence Review</p>
<h1>{html.escape(title)}</h1>
<p>Case {html.escape(_value(pack.get('case_id')))}</p>
</header>
<p class="notice"><strong>HUMAN REVIEW REQUIRED.</strong>
This is a controlled evidence-review candidate, not an authorised
external deliverable.</p>
{''.join(body)}
</main>
</body>
</html>"""
    return content.encode("utf-8")


def render_docx(title: str, pack: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph("Controlled evidence review | Human decision required")
    doc.add_paragraph(f"Case: {_value(pack.get('case_id'))}")

    for heading, lines in report_sections(title, pack):
        doc.add_heading(heading, level=1)
        for line in lines:
            doc.add_paragraph(line)

    for section in doc.sections:
        section.header.paragraphs[0].text = "DIO | Internal review candidate"

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)

    output = BytesIO()
    doc.save(output)
    # Stable package metadata makes repeated proof generation reproducible.
    import zipfile
    normalized = BytesIO()
    with zipfile.ZipFile(BytesIO(output.getvalue())) as source, zipfile.ZipFile(normalized, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0)); info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return normalized.getvalue()


def render_pdf(title: str, pack: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, KeepTogether
    )

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4, invariant=1,
        leftMargin=46,
        rightMargin=46,
        topMargin=52,
        bottomMargin=52,
    )

    title_style = ParagraphStyle(
        "ReportTitle", fontName="Helvetica-Bold",
        fontSize=20, leading=25, spaceAfter=18,
        textColor=colors.HexColor("#183458"),
    )
    heading_style = ParagraphStyle(
        "SectionHeading", fontName="Helvetica-Bold",
        fontSize=13, leading=17, spaceBefore=14,
        spaceAfter=8, textColor=colors.HexColor("#24527c"),
    )
    body_style = ParagraphStyle(
        "ReportBody", fontName="Helvetica",
        fontSize=9, leading=14, spaceAfter=5,
        alignment=TA_LEFT, splitLongWords=True,
    )

    story = [
        Paragraph(html.escape(title), title_style),
        Paragraph(
            "CONTROLLED EVIDENCE REVIEW | HUMAN DECISION REQUIRED",
            body_style,
        ),
        Spacer(1, 10),
    ]

    for heading, lines in report_sections(title, pack):
        heading_block = Paragraph(html.escape(heading), heading_style)
        if heading == "Authority and release boundary":
            story.append(KeepTogether([heading_block] + [Paragraph(html.escape(line), body_style) for line in lines]))
        elif heading == "Requirements and obligation review":
            import re
            groups = []
            for line in lines:
                if not groups or re.match(r"^\d+\. \[", line): groups.append([])
                groups[-1].append(Paragraph(html.escape(line), body_style))
            for index, group in enumerate(groups):
                story.append(KeepTogether(([heading_block] if index == 0 else []) + group))
        else:
            story.append(KeepTogether([heading_block, Paragraph(html.escape(lines[0]), body_style)]))
            story.extend(Paragraph(html.escape(line), body_style) for line in lines[1:])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#667085'))
        canvas.drawString(46, 30, 'DIO | Internal review candidate')
        canvas.drawRightString(A4[0]-46, 30, f'Page {doc.page}')
        canvas.restoreState()
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
